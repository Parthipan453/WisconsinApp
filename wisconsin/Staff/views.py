from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from .models import StaffProfile, FacultyProfile
from .models import (
    StaffLeave,
    StaffTraining,
    StaffAccessRole,
)
from Staff.models import (
    StaffProfile,
    StaffPosition,
    StaffAddress,
    StaffEmergencyContact,
    StaffEducation,
    StaffCertification,
    StaffAccessRole,
    StaffDocument,
    AdvisorAssignment,
)
from django.db.models import Count, Q
import re
from datetime import date
from django.db import transaction
from django.http import JsonResponse
from django.views.decorators.http import require_http_methods
from django.core.paginator import Paginator
from django.db.models import Q, Count
from Admin.bela_admin.models import Department, Course
from Students.models import Schedule
from Staff.utils import notify_submitter_ticket_resolved, notify_submitter_ticket_in_progress
from Staff.utils import notify_submitter_ticket_resolved, notify_submitter_ticket_in_progress
from Staff.utils import (
    notify_submitter_ticket_resolved,
    notify_submitter_ticket_in_progress,
    notify_staff_new_financial_aid,
    notify_submitter_financial_aid_status,
    notify_students_exam_scheduled,
    notify_students_exam_room_assigned,
    notify_faculty_invigilation_assigned,
)

from Students.models import CourseSection, Semester, StudentEnrollment
from .forms import CourseSectionForm
from Faculty.models import FacultyCourseAssignment
########### Rixie Code ######
from datetime import datetime, date
from django.db.models import Count


from django.db.models import Count, Q
from datetime import timedelta
from Staff.models import StaffLeaveRequest, StaffLeaveBalance, MessageAttachment
from Students.models import StudentFinancialAid, StudentLeaveRequest, SupportTicket
from Faculty.models import LeaveRequest as FacultyLeaveRequest

from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.utils import timezone
from Students.models import StudentLeaveRequest
from Staff.utils import notify_student_leave_approved, notify_student_leave_rejected
import json, pyotp, qrcode

from datetime import date, timedelta
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from django.db.models import Q

from Admin.bela_admin.models import Department
from Staff.models import StaffLeaveRequest, StaffLeaveBalance, Announcement
from Students.models import StudentFinancialAid, StudentLeaveRequest, SupportTicket
from Faculty.models import LeaveRequest as FacultyLeaveRequest, FacultyCourseAssignment
from django.http import JsonResponse
from django.template.loader import render_to_string
from Staff.utils import create_resource_alert, get_resource_alerts_context, mark_resource_alerts_seen

from Admin.audit import AuditLogger


# from Staff.utils import notify_staff_new_leave_request 

# from Staff.utils import (
#     notify_student_leave_approved,
#     notify_student_leave_rejected,
#     notify_staff_new_financial_aid,
#     notify_submitter_financial_aid_status,
# )




@login_required
def staff_dashboard(request, uuid):
    
    #  CRITICAL CHECKS
    if not request.user.is_authenticated:
        return redirect('/login/')
    
    if not request.user.is_staff:
        return redirect('/student/dashboard/' + str(request.user.uuid) + '/')
    
    if str(request.user.uuid) != str(uuid):
        return redirect('/staff/' + str(request.user.uuid) + '/')
    
    #  Get staff profile
    try:
        sp = request.user.staff_profile
    except Exception as e:
        print(f"ERROR: {e}")
        return redirect('/login/')


    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    try:
        sp = request.user.staff_profile
    except Exception:
        return redirect("login")

    current_position = sp.positions.filter(position_status="ACTIVE").first()

    department = None
    if current_position and current_position.department_id:
        department = Department.objects.filter(
            department_id=current_position.department_id
        ).first()

    hour = timezone.localtime().hour
    if hour < 12:
        greeting = "Good Morning"
    elif hour < 17:
        greeting = "Good Afternoon"
    else:
        greeting = "Good Evening"

    today = timezone.localdate()
    month_start = today.replace(day=1)

    pending_student_leave = StudentLeaveRequest.objects.filter(status="pending")
    pending_faculty_leave = FacultyLeaveRequest.objects.filter(status="pending")
    pending_financial_aid = StudentFinancialAid.objects.filter(
        status__in=["PENDING", "UNDER_REVIEW"]
    )
    open_tickets = SupportTicket.objects.exclude(status="RESOLVED")

    requests_count = (
        pending_student_leave.count()
        + pending_faculty_leave.count()
        + pending_financial_aid.count()
    )
    open_tasks_count = requests_count + open_tickets.count()

    resolved_tickets_month = SupportTicket.objects.filter(
        status="RESOLVED", resolved_at__gte=month_start
    )
    awarded_aid_month = StudentFinancialAid.objects.filter(
        status="AWARDED", updated_date__gte=month_start
    )
    approved_leave_month = StudentLeaveRequest.objects.filter(
        status="approved", approved_date__gte=month_start
    )

    assisted_student_ids = set()
    assisted_student_ids.update(
        resolved_tickets_month.values_list("submitted_by_id", flat=True)
    )
    assisted_student_ids.update(
        awarded_aid_month.values_list("student__user_id", flat=True)
    )
    assisted_student_ids.update(
        approved_leave_month.values_list("student_id", flat=True)
    )
    students_assisted = len(assisted_student_ids)

    leave_balance, _ = StaffLeaveBalance.objects.get_or_create(
        staff=request.user,
        defaults={
            "annual_leave_balance": 30.0,
            "sick_leave_balance": 15.0,
            "casual_leave_balance": 10.0,
            "earned_leave_balance": 0.0,
            "compensatory_leave_balance": 0.0,
        }
    )
    leave_days = leave_balance.annual_leave_balance

    training_completed = sp.training_records.filter(training_status="COMPLETED").count()
    training_total = sp.training_records.exclude(training_status="CANCELLED").count()

    ticket_completed = SupportTicket.objects.filter(status="RESOLVED").count()
    ticket_in_progress = SupportTicket.objects.filter(status="IN_PROGRESS").count()
    ticket_pending = SupportTicket.objects.filter(status="OPEN").count()
    donut_total = ticket_completed + ticket_in_progress + ticket_pending

    CIRCUMFERENCE = 314.16
    if donut_total:
        completed_len = round((ticket_completed / donut_total) * CIRCUMFERENCE, 2)
        inprogress_len = round((ticket_in_progress / donut_total) * CIRCUMFERENCE, 2)
        pending_len = round((ticket_pending / donut_total) * CIRCUMFERENCE, 2)
    else:
        completed_len = inprogress_len = pending_len = 0

    donut_data = {
        "total": donut_total,
        "completed": {
            "count": ticket_completed,
            "pct": round((ticket_completed / donut_total) * 100) if donut_total else 0,
            "dash": completed_len,
            "offset": 0,
        },
        "in_progress": {
            "count": ticket_in_progress,
            "pct": round((ticket_in_progress / donut_total) * 100) if donut_total else 0,
            "dash": inprogress_len,
            "offset": -completed_len,
        },
        "pending": {
            "count": ticket_pending,
            "pct": round((ticket_pending / donut_total) * 100) if donut_total else 0,
            "dash": pending_len,
            "offset": -(completed_len + inprogress_len),
        },
    }

    category_counts = [
        {"label": "Enrollment", "count": 0, "color": "#dc2626", "url_type": "all"},
        {"label": "Financial Aid", "count": pending_financial_aid.count(), "color": "#3b82f6", "url_type": "financial_aid"},
        {"label": "Student Leave", "count": pending_student_leave.count(), "color": "#22c55e", "url_type": "student_leave"},
        {"label": "Faculty Leave", "count": pending_faculty_leave.count(), "color": "#8b5cf6", "url_type": "faculty_leave"},
    ]
    max_count = max((c["count"] for c in category_counts), default=1) or 1
    for c in category_counts:
        c["pct"] = round((c["count"] / max_count) * 100) if max_count else 0

    months = []
    month_labels = []
    cursor = today.replace(day=1)
    for i in range(5, -1, -1):
        year = cursor.year
        month = cursor.month - i
        while month <= 0:
            month += 12
            year -= 1
        months.append((year, month))
        month_labels.append(date(year, month, 1).strftime("%b").upper())

    workload_data = []
    for year, month in months:
        count = (
            SupportTicket.objects.filter(
                status="RESOLVED", resolved_at__year=year, resolved_at__month=month
            ).count()
            + StudentFinancialAid.objects.filter(
                status="AWARDED", updated_date__year=year, updated_date__month=month
            ).count()
            + StudentLeaveRequest.objects.filter(
                status="approved", approved_date__year=year, approved_date__month=month
            ).count()
        )
        workload_data.append(count)

    workload_max = max(workload_data) if any(workload_data) else 10
    workload_axis_max = int(workload_max * 1.25) + 1

    def _as_date(value):
        if value is None:
            return date.min
        if hasattr(value, "date") and callable(value.date):
            return value.date()
        return value  

    upcoming_tasks = []
    for aid in pending_financial_aid.select_related("student__user").order_by("applied_date")[:2]:
        upcoming_tasks.append({
            "title": f"Process {aid.get_aid_type_display()} Application",
            "student_id": aid.student.student_number,
            "priority": "High" if aid.status == "PENDING" else "Medium",
            "date": aid.applied_date,
            "url_name": "staff_financial_aid_detail",
            "url_kwargs": {"uuid": uuid, "aid_id": aid.id},
        })
    for leave in pending_student_leave.select_related("student", "leave_type").order_by("start_date")[:2]:
        student_number = "N/A"
        if hasattr(leave.student, "student_profile"):
            student_number = getattr(leave.student.student_profile, "student_number", "N/A")
        upcoming_tasks.append({
            "title": f"Review {leave.leave_type.name} Leave Request",
            "student_id": student_number,
            "priority": "Medium",
            "date": leave.start_date,
            "url_name": "staff_student_leave_detail",
            "url_kwargs": {"uuid": uuid, "leave_id": leave.id},
        })
    from django.urls import reverse
    for task in upcoming_tasks:
        try:
            task["href"] = reverse(task.pop("url_name"), kwargs=task.pop("url_kwargs"))
        except Exception:
            task["href"] = "#"
            task.pop("url_name", None)
            task.pop("url_kwargs", None)

    upcoming_tasks = sorted(upcoming_tasks, key=lambda t: _as_date(t["date"]))[:4]

    recent_requests = []
    for aid in StudentFinancialAid.objects.exclude(status="PENDING").select_related(
        "student__user"
    ).order_by("-updated_date")[:2]:
        recent_requests.append({
            "title": f"{aid.get_aid_type_display()} Application",
            "student_id": aid.student.student_number,
            "status": aid.get_status_display(),
            "date": aid.updated_date,
            "url_name": "staff_financial_aid_detail",
            "url_kwargs": {"uuid": uuid, "aid_id": aid.id},
        })
    for ticket in SupportTicket.objects.exclude(status="OPEN").select_related(
        "submitted_by"
    ).order_by("-created_at")[:2]:
        student_number = "N/A"
        if hasattr(ticket.submitted_by, "student_profile"):
            student_number = getattr(ticket.submitted_by.student_profile, "student_number", "N/A")
        recent_requests.append({
            "title": ticket.subject,
            "student_id": student_number,
            "status": ticket.get_status_display(),
            "date": ticket.created_at,
            "url_name": "support_tickets_list",
            "url_kwargs": {"uuid": uuid},
        })
    for req in recent_requests:
        try:
            req["href"] = reverse(req.pop("url_name"), kwargs=req.pop("url_kwargs"))
        except Exception:
            req["href"] = "#"
            req.pop("url_name", None)
            req.pop("url_kwargs", None)

    recent_requests = sorted(recent_requests, key=lambda r: _as_date(r["date"]), reverse=True)[:4]

    announcements = Announcement.objects.filter(
        Q(expires_at__isnull=True) | Q(expires_at__gte=today)
    ).order_by("-posted_at")[:4]

    context = {
        "user": request.user,
        "today": today,
        "greeting": greeting,
        "position": current_position,
        "department": department,
        "open_tasks_count": open_tasks_count,
        "requests_count": requests_count,
        "students_assisted": students_assisted,
        "leave_days": leave_days,
        "training_completed": training_completed,
        "training_total": training_total or 1,
        "donut_data": donut_data,
        "category_counts": category_counts,
        "month_labels": month_labels,
        "workload_data": workload_data,
        "workload_axis_max": workload_axis_max,
        "upcoming_tasks": upcoming_tasks,
        "recent_requests": recent_requests,
        "announcements": announcements,
    }

    return render(request, "staff_dashboard.html", context)



from datetime import timedelta
from django.urls import reverse
from Staff.models import StaffLeaveRequest, Resource
from Faculty.models import LeaveRequest as FacultyLeaveRequest
from Students.models import StudentLeaveRequest


@login_required
def staff_work_overview(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    try:
        sp = request.user.staff_profile
    except Exception:
        return redirect("login")

    today = timezone.localdate()
    soon_cutoff = today + timedelta(days=60)

    training_qs = sp.training_records.all()
    training_completed = training_qs.filter(training_status="COMPLETED").count()
    training_in_progress = training_qs.filter(training_status="IN_PROGRESS").count()
    recent_training = training_qs.order_by("-completion_date")[:5]

    my_leave_qs = StaffLeaveRequest.objects.filter(staff=request.user).select_related("leave_type")
    leave_pending = my_leave_qs.filter(status="pending").count()
    leave_approved = my_leave_qs.filter(status="approved").count()
    recent_leave = my_leave_qs.order_by("-created_at")[:5]

    pending_faculty_leave = (
        FacultyLeaveRequest.objects.filter(status="pending")
        .select_related("faculty", "leave_type")
        .order_by("start_date")[:5]
    )
    pending_student_leave = (
        StudentLeaveRequest.objects.filter(status="pending")
        .select_related("student", "leave_type")
        .order_by("start_date")[:5]
    )
    approvals_needed_count = (
        FacultyLeaveRequest.objects.filter(status="pending").count()
        + StudentLeaveRequest.objects.filter(status="pending").count()
    )

    # ---------------- Certifications ----------------
    cert_qs = sp.certifications.all()
    certs_active = cert_qs.filter(
        Q(expiration_date__isnull=True) | Q(expiration_date__gte=today)
    ).count()
    upcoming_cert_expirations = cert_qs.filter(
        expiration_date__gte=today
    ).order_by("expiration_date")[:5]
    for c in upcoming_cert_expirations:
        c.is_urgent = c.expiration_date <= soon_cutoff
    certs_expiring_soon = cert_qs.filter(
        expiration_date__gte=today, expiration_date__lte=soon_cutoff
    ).count()

    recent_resources = Resource.objects.order_by("-uploaded_at")[:5]
    total_resources = Resource.objects.count()

    access_roles = sp.access_roles.filter(active=True).order_by("-assigned_date")

    latest_review = sp.performance_reviews.order_by("-review_date").first()

    on_leave_today_user_ids = set(
        StaffLeaveRequest.objects.filter(
            status="approved", start_date__lte=today, end_date__gte=today
        ).values_list("staff_id", flat=True)
    )

    team_members_active = StaffProfile.objects.filter(employment_status="ACTIVE").count()
    team_total = StaffProfile.objects.count()
    team_on_leave = len(on_leave_today_user_ids)

    context = {
        "today": today,

        "training_completed": training_completed,
        "training_in_progress": training_in_progress,
        "recent_training": recent_training,

        "leave_pending": leave_pending,
        "leave_approved": leave_approved,
        "recent_leave": recent_leave,

        "pending_faculty_leave": pending_faculty_leave,
        "pending_student_leave": pending_student_leave,
        "approvals_needed_count": approvals_needed_count,

        "certs_active": certs_active,
        "certs_expiring_soon": certs_expiring_soon,
        "upcoming_cert_expirations": upcoming_cert_expirations,

        "recent_resources": recent_resources,
        "total_resources": total_resources,

        "access_roles": access_roles,
        "latest_review": latest_review,

        "team_members_active": team_members_active,
        "team_total": team_total,
        "team_on_leave": team_on_leave,
    }

    return render(request, "Administrator/work_overview.html", context)

#<--------------------------BLAZE CODE START APPROVE AND REJECT(02.07.2024) ---------------------->


from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.db.models import Count, Sum, Q
from django.utils import timezone  
from decimal import Decimal
from datetime import datetime
from Staff.models import StaffLeave, StaffProfile
from Students.models import StudentFinancialAid, StudentProfile
from Students.models import StudentLeaveRequest, StudentLeaveType, StudentLeaveBalance
from Faculty.models import LeaveRequest as FacultyLeaveRequest
from Faculty.models import LeaveType
from django.http import JsonResponse, HttpResponse
import json
import csv


def get_status_color(status):
    colors = {
        'PENDING': '#c5850c',
        'UNDER_REVIEW': '#185fa5',
        'APPROVED': '#1a7a1a',
        'AWARDED': '#1a7a1a',
        'REJECTED': '#c5050c',
        'CANCELLED': '#888',
        'EXPIRED': '#888',
    }
    return colors.get(status, '#888')


def get_aid_icon(aid_type):
    icons = {
        'SCHOLARSHIP': 'ti ti-school',
        'GRANT': 'ti ti-clipboard-text',
        'FELLOWSHIP': 'ti ti-star',
        'ASSISTANTSHIP': 'ti ti-briefcase',
        'LOAN': 'ti ti-credit-card',
    }
    return icons.get(aid_type, 'ti ti-wallet')


def get_aid_status_color(status):
    colors = {
        'PENDING': '#c5850c',
        'UNDER_REVIEW': '#185fa5',
        'AWARDED': '#1a7a1a',
        'REJECTED': '#c5050c',
        'CANCELLED': '#888',
        'EXPIRED': '#888',
    }
    return colors.get(status, '#888')


@login_required
def staff_requests(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    
    try:
        sp = request.user.staff_profile
    except Exception:
        return redirect("login")
    
    student_leave_requests = StudentLeaveRequest.objects.select_related(
        'student', 'student__student_profile', 'leave_type'
    ).order_by('-created_at')
    
    total_student_leave = student_leave_requests.count()
    pending_student_leave = student_leave_requests.filter(status='pending').count()
    approved_student_leave = student_leave_requests.filter(status='approved').count()
    rejected_student_leave = student_leave_requests.filter(status='rejected').count()
    cancelled_student_leave = student_leave_requests.filter(status='cancelled').count()
    
    leave_requests = StaffLeave.objects.filter(staff=sp)
    total_leave = leave_requests.count()
    pending_leave = leave_requests.filter(approval_status="PENDING").count()
    approved_leave = leave_requests.filter(approval_status="APPROVED").count()
    rejected_leave = leave_requests.filter(approval_status="REJECTED").count()
    
    faculty_leave_requests = FacultyLeaveRequest.objects.all().order_by('-created_at')
    total_faculty_leave = faculty_leave_requests.count()
    pending_faculty_leave = faculty_leave_requests.filter(status='pending').count()
    approved_faculty_leave = faculty_leave_requests.filter(status='approved').count()
    rejected_faculty_leave = faculty_leave_requests.filter(status='rejected').count()
    cancelled_faculty_leave = faculty_leave_requests.filter(status='cancelled').count()
    
    from django.db import connection
    from decimal import Decimal
    
    try:
        with connection.cursor() as cursor:
            cursor.execute("""
                SELECT 
                    fa.id,
                    fa.aid_type,
                    fa.academic_year,
                    fa.reason,
                    fa.award_amount,
                    fa.status,
                    fa.applied_date,
                    fa.updated_date,
                    fa.student_id,
                    COALESCE(u.first_name, '') as first_name,
                    COALESCE(u.last_name, '') as last_name,
                    COALESCE(sp.student_number, 'N/A') as student_number
                FROM student_financial_aids fa
                LEFT JOIN student_profiles sp ON fa.student_id = sp.id
                LEFT JOIN users u ON sp.user_id = u.id
                ORDER BY fa.applied_date DESC
            """)
            rows = cursor.fetchall()
        
        financial_aid_list = []
        total_awarded_amount = Decimal("0.00")
        
        for row in rows:
            try:
                award_amount = Decimal("0.00")
                if row[4] is not None:
                    try:
                        award_amount = Decimal(str(row[4]))
                    except:
                        award_amount = Decimal("0.00")
                
                aid_data = {
                    'id': row[0],
                    'aid_type': row[1] or 'N/A',
                    'academic_year': row[2] or 'N/A',
                    'reason': row[3] or '',
                    'award_amount': award_amount,
                    'status': row[5] or 'PENDING',
                    'applied_date': row[6],
                    'updated_date': row[7],
                    'student_id': row[8],
                    'student_name': f"{row[9]} {row[10]}".strip() or "N/A",
                    'student_number': row[11] or "N/A",
                }
                financial_aid_list.append(aid_data)
                
                if aid_data['status'] == "AWARDED":
                    total_awarded_amount += aid_data['award_amount']
            except Exception as e:
                print(f"Error processing row {row[0]}: {e}")
                continue
        
        total_financial_aid = len(financial_aid_list)
        pending_financial_aid = len([a for a in financial_aid_list if a['status'] == "PENDING"])
        under_review_financial_aid = len([a for a in financial_aid_list if a['status'] == "UNDER_REVIEW"])
        awarded_financial_aid = len([a for a in financial_aid_list if a['status'] == "AWARDED"])
        rejected_financial_aid = len([a for a in financial_aid_list if a['status'] == "REJECTED"])
        cancelled_financial_aid = len([a for a in financial_aid_list if a['status'] == "CANCELLED"])
        
    except Exception as e:
        print(f"Error fetching financial aid: {e}")
        financial_aid_list = []
        total_financial_aid = 0
        pending_financial_aid = 0
        under_review_financial_aid = 0
        awarded_financial_aid = 0
        rejected_financial_aid = 0
        cancelled_financial_aid = 0
        total_awarded_amount = Decimal("0.00")
    
    request_type = request.GET.get('type', 'all')
    status_filter = request.GET.get('status', '')
    search_query = request.GET.get('search', '')
    
    requests_list = []
    
    for leave in student_leave_requests:
        requests_list.append({
            'id': leave.id,
            'type': 'student_leave',
            'type_display': 'Student Leave',
            'title': f"{leave.leave_type.name} Leave",
            'description': f"{leave.start_date} to {leave.end_date} ({leave.days_count} days)",
            'status': leave.status,
            'status_display': leave.get_status_display(),
            'created_date': leave.created_at,
            'student_name': leave.student.get_full_name(),
            'student_id': leave.student.student_profile.student_number if hasattr(leave.student, 'student_profile') else 'N/A',
            'url': f'/staff/{uuid}/student-leave-detail/{leave.id}/',
            'icon': 'ti ti-calendar',
            'color': '#c5050c',
            'category': 'student_leave',
            'is_faculty': False,
            'days_count': leave.days_count,
            'half_day': leave.get_half_day_type_display() if leave.half_day_type != 'full' else 'Full Day',
            'reason': leave.reason,
            'leave_type_name': leave.leave_type.name,
            'leave_type_color': leave.leave_type.color,
            'department': 'Student'
        })
    
    for leave in leave_requests:
        requests_list.append({
            'id': leave.id,
            'type': 'staff_leave',
            'type_display': 'Staff Leave',
            'title': f"{leave.get_leave_type_display()} Leave",
            'description': f"{leave.start_date} to {leave.end_date}",
            'status': leave.approval_status,
            'status_display': leave.get_approval_status_display(),
            'created_date': leave.start_date,
            'student_name': leave.staff.user.get_full_name() if leave.staff else 'N/A',
            'student_id': leave.staff.employee_id if leave.staff else 'N/A',
            'url': f'/staff/{uuid}/staff-leave/{leave.id}/',
            'icon': 'ti ti-calendar',
            'color': '#185fa5',
            'category': 'staff_leave',
            'is_faculty': False,
            'department': 'Staff'
        })
    
    for leave in faculty_leave_requests:
        requests_list.append({
            'id': leave.id,
            'type': 'faculty_leave',
            'type_display': 'Faculty Leave',
            'title': f"{leave.leave_type.name} Leave",
            'description': f"{leave.start_date} to {leave.end_date} ({leave.days_count} days)",
            'status': leave.status,
            'status_display': leave.display_status,
            'created_date': leave.created_at,
            'student_name': leave.faculty.get_full_name() if leave.faculty else 'N/A',
            'student_id': str(leave.faculty.uuid) if leave.faculty else 'N/A',
            'url': f'/staff/{uuid}/faculty-leave/{leave.id}/',
            'icon': 'ti ti-calendar',
            'color': '#7c3aed',
            'category': 'faculty_leave',
            'is_faculty': True,
            'faculty_id': leave.faculty.id if leave.faculty else None,
            'days_count': leave.days_count,
            'half_day': leave.get_half_day_type_display() if leave.half_day_type != 'full' else 'Full Day',
            'reason': leave.reason,
            'leave_type_name': leave.leave_type.name,
            'department': 'Faculty'
        })
    
    for aid in financial_aid_list:
        requests_list.append({
            'id': aid['id'],
            'type': 'financial_aid',
            'type_display': 'Financial Aid',
            'title': aid['aid_type'],
            'description': f"Academic Year: {aid['academic_year']}",
            'status': aid['status'],
            'status_display': aid['status'],
            'created_date': aid['applied_date'],
            'student_name': aid['student_name'],
            'student_id': aid['student_number'],
            'award_amount': aid['award_amount'],
            'reason': aid['reason'],
            'url': f'/staff/{uuid}/financial-aid/{aid["id"]}/',
            'icon': 'ti ti-wallet',
            'color': '#c5050c',
            'category': 'financial_aid',
            'is_faculty': False,
            'department': 'Financial Aid'
        })
    
    from django.utils import timezone
    from datetime import datetime
    
    def get_safe_date(item):
        date_val = item.get('created_date')
        if date_val is None:
            return timezone.make_aware(datetime(1900, 1, 1))
        
        if hasattr(date_val, 'year') and not hasattr(date_val, 'hour'):
            dt = datetime(date_val.year, date_val.month, date_val.day)
            return timezone.make_aware(dt)
        
        if timezone.is_naive(date_val):
            return timezone.make_aware(date_val)
        
        return date_val
    
    requests_list.sort(key=get_safe_date, reverse=True)
    
    if request_type != 'all':
        if request_type == 'leave':
            requests_list = [r for r in requests_list if r['type'] in ['staff_leave', 'faculty_leave', 'student_leave']]
        else:
            requests_list = [r for r in requests_list if r['type'] == request_type]
    
    if status_filter:
        requests_list = [r for r in requests_list if r['status'] == status_filter]
    
    if search_query:
        search_query = search_query.lower()
        requests_list = [
            r for r in requests_list 
            if search_query in r['student_name'].lower() 
            or search_query in r['student_id'].lower()
            or search_query in r['title'].lower()
        ]
    
    total_requests = len(requests_list)
    
    status_counts = {}
    for req in requests_list:
        status_counts[req['status']] = status_counts.get(req['status'], 0) + 1
    
    recent_approvals = []
    
    def get_safe_approval_date(item):
        date_val = item.get('date')
        if date_val is None:
            return timezone.make_aware(datetime(1900, 1, 1))
        
        if hasattr(date_val, 'year') and not hasattr(date_val, 'hour'):
            dt = datetime(date_val.year, date_val.month, date_val.day)
            return timezone.make_aware(dt)
        
        if timezone.is_naive(date_val):
            return timezone.make_aware(date_val)
        
        return date_val
    
    recent_student_leave = student_leave_requests.filter(
        status="approved"
    ).order_by("-approved_date")[:3]
    
    for leave in recent_student_leave:
        recent_approvals.append({
            'type': 'Student Leave',
            'title': f"{leave.leave_type.name} Leave",
            'student': leave.student.get_full_name(),
            'date': leave.approved_date or leave.updated_at,
        })
    
    recent_leave = leave_requests.filter(
        approval_status="APPROVED"
    ).order_by("-end_date")[:3]
    
    for leave in recent_leave:
        recent_approvals.append({
            'type': 'Staff Leave',
            'title': f"{leave.get_leave_type_display()} Leave",
            'student': leave.staff.user.get_full_name() if leave.staff else 'N/A',
            'date': leave.end_date,
        })
    
    recent_faculty_leave = faculty_leave_requests.filter(
        status="approved"
    ).order_by("-approved_date")[:3]
    
    for leave in recent_faculty_leave:
        recent_approvals.append({
            'type': 'Faculty Leave',
            'title': f"{leave.leave_type.name} Leave",
            'student': leave.faculty.get_full_name() if leave.faculty else 'N/A',
            'date': leave.approved_date or leave.updated_at,
        })
    
    awarded_aid = [a for a in financial_aid_list if a['status'] == "AWARDED"]
    awarded_aid.sort(
        key=lambda x: get_safe_approval_date({'date': x['updated_date']}), 
        reverse=True
    )
    recent_aid = awarded_aid[:3]
    
    for aid in recent_aid:
        recent_approvals.append({
            'type': 'Financial Aid',
            'title': f"{aid['aid_type']} - ${aid['award_amount']}",
            'student': aid['student_name'],
            'date': aid['updated_date'],
        })
    
    recent_approvals.sort(key=get_safe_approval_date, reverse=True)
    recent_approvals = recent_approvals[:5]
    
    request_types = [
        {
            'type': 'student_leave',
            'display': 'Student Leave',
            'count': total_student_leave,
            'pending': pending_student_leave,
            'icon': 'ti ti-users',
            'color': '#c5050c',
        },
        {
            'type': 'faculty_leave',
            'display': 'Faculty Leave',
            'count': total_faculty_leave,
            'pending': pending_faculty_leave,
            'icon': 'ti ti-school',
            'color': '#7c3aed',
        },
        {
            'type': 'staff_leave',
            'display': 'Staff Leave',
            'count': total_leave,
            'pending': pending_leave,
            'icon': 'ti ti-user',
            'color': '#185fa5',
        },
        {
            'type': 'financial_aid',
            'display': 'Financial Aid',
            'count': total_financial_aid,
            'pending': pending_financial_aid + under_review_financial_aid,
            'icon': 'ti ti-wallet',
            'color': '#c5050c',
        },
    ]
    
    context = {
        "user": request.user,
        "uuid": uuid,
        "total_requests": total_requests,
        "total_student_leave": total_student_leave,
        "pending_student_leave": pending_student_leave,
        "approved_student_leave": approved_student_leave,
        "rejected_student_leave": rejected_student_leave,
        "cancelled_student_leave": cancelled_student_leave,
        "total_leave": total_leave,
        "pending_leave": pending_leave,
        "approved_leave": approved_leave,
        "rejected_leave": rejected_leave,
        "total_faculty_leave": total_faculty_leave,
        "pending_faculty_leave": pending_faculty_leave,
        "approved_faculty_leave": approved_faculty_leave,
        "rejected_faculty_leave": rejected_faculty_leave,
        "cancelled_faculty_leave": cancelled_faculty_leave,
        "total_financial_aid": total_financial_aid,
        "pending_financial_aid": pending_financial_aid,
        "under_review_financial_aid": under_review_financial_aid,
        "awarded_financial_aid": awarded_financial_aid,
        "rejected_financial_aid": rejected_financial_aid,
        "cancelled_financial_aid": cancelled_financial_aid,
        "total_awarded_amount": total_awarded_amount,
        "requests": requests_list,
        "status_counts": status_counts,
        "request_types": request_types,
        "recent_approvals": recent_approvals,
        "current_type": request_type,
        "current_status": status_filter,
        "search_query": search_query,
        "status_choices": [
            ('PENDING', 'Pending'),
            ('UNDER_REVIEW', 'Under Review'),
            ('APPROVED', 'Approved'),
            ('AWARDED', 'Awarded'),
            ('REJECTED', 'Rejected'),
            ('CANCELLED', 'Cancelled'),
        ],
    }
    
    return render(request, "Administrator/requests.html", context)


# ===================== HELPER FUNCTIONS =====================

def get_status_color(status):
    colors = {
        'PENDING': '#c5850c',
        'UNDER_REVIEW': '#185fa5',
        'APPROVED': '#1a7a1a',
        'AWARDED': '#1a7a1a',
        'REJECTED': '#c5050c',
        'CANCELLED': '#888',
        'EXPIRED': '#888',
    }
    return colors.get(status, '#888')


def get_aid_icon(aid_type):
    icons = {
        'SCHOLARSHIP': 'ti ti-school',
        'GRANT': 'ti ti-clipboard-text',
        'FELLOWSHIP': 'ti ti-star',
        'ASSISTANTSHIP': 'ti ti-briefcase',
        'LOAN': 'ti ti-credit-card',
    }
    return icons.get(aid_type, 'ti ti-wallet')


def get_aid_status_color(status):
    colors = {
        'PENDING': '#c5850c',
        'UNDER_REVIEW': '#185fa5',
        'AWARDED': '#1a7a1a',
        'REJECTED': '#c5050c',
        'CANCELLED': '#888',
        'EXPIRED': '#888',
    }
    return colors.get(status, '#888')


def update_student_leave_balance(leave):
    try:
        balance, created = StudentLeaveBalance.objects.get_or_create(
            student=leave.student,
            defaults={
                'annual_leave_balance': 30.0,
                'sick_leave_balance': 15.0,
                'casual_leave_balance': 10.0,
                'earned_leave_balance': 0.0,
                'compensatory_leave_balance': 0.0,
            }
        )
        
        leave_type_code = leave.leave_type.code.upper()
        days_used = leave.days_count
        
        if leave_type_code == 'ANNUAL':
            balance.annual_leave_balance = max(0, balance.annual_leave_balance - days_used)
        elif leave_type_code == 'SICK':
            balance.sick_leave_balance = max(0, balance.sick_leave_balance - days_used)
        elif leave_type_code == 'CASUAL':
            balance.casual_leave_balance = max(0, balance.casual_leave_balance - days_used)
        elif leave_type_code == 'EARNED':
            balance.earned_leave_balance = max(0, balance.earned_leave_balance - days_used)
        elif leave_type_code == 'COMP':
            balance.compensatory_leave_balance = max(0, balance.compensatory_leave_balance - days_used)
        
        balance.save()
        
    except Exception as e:
        print(f"Error updating balance: {e}")


#<--------------------Blaze Code Start (24.07.26)---------------------->

# Staff/views.py

from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.db.models import Count, Sum, Q
from django.utils import timezone  
from decimal import Decimal
from datetime import datetime
from Staff.models import StaffLeave, StaffProfile, StaffLeaveBalance
from Students.models import StudentFinancialAid, StudentProfile
from Students.models import StudentLeaveRequest, StudentLeaveType, StudentLeaveBalance
from Faculty.models import LeaveRequest as FacultyLeaveRequest
from Faculty.models import LeaveType
from django.http import JsonResponse, HttpResponse
import json
import csv

from Admin.notifications import (
    notify_faculty_leave_approved, 
    notify_faculty_leave_rejected,
    notify_staff_new_faculty_leave_request
)


# ... rest of your staff_requests function ...


@login_required
def staff_faculty_leave_detail(request, uuid, leave_id):
    """
    View and approve/reject faculty leave requests
    """
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'message': 'Access denied'})
    
    leave = get_object_or_404(FacultyLeaveRequest, id=leave_id)
    
    # Check if this is an AJAX request
    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
    
    if request.method == 'POST':
        action = request.POST.get('action')
        remarks = request.POST.get('remarks', '')
        
        if action == 'approve':
            before_data = AuditLogger.model_to_dict(
                leave,
                fields=[
                    "leave_type",
                    "start_date",
                    "end_date",
                    "reason",
                    "status",
                    "approval_remarks",
                ]
            )
            if leave.status != 'pending':
                if is_ajax:
                    return JsonResponse({'success': False, 'message': 'This request has already been processed.'})
                messages.error(request, 'This request has already been processed.')
            else:
                leave.status = 'approved'
                leave.approved_by = request.user
                leave.approved_date = timezone.now()
                leave.approval_remarks = remarks
                leave.save()
                # =========================================
                # AUDIT AFTER DATA
                # =========================================
                after_data = AuditLogger.model_to_dict(
                    leave,
                    fields=[
                        "leave_type",
                        "start_date",
                        "end_date",
                        "reason",
                        "status",
                        "approval_remarks",
                        "approved_by",
                        "approved_date",
                    ]
                )
                # =========================================
                # AUDIT LOG
                # =========================================

                AuditLogger.log(
                    request=request,
                    action="APPROVE",
                    module="Staff",
                    object_type="Faculty Leave Request",
                    object_id=leave.id,
                    description=(
                        f'Staff approved the faculty leave request '
                        f'from {leave.start_date} to {leave.end_date}.'
                    ),
                    before_data=before_data,
                    after_data=after_data,
                    status="SUCCESS",
                )

                # ✅ Send notification to FACULTY (not staff)
                try:
                    notify_faculty_leave_approved(leave)
                    print(f"✅ Approval notification sent to: {leave.faculty.get_full_name()}")
                except Exception as e:
                    print(f"❌ Notification error: {e}")
                
                if is_ajax:
                    return JsonResponse({
                        'success': True,
                        'message': f"✅ Leave approved for {leave.faculty.get_full_name()}"
                    })
                messages.success(request, f"✅ Leave approved for {leave.faculty.get_full_name()}")
                return redirect('staff_requests', uuid=uuid)
        
        elif action == 'reject':

            # =========================================
            # AUDIT BEFORE DATA
            # =========================================
            before_data = AuditLogger.model_to_dict(
                leave,
                fields=[
                    "leave_type",
                    "start_date",
                    "end_date",
                    "reason",
                    "status",
                    "approval_remarks",
                ]
            )
            if leave.status != 'pending':
                if is_ajax:
                    return JsonResponse({'success': False, 'message': 'This request has already been processed.'})
                messages.error(request, 'This request has already been processed.')
            else:
                leave.status = 'rejected'
                leave.approved_by = request.user
                leave.approved_date = timezone.now()
                leave.approval_remarks = remarks
                leave.save()
                # =========================================
                # AUDIT AFTER DATA
                # =========================================
                after_data = AuditLogger.model_to_dict(
                    leave,
                    fields=[
                        "leave_type",
                        "start_date",
                        "end_date",
                        "reason",
                        "status",
                        "approval_remarks",
                        "approved_by",
                        "approved_date",
                    ]
                )
                # =========================================
                # AUDIT LOG
                # =========================================
                AuditLogger.log(
                    request=request,
                    action="REJECT",
                    module="Staff",
                    object_type="Faculty Leave Request",
                    object_id=leave.id,
                    description=(
                        f'Staff rejected the faculty leave request '
                        f'from {leave.start_date} to {leave.end_date}.'
                    ),
                    before_data=before_data,
                    after_data=after_data,
                    status="SUCCESS",
                )

                # ✅ Send notification to FACULTY (not staff)
                try:
                    notify_faculty_leave_rejected(leave)
                    print(f"❌ Rejection notification sent to: {leave.faculty.get_full_name()}")
                except Exception as e:
                    print(f"❌ Notification error: {e}")
                
                if is_ajax:
                    return JsonResponse({
                        'success': True,
                        'message': f"❌ Leave rejected for {leave.faculty.get_full_name()}"
                    })
                messages.success(request, f"❌ Leave rejected for {leave.faculty.get_full_name()}")
                return redirect('staff_requests', uuid=uuid)
        
        elif action == 'under_review':
            if leave.status != 'pending':
                if is_ajax:
                    return JsonResponse({'success': False, 'message': 'This request has already been processed.'})
                messages.error(request, 'This request has already been processed.')
            else:
                if remarks:
                    leave.approval_remarks = remarks
                leave.save()
                
                if is_ajax:
                    return JsonResponse({
                        'success': True,
                        'message': f"📝 Leave marked for review - {leave.faculty.get_full_name()}"
                    })
                messages.info(request, f"📝 Leave marked for review - {leave.faculty.get_full_name()}")
                return redirect('staff_faculty_leave_detail', uuid=uuid, leave_id=leave_id)
        
        if is_ajax:
            return JsonResponse({'success': False, 'message': 'Invalid action'})
        return redirect('staff_faculty_leave_detail', uuid=uuid, leave_id=leave_id)
    
    # For GET request - render detail page
    context = {
        'leave': leave,
        'uuid': uuid,
    }
    return render(request, 'Administrator/faculty_leave_detail.html', context)

#<--------------------Blaze Code End (24.07.26)---------------------->

#<--------------------Blaze Code Start (27.07.26)---------------------->

@login_required
def staff_student_leave_detail(request, uuid, leave_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'message': 'Access denied'})
    
    leave = get_object_or_404(
        StudentLeaveRequest.objects.select_related(
            'student', 'student__student_profile', 'leave_type'
        ),
        id=leave_id
    )
    
    if request.method == 'POST':
        action = request.POST.get('action')
        remarks = request.POST.get('remarks', '')
        
        if action == 'approve':

            before_data = AuditLogger.model_to_dict(
                leave,
                fields=[
                    "leave_type",
                    "start_date",
                    "end_date",
                    "reason",
                    "status",
                    "approval_remarks",
                ]
            )
            if leave.status != 'pending':
                return JsonResponse({'success': False, 'message': 'This request has already been processed.'})
            
            leave.status = 'approved'
            leave.approved_by = request.user
            leave.approved_date = timezone.now()
            leave.approval_remarks = remarks
            leave.save()
            after_data = AuditLogger.model_to_dict(
                leave,
                fields=[
                    "leave_type",
                    "start_date",
                    "end_date",
                    "reason",
                    "status",
                    "approval_remarks",
                ]
            )
            AuditLogger.log(
                request=request,
                action="APPROVE",
                module="Staff",
                object_type="Student Leave Request",
                object_id=leave.id,
                description=(
                    f'Staff approved a student leave request '
                    f'from {leave.start_date} to {leave.end_date}.'
                ),
                before_data=before_data,
                after_data=after_data,
                status="SUCCESS",
            )
            
            update_student_leave_balance(leave)
            
            notify_student_leave_approved(leave)
            

            return JsonResponse({
                'success': True, 
                'message': f'Leave approved for {leave.student.get_full_name()}'
            })
        
        elif action == 'reject':

            before_data = AuditLogger.model_to_dict(
                leave,
                fields=[
                    "leave_type",
                    "start_date",
                    "end_date",
                    "reason",
                    "status",
                    "approval_remarks",
                ]
            )
            if leave.status != 'pending':
                return JsonResponse({'success': False, 'message': 'This request has already been processed.'})
            
            leave.status = 'rejected'
            leave.approved_by = request.user
            leave.approved_date = timezone.now()
            leave.approval_remarks = remarks
            leave.save()
            after_data = AuditLogger.model_to_dict(
                leave,
                fields=[
                    "leave_type",
                    "start_date",
                    "end_date",
                    "reason",
                    "status",
                    "approval_remarks",
                ]
            )
            AuditLogger.log(
                request=request,
                action="REJECT",
                module="Staff",
                object_type="Student Leave Request",
                object_id=leave.id,
                description=(
                    f'Staff rejected a student leave request '
                    f'from {leave.start_date} to {leave.end_date}.'
                ),
                before_data=before_data,
                after_data=after_data,
                status="SUCCESS",
            )

            notify_student_leave_rejected(leave)
            
            return JsonResponse({
                'success': True, 
                'message': f'Leave rejected for {leave.student.get_full_name()}'
            })
        
        elif action == 'under_review':
            if leave.status != 'pending':
                return JsonResponse({'success': False, 'message': 'This request has already been processed.'})
            
            if remarks:
                leave.approval_remarks = remarks
            leave.save()
            
            return JsonResponse({
                'success': True, 
                'message': f' Leave marked for review - {leave.student.get_full_name()}'
            })
        
        return JsonResponse({'success': False, 'message': 'Invalid action'})
    
    context = {
        'leave': leave,
        'uuid': uuid,
    }
    return render(request, 'Administrator/student_leave_detail.html', context)


#<--------------------Blaze Code End (27.07.26)---------------------->


#<--------------------Blaze Code Start (Financial Aid)---------------------->

@login_required
def staff_financial_aid_detail(request, uuid, aid_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'message': 'Access denied'})
    
    aid = get_object_or_404(StudentFinancialAid, id=aid_id)
    
    if request.method == 'POST':
        action = request.POST.get('action')
        remarks = request.POST.get('remarks', '')
        award_amount = request.POST.get('award_amount', 0)
        
        if action == 'approve':
            if aid.status != 'PENDING':
                return JsonResponse({'success': False, 'message': 'This request has already been processed.'})
            
            aid.status = 'AWARDED'
            aid.award_amount = float(award_amount) if award_amount else 0
            aid.remarks = remarks
            aid.save()
            
            return JsonResponse({
                'success': True,
                'message': f"Financial aid approved for ${aid.award_amount}"
            })
        
        elif action == 'reject':
            if aid.status != 'PENDING':
                return JsonResponse({'success': False, 'message': 'This request has already been processed.'})
            
            aid.status = 'REJECTED'
            aid.remarks = remarks
            aid.save()
            
            return JsonResponse({
                'success': True,
                'message': f" Financial aid rejected"
            })
        
        elif action == 'under_review':
            if aid.status != 'PENDING':
                return JsonResponse({'success': False, 'message': 'This request has already been processed.'})
            
            aid.status = 'UNDER_REVIEW'
            aid.remarks = remarks
            aid.save()
            
            return JsonResponse({
                'success': True,
                'message': f" Financial aid marked for review"
            })
        
        return JsonResponse({'success': False, 'message': 'Invalid action'})
    
    context = {
        'aid': aid,
        'uuid': uuid,
    }
    return render(request, 'Administrator/financial_aid_detail.html', context)


#<--------------------Blaze Code End (Financial Aid)---------------------->

@login_required
def staff_financial_aid_approve(request, uuid, aid_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'message': 'Access denied'})
    
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid request method'})
    
    aid = get_object_or_404(StudentFinancialAid, id=aid_id)
    # ==========================================
    # AUDIT: BEFORE DATA
    # ==========================================
    before_data = AuditLogger.model_to_dict(
        aid,
        fields=[
            "status",
            "award_amount",
            "remarks",
        ]
    )
    award_amount_raw = request.POST.get('award_amount', '')
    remarks = request.POST.get('remarks', '')
    
    try:
        if not award_amount_raw:
            return JsonResponse({'success': False, 'message': 'Award amount is required'})
        
        award_amount = Decimal(award_amount_raw)
        
        if award_amount <= 0:
            return JsonResponse({'success': False, 'message': 'Award amount must be greater than zero'})
        
        if award_amount > 999999.99:
            return JsonResponse({'success': False, 'message': 'Award amount is too high'})
            
    except (InvalidOperation, ValueError, TypeError):
        return JsonResponse({'success': False, 'message': 'Please enter a valid numeric award amount'})
    
    aid.status = 'AWARDED'
    aid.award_amount = award_amount
    if remarks:
        aid.remarks = remarks
    aid.save()
    # ==========================================
    # AUDIT: AFTER DATA
    # ==========================================

    after_data = AuditLogger.model_to_dict(
        aid,
        fields=[
            "status",
            "award_amount",
            "remarks",
        ]
    )
    # ==========================================
    # AUDIT LOG
    # ==========================================
    AuditLogger.log(
        request=request,
        action="APPROVE",
        module="Staff",
        object_type="Student Financial Aid Request",
        object_id=aid.id,
        description=(
            f'Staff approved financial aid request '
            f'#{aid.id} with an award amount of '
            f'${award_amount:,.2f}.'
        ),
        before_data=before_data,
        after_data=after_data,
        status="SUCCESS",
    )

    notify_submitter_financial_aid_status(aid)
    
    return JsonResponse({
        'success': True,
        'message': f'Financial aid approved successfully! Amount: ${award_amount:,.2f}'
    })


@login_required
def staff_financial_aid_reject(request, uuid, aid_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'message': 'Access denied'})
    
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid request method'})
    
    aid = get_object_or_404(StudentFinancialAid, id=aid_id)
    # ==========================================
    # AUDIT: BEFORE DATA
    # ==========================================
    before_data = AuditLogger.model_to_dict(
        aid,
        fields=[
            "status",
            "award_amount",
            "remarks",
        ]
    )
    remarks = request.POST.get('remarks', '')
    
    aid.status = 'REJECTED'
    if remarks:
        aid.remarks = remarks
    aid.save()
    # ==========================================
    # AUDIT: AFTER DATA
    # ==========================================

    after_data = AuditLogger.model_to_dict(
        aid,
        fields=[
            "status",
            "award_amount",
            "remarks",
        ]
    )


    # ==========================================
    # AUDIT LOG
    # ==========================================

    AuditLogger.log(
        request=request,
        action="REJECT",
        module="Staff",
        object_type="Student Financial Aid Request",
        object_id=aid.id,
        description=(
            f'Staff rejected financial aid request '
            f'#{aid.id}.'
        ),
        before_data=before_data,
        after_data=after_data,
        status="SUCCESS",
    )


    notify_submitter_financial_aid_status(aid)
    
    return JsonResponse({
        'success': True,
        'message': 'Financial aid rejected'
    })


@login_required
def staff_financial_aid_under_review(request, uuid, aid_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'message': 'Access denied'})
    
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid request method'})
    
    aid = get_object_or_404(StudentFinancialAid, id=aid_id)
    # ==========================================
    # AUDIT: BEFORE DATA
    # ==========================================
    before_data = AuditLogger.model_to_dict(
        aid,
        fields=[
            "status",
            "award_amount",
            "remarks",
        ]
    )
    remarks = request.POST.get('remarks', '')
    
    aid.status = 'UNDER_REVIEW'
    if remarks:
        aid.remarks = remarks
    aid.save()
    # ==========================================
    # AUDIT: AFTER DATA
    # ==========================================
    after_data = AuditLogger.model_to_dict(
        aid,
        fields=[
            "status",
            "award_amount",
            "remarks",
        ]
    )
    # ==========================================
    # AUDIT LOG
    # ==========================================
    AuditLogger.log(
        request=request,
        action="STATUS_CHANGE",
        module="Staff",
        object_type="Student Financial Aid Request",
        object_id=aid.id,
        description=(
            f'Staff marked financial aid request '
            f'#{aid.id} as under review.'
        ),
        before_data=before_data,
        after_data=after_data,
        status="SUCCESS",
    )
    notify_submitter_financial_aid_status(aid)
    
    return JsonResponse({
        'success': True,
        'message': 'Financial aid marked as under review'
    })


#<--------------------Blaze Code End (Financial Aid)---------------------->

# <---------------------Blaze Code end----------------------->(02.07.2024)


login_required
def staff_forms(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
 
    try:
        request.user.staff_profile
    except Exception:
        return redirect("login")
 
    all_forms = Resource.objects.filter(resource_type='form')
 
    search_query = request.GET.get('search', '').strip()
    category_filter = request.GET.get('category', 'all')
 
    forms_qs = all_forms.order_by('-updated_at')
 
    if search_query:
        forms_qs = forms_qs.filter(
            Q(title__icontains=search_query) |
            Q(description__icontains=search_query) |
            Q(department_name__icontains=search_query)
        )
    if category_filter != 'all':
        forms_qs = forms_qs.filter(form_category=category_filter)
 
    total_forms = all_forms.count()
 
    thirty_days_ago = timezone.now() - timedelta(days=30)
    added_this_month = all_forms.filter(uploaded_at__gte=thirty_days_ago).count()
 
    week_ago = timezone.now() - timedelta(days=7)
    updated_this_week = all_forms.filter(updated_at__gte=week_ago).count()
 
    pending_signatures = all_forms.filter(status='pending_signature').count()
    archived_count = all_forms.filter(status='archived').count()
 
    today = timezone.localdate()
    expiring_cutoff = today + timedelta(days=30)
    expiring_qs = all_forms.filter(
        expires_at__isnull=False, expires_at__gte=today, expires_at__lte=expiring_cutoff
    )
    expiring_soon = expiring_qs.count()
 
    # Real category breakdown (drives tabs, quick-access cards, and the
    # "By Category" donut/bar widget)
    category_counts = []
    for code, label in Resource.FORM_CATEGORY_CHOICES:
        count = all_forms.filter(form_category=code).count()
        category_counts.append({'code': code, 'label': label, 'count': count})
    for c in category_counts:
        c['pct'] = round((c['count'] / total_forms) * 100) if total_forms else 0
 
    # "Needs Attention" widget = pending signature + expiring soon
    needs_attention = (
        list(all_forms.filter(status='pending_signature').order_by('-updated_at')[:3])
        + list(expiring_qs.order_by('expires_at')[:3])
    )
 
    recently_updated = all_forms.order_by('-updated_at')[:5]
 
    # Preserve filters across pagination links
    querydict = request.GET.copy()
    querydict.pop('page', None)
    querystring = querydict.urlencode()
 
    paginator = Paginator(forms_qs, 8)
    page_obj = paginator.get_page(request.GET.get('page', 1))
 
    context = {
        "user": request.user,
        "uuid": uuid,
        "forms": page_obj,
        "total_forms": total_forms,
        "added_this_month": added_this_month,
        "updated_this_week": updated_this_week,
        "pending_signatures": pending_signatures,
        "expiring_soon": expiring_soon,
        "archived_count": archived_count,
        "category_counts": category_counts,
        "needs_attention": needs_attention,
        "recently_updated": recently_updated,
        "search_query": search_query,
        "category_filter": category_filter,
        "form_categories": Resource.FORM_CATEGORY_CHOICES,
        "querystring": querystring,
    }
    return render(request, "Administrator/forms.html", context)

@login_required
@require_http_methods(["POST"])
def update_form_resource(request, uuid, resource_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False, "message": "Unauthorized"}, status=403)
 
    resource = get_object_or_404(Resource, id=resource_id, resource_type='form')
 
    try:
        data = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({"success": False, "message": "Invalid JSON data"}, status=400)
 
    resource.title = data.get('title', resource.title)
    resource.description = data.get('description', resource.description)
    resource.form_category = data.get('form_category', resource.form_category)
    resource.department_name = data.get('department_name', resource.department_name)
    resource.version = data.get('version', resource.version)
    resource.status = data.get('status', resource.status)
    resource.expires_at = data.get('expires_at') or None
    resource.save()
 
    return JsonResponse({"success": True, "message": "Form updated successfully!"})

@login_required
def staff_communications(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    return render(request, "Administrator/communications.html", {"user": request.user})


#<--------------------Blaze Code Start----------------------->

from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse
from django.db.models import Q, Count, Avg
from django.utils import timezone
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from Admin.models import User
import json
import logging

logger = logging.getLogger(__name__)

def staff_student_records(request, uuid):
    """Main student records view with dynamic data and error handling"""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    
    try:
        # Get all students with related data
        students = User.objects.filter(
            is_student=True
        ).select_related(
            'student_profile',
            'student_profile__academic_profile',
            'student_profile__academic_profile__program'
        ).order_by('first_name', 'last_name')
        
        print(f"Students count: {students.count()}")
        
        total_students = students.count()
        
        active_students = students.filter(account_status='ACTIVE').count()
        
        graduated = students.filter(account_status='GRADUATED').count()
        
        on_leave = students.filter(account_status='LEAVE').count()
        
        suspended = students.filter(account_status='SUSPENDED').count()
        
        if active_students == 0 and total_students > 0:
            # Check if account_status field exists
            try:
                from django.db import connection
                with connection.cursor() as cursor:
                    cursor.execute("PRAGMA table_info(users)")
                    columns = [col[1] for col in cursor.fetchall()]
                    if 'account_status' not in columns:
                        # Use is_active as fallback
                        active_students = students.filter(is_active=True).count()
                        graduated = 0
                        on_leave = 0
                        suspended = 0
            except:
                pass
       
        
        # Get filter parameters
        search_query = request.GET.get('search', '').strip()
        status_filter = request.GET.get('status', '').strip().upper()
        program_filter = request.GET.get('program', '').strip()
        page = request.GET.get('page', 1)
        
        # Apply search filter
        if search_query:
            students = students.filter(
                Q(first_name__icontains=search_query) |
                Q(last_name__icontains=search_query) |
                Q(email__icontains=search_query) |
                Q(student_profile__student_number__icontains=search_query)
            )
        
        # Apply status filter
        if status_filter:
            # Check if account_status field exists
            try:
                from django.db import connection
                with connection.cursor() as cursor:
                    cursor.execute("PRAGMA table_info(users)")
                    columns = [col[1] for col in cursor.fetchall()]
                    if 'account_status' in columns:
                        students = students.filter(account_status=status_filter)
                    else:
                        # Use is_active for filtering
                        if status_filter == 'ACTIVE':
                            students = students.filter(is_active=True)
                        elif status_filter == 'INACTIVE':
                            students = students.filter(is_active=False)
            except:
                pass
        
        # Apply program filter
        if program_filter and program_filter.isdigit():
            try:
                students = students.filter(
                    student_profile__academic_profile__program_id=int(program_filter)
                )
            except:
                pass
        
        # Pagination
        paginator = Paginator(students, 10)
        try:
            page_obj = paginator.page(page)
        except PageNotAnInteger:
            page_obj = paginator.page(1)
        except EmptyPage:
            page_obj = paginator.page(paginator.num_pages)
        
        # Get programs for filter dropdown
        try:
            from Students.models import AcademicProgram
            programs = AcademicProgram.objects.all().order_by('program_name')
        except (ImportError, AttributeError):
            programs = []
        
        context = {
            'user': request.user,
            'students': page_obj,
            'total_students': total_students,
            'active_students': active_students,
            'graduated': graduated,
            'on_leave': on_leave,
            'suspended': suspended,
            'programs': programs,
            'today': timezone.now().date(),
            'search_query': search_query,
            'status_filter': status_filter,
            'program_filter': program_filter,
        }
        
        return render(request, "Student_Services/Student_records/student_records.html", context)
        
    except Exception as e:
    
        import traceback
        traceback.print_exc()
        
        context = {
            'user': request.user,
            'students': [],
            'total_students': 0,
            'active_students': 0,
            'graduated': 0,
            'on_leave': 0,
            'suspended': 0,
            'programs': [],
            'today': timezone.now().date(),
            'search_query': '',
            'status_filter': '',
            'program_filter': '',
            'error': str(e),
        }
        return render(request, "Student_Services/Student_records/student_records.html", context)


def get_student_detail(request, student_id):
    """Get detailed student information with proper error handling"""
    try:
        student = get_object_or_404(
            User.objects.filter(is_student=True),
            id=student_id
        )
        
        # Safely get profile
        profile = None
        if hasattr(student, 'student_profile'):
            profile = student.student_profile
        
        # Safely get academic profile
        academic = None
        if profile and hasattr(profile, 'academic_profile'):
            academic = profile.academic_profile
        
        # Build response data with defaults
        data = {
            'id': student.id,
            'uuid': str(student.uuid) if student.uuid else '',
            'name': student.get_full_name() or 'Unknown Student',
            'first_name': student.first_name or '',
            'last_name': student.last_name or '',
            'email': student.email or 'No Email',
            'student_id': getattr(profile, 'student_number', 'N/A') or 'N/A',
            'date_of_birth': student.date_of_birth.strftime('%Y-%m-%d') if student.date_of_birth else None,
            'gender': get_gender_display(student) or 'N/A',
            'mobile': student.mobile_number or 'N/A',
            'status': student.account_status or 'ACTIVE',
            'status_display': get_status_display(student) or student.account_status,
            'profile_photo': student.profile_photo.url if student.profile_photo else None,
            
            # Academic Info
            'program': get_program_name(academic),
            'program_id': getattr(academic, 'program_id', None),
            'major': getattr(academic, 'major', 'N/A') or 'N/A',
            'minor': getattr(academic, 'minor', 'N/A') or 'N/A',
            'gpa': str(get_gpa(academic)) if get_gpa(academic) is not None else '0.00',
            'advisor': get_advisor_name(academic),
            'catalog_year': getattr(academic, 'catalog_year', 'N/A') or 'N/A',
            
            # Enrollments
            'enrollments': get_enrollments(profile),
            
            # Addresses
            'addresses': get_addresses(profile),
            
            # Emergency Contacts
            'emergency_contacts': get_emergency_contacts(profile),
            
            # Financial Aid
            'financial_aid': get_financial_aid(profile),
            
            # Research
            'research': get_research(profile),
            
            # Organizations
            'organizations': get_organizations(profile),
            
            # Career
            'career': get_career(profile),
        }
        
        return JsonResponse({'success': True, 'data': data})
        
    except Exception as e:
        logger.error(f"Error getting student detail: {str(e)}")
        return JsonResponse({
            'success': False, 
            'error': 'Student not found or error loading details'
        }, status=500)


# Helper functions 
def get_gender_display(student):
    try:
        return student.get_gender_display() if student.gender else None
    except:
        return None

def get_status_display(student):
    try:
        if hasattr(student, 'get_account_status_display'):
            return student.get_account_status_display()
        return student.account_status
    except:
        return student.account_status

def get_program_name(academic):
    try:
        if academic and hasattr(academic, 'program') and academic.program:
            return academic.program.name
        return 'N/A'
    except:
        return 'N/A'

def get_advisor_name(academic):
    try:
        if academic and hasattr(academic, 'advisor') and academic.advisor:
            return academic.advisor.get_full_name()
        return 'N/A'
    except:
        return 'N/A'

def get_gpa(academic):
    try:
        if academic and hasattr(academic, 'cumulative_gpa'):
            return academic.cumulative_gpa
        return None
    except:
        return None

def get_enrollments(profile):
    try:
        if not profile:
            return []
        enrollments = []
        for e in profile.enrollments.all()[:5]:
            enrollments.append({
                'term': getattr(e, 'term', {}).name if hasattr(e, 'term') and e.term else 'N/A',
                'status': getattr(e, 'enrollment_status', 'N/A') or 'N/A',
                'credit_load': str(getattr(e, 'credit_load', 0) or 0),
                'academic_standing': getattr(e, 'academic_standing', 'N/A') or 'N/A',
            })
        return enrollments
    except:
        return []

def get_addresses(profile):
    try:
        if not profile:
            return []
        addresses = []
        for addr in profile.addresses.all()[:2]:
            addresses.append({
                'type': getattr(addr, 'address_type', 'N/A') or 'N/A',
                'address': f"{getattr(addr, 'address_line_1', '')}, {getattr(addr, 'city', '')}, {getattr(addr, 'state', '')} {getattr(addr, 'postal_code', '')}".strip(', '),
            })
        return addresses
    except:
        return []

def get_emergency_contacts(profile):
    try:
        if not profile:
            return []
        contacts = []
        for c in profile.emergency_contacts.all()[:2]:
            contacts.append({
                'name': getattr(c, 'contact_name', 'N/A') or 'N/A',
                'relationship': getattr(c, 'relationship', 'N/A') or 'N/A',
                'phone': getattr(c, 'phone_number', 'N/A') or 'N/A',
                'email': getattr(c, 'email', 'N/A') or 'N/A',
            })
        return contacts
    except:
        return []

def get_financial_aid(profile):
    try:
        if not profile:
            return []
        aid_list = []
        for aid in profile.financial_aid.all()[:3]:
            aid_list.append({
                'type': getattr(aid, 'aid_type', 'N/A') or 'N/A',
                'amount': str(getattr(aid, 'award_amount', 0) or 0),
                'year': getattr(aid, 'academic_year', 'N/A') or 'N/A',
                'status': getattr(aid, 'status', 'N/A') or 'N/A',
            })
        return aid_list
    except:
        return []

def get_research(profile):
    try:
        if not profile:
            return []
        research_list = []
        for r in profile.research_profiles.all()[:3]:
            mentor_name = 'N/A'
            if hasattr(r, 'faculty_mentor') and r.faculty_mentor:
                try:
                    mentor_name = r.faculty_mentor.user.get_full_name() if hasattr(r.faculty_mentor, 'user') else 'N/A'
                except:
                    mentor_name = 'N/A'
            research_list.append({
                'title': getattr(r, 'project_title', 'N/A') or 'N/A',
                'area': getattr(r, 'research_area', 'N/A') or 'N/A',
                'mentor': mentor_name,
                'start': getattr(r, 'participation_start', None).strftime('%Y-%m-%d') if hasattr(r, 'participation_start') and r.participation_start else 'N/A',
            })
        return research_list
    except:
        return []

def get_organizations(profile):
    try:
        if not profile:
            return []
        org_list = []
        for o in profile.organization_memberships.all()[:3]:
            org_list.append({
                'name': getattr(o, 'organization_name', 'N/A') or 'N/A',
                'role': getattr(o, 'role', 'Member') or 'Member',
                'start': getattr(o, 'start_date', None).strftime('%Y-%m-%d') if hasattr(o, 'start_date') and o.start_date else 'N/A',
            })
        return org_list
    except:
        return []

def get_career(profile):
    try:
        if not profile:
            return []
        career_list = []
        for c in profile.career_profiles.all()[:3]:
            career_list.append({
                'company': getattr(c, 'internship_company', 'N/A') or 'N/A',
                'title': getattr(c, 'internship_title', 'N/A') or 'N/A',
                'start': getattr(c, 'internship_start', None).strftime('%Y-%m-%d') if hasattr(c, 'internship_start') and c.internship_start else 'N/A',
                'end': getattr(c, 'internship_end', None).strftime('%Y-%m-%d') if hasattr(c, 'internship_end') and c.internship_end else 'Present',
            })
        return career_list
    except:
        return []


def update_student_status(request):
    """Update student status with error handling"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)
    
    try:
        data = json.loads(request.body)
        student_id = data.get('student_id')
        new_status = data.get('status', '').upper()
        
        student = get_object_or_404(User, id=student_id, is_student=True)
        
        valid_statuses = ['ACTIVE', 'INACTIVE', 'LEAVE', 'GRADUATED', 'SUSPENDED']
        if new_status not in valid_statuses:
            return JsonResponse({'success': False, 'error': 'Invalid status'}, status=400)
        
        student.account_status = new_status
        student.save()
        
        return JsonResponse({
            'success': True,
            'message': f'Student status updated to {new_status}',
            'new_status': new_status
        })
        
    except json.JSONDecodeError:
        return JsonResponse({'success': False, 'error': 'Invalid JSON data'}, status=400)
    except Exception as e:
        logger.error(f"Error updating student status: {str(e)}")
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


def get_student_statistics(request):
    """Get statistics for dashboard with error handling"""
    try:
        stats = {
            'total': User.objects.filter(is_student=True).count(),
            'active': User.objects.filter(is_student=True, account_status='ACTIVE').count(),
            'graduated': User.objects.filter(is_student=True, account_status='GRADUATED').count(),
            'on_leave': User.objects.filter(is_student=True, account_status='LEAVE').count(),
            'suspended': User.objects.filter(is_student=True, account_status='SUSPENDED').count(),
            'inactive': User.objects.filter(is_student=True, account_status='INACTIVE').count(),
        }
        return JsonResponse(stats)
        
    except Exception as e:
        logger.error(f"Error getting statistics: {str(e)}")
        return JsonResponse({
            'error': 'Error loading statistics',
            'total': 0,
            'active': 0,
            'graduated': 0,
            'on_leave': 0,
            'suspended': 0,
            'inactive': 0,
        }, status=500)


def search_students(request):
    """Search students with error handling"""
    try:
        query = request.GET.get('q', '').strip()
        if not query:
            return JsonResponse({'success': True, 'results': []})
        
        students = User.objects.filter(
            is_student=True
        ).filter(
            Q(first_name__icontains=query) |
            Q(last_name__icontains=query) |
            Q(email__icontains=query) |
            Q(student_profile__student_number__icontains=query)
        )[:10]
        
        results = []
        for student in students:
            profile = getattr(student, 'student_profile', None)
            results.append({
                'id': student.id,
                'name': student.get_full_name() or 'Unknown',
                'student_id': getattr(profile, 'student_number', 'N/A') or 'N/A',
                'email': student.email or 'No Email',
                'profile_photo': student.profile_photo.url if student.profile_photo else None,
            })
        
        return JsonResponse({'success': True, 'results': results})
        
    except Exception as e:
        logger.error(f"Error searching students: {str(e)}")
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


def export_students(request):
    """Export student data to CSV"""
    try:
        import csv
        from django.http import HttpResponse
        
        students = User.objects.filter(is_student=True).select_related('student_profile')
        student_count = students.count()
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="students_export.csv"'
        
        writer = csv.writer(response)
        writer.writerow([
            'Student ID', 'Name', 'Email', 'Program', 'GPA', 
            'Status', 'Date of Birth', 'Phone'
        ])
        
        for student in students:
            profile = getattr(student, 'student_profile', None)
            academic = getattr(profile, 'academic_profile', None) if profile else None
            
            writer.writerow([
                getattr(profile, 'student_number', 'N/A') or 'N/A',
                student.get_full_name() or 'Unknown',
                student.email or 'No Email',
                get_program_name(academic),
                str(get_gpa(academic)) if get_gpa(academic) is not None else '0.00',
                student.account_status or 'ACTIVE',
                student.date_of_birth.strftime('%Y-%m-%d') if student.date_of_birth else 'N/A',
                student.mobile_number or 'N/A',
            ])

        # ==========================================
        # AUDIT LOG
        # ==========================================
        AuditLogger.log(
            request=request,
            action="EXPORT",
            module="Staff",
            object_type="Student Records",
            object_id="",
            description=(
                f'Staff exported '
                f'{student_count} student record(s) '
                f'to CSV.'
            ),
            status="SUCCESS",
        )
        
        return response
        
    except Exception as e:
        logger.error(f"Error exporting students: {str(e)}")
        AuditLogger.log(
            request=request,
            action="EXPORT",
            module="Staff",
            object_type="Student Records",
            object_id="",
            description=(
                f'Failed to export student records. '
                f'Error: {str(e)}'
            ),
            status="FAILED",
        )
        return JsonResponse({'success': False, 'error': str(e)}, status=500)
    


from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from Admin.models import User

@login_required
def edit_student(request, student_id):
    """Edit student page"""
    # Get the student
    student = get_object_or_404(User, id=student_id, is_student=True)
    
    if request.method == 'POST':
        # ==========================================
        # AUDIT: CAPTURE BEFORE DATA
        # ==========================================
        before_data = AuditLogger.model_to_dict(
            student,
            fields=[
                "first_name",
                "last_name",
                "email",
                "mobile_number",
                "account_status",
            ]
        )
        # Update student
        student.first_name = request.POST.get('first_name', student.first_name)
        student.last_name = request.POST.get('last_name', student.last_name)
        student.email = request.POST.get('email', student.email)
        student.mobile_number = request.POST.get('mobile_number', student.mobile_number)
        student.account_status = request.POST.get('status', student.account_status)
        student.save()
         # ==========================================
        # AUDIT: CAPTURE AFTER DATA
        # ==========================================
        after_data = AuditLogger.model_to_dict(
            student,
            fields=[
                "first_name",
                "last_name",
                "email",
                "mobile_number",
                "account_status",
            ]
        )
        # ==========================================
        # AUDIT LOG
        # ==========================================

        AuditLogger.log(
            request=request,
            action="UPDATE",
            module="Staff",
            object_type="Student",
            object_id=student.id,
            description=(
                f'Staff updated student record for '
                f'"{student.get_full_name()}".'
            ),
            before_data=before_data,
            after_data=after_data,
            status="SUCCESS",
        )
        # Redirect back to records page
        return redirect('staff_student_records', uuid=request.user.uuid)
    
    context = {
        'student': student,
        'user': request.user,
    }
    return render(request, 'Student_Services/Student_records/edit_student.html', context)  

#<--------------------Blaze Code end staff stu rec----------------------->

    
#<--------------------Blaze Code start enrollment(29-06-26)----------------------->


from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse
from django.utils import timezone
from django.db.models import Count, Avg, Q
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from Admin.models import User

# Correct imports from Student app
from Students.models import (
    StudentProfile,
    StudentEnrollment, 
    CourseSection,
    Semester,
    StudentAcademicProfile,
)


@login_required
def staff_enrollment(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    
    try:
        # Get all enrollments with related data
        enrollments = StudentEnrollment.objects.select_related(
            'student',
            'student__user',
            'section_id',
            'section_id__course',
            'section_id__semester_id',
        ).all().order_by('-section_id__semester_id__start_date', 'student__user__first_name')
        
        # Statistics
        total_enrollments = enrollments.count()
        active_students = User.objects.filter(is_student=True, account_status='ACTIVE').count()
        
        # Waitlisted - StudentEnrollment doesn't have waitlist status, using default
        waitlisted = 0
        
        # Calculate average credit load
        avg_credits = enrollments.aggregate(avg=Avg('credit_load'))['avg'] or 0
        
        # Get filter parameters
        search_query = request.GET.get('search', '').strip()
        status_filter = request.GET.get('status', 'all')
        
        # Apply search filter
        if search_query:
            enrollments = enrollments.filter(
                Q(student__user__first_name__icontains=search_query) |
                Q(student__user__last_name__icontains=search_query) |
                Q(student__user__email__icontains=search_query) |
                Q(student__student_number__icontains=search_query) |
                Q(section_id__course__course_code__icontains=search_query) |
                Q(section_id__course__course_name__icontains=search_query)
            )
        
        # Apply status filter (enrollment_status)
        if status_filter and status_filter != 'all':
            enrollments = enrollments.filter(enrollment_status=status_filter.upper())
        
        # A student can have multiple enrollment rows (one for each course).
        # Keep one row per student and expose the number of distinct courses.
        records_by_student = {}
        for enrollment in enrollments:
            records_by_student.setdefault(enrollment.student_id, []).append(enrollment)

        student_records = []
        for student_enrollments in records_by_student.values():
            representative = student_enrollments[0]  # queryset is newest first
            courses = {
                enrollment.section_id.course_id: enrollment.section_id.course
                for enrollment in student_enrollments
                if enrollment.section_id and enrollment.section_id.course_id
            }
            representative.course_count = len(courses)
            representative.total_credits = sum(
                (course.credits or 0) for course in courses.values()
            )
            student_records.append(representative)

        # Pagination
        paginator = Paginator(student_records, 10)
        page = request.GET.get('page', 1)
        page_obj = paginator.get_page(page)
        
        # Get departments list with dynamic counts
        from Admin.bela_admin.models import Department, Course
        from Admin.Colleges.models import AcademicProgram

        departments_list = []
        color_palette = ['red', 'blue', 'green', 'purple', 'orange', 'teal', 'pink', 'brown']
        
        all_departments = Department.objects.all().order_by('department_name')
        for idx, dept in enumerate(all_departments):
            color = color_palette[idx % len(color_palette)]
            
            # Counts
            student_count = StudentAcademicProfile.objects.filter(department=dept).count()
            program_count = AcademicProgram.objects.filter(department=dept).count()
            course_count = Course.objects.filter(department=dept).count()
            
            departments_list.append({
                'id': dept.department_id,
                'code': dept.department_code,
                'name': dept.department_name,
                'short_name': dept.short_name,
                'color': color,
                'student_count': student_count,
                'program_count': program_count,
                'course_count': course_count,
            })
            
        context = {
            'user': request.user,
            'enrollments': page_obj,
            'total_enrollments': total_enrollments,
            'active_students': active_students,
            'waitlisted': waitlisted,
            'avg_credits': round(avg_credits, 1) if avg_credits else 0,
            'today': timezone.now().date(),
            'search_query': search_query,
            'status_filter': status_filter,
            'departments': departments_list,
        }
        
        return render(request, "Student_Services/Student_enrollment/enrollment.html", context)
        
    except Exception as e:
        print(f"Error in staff_enrollment: {e}")  # For debugging
        context = {
            'user': request.user,
            'enrollments': [],
            'total_enrollments': 0,
            'active_students': 0,
            'waitlisted': 0,
            'avg_credits': 0,
            'today': timezone.now().date(),
            'search_query': '',
            'status_filter': 'all',
            'error': str(e),
        }
        return render(request, "Student_Services/Student_enrollment/enrollment.html", context)
@login_required
def staff_department_enrollment(request, uuid, department_id):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
        
    try:
        from django.shortcuts import get_object_or_404
        from Admin.bela_admin.models import Department, Course
        from Admin.Colleges.models import AcademicProgram
        from Students.models import StudentAcademicProfile
        
        department = get_object_or_404(Department, department_id=department_id)
        
        # Dept-level stats
        students_count = StudentAcademicProfile.objects.filter(department=department).count()
        programs_count = AcademicProgram.objects.filter(department=department).count()
        courses_count = Course.objects.filter(department=department).count()
        
        # Fetch programs in this department
        programs = AcademicProgram.objects.filter(department=department).order_by('program_name')
        
        programs_list = []
        for prog in programs:
            p_students = StudentAcademicProfile.objects.filter(program=prog).count()
            p_courses = Course.objects.filter(academic_program=prog).count()
            
            # Determine degree color & description if not set
            name = prog.program_name
            color = 'orange'
            if name.startswith('BS') or name.startswith('B.Tech') or name.startswith('BBA') or name.startswith('B.Sc') or 'Bachelor' in name:
                color = 'red'
            elif name.startswith('MS') or name.startswith('M.Tech') or name.startswith('MBA') or name.startswith('M.Sc') or 'Master' in name:
                color = 'blue'
            elif name.startswith('PhD') or 'Doctor' in name:
                color = 'green'
                
            desc = prog.description or f"Academic program under the {department.department_name} department focusing on specialized concepts and practical applications."
            
            programs_list.append({
                'id': prog.program_id,
                'name': name,
                'student_count': p_students,
                'course_count': p_courses,
                'color': color,
                'description': desc,
            })
            
        context = {
            'user': request.user,
            'department': department,
            'students_count': students_count,
            'programs_count': programs_count,
            'courses_count': courses_count,
            'programs': programs_list,
            'today': timezone.now().date(),
        }
        
        return render(request, "Student_Services/Student_enrollment/department_enrollment.html", context)
        
    except Exception as e:
        print(f"Error in staff_department_enrollment: {e}")
        return redirect('staff_enrollment', uuid=uuid)


@login_required
def staff_program_bulk_enrollment(request, uuid, program_id):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
        
    try:
        from django.shortcuts import get_object_or_404
        from Admin.bela_admin.models import Department, Course
        from Admin.Colleges.models import AcademicProgram, ProgramCourse
        from Students.models import CourseSection, Semester
        
        program = get_object_or_404(AcademicProgram, program_id=program_id)
        department = program.department
        
        # ---- Each ProgramCourse record becomes its own row so that the same
        # course appearing in e.g. Year 2 Fall and Year 3 Spring shows as two
        # distinct rows, each with its own curriculum year / term. ----
        current_semester = Semester.objects.filter(is_current=True).first()

        pcs = ProgramCourse.objects.filter(program=program).select_related('course')
        section_cache = {}  # course_id -> section count for current semester
        def _section_count(course_id):
            if course_id not in section_cache:
                qs = CourseSection.objects.filter(course_id=course_id, is_active=True)
                if current_semester:
                    qs = qs.filter(semester_id=current_semester)
                section_cache[course_id] = qs.count()
            return section_cache[course_id]

        courses_list = []
        for pc in pcs:
            course = pc.course
            code = course.course_code
            if code.isdigit() and department:
                code = f"{department.department_code}{code}"
            courses_list.append({
                'id': course.course_id,
                'code': code,
                'name': course.course_name,
                'credits': course.credits,
                'sections_count': _section_count(course.course_id),
                'study_years': [pc.study_year],
                'terms': [pc.term],
                'statuses': [pc.status],
                'has_elective': pc.is_elective,
                'has_core': not pc.is_elective,
            })
            
        context = {
            'user': request.user,
            'program': program,
            'department': department,
            'courses': courses_list,
            'study_year_choices': ProgramCourse.YEAR_CHOICES,
            'semester_choices': ProgramCourse.TERM_CHOICES,
            'status_choices': ProgramCourse.STATUS_CHOICES,
            'today': timezone.now().date(),
        }
        
        return render(request, "Student_Services/Student_enrollment/program_bulk_enrollment.html", context)
        
    except Exception as e:
        print(f"Error in staff_program_bulk_enrollment: {e}")
        return redirect('staff_enrollment', uuid=uuid)

def _program_bulk_course_dict(course, department):
    """Build a lightweight course dict for the program bulk enrollment flow."""
    code = str(course.course_code)
    if code.isdigit() and department:
        code = f"{department.department_code}{code}"
    upper = code.upper()
    if upper.startswith('CS'):
        theme = 'purple'
    elif upper.startswith('MATH') or upper.startswith('MTH'):
        theme = 'green'
    elif upper.startswith('ENG') or upper.startswith('ENGL'):
        theme = 'orange'
    else:
        theme = 'blue'
    return {
        'id': course.course_id,
        'code': code,
        'name': course.course_name,
        'credits': course.credits,
        'theme': theme,
    }


def _program_bulk_course_ids(course_ids_str):
    return [int(x) for x in (course_ids_str or '').split(',') if str(x).strip().isdigit()]


def _program_bulk_default_term(program, course_ids_str):
    """Pick a default semester term for the selected courses.

    Prefers the term type that actually has CourseSections created for the
    selected courses (matching real Semester data), falling back to the most
    common ACTIVE curriculum term so the wizard always carries a usable term.
    """
    from collections import Counter
    from Admin.Colleges.models import ProgramCourse
    from Students.models import CourseSection

    term_by_type = {'FA': 'FALL', 'SP': 'SPRING', 'SU': 'SUMMER', 'WI': 'WINTER'}
    ids = _program_bulk_course_ids(course_ids_str)

    if ids:
        counts = Counter(
            term_by_type.get(st)
            for st in CourseSection.objects.filter(
                course_id__in=ids, is_active=True
            ).exclude(course_id__isnull=True).values_list('semester_id__semester_type', flat=True)
            if term_by_type.get(st)
        )
        if counts:
            return counts.most_common(1)[0][0]

    counts = Counter()
    for pc in ProgramCourse.objects.filter(program=program, course_id__in=ids, status='ACTIVE'):
        counts[pc.term] += 1
    return counts.most_common(1)[0][0] if counts else ''


def _program_bulk_courses_with_sections(program, department, course_ids_str, section_ids_str):
    """Courses for the selected ids, each with its matching selected section (if any)."""
    from Students.models import CourseSection, StudentEnrollment

    course_ids = _program_bulk_course_ids(course_ids_str)
    section_ids = _program_bulk_course_ids(section_ids_str)

    sections_by_id = {}
    if section_ids:
        for s in CourseSection.objects.filter(section_id__in=section_ids):
            sections_by_id[s.section_id] = s

    courses = []
    for course in Course.objects.filter(course_id__in=course_ids, academic_program=program).order_by('course_code'):
        c = _program_bulk_course_dict(course, department)
        sec = None
        for sid in section_ids:
            s = sections_by_id.get(sid)
            if s and s.course_id == course.course_id:
                sec = s
                break
        if sec is not None:
            c['section'] = {
                'id': sec.section_id,
                'number': sec.section_number,
                'faculty_name': sec.faculty_name or 'Unassigned',
                'capacity': sec.capacity,
                'enrolled_count': StudentEnrollment.objects.filter(section_id=sec).count(),
            }
        else:
            c['section'] = None
        courses.append(c)
    return courses


def _program_bulk_section_cards(program, department, course_ids_str, section_ids_str):
    """Flat per-section cards (one card per involved section) for the bulk flow."""
    from Students.models import CourseSection, StudentEnrollment

    course_ids = _program_bulk_course_ids(course_ids_str)
    section_ids = _program_bulk_course_ids(section_ids_str)

    sections_by_id = {}
    if section_ids:
        for s in CourseSection.objects.filter(section_id__in=section_ids).select_related('course'):
            sections_by_id[s.section_id] = s

    cards = []
    for course in Course.objects.filter(course_id__in=course_ids, academic_program=program).order_by('course_code'):
        for sid in section_ids:
            s = sections_by_id.get(sid)
            if not s or s.course_id != course.course_id:
                continue
            cc = _program_bulk_course_dict(course, department)
            cc['section'] = {
                'id': s.section_id,
                'number': s.section_number,
                'faculty_name': s.faculty_name or 'Unassigned',
                'capacity': s.capacity,
                'enrolled_count': StudentEnrollment.objects.filter(section_id=s).count(),
            }
            cards.append({'course': cc})
    return cards


@login_required
def staff_program_bulk_enrollment_sections(request, uuid, program_id):
    """Step 2 (standalone): select one section per course before enrolling students."""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    try:
        from django.shortcuts import get_object_or_404
        from Admin.Colleges.models import AcademicProgram, ProgramCourse

        program = get_object_or_404(AcademicProgram, program_id=program_id)
        department = program.department
        course_ids_str = request.GET.get('course_ids', '')

        courses = []
        for course in Course.objects.filter(
            course_id__in=_program_bulk_course_ids(course_ids_str),
            academic_program=program,
        ).order_by('course_code'):
            courses.append(_program_bulk_course_dict(course, department))

        context = {
            'user': request.user,
            'program': program,
            'department': department,
            'courses': courses,
            'courses_json': json.dumps(courses),
            'semester_choices': ProgramCourse.TERM_CHOICES,
            'default_term': _program_bulk_default_term(program, course_ids_str),
            'today': timezone.now().date(),
        }
        return render(request, "Student_Services/Student_enrollment/program_bulk_enrollment_sections.html", context)

    except Exception as e:
        print(f"Error in staff_program_bulk_enrollment_sections: {e}")
        return redirect('staff_program_bulk_enrollment', uuid=uuid, program_id=program_id)


@login_required
def staff_program_bulk_enrollment_students(request, uuid, program_id):
    """Step 3 (standalone): select the students to enroll in the chosen sections."""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    try:
        from django.shortcuts import get_object_or_404
        from Admin.Colleges.models import AcademicProgram, ProgramCourse

        program = get_object_or_404(AcademicProgram, program_id=program_id)
        department = program.department
        course_ids_str = request.GET.get('course_ids', '')
        section_ids_str = request.GET.get('section_ids', '')
        term = request.GET.get('term', '')

        courses = _program_bulk_courses_with_sections(program, department, course_ids_str, section_ids_str)
        section_cards = _program_bulk_section_cards(program, department, course_ids_str, section_ids_str)

        course_ids = [c['id'] for c in courses]
        section_ids = [c['section']['id'] for c in courses if c['section']]
        resolved_semester = _resolve_semester(term) if term else None

        context = {
            'user': request.user,
            'program': program,
            'department': department,
            'courses': courses,
            'courses_json': json.dumps(courses),
            'section_cards': section_cards,
            'course_ids': course_ids,
            'section_ids': section_ids,
            'semester_id': resolved_semester.semester_id if resolved_semester else term,
            'term': term,
            'term_label': dict(ProgramCourse.TERM_CHOICES).get(term, term or '—'),
            'today': timezone.now().date(),
        }
        return render(request, "Student_Services/Student_enrollment/program_bulk_enrollment_students.html", context)

    except Exception as e:
        print(f"Error in staff_program_bulk_enrollment_students: {e}")
        return redirect('staff_program_bulk_enrollment_sections', uuid=uuid, program_id=program_id)

 # speed code start
def _student_admission_year(student_profile):
    """Best-effort cohort year: prefer admission_date, fall back to the first
    4-digit group found in the student number."""
    sp = student_profile
    if sp is None:
        return None
    if getattr(sp, 'admission_date', None):
        try:
            return int(sp.admission_date.year)
        except (TypeError, ValueError):
            pass
    match = re.search(r'(\d{4})', str(sp.student_number or ''))
    return int(match.group(1)) if match else None


def _cohort_study_years(course_ids, program_id=None):
    """Study years a set of courses belongs to in the program curriculum."""
    from Admin.Colleges.models import ProgramCourse

    pc_qs = ProgramCourse.objects.filter(
        course_id__in=list(set(course_ids)),
        status='ACTIVE',
    )
    if program_id:
        pc_qs = pc_qs.filter(program_id=program_id)
    return sorted({int(y) for y in pc_qs.values_list('study_year', flat=True) if y})


def _cohort_matches(admission_year, study_years, academic_year):
    """True when admission_year + (study_year - 1) equals the semester's
    academic year for any of the curriculum study years.  None means we cannot
    judge (missing admission year or curriculum mapping) so treat as eligible."""
    if not admission_year or not academic_year or not study_years:
        return None
    return any(int(admission_year) + (int(y) - 1) == int(academic_year) for y in study_years)




@login_required
def staff_program_bulk_enrollment_review(request, uuid, program_id):
    """Step 4 (standalone): review the selection and submit the enrollments."""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    try:
        from django.shortcuts import get_object_or_404
        from Admin.Colleges.models import AcademicProgram, ProgramCourse
        from Admin.models import User

        program = get_object_or_404(AcademicProgram, program_id=program_id)
        department = program.department
        course_ids_str = request.GET.get('course_ids', '')
        section_ids_str = request.GET.get('section_ids', '')
        student_ids_str = request.GET.get('student_ids', '')
        section_students_str = request.GET.get('section_students', '')
        term = request.GET.get('term', '')

        courses = _program_bulk_courses_with_sections(program, department, course_ids_str, section_ids_str)
        resolved_semester = _resolve_semester(term) if term else None
        cohort_override = request.GET.get('cohort_override') == '1'

        uid_list = [int(x) for x in (student_ids_str or '').split(',') if str(x).strip().isdigit()]
        semester_academic_year = resolved_semester.academic_year if resolved_semester else None
        try:
            cohort_study_years = _cohort_study_years(
                [c['id'] for c in courses], program.program_id
            )
        except Exception:
            cohort_study_years = []
        user_by_id = {}
        for user_obj in User.objects.filter(id__in=uid_list, is_student=True).select_related('student_profile'):
            sp = getattr(user_obj, 'student_profile', None)
            admission_year = _student_admission_year(sp)
            user_by_id[user_obj.id] = {
                'id': user_obj.id,
                'first_name': user_obj.first_name or '',
                'last_name': user_obj.last_name or '',
                'student_number': sp.student_number if sp else '—',
                'admission_year': admission_year,
                'cohort_match': _cohort_matches(
                    admission_year, cohort_study_years, semester_academic_year
                ),
            }

        def _section_dict(s):
            from Students.models import StudentEnrollment
            return {
                'id': s.section_id,
                'number': s.section_number,
                'faculty_name': s.faculty_name or 'Unassigned',
                'capacity': s.capacity,
                'enrolled_count': StudentEnrollment.objects.filter(section_id=s).count(),
            }

        # The submit action persists StudentEnrollment rows.  Look them up again
        # when the review page is rendered so a browser refresh reflects the
        # saved state instead of relying on the previous page's JavaScript.
        from Students.models import StudentEnrollment
        selected_section_ids = [
            int(section_id) for section_id in (section_ids_str or '').split(',')
            if str(section_id).strip().isdigit()
        ]
        enrollment_keys = set()
        if resolved_semester and uid_list and selected_section_ids:
            enrollment_keys = set(
                StudentEnrollment.objects.filter(
                    student__user_id__in=uid_list,
                    semester=resolved_semester,
                    section_id__in=selected_section_ids,
                ).values_list('student__user_id', 'section_id')
            )

        def _group_users(ids, section_id):
            group_users = []
            for uid in ids:
                if uid not in user_by_id:
                    continue
                student = dict(user_by_id[uid])
                student['enrollment_state'] = (
                    'enrolled' if (uid, section_id) in enrollment_keys else 'pending'
                )
                group_users.append(student)
            return group_users

        groups = []
        if section_students_str:
            from Students.models import CourseSection
            section_students = {}
            for part in (section_students_str or '').split(';'):
                if ':' not in part:
                    continue
                sec_str, ids_str = part.split(':', 1)
                ids = [int(x) for x in ids_str.split(',') if str(x).strip().isdigit()]
                if ids:
                    section_students[int(sec_str)] = ids

            course_by_id = {c['id']: c for c in courses}
            sections_by_id = {}
            if section_students:
                for s in CourseSection.objects.filter(section_id__in=list(section_students.keys())).select_related('course'):
                    sections_by_id[s.section_id] = s

            for sec_id, ids in section_students.items():
                sec = sections_by_id.get(sec_id)
                if not sec or not sec.course_id:
                    continue
                c = course_by_id.get(sec.course_id)
                if not c:
                    continue
                group_users = _group_users(ids, sec_id)
                if not group_users:
                    continue
                group_course = dict(c)
                group_course['section'] = _section_dict(sec)
                groups.append({
                    'course': group_course,
                    'section_id': sec_id,
                    'student_ids': ids,
                    'students': group_users,
                })
        else:
            for c in courses:
                sec = c.get('section')
                if not sec:
                    continue
                group_users = _group_users(uid_list, sec['id'])
                if not group_users:
                    continue
                groups.append({
                    'course': c,
                    'section_id': sec['id'],
                    'student_ids': uid_list,
                    'students': group_users,
                })

        students = []
        seen_ids = set()
        for g in groups:
            for s in g['students']:
                if s['id'] in seen_ids:
                    continue
                seen_ids.add(s['id'])
                students.append(s)

        cohort_mismatch_count = sum(1 for s in students if s.get('cohort_match') is False)

        course_ids = [c['id'] for c in courses]
        section_ids = [g['section_id'] for g in groups]
        context = {
            'user': request.user,
            'program': program,
            'department': department,
            'courses': courses,
            'courses_json': json.dumps(courses),
            'groups': groups,
            'groups_json': json.dumps(groups),
            'students': students,
            'students_json': json.dumps(students),
            'cohort_override': cohort_override,
            'cohort_info': {
                'academic_year': semester_academic_year,
                'study_years': cohort_study_years,
            },
            'cohort_mismatch_count': cohort_mismatch_count,
            'course_ids': course_ids,
            'section_ids': section_ids,
            'student_ids': uid_list,
            'section_students': section_students_str,
            'semester_id': resolved_semester.semester_id if resolved_semester else '',
            'term': term,
            'term_label': dict(ProgramCourse.TERM_CHOICES).get(term, term or '—'),
            'today': timezone.now().date(),
        }
        return render(request, "Student_Services/Student_enrollment/program_bulk_enrollment_review.html", context)

    except Exception as e:
        print(f"Error in staff_program_bulk_enrollment_review: {e}")
        return redirect('staff_program_bulk_enrollment_students', uuid=uuid, program_id=program_id)



@login_required
def get_enrollment_detail(request, enrollment_id):
    """Get enrollment details via AJAX"""
    try:
        enrollment = get_object_or_404(StudentEnrollment, id=enrollment_id)
        
        # Get student data
        student = enrollment.student
        user = student.user
        
        # Get course data
        course_section = enrollment.section_id if hasattr(enrollment, 'section_id') else None
        course = course_section.course if course_section else None
        semester = course_section.semester_id if course_section else None
        
        academic_profile = getattr(student, 'academic_profile', None)   #speed add this 
        department = academic_profile.department if academic_profile else None      #speed add this 
        program = academic_profile.program if academic_profile else None

        status_choices = []
        for key, value in StudentEnrollment.ENROLLMENT_STATUS_CHOICES:
            status_choices.append({'value': key, 'label': value})
        #speed add this code #########
        department_course_ids = []
        if department:
            department_course_ids = list(
                Course.objects.filter(department=department).values_list('course_id', flat=True)
            )

        course_ids = Q(course_id__in=department_course_ids)
        if course:
            course_ids |= Q(course_id=course.course_id)
        course_qs = Course.objects.filter(course_ids) if department_course_ids or course else Course.objects.none()
        # A program can have shared department courses, so include courses with no
        # program as well as those assigned to the student's program.
        if program:
            course_qs = course_qs.filter(
                Q(academic_program=program) | Q(academic_program__isnull=True) |
                Q(course_id=course.course_id if course else None)
            )
        # Keep the current course selectable even when older enrollment data does
        # not match the student's latest department/program assignment.
        course_qs = course_qs.order_by('course_code')
        courses = []
        for c in course_qs:
            courses.append({
                'id': c.course_id,
                'code': c.course_code,
                'name': c.course_name,
                'credits': c.credits,
            })
        #speed  code ends #########
        semesters = []
        for s in Semester.objects.all().order_by('-academic_year'):
            semesters.append({
                'id': s.semester_id,
                'name': s.__str__(),
                'semester_type': s.semester_type,
                'is_current': s.is_current,
            })

        #speed code starts (same as new_enrollment) - Program Curriculum term -> Semester auto-fill
        from Admin.Colleges.models import ProgramCourse
        program_course_terms = {}
        for pc in ProgramCourse.objects.filter(status='ACTIVE').values('course_id', 'program_id', 'term'):
            key = f"{pc['course_id']}:{pc['program_id']}"
            program_course_terms.setdefault(key, set()).add(pc['term'])
        program_course_terms = {k: sorted(v) for k, v in program_course_terms.items()}
        #speed code ends

        linked_courses = []
        for e in StudentEnrollment.objects.filter(student=student).select_related('section_id__course'):
            sec = e.section_id
            c = sec.course if sec else None
            linked_courses.append({
                'enrollment_id': e.id,
                'code': c.course_code if c else 'N/A',
                'name': c.course_name if c else 'No Course',
            })

        data = {
            'id': enrollment.id,
            'student_id': user.id,
            'student_name': user.get_full_name() or 'Unknown',
            'student_initials': (user.first_name[0] if user.first_name else '') + 
                               (user.last_name[0] if user.last_name else '') or '?',
            'student_id_number': student.student_number if student.student_number else 'N/A',
            'email': user.email or 'No Email',
            'department_name': department.department_name if department else 'N/A', #speed add this code
            'program_id': program.program_id if program else None,
            'program_name': program.program_name if program else 'N/A',
            'course_id': course.course_id if course else None,
            'course_code': course.course_code if course else 'N/A',
            'course_name': course.course_name if course else 'No Course',
            'section_id': course_section.section_id if course_section else None,
            'section_number': course_section.section_number if course_section else 'N/A',
            'semester_id': semester.semester_id if semester else None,
            'semester_name': semester.__str__() if semester else 'N/A',
            'credit_load': course.credits if course else (enrollment.credit_load or 0),
            'status': enrollment.enrollment_status or 'PENDING',
            'status_display': dict(StudentEnrollment.ENROLLMENT_STATUS_CHOICES).get(
                enrollment.enrollment_status, enrollment.enrollment_status
            ) if hasattr(enrollment, 'enrollment_status') else 'N/A',
            'academic_standing': enrollment.academic_standing or 'GOOD_STANDING',
            'status_choices': status_choices,
            'courses': courses,
            'semesters': semesters,
            'linked_courses': linked_courses,
            'program_course_terms': program_course_terms,
        }
        
        
        return JsonResponse({'success': True, 'data': data})
        
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)
        

# Staff/views.py
from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse
from django.utils import timezone
from django.db.models import Count, Avg, Q
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from Admin.models import User

# Import from Student_Services
from Students.models import (
    StudentProfile,
    StudentEnrollment,
    Course,
    CourseSection,
    Semester,
    StudentAcademicProfile,
)

@login_required
def new_enrollment(request, uuid):
    """New Enrollment Form Page"""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
   
    try:
        # Get all students, courses, semesters for dropdown
        students = User.objects.filter(is_student=True, account_status='ACTIVE').order_by('first_name', 'last_name')
        courses = Course.objects.all().order_by('course_code')
        semesters = Semester.objects.filter(is_current=True).order_by('-academic_year')
       
        # If no current semester, get the latest one
        if not semesters.exists():
            semesters = Semester.objects.all().order_by('-academic_year')[:1]
        #speed code starts
        departments = Department.objects.all().order_by('department_name')
        academic_programs = AcademicProgram.objects.all().order_by('program_name')
       #speed code starts
        from Admin.Colleges.models import ProgramCourse
        program_course_terms = {}
        for pc in ProgramCourse.objects.filter(status='ACTIVE').values('course_id', 'program_id', 'term'):
            key = f"{pc['course_id']}:{pc['program_id']}"
            program_course_terms.setdefault(key, set()).add(pc['term'])
        program_course_terms = {k: sorted(v) for k, v in program_course_terms.items()}
        context = {
            'user': request.user,
            'students': students,
            'courses': courses,
            'semesters': semesters,
            'departments': departments, #speed code
            'academic_programs': academic_programs, #speed code
            'program_course_terms_json': json.dumps(program_course_terms),
            'today': timezone.now().date(),
        }
        return render(request, "Student_Services/Student_enrollment/new_enrollment.html", context)
       
    except Exception as e:
        messages.error(request, f'Error loading form: {str(e)}')
        return redirect('staff_enrollment', uuid=uuid)



from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse
from django.utils import timezone
from django.db.models import Count, Avg, Q
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from Admin.models import User
import json

# Import from Students app
from Students.models import (
    StudentProfile,
    StudentEnrollment,
    Course,
    CourseSection,
    Semester,
    StudentAcademicProfile,
)


@login_required
def create_enrollment(request):
    """Create new enrollment via AJAX"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)
    
    try:
        data = json.loads(request.body)
        
        # Get data from request
        student_id = data.get('student_id')
        course_section_id = data.get('course_section_id')
        semester_id = data.get('semester_id')
        enrollment_status = data.get('enrollment_status', 'FULL_TIME')
        credit_load = data.get('credit_load', 3)
        
        print(f" Creating enrollment: student={student_id}, section={course_section_id}, semester={semester_id}")
        
        # Validate
        if not student_id:
            return JsonResponse({'success': False, 'error': 'Student ID required'}, status=400)
        if not course_section_id:
            return JsonResponse({'success': False, 'error': 'Course Section ID required'}, status=400)
        if not semester_id:
            return JsonResponse({'success': False, 'error': 'Semester ID required'}, status=400)
        
        semester = _resolve_semester(semester_id)
        if not semester:
            return JsonResponse({'success': False, 'error': 'Invalid semester or term.'}, status=400)
        
        # Get student profile
        student_user = get_object_or_404(User, id=student_id, is_student=True)
        student_profile = get_object_or_404(StudentProfile, user=student_user)
        
        # print(f"Found student: {student_user.get_full_name()}")  # Debug
        
        # Get course section
        section = get_object_or_404(CourseSection, section_id=course_section_id)
        
        # When a term code is used, the section must belong to that term
        if str(semester_id).isalpha() and _term_of_semester(section.semester_id) != str(semester_id).upper():
            return JsonResponse({
                'success': False,
                'error': 'Please select a section for the selected term.',
                'field': 'section',
            }, status=400)
        
        
        # A student may take a course only once per semester, irrespective of
        # section.  Check this explicitly because legacy databases may not have
        # the model's historical unique constraint applied.
        existing_course = StudentEnrollment.objects.filter(
            student=student_profile,
            semester_id=semester.semester_id,
            section_id__course=section.course,
        ).exists()
        if existing_course:
            return JsonResponse({
                'success': False,
                'error': 'Student is already enrolled in this course.',
                'field': 'student',
            }, status=400)

        if section.capacity and StudentEnrollment.objects.filter(
            section_id=section, semester_id=semester.semester_id
        ).count() >= section.capacity:
            return JsonResponse({
                'success': False,
                'error': 'This section is already full.',
                'field': 'section',
            }, status=400)
        
        # Create enrollment
        enrollment = StudentEnrollment.objects.create(
            student=student_profile,
            semester=semester,
            section_id=section,
            enrollment_status=enrollment_status,
            credit_load=section.course.credits if section.course else credit_load,
            academic_standing='GOOD_STANDING',
        )

        # ==========================================
        # AUDIT: AFTER DATA
        # ==========================================

        after_data = AuditLogger.model_to_dict(
            enrollment,
            fields=[
                "student",
                "semester",
                "section_id",
                "enrollment_status",
                "credit_load",
                "academic_standing",
            ]
        )
        # ==========================================
        # AUDIT LOG: CREATE ENROLLMENT
        # ==========================================
        AuditLogger.log(
            request=request,
            action="CREATE",
            module="Academic",
            object_type="Student Enrollment",
            object_id=enrollment.id,
            description=(
                f'Created enrollment for student '
                f'"{student_user.get_full_name()}" '
                f'in course '
                f'"{section.course.course_code if section.course else "N/A"}".'
            ),
            before_data=None,
            after_data=after_data,
            status="SUCCESS",
        )
        return JsonResponse({
            'success': True,
            'message': 'Enrollment created successfully!',
            'data': {
                'id': enrollment.id,
                'student_name': student_user.get_full_name(),
                'course_code': section.course.course_code if section.course else 'N/A',
                'status': enrollment.enrollment_status,
            }
        })
        
    except json.JSONDecodeError as e:
        print(f"JSON Error: {e}")
        return JsonResponse({'success': False, 'error': 'Invalid JSON data'}, status=400)
    except Exception as e:
        print(f"Error in create_enrollment: {e}")
        return JsonResponse({'success': False, 'error': str(e)}, status=500)



import traceback

@login_required
def get_course_sections(request):
    try:
        course_id = request.GET.get('course_id')
        semester_id = request.GET.get('semester_id')

        sections = CourseSection.objects.filter(
            course_id=course_id
        ).select_related('semester_id')

        sem_ids = _resolve_semester_ids(semester_id)
        if sem_ids is not None:
            sections = sections.filter(semester_id__in=sem_ids)

        data = []
        for section in sections:
            enrollment_filter = {'section_id': section.section_id}
            if sem_ids is not None:
                enrollment_filter['semester_id__in'] = sem_ids
            enrolled_count = StudentEnrollment.objects.filter(**enrollment_filter).count()
            enrolled_students = StudentEnrollment.objects.filter(
                **enrollment_filter
            ).select_related('student__user')

            data.append({
                'id': section.section_id,
                'section_number': section.section_number,
                'semester': str(section.semester_id) if section.semester_id else 'N/A',
                'faculty_name': section.faculty_name or 'Unassigned',
                'capacity': section.capacity,
                'enrolled_count': enrolled_count,
                'students': [{
                    'id': enrollment.student.user_id,
                    'first_name': enrollment.student.user.first_name,
                    'last_name': enrollment.student.user.last_name,
                    'student_number': enrollment.student.student_number,
                } for enrollment in enrolled_students],
            })

        return JsonResponse({'success': True, 'sections': data})

    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'success': False, 'error': str(e)}, status=500)

# Staff/views.py
@login_required
def update_enrollment(request):
    """Update enrollment via AJAX"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)
   
    try:
        data = json.loads(request.body)
       
        enrollment_id = data.get('enrollment_id')
        credit_load = data.get('credit_load')
        status = data.get('status')
        academic_standing = data.get('academic_standing')
        section_id = data.get('section_id')
        course_id = data.get('course_id')
        semester_id = data.get('semester_id')
       
        if not enrollment_id:
            return JsonResponse({'success': False, 'error': 'Enrollment ID required'}, status=400)
       
        enrollment = get_object_or_404(StudentEnrollment, id=enrollment_id)
        # ==========================================
        # AUDIT: BEFORE DATA
        # ==========================================
        before_data = AuditLogger.model_to_dict(
            enrollment,
            fields=[
                "student",
                "semester",
                "section_id",
                "enrollment_status",
                "credit_load",
                "academic_standing",
            ]
        )
        semester = _resolve_semester(semester_id) or enrollment.semester
        selected_semester_id = semester.semester_id if semester else None
 
        if not section_id:
            return JsonResponse({'success': False, 'error': 'Please select a section.', 'field': 'section'}, status=400)
 
        section = get_object_or_404(CourseSection, section_id=section_id)
        if course_id and str(section.course_id) != str(course_id):
            return JsonResponse({'success': False, 'error': 'The selected section does not belong to this course.', 'field': 'section'}, status=400)
        if semester_id:
            s = str(semester_id)
            if s.isdigit():
                if str(section.semester_id_id) != s:
                    return JsonResponse({'success': False, 'error': 'Please select a section for the selected semester.', 'field': 'section'}, status=400)
            else:
                if _term_of_semester(section.semester_id) != s.upper():
                    return JsonResponse({'success': False, 'error': 'Please select a section for the selected term.', 'field': 'section'}, status=400)
        elif selected_semester_id and str(section.semester_id_id) != str(selected_semester_id):
            return JsonResponse({'success': False, 'error': 'Please select a section for the selected semester.', 'field': 'section'}, status=400)
 
        other_enrollments = StudentEnrollment.objects.filter(
            student=enrollment.student,
            semester_id=selected_semester_id,
        ).exclude(id=enrollment.id)
        if other_enrollments.filter(section_id__course=section.course).exists():
            return JsonResponse({
                'success': False,
                'error': 'Student is already enrolled in this course.',
                'field': 'student',
            }, status=400)
        if section.capacity and StudentEnrollment.objects.filter(
            section_id=section, semester_id=selected_semester_id
        ).exclude(id=enrollment.id).count() >= section.capacity:
            return JsonResponse({
                'success': False,
                'error': 'This section is already full.',
                'field': 'section',
            }, status=400)
       
        # Credit load is determined by the selected course; do not rely on the
        # client value, which can be stale or omitted by a disabled field.
        enrollment.credit_load = section.course.credits if section.course else (int(credit_load) if credit_load is not None else enrollment.credit_load)
        if status:
            enrollment.enrollment_status = status
        if academic_standing:
            enrollment.academic_standing = academic_standing
        enrollment.section_id = section
        if semester: #speed add this
            enrollment.semester = semester
       
        enrollment.save()
        # ==========================================
        # AUDIT: AFTER DATA
        # ==========================================
        after_data = AuditLogger.model_to_dict(
            enrollment,
            fields=[

                "student",
                "semester",
                "section_id",
                "enrollment_status",
                "credit_load",
                "academic_standing",
            ]
        )
        # ==========================================
        # AUDIT LOG: UPDATE
        # ==========================================
        AuditLogger.log(
            request=request,
            action="UPDATE",
            module="Academic",
            object_type="Student Enrollment",
            object_id=enrollment.id,
            description=(
                f'Updated enrollment for student '
                f'"{enrollment.student.user.get_full_name()}" '
                f'in course '
                f'"{section.course.course_code if section.course else "N/A"}".'
            ),
            before_data=before_data,
            after_data=after_data,
            status="SUCCESS",
        )
        return JsonResponse({
            'success': True,
            'message': 'Enrollment updated successfully!',
            'data': {
                'id': enrollment.id,
                'credit_load': enrollment.credit_load,
                'status': enrollment.enrollment_status,
            }
        })
   
       
    except json.JSONDecodeError as e:
 
        return JsonResponse({'success': False, 'error': 'Invalid JSON data'}, status=400)
    except Exception as e:
 
        return JsonResponse({'success': False, 'error': str(e)}, status=500)
    


#<--------------------Blaze Code End enrollment(29-06-26)----------------------->




#<--------------------Blaze Code start Finance Sup(30-06-26)----------------------->
 

from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse
from django.utils import timezone
from django.db.models import Count, Avg, Sum, Q
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from Admin.models import User
from Students.models import StudentProfile, StudentFinancialAid
import json


@login_required
def staff_financial_aid(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    
    try:
        # Get all financial aid records with related data
        aid_records = StudentFinancialAid.objects.select_related(
            'student',
            'student__user'
        ).all().order_by('-id')
        
        # Statistics
        total_applications = aid_records.count()
        avg_award = aid_records.aggregate(avg=Avg('award_amount'))['avg'] or 0
        pending_reviews = aid_records.filter(status='PENDING').count()
        scholarships_offered = aid_records.filter(aid_type='SCHOLARSHIP').count()
        
        # Get filter parameters
        search_query = request.GET.get('search', '').strip()
        status_filter = request.GET.get('status', 'all')
        type_filter = request.GET.get('type', 'all')
        page = request.GET.get('page', 1)
        
        # Apply filters
        if search_query:
            aid_records = aid_records.filter(
                Q(student__user__first_name__icontains=search_query) |
                Q(student__user__last_name__icontains=search_query) |
                Q(student__user__email__icontains=search_query) |
                Q(student__student_number__icontains=search_query)
            )
        
        if status_filter and status_filter != 'all':
            aid_records = aid_records.filter(status=status_filter.upper())
        
        if type_filter and type_filter != 'all':
            aid_records = aid_records.filter(aid_type=type_filter.upper())
        
        # Pagination
        paginator = Paginator(aid_records, 10)
        page_obj = paginator.get_page(page)
        
        # Get students for dropdown
        students = User.objects.filter(is_student=True, account_status='ACTIVE').order_by('first_name', 'last_name')
        
        context = {
            'user': request.user,
            'aid_records': page_obj,
            'total_applications': total_applications,
            'avg_award': round(avg_award, 2) if avg_award else 0,
            'pending_reviews': pending_reviews,
            'scholarships_offered': scholarships_offered,
            'today': timezone.now().date(),
            'search_query': search_query,
            'status_filter': status_filter,
            'type_filter': type_filter,
            'students': students,
        }
        
        return render(request, "Student_Services/Financial/financial_aid_support.html", context)
        
    except Exception as e:
        print(f"Error in staff_financial_aid: {e}")
        context = {
            'user': request.user,
            'aid_records': [],
            'total_applications': 0,
            'avg_award': 0,
            'pending_reviews': 0,
            'scholarships_offered': 0,
            'today': timezone.now().date(),
            'search_query': '',
            'status_filter': 'all',
            'type_filter': 'all',
            'students': [],
            'error': str(e),
        }
        return render(request, "Student_Services/Financial/financial_aid_support.html", context)


from decimal import Decimal, InvalidOperation
from django.db.models import Sum
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib.auth.decorators import login_required
from django.contrib import messages
import json



@login_required
def get_financial_aid_detail(request, aid_id):
    """Get financial aid details via AJAX"""
    try:
        aid = get_object_or_404(StudentFinancialAid, id=aid_id)
        
        student = aid.student
        user = student.user
        
        status_choices = dict(StudentFinancialAid.STATUS_CHOICES)
        status_display = status_choices.get(aid.status, aid.status or 'PENDING')
        
        type_choices = dict(StudentFinancialAid.AID_TYPE_CHOICES)
        type_display = type_choices.get(aid.aid_type, aid.aid_type or 'N/A')
        

        try:
            award_amount = float(aid.award_amount) if aid.award_amount is not None else 0
        except (ValueError, TypeError, InvalidOperation):
            award_amount = 0
        
        # Get all students for dropdown
        all_students = User.objects.filter(is_student=True, account_status='ACTIVE').order_by('first_name', 'last_name')
        students_data = []
        for s in all_students:
            students_data.append({
                'id': s.id,
                'name': s.get_full_name(),
                'student_id': getattr(s.student_profile, 'student_number', 'N/A') if hasattr(s, 'student_profile') else 'N/A',
            })
        
        data = {
            'id': aid.id,
            'student_id': user.id,
            'student_name': user.get_full_name() or 'Unknown',
            'student_initials': (user.first_name[0] if user.first_name else '') + 
                               (user.last_name[0] if user.last_name else '') or '?',
            'student_id_number': student.student_number if student.student_number else 'N/A',
            'email': user.email or 'No Email',
            'aid_type': aid.aid_type or 'N/A',
            'aid_type_display': type_display,
            'award_amount': str(award_amount),
            'academic_year': aid.academic_year or 'N/A',
            'status': aid.status or 'PENDING',
            'status_display': status_display,
            'created_at': aid.created_at.strftime('%Y-%m-%d %H:%M') if hasattr(aid, 'created_at') and aid.created_at else 'N/A',
            'updated_at': aid.updated_at.strftime('%Y-%m-%d %H:%M') if hasattr(aid, 'updated_at') and aid.updated_at else 'N/A',
            'status_choices': [{'value': k, 'label': v} for k, v in StudentFinancialAid.STATUS_CHOICES],
            'type_choices': [{'value': k, 'label': v} for k, v in StudentFinancialAid.AID_TYPE_CHOICES],
            'students': students_data,
        }
        
        return JsonResponse({'success': True, 'data': data})
        
    except Exception as e:
        print(f"Error in get_financial_aid_detail: {str(e)}")
        return JsonResponse({'success': False, 'error': str(e)}, status=500)




@login_required
def create_financial_aid(request):
    """Create new financial aid via AJAX"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)
    
    try:
        data = json.loads(request.body)
        
        student_id = data.get('student_id')
        aid_type = data.get('aid_type')
        award_amount = data.get('award_amount')
        academic_year = data.get('academic_year')
        status = data.get('status', 'PENDING')
        
        # Validate required fields
        if not student_id:
            return JsonResponse({'success': False, 'error': 'Student ID required'}, status=400)
        if not aid_type:
            return JsonResponse({'success': False, 'error': 'Aid type required'}, status=400)
        if not academic_year:
            return JsonResponse({'success': False, 'error': 'Academic year required'}, status=400)
        
        if status == 'AWARDED':
            try:
                amount = float(award_amount) if award_amount else 0
                if amount <= 0:
                    return JsonResponse({
                        'success': False, 
                        'error': 'Award amount must be greater than zero for AWARDED status'
                    }, status=400)
                award_amount = Decimal(str(amount))
            except (ValueError, TypeError, InvalidOperation):
                return JsonResponse({
                    'success': False, 
                    'error': 'Please enter a valid numeric award amount'
                }, status=400)
        else:
            award_amount = Decimal('0.00')
        
        student_user = get_object_or_404(User, id=student_id, is_student=True)
        student_profile = get_object_or_404(StudentProfile, user=student_user)
        
        aid = StudentFinancialAid.objects.create(
            student=student_profile,
            aid_type=aid_type,
            award_amount=award_amount,
            academic_year=academic_year,
            status=status,
        )
        # ==========================================
        # AUDIT: AFTER DATA
        # ==========================================
        after_data = AuditLogger.model_to_dict(
            aid,
            fields=[
                "student",
                "aid_type",
                "award_amount",
                "academic_year",
                "status",
            ]
        )
        # Add readable student information
        after_data["student_name"] = (
            aid.student.user.get_full_name()
        )
        after_data["student_id"] = (
            aid.student.student_number
        )
        # ==========================================
        # AUDIT LOG: CREATE
        # ==========================================
        AuditLogger.log(
            request=request,
            action="CREATE",
            module="Staff",
            object_type="Student Financial Aid",
            object_id=aid.id,
            description=(
                f'Staff created a financial aid record '
                f'for "{aid.student.user.get_full_name()}" '
                f'with status "{aid.status}".'
            ),
            before_data=None,
            after_data=after_data,
            status="SUCCESS",
        )
        
        return JsonResponse({
            'success': True,
            'message': 'Financial aid created successfully!',
            'data': {
                'id': aid.id,
                'student_name': aid.student.user.get_full_name(),
                'aid_type': aid.aid_type,
                'status': aid.status,
            }
        })
        
    except json.JSONDecodeError:
        return JsonResponse({'success': False, 'error': 'Invalid JSON data'}, status=400)
    except Exception as e:
        print(f"Error in create_financial_aid: {str(e)}")
        return JsonResponse({'success': False, 'error': str(e)}, status=500)



@login_required
def update_financial_aid(request):
    """Update financial aid via AJAX"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)
    
    try:
        data = json.loads(request.body)
        
        aid_id = data.get('aid_id')
        student_id = data.get('student_id')
        aid_type = data.get('aid_type')
        award_amount = data.get('award_amount')
        academic_year = data.get('academic_year')
        status = data.get('status')
        
        if not aid_id:
            return JsonResponse({'success': False, 'error': 'Aid ID required'}, status=400)
        
        aid = get_object_or_404(StudentFinancialAid, id=aid_id)

        # ==========================================
        # AUDIT: BEFORE DATA
        # IMPORTANT - BEFORE ANY CHANGES
        # ==========================================
        before_data = AuditLogger.model_to_dict(
            aid,
            fields=[
                "student",
                "aid_type",
                "award_amount",
                "academic_year",
                "status",
            ]
        )
        before_data["student_name"] = (
            aid.student.user.get_full_name()
        )
        before_data["student_id"] = (
            aid.student.student_number
        )
        if status == 'AWARDED':
            try:
                amount = float(award_amount) if award_amount else 0
                if amount <= 0:
                    return JsonResponse({
                        'success': False, 
                        'error': 'Award amount must be greater than zero for AWARDED status'
                    }, status=400)
                award_amount = Decimal(str(amount))
            except (ValueError, TypeError, InvalidOperation):
                return JsonResponse({
                    'success': False, 
                    'error': 'Please enter a valid numeric award amount'
                }, status=400)
        else:
            award_amount = Decimal('0.00')
        
        # Update fields
        if student_id:
            try:
                student_user = User.objects.get(id=student_id, is_student=True)
                student_profile = StudentProfile.objects.get(user=student_user)
                aid.student = student_profile
            except (User.DoesNotExist, StudentProfile.DoesNotExist):
                pass
        
        if aid_type:
            aid.aid_type = aid_type
        if academic_year:
            aid.academic_year = academic_year
        if status:
            aid.status = status
        
        aid.award_amount = award_amount
        aid.save()
        # ==========================================
        # AUDIT: AFTER DATA
        # ==========================================
        after_data = AuditLogger.model_to_dict(
            aid,
            fields=[
                "student",
                "aid_type",
                "award_amount",
                "academic_year",
                "status",
            ]
        )
        after_data["student_name"] = (
            aid.student.user.get_full_name()
        )
        after_data["student_id"] = (
            aid.student.student_number
        )
        # ==========================================
        # AUDIT LOG: UPDATE
        # ==========================================
        AuditLogger.log(
            request=request,
            action="UPDATE",
            module="Staff",
            object_type="Student Financial Aid",
            object_id=aid.id,
            description=(
                f'Staff updated financial aid record '
                f'#{aid.id} for '
                f'"{aid.student.user.get_full_name()}".'
            ),
            before_data=before_data,
            after_data=after_data,
            status="SUCCESS",
        )
        if status:
            notify_submitter_financial_aid_status(aid)
        
        return JsonResponse({
            'success': True,
            'message': 'Financial aid updated successfully!',
            'data': {
                'id': aid.id,
                'student_name': aid.student.user.get_full_name(),
                'aid_type': aid.aid_type,
                'status': aid.status,
            }
        })
        
    except json.JSONDecodeError:
        return JsonResponse({'success': False, 'error': 'Invalid JSON data'}, status=400)
    except Exception as e:
        print(f"Error in update_financial_aid: {str(e)}")
        return JsonResponse({'success': False, 'error': str(e)}, status=500)




@login_required
def staff_financial_aid_detail(request, uuid, aid_id):
    """
    View and manage individual financial aid request
    """
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    
    try:
        sp = request.user.staff_profile
    except Exception:
        return redirect("login")
    
    # Get the financial aid request with error handling
    try:
        aid = get_object_or_404(
            StudentFinancialAid.objects.select_related('student__user'),
            id=aid_id
        )
    except Exception as e:
        messages.error(request, f"Error loading financial aid: {str(e)}")
        return redirect('staff_requests', uuid=uuid)
    
    # Get student profile
    student = aid.student
    
    try:
        if aid.award_amount is None:
            aid.award_amount = Decimal('0.00')
            aid.save(update_fields=['award_amount'])
    except Exception:
        aid.award_amount = Decimal('0.00')
    
    if request.method == 'POST':
        action = request.POST.get('action')
        remarks = request.POST.get('remarks', '')
        
        if action == 'approve':
            award_amount_raw = request.POST.get('award_amount', 0)
            
            try:
                if not award_amount_raw:
                    messages.error(request, 'Award amount is required to approve financial aid.')
                    return redirect('staff_financial_aid_detail', uuid=uuid, aid_id=aid_id)
                
                award_amount = Decimal(str(award_amount_raw))
                
                if award_amount <= 0:
                    messages.error(request, 'Award amount must be greater than zero.')
                    return redirect('staff_financial_aid_detail', uuid=uuid, aid_id=aid_id)
                    
            except (InvalidOperation, ValueError, TypeError):
                messages.error(request, 'Please enter a valid numeric award amount.')
                return redirect('staff_financial_aid_detail', uuid=uuid, aid_id=aid_id)
            
            aid.status = 'AWARDED'
            aid.award_amount = award_amount
            if remarks:
                aid.remarks = remarks
            aid.save()

            notify_submitter_financial_aid_status(aid)
            
            messages.success(
                request,
                f"Financial aid approved for {student.user.get_full_name()} - ${award_amount}"
            )
            
        elif action == 'reject':
            aid.status = 'REJECTED'
            if remarks:
                aid.remarks = remarks
            aid.save()
            
            messages.success(
                request,
                f"Financial aid rejected for {student.user.get_full_name()}"
            )
            
        elif action == 'under_review' or action == 'review':
            aid.status = 'UNDER_REVIEW'
            if remarks:
                aid.remarks = remarks
            aid.save()
            
            messages.success(
                request,
                f"Financial aid marked as under review"
            )
            
        elif action == 'cancel':
            aid.status = 'CANCELLED'
            if remarks:
                aid.remarks = remarks
            aid.save()
            
            messages.success(
                request,
                f"Financial aid cancelled"
            )
        
        return redirect('staff_financial_aid_detail', uuid=uuid, aid_id=aid_id)
    
    # GET request - show detail
    # Get student's fee summary
    total_fees = student.fees.aggregate(
        total=Sum('amount')
    )['total'] or Decimal('0.00')
    
    total_paid = student.fee_payments.filter(
        status='SUCCESS'
    ).aggregate(
        total=Sum('amount')
    )['total'] or Decimal('0.00')
    
    student_balance = total_fees - total_paid
    
    context = {
        'aid': aid,
        'student': student,
        'total_fees': total_fees,
        'total_paid': total_paid,
        'student_balance': student_balance,
        'status_choices': StudentFinancialAid.STATUS_CHOICES,
    }
    
    return render(request, "Administrator/financial_aid_detail.html", context)

#<--------------------Blaze Code End Finance Sup(30-06-26)----------------------->


#<--------------------Blaze Code Start Finance bulk upload(08-08-26)----------------------->


from django.http import JsonResponse
from django.contrib.admin.views.decorators import staff_member_required
from django.views.decorators.csrf import csrf_exempt
from django.db import transaction
from decimal import Decimal
from django.utils import timezone
import json


@staff_member_required
def get_students_list(request):
    try:
        try:
            from Staff.models import StudentProfile
        except ImportError:
            from Students.models import StudentProfile
        
        students = StudentProfile.objects.select_related('user').all()
        
        data = []
        for student in students:
            user = student.user
            if user:
                full_name = user.get_full_name() or user.username or 'Unknown'
                student_number = student.student_number or 'N/A'
                email = user.email or 'No Email'
                
                data.append({
                    'id': student.id,
                    'full_name': full_name,
                    'student_number': student_number,
                    'email': email,
                })
        
        data.sort(key=lambda x: x['full_name'])
        
        return JsonResponse({
            'success': True,
            'data': data,
            'count': len(data)
        })
        
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e),
            'data': []
        }, status=500)


@staff_member_required
@csrf_exempt
def bulk_create_financial_aid(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Method not allowed'}, status=405)
    
    try:
        data = json.loads(request.body)
        
        student_ids = data.get('student_ids', [])
        aid_type = data.get('aid_type')
        award_amount = data.get('award_amount', 0)
        academic_year = data.get('academic_year')
        status = data.get('status', 'PENDING')
        
        if not student_ids:
            return JsonResponse({'success': False, 'error': 'No students selected'}, status=400)
        
        if not aid_type:
            return JsonResponse({'success': False, 'error': 'Aid type is required'}, status=400)
        
        if award_amount <= 0:
            return JsonResponse({'success': False, 'error': 'Invalid award amount'}, status=400)
        
        if not academic_year:
            return JsonResponse({'success': False, 'error': 'Academic year is required'}, status=400)
        
        try:
            from Staff.models import StudentProfile, StudentFinancialAid
        except ImportError:
            from Students.models import StudentProfile, StudentFinancialAid
        
        students = StudentProfile.objects.filter(id__in=student_ids)
        
        if not students:
            return JsonResponse({'success': False, 'error': 'No valid students found'}, status=404)
        
        created_count = 0
        errors = []
        
        with transaction.atomic():
            for student in students:
                try:
                    existing = StudentFinancialAid.objects.filter(
                        student=student,
                        academic_year=academic_year,
                        aid_type=aid_type
                    ).first()
                    
                    if existing:
                        existing.award_amount = Decimal(str(award_amount))
                        existing.status = status
                        existing.updated_date = timezone.now()
                        existing.save()
                    else:
                        StudentFinancialAid.objects.create(
                            student=student,
                            aid_type=aid_type,
                            award_amount=Decimal(str(award_amount)),
                            academic_year=academic_year,
                            status=status,
                            reason='Bulk add from student selection',
                            applied_date=timezone.now()
                        )
                    
                    created_count += 1
                    
                except Exception as e:
                    errors.append({
                        'student_id': student.student_number if student.student_number else 'Unknown',
                        'error': str(e)
                    })
        
        return JsonResponse({
            'success': True,
            'message': f'Created/Updated {created_count} records',
            'created': created_count,
            'errors': errors
        })
        
    except json.JSONDecodeError:
        return JsonResponse({'success': False, 'error': 'Invalid JSON data'}, status=400)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)

#<--------------------Blaze Code End Finance bulk upload(08-08-26)----------------------->


#<--------------------Blaze Code Start Finance Reports(10-08-26)----------------------->

from django.shortcuts import render, redirect
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.db.models import Count, Sum, Avg
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from django.http import HttpResponse
import json
import pandas as pd
from io import BytesIO
from decimal import Decimal
import traceback

@login_required
def financial_aid_reports(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    
    from Students.models import StudentFinancialAid
    
    selected_year = request.GET.get('year', 'all')
    selected_status = request.GET.get('status', 'all')
    selected_type = request.GET.get('type', 'all')
    export_format = request.GET.get('export', None)
    
    aid_records = StudentFinancialAid.objects.select_related('student', 'student__user').all()
    
    if selected_year != 'all': aid_records = aid_records.filter(academic_year=selected_year)
    if selected_status != 'all': aid_records = aid_records.filter(status=selected_status.upper())
    if selected_type != 'all': aid_records = aid_records.filter(aid_type=selected_type.upper())
    

    if export_format == 'excel':
        return generate_excel_response(aid_records)
    elif export_format == 'pdf':
        return generate_pdf_response(aid_records)
    
    # --- STATS & PAGINATION ---
    total_applications = aid_records.count()
    total_awarded = aid_records.filter(status='AWARDED').aggregate(total=Sum('award_amount'))['total'] or Decimal('0.00')
    avg_award = aid_records.aggregate(avg=Avg('award_amount'))['avg'] or 0
    awarded_count = aid_records.filter(status='AWARDED').count()
    approval_rate = (awarded_count / total_applications * 100) if total_applications > 0 else 0
    pending_count = aid_records.filter(status='PENDING').count()
    scholarship_count = aid_records.filter(aid_type='SCHOLARSHIP').count()

    type_distribution = aid_records.values('aid_type').annotate(count=Count('id'))
    type_labels = [dict(StudentFinancialAid.AID_TYPE_CHOICES).get(i['aid_type'], i['aid_type']) for i in type_distribution]
    type_data = [i['count'] for i in type_distribution]
    
    status_distribution = aid_records.values('status').annotate(count=Count('id'))
    status_labels = [dict(StudentFinancialAid.STATUS_CHOICES).get(i['status'], i['status']) for i in status_distribution]
    status_data = [i['count'] for i in status_distribution]

    aid_records = aid_records.order_by('-applied_date')
    paginator = Paginator(aid_records, 15)
    page = request.GET.get('page', 1)
    page_obj = paginator.get_page(page)

    start_index = (page_obj.number - 1) * paginator.per_page + 1
    end_index = min(page_obj.number * paginator.per_page, total_applications)

    context = {
        'uuid': uuid, 'aid_records': page_obj,
        'total_applications': total_applications, 'total_awarded': total_awarded,
        'avg_award': avg_award, 'approval_rate': approval_rate,
        'pending_count': pending_count, 'scholarship_count': scholarship_count,
        'type_labels': json.dumps(type_labels), 'type_data': json.dumps(type_data),
        'status_labels': json.dumps(status_labels), 'status_data': json.dumps(status_data),
        'selected_year': selected_year, 'selected_status': selected_status, 'selected_type': selected_type,
        'start_index': start_index, 'end_index': end_index, 'today': timezone.now(),
    }
    return render(request, 'Student_Services/Financial/financial_aid_reports.html', context)



def generate_excel_response(request, queryset):
    try:
        data = []
        for aid in queryset:
            data.append({
                'Student Name': aid.student.user.get_full_name() if aid.student.user else 'Unknown',
                'Student ID': aid.student.student_number or 'N/A',
                'Email': aid.student.user.email if aid.student.user else 'N/A',
                'Aid Type': aid.get_aid_type_display() or aid.aid_type,
                'Award Amount': float(aid.award_amount) if aid.award_amount else 0,
                'Academic Year': aid.academic_year or 'N/A',
                'Status': aid.get_status_display() or aid.status,
                'Applied Date': aid.applied_date.strftime('%Y-%m-%d') if aid.applied_date else 'N/A',
            })
        record_count = queryset.count()
        df = pd.DataFrame(data)
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Financial Aid Report', index=False)
            # Auto adjust width
            worksheet = writer.sheets['Financial Aid Report']
            for i, col in enumerate(df.columns):
                max_length = max(df[col].astype(str).map(len).max(), len(col)) + 2
                worksheet.column_dimensions[chr(65 + i)].width = min(max_length, 50)
        
        output.seek(0)
        filename = f"Financial_Aid_Report_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
        response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        # ==========================================
        # AUDIT LOG - EXCEL EXPORT
        # ==========================================
        AuditLogger.log(
            request=request,
            action="EXPORT",
            module="Staff",
            object_type="Financial Aid Report",
            object_id="",
            description=(
                f'Staff exported {record_count} '
                f'financial aid record(s) to Excel.'
            ),
            status="SUCCESS",
        )
        return response
    except Exception as e:
        print(f"Excel Error: {e}")
        traceback.print_exc()
        # ==========================================
        # AUDIT LOG - FAILED EXPORT
        # ==========================================
        AuditLogger.log(
            request=request,
            action="EXPORT",
            module="Staff",
            object_type="Financial Aid Report",
            object_id="",
            description=(
                f'Failed to export financial aid '
                f'report to Excel. Error: {str(e)}'
            ),
            status="FAILED",
        )
        return HttpResponse(f"Error: {str(e)}", status=500)



def generate_pdf_response(request, queryset):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        record_count = queryset.count()
        response = HttpResponse(content_type='application/pdf')
        filename = f"Financial_Aid_Report_{timezone.now().strftime('%Y%m%d_%H%M')}.pdf"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        doc = SimpleDocTemplate(response, pagesize=landscape(letter))
        elements = []
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=14, textColor=colors.HexColor('#1a1a1a'), spaceAfter=12)
        elements.append(Paragraph("Financial Aid Report", title_style))
        elements.append(Spacer(1, 0.2*inch))
        
        subtitle = f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')} | Records: {queryset.count()}"
        elements.append(Paragraph(subtitle, styles['Normal']))
        elements.append(Spacer(1, 0.3*inch))
        
        table_data = [['Student', 'ID', 'Type', 'Amount ($)', 'Year', 'Status', 'Date']]
        for aid in queryset[:100]:  # Limit to 100 rows for PDF performance
            table_data.append([
                (aid.student.user.get_full_name() or 'Unknown')[:20],
                (aid.student.student_number or 'N/A')[:12],
                (aid.get_aid_type_display() or aid.aid_type or 'N/A')[:12],
                f"{float(aid.award_amount):,.2f}" if aid.award_amount else '0.00',
                (aid.academic_year or 'N/A')[:9],
                (aid.get_status_display() or aid.status or 'N/A')[:12],
                aid.applied_date.strftime('%Y-%m-%d') if aid.applied_date else 'N/A',
            ])
        
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f8f8')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
            ('ALIGN', (3, 0), (3, -1), 'RIGHT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(table)
        
        elements.append(Spacer(1, 0.3*inch))
        footer_style = ParagraphStyle('Footer', parent=styles['Normal'], fontSize=8, textColor=colors.HexColor('#6b7280'), alignment=1)
        elements.append(Paragraph(f"Total Records: {queryset.count()} | University of Wisconsin", footer_style))
        
        doc.build(elements)
        # ==========================================
        # AUDIT LOG - PDF EXPORT
        # ==========================================
        AuditLogger.log(
            request=request,
            action="EXPORT",
            module="Staff",
            object_type="Financial Aid Report",
            object_id="",
            description=(
                f'Staff exported {record_count} '
                f'financial aid record(s) to PDF.'
            ),
            status="SUCCESS",
        )
        return response
    except Exception as e:
        print(f"PDF Error: {e}")
        traceback.print_exc()
        # ==========================================
        # AUDIT LOG - FAILED EXPORT
        # ==========================================
        AuditLogger.log(
            request=request,
            action="EXPORT",
            module="Staff",
            object_type="Financial Aid Report",
            object_id="",
            description=(
                f'Failed to export financial aid '
                f'report to PDF. Error: {str(e)}'
            ),
            status="FAILED",
        )
        return HttpResponse(f"Error: {str(e)}", status=500)

#<--------------------Blaze Code End Finance Reports(10-08-26)----------------------->



@login_required
def staff_advising(request, uuid):

    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    # =====================================================
    # SEARCH
    # =====================================================

    search = request.GET.get("q", "").strip()


    # =====================================================
    # ACTIVE DEPARTMENTS
    # =====================================================

    departments = Department.objects.filter(
        status="ACTIVE"
    ).annotate(
        active_program_count=Count(
            "programs",
            filter=Q(
                programs__status="ACTIVE"
            ),
            distinct=True
        )
    ).order_by(
        "department_name"
    )


    # Department search

    if search:

        departments = departments.filter(
            Q(department_name__icontains=search) |
            Q(department_code__icontains=search) |
            Q(short_name__icontains=search)
        )


    # =====================================================
    # STATISTICS
    # =====================================================

    total_departments = Department.objects.filter(
        status="ACTIVE"
    ).count()


    total_programs = AcademicProgram.objects.filter(
        status="ACTIVE",
        department__status="ACTIVE"
    ).count()


    total_active_students = StudentAcademicProfile.objects.filter(
        student__current_status="ACTIVE",
        program__status="ACTIVE"
    ).count()


    total_active_advisors = FacultyProfile.objects.filter(
        employment_status="ACTIVE"
    ).count()


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        "user": request.user,

        "departments": departments,

        "search": search,

        # Stats
        "total_departments": total_departments,
        "total_programs": total_programs,
        "total_active_students": total_active_students,
        "total_active_advisors": total_active_advisors,

    }


    return render(
        request,
        "Student_Services/advising_support.html",
        context
    )


AVATAR_PALETTE = [
    ("#eff6ff", "#2563eb"),
    ("#f5f3ff", "#7c3aed"),
    ("#fffbeb", "#d97706"),
    ("#f0fdf4", "#16a34a"),
    ("#fff0f0", "#C5050C"),
]


def _avatar_colors(seed_text):
    idx = sum(ord(c) for c in (seed_text or "?")) % len(AVATAR_PALETTE)
    return AVATAR_PALETTE[idx]


def _normalize_staff(sp, dept_lookup):
    pos = sp.positions.filter(position_status="ACTIVE").first() or sp.positions.first()
    fname = sp.user.first_name or ""
    lname = sp.user.last_name or ""
    bg, fg = _avatar_colors(sp.user.get_full_name() or sp.employee_id)
    return {
        "kind": "staff",
        "id": sp.id,
        "full_name": sp.preferred_name or sp.user.get_full_name() or sp.user.username,
        "initials": (fname[:1] + lname[:1]).upper() or "?",
        "photo_url": sp.user.profile_photo.url if sp.user.profile_photo else None,   
        "email": sp.work_email,
        "employee_id": sp.employee_id,
        "department_name": dept_lookup.get(pos.department_id, "—") if pos else "—",
        "job_title": pos.job_title if pos else "—",
        "employment_type_display": sp.get_employment_type_display() or "—",
        "employment_status": sp.employment_status,
        "employment_status_display": sp.get_employment_status_display(),
        "phone": sp.office_phone,
        "avatar_bg": bg,
        "avatar_fg": fg,
        "search_blob": f"{sp.user.get_full_name()} {sp.employee_id} {sp.work_email} {pos.job_title if pos else ''}".lower(),
    }


def _normalize_faculty(fp, dept_lookup):
    fname = fp.user.first_name or ""
    lname = fp.user.last_name or ""
    bg, fg = _avatar_colors(fp.user.get_full_name() or fp.employee_id)
    return {
        "kind": "faculty",
        "id": fp.id,
        "full_name": fp.preferred_name or fp.user.get_full_name() or fp.user.username,
        "initials": (fname[:1] + lname[:1]).upper() or "?",
        "photo_url": fp.user.profile_photo.url if fp.user.profile_photo else (
            fp.profile_photo.url if getattr(fp, "profile_photo", None) else None
        ),   
        "email": fp.email,
        "employee_id": fp.employee_id,
        "department_name": dept_lookup.get(fp.department_id, "—"),
        "job_title": fp.faculty_rank.rank_name if fp.faculty_rank else "Faculty",
        "employment_type_display": "Faculty",
        "employment_status": fp.employment_status,
        "employment_status_display": fp.get_employment_status_display(),
        "phone": fp.phone,
        "avatar_bg": bg,
        "avatar_fg": fg,
        "search_blob": f"{fp.user.get_full_name()} {fp.employee_id} {fp.email} {fp.faculty_rank.rank_name if fp.faculty_rank else ''}".lower(),
    }


@login_required
def staff_directory(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    dept_lookup = dict(Department.objects.values_list("department_id", "department_name"))

    staff_qs = StaffProfile.objects.select_related("user").prefetch_related("positions")
    faculty_qs = FacultyProfile.objects.select_related("user", "faculty_rank")

    members = [_normalize_staff(sp, dept_lookup) for sp in staff_qs]
    members += [_normalize_faculty(fp, dept_lookup) for fp in faculty_qs]
    members.sort(key=lambda m: m["full_name"])

    total_staff = len(members)
    faculty_count = sum(1 for m in members if m["kind"] == "faculty")
    staff_count = total_staff - faculty_count
    active_count = sum(1 for m in members if m["employment_status"] == "ACTIVE")
    leave_count = sum(1 for m in members if m["employment_status"] == "LEAVE")
    departments_count = Department.objects.filter(status="ACTIVE").count()

    search_query  = request.GET.get("search", "").strip().lower()
    status_filter = request.GET.get("status", "").strip().upper()
    dept_filter   = request.GET.get("department", "").strip()
    tab_filter    = request.GET.get("tab", "all").strip()

    filtered = members
    if search_query:
        filtered = [m for m in filtered if search_query in m["search_blob"]]
    if status_filter:
        filtered = [m for m in filtered if m["employment_status"] == status_filter]
    if dept_filter.isdigit():
        dept_name = dept_lookup.get(int(dept_filter))
        filtered = [m for m in filtered if m["department_name"] == dept_name]
    if tab_filter == "faculty":
        filtered = [m for m in filtered if m["kind"] == "faculty"]
    elif tab_filter == "staff":
        filtered = [m for m in filtered if m["kind"] == "staff"]
    elif tab_filter == "active":
        filtered = [m for m in filtered if m["employment_status"] == "ACTIVE"]
    elif tab_filter == "leave":
        filtered = [m for m in filtered if m["employment_status"] == "LEAVE"]

    paginator = Paginator(filtered, 10)
    page_obj = paginator.get_page(request.GET.get("page", 1))

    dept_counts = {}
    for m in members:
        dept_counts[m["department_name"]] = dept_counts.get(m["department_name"], 0) + 1
    top_departments = sorted(dept_counts.items(), key=lambda kv: -kv[1])[:5]
    max_total = max([c for _, c in top_departments], default=1)
    top_departments = [
        {"name": name, "total": count, "pct": round((count / max_total) * 100)}
        for name, count in top_departments
    ]

    context = {
        "user": request.user,
        "directory_members": page_obj,
        "total_staff": total_staff,
        "faculty_count": faculty_count,
        "staff_count": staff_count,
        "departments_count": departments_count,
        "active_count": active_count,
        "leave_count": leave_count,
        "search_query": search_query,
        "status_filter": status_filter,
        "dept_filter": dept_filter,
        "tab_filter": tab_filter,
        "all_departments": Department.objects.filter(status="ACTIVE").order_by("department_name"),
        "top_departments": top_departments,
    }
    return render(request, "Academicsupport/directory.html", context)


@login_required
def staff_course_support(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    from Admin.bela_admin.models import Department, Course
    from Students.models import (
        Course as StudentCourse, CourseSection, StudentEnrollment,
        CourseMaterial, SupportTicket,
    )

    courses_qs = Course.objects.select_related('department').order_by('course_code')

    search_query  = request.GET.get('search', '').strip()
    status_filter = request.GET.get('status', 'all').strip().lower()
    dept_filter   = request.GET.get('department', '').strip()

    if search_query:
        courses_qs = courses_qs.filter(
            Q(course_name__icontains=search_query) |
            Q(course_code__icontains=search_query) |
            Q(department__department_name__icontains=search_query)
        )
    if dept_filter.isdigit():
        courses_qs = courses_qs.filter(department_id=int(dept_filter))

    total_courses = courses_qs.count()

    student_course_by_code = {c.course_code: c for c in StudentCourse.objects.all()}

    courses_data = []
    unassigned_count = 0
    total_enrolled_all = 0
    dept_enroll_counts = {}   # department_name -> enrolled total, for the widget
    dept_capacity_counts = {} # department_name -> capacity total, for pct

    for course in courses_qs:
        student_course = student_course_by_code.get(course.course_code)

        section = None
        if student_course:
            section = (
                CourseSection.objects.filter(course=course)
                .select_related('semester_id')
                .order_by('-semester_id__start_date')
                .first()
            )

        instructor_name = getattr(section, 'faculty_name', None) if section else None
        schedule        = getattr(section, 'schedule', None) if section else None
        capacity        = getattr(section, 'capacity', None) if section else 0

        enrolled_count = 0
        if section:
            enrolled_count = StudentEnrollment.objects.filter(
                section_id=section
            ).exclude(enrollment_status__in=['DROPPED', 'CANCELLED']).count()

        capacity = capacity or 0
        fill_pct = round((enrolled_count / capacity) * 100) if capacity else 0

        dept_name = course.department.department_name if course.department else 'Unassigned Dept'
        dept_enroll_counts[dept_name] = dept_enroll_counts.get(dept_name, 0) + enrolled_count
        dept_capacity_counts[dept_name] = dept_capacity_counts.get(dept_name, 0) + capacity

        is_unassigned = not instructor_name
        if is_unassigned:
            unassigned_count += 1
        total_enrolled_all += enrolled_count

        row_status = 'unassigned' if is_unassigned else (course.status or 'ACTIVE').lower()
        initials = ''.join(p[0] for p in instructor_name.split()[:2]).upper() if instructor_name else '?'

        courses_data.append({
            'course_id': course.course_id,
            'course_uuid': course.course_uuid,
            'name': course.course_name,
            'code': course.course_code,
            'credits': course.credits,
            'department': dept_name,
            'instructor': instructor_name or 'Unassigned',
            'instructor_initials': initials,
            'enrolled': enrolled_count,
            'capacity': capacity,
            'fill_pct': fill_pct,
            'schedule': schedule or 'TBD',
            'status': row_status,
            'status_display': 'Unassigned' if is_unassigned else course.get_status_display(),
        })

    if status_filter and status_filter != 'all':
        courses_data = [c for c in courses_data if c['status'] == status_filter]

    querydict = request.GET.copy()
    querydict.pop('page', None)
    querystring = querydict.urlencode()

    paginator = Paginator(courses_data, 10)
    page_obj  = paginator.get_page(request.GET.get('page', 1))

    # --- Smart Pages for pagination ---
    current = page_obj.number
    total   = paginator.num_pages

    def smart_range(current, total):
        pages = set()
        pages.update([1, 2])
        pages.update([total - 1, total])
        for i in range(current - 2, current + 3):
            if 1 <= i <= total:
                pages.add(i)
        return sorted(pages)

    page_range = smart_range(current, total)

    smart_pages = []
    prev = None
    for p in page_range:
        if prev is not None and p - prev > 1:
            smart_pages.append(None)
        smart_pages.append(p)
        prev = p

    # --- Enrollment by Department widget ---
    enrollment_by_dept = []
    for dept_name, enrolled in dept_enroll_counts.items():
        cap = dept_capacity_counts.get(dept_name, 0)
        pct = round((enrolled / cap) * 100) if cap else 0
        enrollment_by_dept.append({'name': dept_name, 'pct': pct})
    enrollment_by_dept.sort(key=lambda d: -d['pct'])
    enrollment_by_dept = enrollment_by_dept[:5]

    # --- Recent Uploads widget ---
    recent_uploads = CourseMaterial.objects.select_related(
        'course_section__course', 'uploaded_by'
    ).order_by('-uploaded_at')[:4]
    recent_uploads_data = [{
        'title': f"{m.course_section.course.course_code} — {m.title}",
        'uploader': m.uploaded_by.get_full_name() if m.uploaded_by else 'Unknown',
        'date': m.uploaded_at,
    } for m in recent_uploads]

    # --- Open Support Tickets widget ---
    open_tickets = SupportTicket.objects.select_related(
        'course_section__course', 'submitted_by'
    ).exclude(status='RESOLVED').order_by('-priority', '-created_at')[:5]
    open_tickets_data = [{
        'id': t.id,
        'title': f"{t.course_section.course.course_code if t.course_section else 'General'} — {t.subject}",
        'submitted_by': t.submitted_by.get_full_name() if t.submitted_by else 'Unknown',
        'date': t.created_at,
        'priority': t.priority,
    } for t in open_tickets]

    support_tickets_open_count = SupportTicket.objects.exclude(status='RESOLVED').count()
    support_tickets_urgent_count = SupportTicket.objects.filter(priority='URGENT').exclude(status='RESOLVED').count()

    context = {
        'user': request.user,
        'courses': page_obj,
        'total_courses': total_courses,
        'assigned_courses': total_courses - unassigned_count,
        'unassigned_courses': unassigned_count,
        'total_enrolled': total_enrolled_all,
        'support_tickets_open': support_tickets_open_count,
        'support_tickets_urgent': support_tickets_urgent_count,
        'enrollment_by_dept': enrollment_by_dept,
        'recent_uploads': recent_uploads_data,
        'open_tickets': open_tickets_data,
        'departments': Department.objects.filter(status='ACTIVE').order_by('department_name'),
        'search_query': search_query,
        'status_filter': status_filter,
        'dept_filter': dept_filter,
        'querystring': querystring,
        'smart_pages': smart_pages,
        'page_obj': page_obj,
    }
    return render(request, "Academicsupport/coursesupport.html", context)


 
@login_required
def staff_course_view(request, uuid, course_uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
 
    from Admin.bela_admin.models import Course
    from Students.models import CourseSection
 
    course = get_object_or_404(
        Course.objects.select_related('department'),
        course_uuid=course_uuid
    )
 
    sections = CourseSection.objects.filter(course=course).order_by('section_number')
    total_sections = sections.count()
 
    context = {
        "course": course,
        "sections": sections,
        "total_sections": total_sections,
    }
    return render(request, "Academicsupport/course_view.html", context)
 


@login_required
def staff_scheduling(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
 
    from Admin.bela_admin.models import Course
 
    all_sections = CourseSection.objects.select_related(
        "course", "semester_id"
    ).prefetch_related("schedules")
 
    search = request.GET.get("search", "").strip()
    day = request.GET.get("day", "").strip()
    status_filter = request.GET.get("status", "all").strip().lower()
 
    if search:
        all_sections = all_sections.filter(
            Q(course__course_name__icontains=search) |
            Q(course__course_code__icontains=search) |
            Q(faculty_name__icontains=search) |
            Q(room_number__icontains=search)
        )
 
    if day:
        day_code = day[:3].upper()
        all_sections = all_sections.filter(
            schedules__day_of_week=day_code
        ).distinct()
 
    course_map = {}
    for sec in all_sections:
        course = sec.course
        if not course:
            continue
        cid = course.course_id
 
        if cid not in course_map:
            course_map[cid] = {
                "course_id": cid,
                "course_name": course.course_name,
                "course_code": course.course_code,
                "department": course.department.department_name if course.department else "—",
                "total_sections": 0,
                "total_enrolled": 0,
                "total_capacity": 0,
                "statuses": [],
            }
 
        course_map[cid]["total_sections"] += 1
        cap = sec.capacity or 0
        course_map[cid]["total_capacity"] += cap
 
        enrolled = StudentEnrollment.objects.filter(
            section_id=sec
        ).exclude(enrollment_status__in=["WITHDRAWN"]).count()
        course_map[cid]["total_enrolled"] += enrolled
        course_map[cid]["statuses"].append(sec.status)
 
    rows = []
    for cid, info in course_map.items():
        statuses = info["statuses"]
        if "CONFLICT" in statuses:
            overall_status = "CONFLICT"
        elif all(s == "CONFIRMED" for s in statuses):
            overall_status = "CONFIRMED"
        elif any(s == "PENDING" for s in statuses):
            overall_status = "PENDING"
        elif all(s == "CANCELLED" for s in statuses):
            overall_status = "CANCELLED"
        else:
            overall_status = statuses[0] if statuses else "PENDING"
 
        status_map = {
            "CONFIRMED": "Confirmed",
            "PENDING": "Pending",
            "CONFLICT": "Conflict",
            "CANCELLED": "Cancelled",
        }
 
        rows.append({
            "course_id": info["course_id"],
            "course_name": info["course_name"],
            "course_code": info["course_code"],
            "department": info["department"],
            "total_sections": info["total_sections"],
            "enrolled": info["total_enrolled"],
            "capacity": info["total_capacity"],
            "status": overall_status,
            "status_lower": overall_status.lower(),
            "status_display": status_map.get(overall_status, overall_status),
        })
 
    rows.sort(key=lambda r: r["course_code"])
 
    total_schedules = len(rows)
    confirmed = sum(1 for r in rows if r["status"] == "CONFIRMED")
    pending = sum(1 for r in rows if r["status"] == "PENDING")
    conflicts = sum(1 for r in rows if r["status"] == "CONFLICT")
    cancelled = sum(1 for r in rows if r["status"] == "CANCELLED")
    online_classes = Schedule.objects.filter(is_online=True).count()
   
    rooms_booked = (Schedule.objects.exclude(room__isnull=True).exclude(room="").values("room", "building").distinct().count())
        # Department Schedule Overview
    department_overview = (
        CourseSection.objects
        .values("course__department__department_name")
        .annotate(
            total_classes=Count("section_id")
        )
        .order_by("-total_classes")
    )
 
 
    # Class Type Distribution
    class_type_raw = (
        CourseSection.objects
        .values("section_type")
        .annotate(
            total=Count("section_id")
        )
    )
 
 
    total_class_types = sum(
        item["total"] for item in class_type_raw
    )
 
 
    class_type_data = []
 
    for item in class_type_raw:
 
        percentage = 0
 
        if total_class_types:
            percentage = round(
                (item["total"] / total_class_types) * 100
            )
 
        class_type_data.append({
            "section_type": item["section_type"],
            "total": item["total"],
            "percentage": percentage
        })
 
 
    # Semester Schedule Status
    semester_raw = (
        CourseSection.objects
        .values("semester_id__semester_code")
        .annotate(
            total=Count("section_id")
        )
    )
 
 
    total_semester = sum(
        item["total"] for item in semester_raw
    )
 
 
    semester_status = []
 
    for item in semester_raw:
 
        percentage = 0
 
        if total_semester:
            percentage = round(
                (item["total"] / total_semester) * 100
            )
 
        semester_status.append({
            "semester_id__semester_code":
                item["semester_id__semester_code"],
 
            "total":
                item["total"],
 
            "percentage":
                percentage
        })
    total_rooms = Room.objects.count()
    confirmed_percentage = (
    round((confirmed / total_schedules) * 100)
    if total_schedules else 0
)
    total_schedules_this_semester = total_schedules
    pending_increase = pending
    conflicts_increase = conflicts
    rooms_free_today = max(total_rooms - rooms_booked, 0)
    reschedule_requests = pending
 
    if status_filter and status_filter != "all":
        rows = [r for r in rows if r["status_lower"] == status_filter]
 
    paginator = Paginator(rows, 10)
    page_obj = paginator.get_page(request.GET.get("page", 1))
    department_schedule = (
      CourseSection.objects
      .select_related("course__department")
      .values(
        "course__department__department_name"
     )
      .annotate(
        total_classes=Count("section_id")
     )
     .order_by("-total_classes")
)
 
 
    department_overview = []
 
    for dept in department_schedule:
 
      department_overview.append({
 
        "name": dept["course__department__department_name"]
        or "Unknown",
 
        "total_classes": dept["total_classes"]
 
    })
 
    class_type_data = (
        CourseSection.objects
        .values("section_type")
        .annotate(
            total=Count("section_id")
        )
        .order_by("-total")
    )
 
    semester_status = (
        CourseSection.objects
        .values(
            "semester_id__semester_code"
        )
        .annotate(
            total=Count("section_id")
        )
        .order_by("-total")
    )
 
    context = {
    "user": request.user,
    "schedules": page_obj,
 
    "total_schedules": total_schedules,
    "total_schedules_this_semester": total_schedules_this_semester,
 
    "confirmed": confirmed,
    "confirmed_percentage": confirmed_percentage,
 
    "pending": pending,
    "pending_increase": pending_increase,
 
    "conflicts": conflicts,
    "conflicts_increase": conflicts_increase,
 
    "rooms_booked": rooms_booked,
    "total_rooms": total_rooms,
    "rooms_free_today": rooms_free_today,
 
    "reschedule_requests": reschedule_requests,
 
    "cancelled": cancelled,
 
    "search_query": search,
    "day_filter": day,
    "online_classes": online_classes,
    "status_filter": status_filter,
    "department_overview": department_overview,
 
    "class_type_data": class_type_data,
    "semester_status": semester_status,
    "department_overview": department_overview,
    "class_type_data": class_type_data,
    "semester_status": semester_status,
    }
    return render(request, "Academicsupport/scheduling.html", context)
 

from django.db.models import Q

def detect_conflicts():
    CourseSection.objects.exclude(status="CANCELLED").update(status="CONFIRMED")
    schedules = Schedule.objects.select_related("section_id").all()
    for s in schedules:
        overlaps = Schedule.objects.filter(
            day_of_week=s.day_of_week,
            room=s.room,
            building=s.building,
        ).exclude(id=s.id).filter(
            Q(start_time__lt=s.end_time) & Q(end_time__gt=s.start_time)
        )
        if overlaps.exists():
            s.section_id.status = "CONFLICT"
            s.section_id.save(update_fields=["status"])

# @login_required
# def staff_exams(request, uuid):
#     if str(request.user.uuid) != str(uuid):
#         return redirect("login")

#     return render(request, "Academicsupport/exams.html", {"user": request.user})
 
@login_required
def schedule_view(request, uuid):
    """View all sections and schedule entries for a specific course"""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
 
    from Admin.bela_admin.models import Course
 
    course_id = request.GET.get("course_id")
    if not course_id:
        return redirect("staff_scheduling", uuid=uuid)
 
    try:
        course = get_object_or_404(
            Course.objects.select_related("department"),
            course_id=course_id,
        )
    except Exception:
        return redirect("staff_scheduling", uuid=uuid)
 
    sections = CourseSection.objects.filter(
        course=course
    ).select_related("semester_id").prefetch_related("schedules").order_by("section_number")
 
    sections_data = []
    all_schedules = []
    for sec in sections:
        enrolled = StudentEnrollment.objects.filter(
            section_id=sec
        ).exclude(enrollment_status__in=["WITHDRAWN"]).count()
 
        sections_data.append({
            "section_id": sec.section_id,
            "section_number": sec.section_number,
            "faculty_name": sec.faculty_name or "Unassigned",
            "room_number": sec.room_number or "—",
            "building_name": sec.building_name or "—",
            "capacity": sec.capacity,
            "enrolled": enrolled,
            "semester": str(sec.semester_id) if sec.semester_id else "—",
            "status": sec.status,
            "status_display": sec.get_status_display(),
        })
 
        for sched in sec.schedules.all():
            all_schedules.append({
                "schedule_id": sched.schedule_id,
                "section_number": sec.section_number,
                "day_of_week": sched.get_day_of_week_display(),
                "start_time": sched.start_time.strftime("%I:%M %p"),
                "end_time": sched.end_time.strftime("%I:%M %p"),
                "room": sched.room or "—",
                "building": sched.building or "—",
                "is_online": sched.is_online,
                "dates": sched.dates or [],
                "section_status": sec.status,
                "section_status_display": sec.get_status_display(),
            })
 
    all_schedules.sort(key=lambda s: s["section_number"])
 
    context = {
        "user": request.user,
        "course": course,
        "sections": sections_data,
        "schedules": all_schedules,
    }
    return render(request, "Academicsupport/schedule_view.html", context)
 
 
 

@login_required
def staff_reports(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    return render(request, "Resources/reports.html", {"user": request.user})


@login_required
def staff_knowledge(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    return render(request, "Resources/knowledge.html", {"user": request.user})


@login_required
def staff_downloads(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    return render(request, "Resources/downloads.html")

#<------------------------------Blaze code start(06-07-26)------------------------------>

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse
from django.utils import timezone
from django.db import transaction
from Admin.models import User
from Staff.models import StaffProfile
import json
import logging
import os
from PIL import Image
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth import update_session_auth_hash
from django.core.exceptions import ValidationError
from django.views.decorators.http import require_http_methods

logger = logging.getLogger(__name__)

@login_required
def staff_profile_settings(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    
    try:
        staff_profile = request.user.staff_profile
    except StaffProfile.DoesNotExist:
        messages.error(request, "Staff profile not found.")
        return redirect('staff_dashboard', uuid=uuid)
    
    context = {
        'user': request.user,
        'staff': staff_profile,
        'today': timezone.now().date(),
        'tfa_enabled': getattr(staff_profile, 'tfa_enabled', False),
        'sessions': get_active_sessions(request),
    }
    
    return render(request, "Settings/profile_settings.html", context)

def get_active_sessions(request):
    """Get active sessions for the user"""
    try:
        from Admin.models import UserSession
        sessions = UserSession.objects.filter(
            user=request.user,
            logout_time__isnull=True
        ).order_by('-login_time')
        
        session_list = []
        for session in sessions:
            session_list.append({
                'session_key': session.session_id,
                'device': session.device_type or 'Unknown Device',
                'location': 'Unknown',
                'last_active': session.login_time,
                'is_current': session.session_id == request.session.session_key,
            })
        return session_list
    except:
        return []

@login_required
def update_staff_profile_ajax(request, uuid):
    """Update staff profile via AJAX"""
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'message': 'Unauthorized'}, status=403)
    
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid method'}, status=400)
    
    try:
        # Handle file upload (profile photo) with validation
        if request.FILES.get('profile_photo'):
            file = request.FILES['profile_photo']
            
            # Check file type
            allowed_types = [
                'image/jpeg', 'image/png', 'image/gif', 
                'image/webp', 'image/svg+xml', 'image/bmp'
            ]
            
            if file.content_type not in allowed_types:
                return JsonResponse({
                    'success': False,
                    'message': 'Only image files are allowed (JPG, PNG, GIF, WEBP, SVG, BMP)'
                }, status=400)
            
            # Check file extension
            allowed_extensions = ['.jpg', '.jpeg', '.png', '.gif', '.webp', '.svg', '.bmp']
            ext = os.path.splitext(file.name)[1].lower()
            
            if ext not in allowed_extensions:
                return JsonResponse({
                    'success': False,
                    'message': f'Invalid file extension "{ext}". Allowed: JPG, PNG, GIF, WEBP, SVG, BMP'
                }, status=400)
            
            # Check file size (2MB max)
            max_size = 2 * 1024 * 1024
            if file.size > max_size:
                size_mb = file.size / 1024 / 1024
                return JsonResponse({
                    'success': False,
                    'message': f'File is too large ({size_mb:.2f}MB). Maximum size is 2MB'
                }, status=400)
            
            try:
                img = Image.open(file)
                img.verify()
                
                file.seek(0)
                img = Image.open(file)
                width, height = img.size
                
                if width < 100 or height < 100:
                    return JsonResponse({
                        'success': False,
                        'message': 'Image is too small. Minimum size is 100x100 pixels.'
                    }, status=400)
                
                if width > 4000 or height > 4000:
                    return JsonResponse({
                        'success': False,
                        'message': 'Image is too large. Maximum size is 4000x4000 pixels.'
                    }, status=400)
                    
            except Exception as e:
                return JsonResponse({
                    'success': False,
                    'message': f'Invalid image file: {str(e)}'
                }, status=400)
            
            request.user.profile_photo = file
            request.user.save()
            
            messages.success(request, 'Profile photo uploaded successfully!')
            
            return JsonResponse({
                'success': True, 
                'message': 'Photo uploaded successfully!',
                'photo_url': request.user.profile_photo.url if request.user.profile_photo else None
            })
        
        # Handle JSON data
        data = json.loads(request.body)
        user = request.user
        staff = user.staff_profile
        
        with transaction.atomic():
            if 'first_name' in data:
                user.first_name = data['first_name'].strip()
            if 'last_name' in data:
                user.last_name = data['last_name'].strip()
            if 'work_email' in data:
                staff.work_email = data['work_email'].strip()
            if 'personal_email' in data:
                staff.personal_email = data['personal_email'].strip()
            if 'preferred_name' in data:
                staff.preferred_name = data['preferred_name'].strip()
            if 'office_phone' in data:
                staff.office_phone = data['office_phone'].strip()
            if 'employee_id' in data:
                staff.employee_id = data['employee_id'].strip()
            if 'position' in data:
                staff.position = data['position'].strip()
            if 'department' in data:
                staff.department = data['department'].strip()
            if 'bio' in data:
                staff.bio = data['bio'].strip()
            
            user.save()
            staff.save()
        

        messages.success(request,'Profile updated successfully!')
        
        return JsonResponse({
            'success': True,
            'message': 'Profile updated successfully!',
            'data': {
                'full_name': user.get_full_name(),
                'photo_url': user.profile_photo.url if user.profile_photo else None,
            }
        })
        
    except json.JSONDecodeError:
        return JsonResponse({'success': False, 'message': 'Invalid JSON data'}, status=400)
    except Exception as e:
        logger.error(f"Error updating staff profile: {str(e)}")
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


@login_required
def staff_remove_photo_ajax(request, uuid):
    """Remove profile photo via AJAX"""
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'message': 'Unauthorized'})
    
    try:
        user = request.user
        if user.profile_photo:
            user.profile_photo.delete(save=False)
            user.save()
            
            messages.success(request, 'Profile photo removed successfully!')
            
            return JsonResponse({
                'success': True,
                'message': 'Photo removed successfully'
            })
        else:
            return JsonResponse({
                'success': False,
                'message': 'No photo to remove'
            })
    except Exception as e:
        messages.error(request, f'Error: {str(e)}')
        return JsonResponse({
            'success': False,
            'message': str(e)
        })


@login_required
@require_http_methods(["POST"])
def change_password_ajax(request, uuid):
    """Handle password change via AJAX"""
    try:
        if str(request.user.uuid) != str(uuid):
            return JsonResponse({'success': False, 'message': 'Invalid request'})
        
        data = json.loads(request.body)
        current_password = data.get('current_password', '')
        new_password = data.get('new_password', '')
        confirm_password = data.get('confirm_password', '')
        
        if not current_password or not new_password or not confirm_password:
            return JsonResponse({'success': False, 'message': 'All fields are required'})
        
        if new_password != confirm_password:
            return JsonResponse({'success': False, 'message': 'Passwords do not match'})
        
        if not request.user.check_password(current_password):
            return JsonResponse({'success': False, 'message': 'Current password is incorrect'})
        
        try:
            validate_password(new_password, request.user)
        except ValidationError as e:
            return JsonResponse({'success': False, 'message': ' '.join(e.messages)})
        
        request.user.set_password(new_password)
        request.user.save()
        update_session_auth_hash(request, request.user)
        
        # ADD DJANGO MESSAGE
        messages.success(request, 'Password updated successfully!')
        
        return JsonResponse({'success': True, 'message': 'Password updated successfully'})
    
    except Exception as e:
        messages.error(request, f'Error: {str(e)}')
        return JsonResponse({'success': False, 'message': str(e)})

#<------------------------------Blaze code End(06-07-26)------------------------------>

#<------------------------------Blaze code start(06-07-26)------------------------------>


@login_required
@require_http_methods(["POST"])

def setup_2fa_ajax(request, uuid):
    """Generate 2FA secret key and QR code"""
    try:
        if str(request.user.uuid) != str(uuid):
            return JsonResponse({'success': False, 'message': 'Invalid request'})
        
        secret = pyotp.random_base32()
        request.session['2fa_secret'] = secret
        
        totp = pyotp.TOTP(secret)
        provisioning_uri = totp.provisioning_uri(
            name=request.user.email,
            issuer_name="Staff Portal"
        )
        
        qr = qrcode.QRCode(box_size=10, border=4)
        qr.add_data(provisioning_uri)
        qr.make(fit=True)
        img = qr.make_image(fill_color="black", back_color="white")
        
        buffer = io.BytesIO()
        img.save(buffer, format='PNG')
        qr_base64 = base64.b64encode(buffer.getvalue()).decode()
        
        return JsonResponse({
            'success': True,
            'secret_key': secret,
            'qr_code': f'data:image/png;base64,{qr_base64}'
        })
    
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

@login_required
@require_http_methods(["POST"])
def verify_2fa_ajax(request, uuid):
    """Verify 2FA setup"""
    try:
        if str(request.user.uuid) != str(uuid):
            return JsonResponse({'success': False, 'message': 'Invalid request'})
        
        data = json.loads(request.body)
        code = data.get('code')
        secret = request.session.get('2fa_secret')
        
        if not code or not secret:
            return JsonResponse({'success': False, 'message': 'Invalid setup'})
        
        totp = pyotp.TOTP(secret)
        if totp.verify(code):
            profile = request.user.staff_profile
            profile.tfa_enabled = True
            profile.tfa_secret = secret
            profile.save()
            
            backup_codes = []
            for _ in range(10):
                backup_codes.append(pyotp.random_base32()[:12])
            profile.tfa_backup_codes = json.dumps(backup_codes)
            profile.save()
            
            del request.session['2fa_secret']
            return JsonResponse({'success': True, 'message': '2FA enabled successfully'})
        
        return JsonResponse({'success': False, 'message': 'Invalid verification code'})
    
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

@login_required
@require_http_methods(["POST"])
def disable_2fa_ajax(request, uuid):
    """Disable 2FA"""
    try:
        if str(request.user.uuid) != str(uuid):
            return JsonResponse({'success': False, 'message': 'Invalid request'})
        
        data = json.loads(request.body)
        password = data.get('password')
        
        if not request.user.check_password(password):
            return JsonResponse({'success': False, 'message': 'Invalid password'})
        
        profile = request.user.staff_profile
        profile.tfa_enabled = False
        profile.tfa_secret = None
        profile.tfa_backup_codes = None
        profile.save()
        
        return JsonResponse({'success': True, 'message': '2FA disabled'})
    
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

@login_required
@require_http_methods(["GET"])
def get_backup_codes_ajax(request, uuid):
    """Get backup codes"""
    try:
        if str(request.user.uuid) != str(uuid):
            return JsonResponse({'success': False, 'message': 'Invalid request'})
        
        profile = request.user.staff_profile
        if not profile.tfa_enabled:
            return JsonResponse({'success': False, 'message': '2FA not enabled'})
        
        codes = json.loads(profile.tfa_backup_codes) if profile.tfa_backup_codes else []
        return JsonResponse({'success': True, 'codes': codes})
    
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})


@login_required
@require_http_methods(["POST"])
def revoke_session_ajax(request, uuid):
    """Revoke a specific session"""
    try:
        if str(request.user.uuid) != str(uuid):
            return JsonResponse({'success': False, 'message': 'Invalid request'})
        
        data = json.loads(request.body)
        session_key = data.get('session_key')
        
        if not session_key:
            return JsonResponse({'success': False, 'message': 'Session key required'})
        
        from django.contrib.sessions.models import Session
        try:
            session = Session.objects.get(session_key=session_key)
            session.delete()
            return JsonResponse({'success': True, 'message': 'Session revoked'})
        except Session.DoesNotExist:
            return JsonResponse({'success': False, 'message': 'Session not found'})
    
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

@login_required
@require_http_methods(["POST"])
def revoke_all_sessions_ajax(request, uuid):
    """Revoke all sessions except current"""
    try:
        if str(request.user.uuid) != str(uuid):
            return JsonResponse({'success': False, 'message': 'Invalid request'})
        
        from django.contrib.sessions.models import Session
        from django.utils import timezone
        
        current_session_key = request.session.session_key
        user_sessions = Session.objects.filter(expire_date__gte=timezone.now())
        
        for session in user_sessions:
            try:
                data = session.get_decoded()
                if data.get('_auth_user_id') == str(request.user.pk):
                    if session.session_key != current_session_key:
                        session.delete()
            except:
                continue
        
        return JsonResponse({'success': True, 'message': 'All other sessions revoked'})
    
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})



@login_required
@require_http_methods(["POST"])
def deactivate_account_ajax(request, uuid):
    """Deactivate account"""
    try:
        if str(request.user.uuid) != str(uuid):
            return JsonResponse({'success': False, 'message': 'Invalid request'})
        
        data = json.loads(request.body)
        password = data.get('password')
        
        if not request.user.check_password(password):
            return JsonResponse({'success': False, 'message': 'Invalid password'})
        
        request.user.is_active = False
        request.user.save()
        
        return JsonResponse({'success': True, 'message': 'Account deactivated'})
    
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

@login_required
@require_http_methods(["POST"])
def delete_account_ajax(request, uuid):
    """Permanently delete account"""
    try:
        if str(request.user.uuid) != str(uuid):
            return JsonResponse({'success': False, 'message': 'Invalid request'})
        
        data = json.loads(request.body)
        password = data.get('password')
        
        if not request.user.check_password(password):
            return JsonResponse({'success': False, 'message': 'Invalid password'})
        
        request.user.delete()
        return JsonResponse({'success': True, 'message': 'Account deleted permanently'})
    
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})    


#<------------------------------Blaze code End(06-07-26)------------------------------>



#<------------------------------Blaze code Start(09-07-26)------------------------------>
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse
from django.utils import timezone
from django.contrib.auth import update_session_auth_hash
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from Admin.models import User, UserSession, UserAuditLog
from Staff.models import StaffProfile
import json
import logging
import pytz

logger = logging.getLogger(__name__)

# HELPER FUNCTIONS----> by Blaze

def get_active_sessions(request):
    """Get active sessions for the user"""
    try:
        sessions = UserSession.objects.filter(
            user=request.user,
            logout_time__isnull=True
        ).order_by('-login_time')
        
        session_list = []
        for session in sessions:
            session_list.append({
                'id': session.id,
                'session_key': session.session_id,
                'device': session.device_type or 'Unknown Device',
                'browser': session.browser or 'Unknown Browser',
                'ip': session.ip_address or 'Unknown IP',
                'location': getattr(session, 'location', 'Unknown Location'),
                'login_time': session.login_time,
                'is_current': session.session_id == request.session.session_key,
            })
        return session_list
    except:
        return []


def get_recent_activities(request, limit=10):
    """Get recent activity log for the user"""
    try:
        activities = UserAuditLog.objects.filter(
            user=request.user
        ).order_by('-timestamp')[:limit]
        
        activity_list = []
        for activity in activities:
            activity_list.append({
                'id': activity.id,
                'action': activity.action,
                'description': activity.description or activity.action,
                'timestamp': activity.timestamp,
                'ip_address': activity.ip_address or 'N/A',
                'status': activity.status,
                'module': activity.module or '',
            })
        return activity_list
    except:
        return []


def get_available_timezones():
    """Get list of common timezones"""
    return [
        {'value': 'America/New_York', 'label': 'Eastern Time (UTC-5/UTC-4)'},
        {'value': 'America/Chicago', 'label': 'Central Time (UTC-6/UTC-5)'},
        {'value': 'America/Denver', 'label': 'Mountain Time (UTC-7/UTC-6)'},
        {'value': 'America/Los_Angeles', 'label': 'Pacific Time (UTC-8/UTC-7)'},
        {'value': 'America/Anchorage', 'label': 'Alaska Time (UTC-9/UTC-8)'},
        {'value': 'America/Honolulu', 'label': 'Hawaii Time (UTC-10)'},
        {'value': 'Europe/London', 'label': 'London (UTC+0/UTC+1)'},
        {'value': 'Europe/Paris', 'label': 'Paris (UTC+1/UTC+2)'},
        {'value': 'Asia/Dubai', 'label': 'Dubai (UTC+4)'},
        {'value': 'Asia/Kolkata', 'label': 'India (UTC+5:30)'},
        {'value': 'Asia/Shanghai', 'label': 'China (UTC+8)'},
        {'value': 'Asia/Tokyo', 'label': 'Japan (UTC+9)'},
        {'value': 'Australia/Sydney', 'label': 'Sydney (UTC+10/UTC+11)'},
        {'value': 'Pacific/Auckland', 'label': 'Auckland (UTC+12/UTC+13)'},
    ]


@login_required
def staff_account_security(request, uuid):
    """Staff Account & Security - Full Dynamic with Timezone"""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    
    try:
        staff_profile = request.user.staff_profile
    except StaffProfile.DoesNotExist:
        messages.error(request, "Staff profile not found.")
        return redirect('staff_dashboard', uuid=uuid)
    
    # Get user's timezone preference
    user_timezone = request.user.timezone or 'America/Chicago'
    
    # Get current time in user's timezone
    try:
        user_tz = pytz.timezone(user_timezone)
        current_time = timezone.now().astimezone(user_tz)
    except:
        user_tz = pytz.timezone('America/Chicago')
        current_time = timezone.now().astimezone(user_tz)
    
    # Get sessions and activities
    sessions = get_active_sessions(request)
    activities = get_recent_activities(request, limit=10)
    total_activities = UserAuditLog.objects.filter(user=request.user).count()
    
    context = {
        'user': request.user,
        'staff': staff_profile,
        'today': current_time.date(),
        'user_timezone': user_timezone,
        'available_timezones': get_available_timezones(),
        'sessions': sessions,
        'activities': activities,
        'total_activities': total_activities,
    }
    
    return render(request, "Settings/accounts_privacy.html", context)

#Blaze code start (09-07-26)--------------------->    

@login_required
@require_http_methods(["POST"])
def update_account_settings(request, uuid):
    """Update account settings via AJAX"""
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'message': 'Unauthorized'}, status=403)
    
    try:
        data = json.loads(request.body)
        user = request.user
        staff = user.staff_profile
        
        # Update fields as needed
        if 'timezone' in data:
            user.timezone = data['timezone']
            user.save()
        
        # ... other updates
        
        return JsonResponse({'success': True, 'message': 'Settings updated successfully'})
        
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)



# ************************************************ Arun Code ********************************************************


def generate_staff_employee_id(hire_year: int) -> str:
    prefix = f"STF{hire_year}"

    existing = (
        StaffProfile.objects
        .filter(employee_id__startswith=prefix)
        .values_list("employee_id", flat=True)
    )

    max_seq = 0
    pattern = re.compile(rf"^{prefix}(\d+)$")
    for emp_id in existing:
        m = pattern.match(emp_id)
        if m:
            seq = int(m.group(1))
            if seq > max_seq:
                max_seq = seq

    return f"{prefix}{max_seq + 1:03d}"


def suggested_staff_id_for_context(hire_date_str=None):
    year = date.today().year
    if hire_date_str:
        try:
            year = int(hire_date_str[:4])
        except (ValueError, TypeError):
            pass
    return generate_staff_employee_id(year)


@login_required
@require_http_methods(["GET"])
def next_staff_id_json(request):
    year_param = request.GET.get("year", "").strip()
    try:
        year = int(year_param) if year_param else date.today().year
    except ValueError:
        year = date.today().year
    return JsonResponse({"employee_id": generate_staff_employee_id(year)})

def save_staff_details(request, user):
    post = request.POST
    files = request.FILES

    hire_date = post.get("staff_hire_date") or None

    employee_id = (post.get("staff_employee_id") or "").strip()
    if not employee_id:
        year = date.today().year
        if hire_date:
            try:
                year = int(hire_date[:4])
            except ValueError:
                pass
        employee_id = generate_staff_employee_id(year)

    supervisor_obj = None
    supervisor_id = post.get("staff_supervisor_id")
    if supervisor_id:
        try:
            supervisor_obj = StaffProfile.objects.get(id=supervisor_id)
        except (StaffProfile.DoesNotExist, ValueError):
            supervisor_obj = None

    with transaction.atomic():
        # ── StaffProfile ──
        profile = StaffProfile.objects.create(
            user=user,
            employee_id=employee_id,
            work_email=post.get("staff_work_email") or user.email,
            personal_email=post.get("staff_personal_email", ""),
            office_phone=post.get("staff_office_phone", ""),
            preferred_name=post.get("staff_preferred_name", ""),
            hire_date=hire_date,
            employment_status=post.get("staff_employment_status") or "ACTIVE",
            employment_type=post.get("staff_employment_type") or None,
            supervisor=supervisor_obj,
        )

        if post.get("staff_job_title"):
            StaffPosition.objects.create(
                staff=profile,
                job_title=post.get("staff_job_title", ""),
                job_code=post.get("staff_job_code", ""),
                department_id=post.get("staff_department_id") or None,
                unit_id=post.get("staff_unit_id") or None,
                position_start_date=post.get("staff_position_start_date") or date.today(),
                position_end_date=post.get("staff_position_end_date") or None,
                position_status=post.get("staff_position_status") or "ACTIVE",
            )

        if post.get("staff_address_line_1"):
            StaffAddress.objects.create(
                staff=profile,
                address_type=post.get("staff_address_type") or "PERMANENT",
                address_line_1=post.get("staff_address_line_1", ""),
                address_line_2=post.get("staff_address_line_2", ""),
                city=post.get("staff_city", ""),
                state=post.get("staff_state", ""),
                postal_code=post.get("staff_postal_code", ""),
                country=post.get("staff_country", ""),
            )

        if post.get("staff_emergency_name"):
            StaffEmergencyContact.objects.create(
                staff=profile,
                contact_name=post.get("staff_emergency_name", ""),
                relationship=post.get("staff_emergency_relationship", ""),
                phone_number=post.get("staff_emergency_phone", ""),
                email=post.get("staff_emergency_email", ""),
                priority=post.get("staff_emergency_priority") or 1,
            )

        if post.get("staff_degree"):
            StaffEducation.objects.create(
                staff=profile,
                degree=post.get("staff_degree", ""),
                institution_name=post.get("staff_institution_name", ""),
                field_of_study=post.get("staff_field_of_study", ""),
                graduation_year=post.get("staff_graduation_year") or None,
            )

        cert_names = post.getlist("staff_cert_name[]")
        cert_orgs = post.getlist("staff_cert_org[]")
        cert_issue_dates = post.getlist("staff_cert_issue_date[]")
        cert_expiry_dates = post.getlist("staff_cert_expiry_date[]")

        for i in range(len(cert_names)):
            name = cert_names[i] if i < len(cert_names) else ""
            if not name:
                continue  
            StaffCertification.objects.create(
                staff=profile,
                certification_name=name,
                issuing_organization=cert_orgs[i] if i < len(cert_orgs) else "",
                issue_date=cert_issue_dates[i] if i < len(cert_issue_dates) and cert_issue_dates[i] else date.today(),
                expiration_date=cert_expiry_dates[i] if i < len(cert_expiry_dates) and cert_expiry_dates[i] else None,
            )

        if post.get("staff_system_role"):
            StaffAccessRole.objects.create(
                staff=profile,
                system_role=post.get("staff_system_role"),
                assigned_date=post.get("staff_access_assigned_date") or date.today(),
                active=(post.get("staff_access_active") == "true"),
            )

        doc_types = post.getlist("staff_doc_type[]")
        doc_names = post.getlist("staff_doc_name[]")
        doc_files = files.getlist("staff_doc_file[]")

        for i in range(max(len(doc_types), len(doc_names), len(doc_files))):
            doc_file = doc_files[i] if i < len(doc_files) else None
            if not doc_file:
                continue
            doc_type = doc_types[i] if i < len(doc_types) else ""
            doc_name = doc_names[i] if i < len(doc_names) else ""
            StaffDocument.objects.create(
                staff=profile,
                document_type=doc_type or "OTHER",
                file_name=doc_name or doc_file.name,
                file=doc_file,
                verification_status="PENDING",
            )

    return profile



from Students.models import CourseMaterial, SupportTicket, CourseSection

@login_required
def course_materials_list(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    materials = CourseMaterial.objects.select_related(
        'course_section__course', 'uploaded_by'
    ).order_by('-uploaded_at')

    sections = CourseSection.objects.select_related('course').order_by('course__course_code')

    context = {
        'user': request.user,
        'materials': materials,
        'sections': sections,   
    }
    return render(request, "Academicsupport/materials_list.html", context)

@login_required
def upload_course_material(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    if request.method == 'POST':
        material_id = request.POST.get('material_id')  
        section_id = request.POST.get('course_section')
        title = request.POST.get('title')
        material_type = request.POST.get('material_type', 'OTHER')
        file = request.FILES.get('file')

        if material_id:
            material = get_object_or_404(CourseMaterial, id=material_id)
            if section_id:
                material.course_section = get_object_or_404(CourseSection, section_id=section_id)
            if title:
                material.title = title
            material.material_type = material_type
            if file:
                if material.file:
                    material.file.delete(save=False)
                material.file = file
            material.save()
            messages.success(request, "Material updated successfully.")
        else:
            if section_id and title and file:
                section = get_object_or_404(CourseSection, section_id=section_id)
                CourseMaterial.objects.create(
                    course_section=section,
                    title=title,
                    material_type=material_type,
                    file=file,
                    uploaded_by=request.user,
                )
                messages.success(request, "Material uploaded successfully.")
            else:
                messages.error(request, "Please fill in all required fields.")

        return redirect('course_materials_list', uuid=uuid)

    sections = CourseSection.objects.select_related('course').order_by('course__course_code')
    return render(request, "Academicsupport/upload_material.html", {'user': request.user, 'sections': sections})

@login_required
def delete_course_material(request, uuid, material_id):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    if request.method != 'POST':
        return redirect('course_materials_list', uuid=uuid)

    material = get_object_or_404(CourseMaterial, id=material_id)

    if material.file:
        material.file.delete(save=False)

    material.delete()
    messages.success(request, "Material deleted successfully.")
    return redirect('course_materials_list', uuid=uuid)


@login_required
def edit_course_material(request, uuid, material_id):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    material = get_object_or_404(CourseMaterial, id=material_id)

    if request.method == 'POST':
        title = request.POST.get('title')
        material_type = request.POST.get('material_type', material.material_type)
        new_file = request.FILES.get('file')

        if title:
            material.title = title
        material.material_type = material_type

        if new_file:
            if material.file:
                material.file.delete(save=False)
            material.file = new_file

        material.save()
        messages.success(request, "Material updated successfully.")
        return redirect('course_materials_list', uuid=uuid)

    sections = CourseSection.objects.select_related('course').order_by('course__course_code')
    return render(request, "Academicsupport/edit_material.html", {
        'user': request.user,
        'material': material,
        'sections': sections,
    })

@login_required
def support_tickets_list(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    status_filter = request.GET.get('status', 'open')
    tickets = SupportTicket.objects.select_related(
        'course_section__course', 'submitted_by', 'resolved_by'
    ).order_by('-priority', '-created_at')

    if status_filter == 'open':
        tickets = tickets.exclude(status='RESOLVED')
    elif status_filter != 'all':
        tickets = tickets.filter(status=status_filter.upper())

    context = {'user': request.user, 'tickets': tickets, 'status_filter': status_filter}
    return render(request, "Academicsupport/tickets_list.html", context)


@login_required
def resolve_ticket(request, uuid, ticket_id):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    if request.method == 'POST':
        ticket = get_object_or_404(SupportTicket, id=ticket_id)
        ticket.status = 'RESOLVED'
        ticket.resolved_by = request.user
        ticket.resolved_at = timezone.now()
        ticket.save()

        notify_submitter_ticket_resolved(ticket)

        messages.success(request, "Ticket marked as resolved.")

    return redirect('support_tickets_list', uuid=uuid)

@login_required
def staff_ticket_detail(request, uuid, ticket_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False, "error": "Unauthorized"}, status=403)

    ticket = get_object_or_404(
        SupportTicket.objects.select_related('course_section__course', 'submitted_by', 'resolved_by'),
        id=ticket_id
    )

    return JsonResponse({
        "success": True,
        "data": {
            "id": ticket.id,
            "subject": ticket.subject,
            "description": ticket.description,
            "priority_display": ticket.get_priority_display(),
            "status_display": ticket.get_status_display(),
            "course": ticket.course_section.course.course_code if ticket.course_section else None,
            "submitted_by": ticket.submitted_by.full_name if hasattr(ticket.submitted_by, 'full_name') else str(ticket.submitted_by),
            "submitter_role": "Faculty" if ticket.submitted_by.is_faculty else "Student" if ticket.submitted_by.is_student else "Unknown",
            "created_at": ticket.created_at.strftime("%b %d, %Y"),
            "resolved_by": ticket.resolved_by.get_full_name() if ticket.resolved_by else None,
        }
    })

import csv
from django.http import HttpResponse

@login_required
def export_support_tickets(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    status_filter = request.GET.get('status', 'open')
    tickets = SupportTicket.objects.select_related(
        'course_section__course', 'submitted_by', 'resolved_by'
    ).order_by('-priority', '-created_at')

    if status_filter == 'open':
        tickets = tickets.exclude(status='RESOLVED')
    elif status_filter != 'all':
        tickets = tickets.filter(status=status_filter.upper())

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="support_tickets.csv"'

    writer = csv.writer(response)
    writer.writerow(['Subject', 'Description', 'Course', 'Submitted By', 'Role', 'Priority', 'Status', 'Date', 'Resolved By'])

    for t in tickets:
        role = "Faculty" if t.submitted_by.is_faculty else "Student" if t.submitted_by.is_student else "Unknown"
        writer.writerow([
            t.subject,
            t.description,
            t.course_section.course.course_code if t.course_section else "General",
            getattr(t.submitted_by, 'full_name', str(t.submitted_by)),
            role,
            t.get_priority_display(),
            t.get_status_display(),
            t.created_at.strftime("%b %d, %Y"),
            t.resolved_by.get_full_name() if t.resolved_by else "",
        ])

    return response

@login_required
def mark_ticket_in_progress(request, uuid, ticket_id):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    if request.method == 'POST':
        ticket = get_object_or_404(SupportTicket, id=ticket_id)
        ticket.status = 'IN_PROGRESS'
        ticket.save()

        notify_submitter_ticket_in_progress(ticket)

        messages.success(request, "Ticket marked as in progress.")

    return redirect('support_tickets_list', uuid=uuid)

from Admin.bela_admin.models import Course as CatalogCourse
from Students.models import Schedule, CourseSection, Course as StudentCourse, Semester

@login_required
def new_schedule(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
 
    departments = Department.objects.filter(
    status="ACTIVE"
).order_by("department_name")
    semesters = Semester.objects.filter(is_current=True).order_by('-academic_year')
    if not semesters.exists():
        semesters = Semester.objects.all().order_by('-academic_year')[:1]
 
    context = {
    'user': request.user,
    'departments': departments,
    'semesters': semesters,
    'day_choices': Schedule.DAYS_OF_WEEK,
    'section_type_choices': CourseSection.SECTION_TYPES,
}
    return render(request, "Academicsupport/new_schedule.html", context)


@login_required
def get_sections_for_course(request):
    """Dropdown data — sections for a chosen course (reuses your existing pattern)"""
    try:
        course_id = request.GET.get('course_id')
        sections = CourseSection.objects.filter(course_id=course_id).select_related('semester_id')
        data = [{
            'id': s.section_id,
            'section_number': s.section_number,
            'semester': str(s.semester_id) if s.semester_id else 'N/A',
            'semester_id': s.semester_id.pk if s.semester_id else None,
            'faculty_name': s.faculty_name or 'N/A',
            'room_number': s.room_number,
            'building_name': s.building_name,
            'capacity': s.capacity,
        } for s in sections]
        return JsonResponse({'success': True, 'sections': data})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


##########  Rixie code start  #########

@login_required
def get_buildings(request):
    """Return all active buildings for the schedule form dropdown."""
    try:
        from Admin.bela_admin.models import Building
        buildings = Building.objects.filter(
            status="ACTIVE",
            building_type="NORMAL"
        ).order_by("building_name")
        data = [{
            'id': b.pk,
            'name': b.building_name,
            'code': b.building_code,
        } for b in buildings]
        return JsonResponse({'success': True, 'buildings': data})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def get_rooms_for_building(request):
    """Return available rooms for a given building (via Floor FK)."""
    try:
        from Admin.bela_admin.models import Room, Floor
        building_id = request.GET.get('building_id')
        if not building_id:
            return JsonResponse({'success': True, 'rooms': []})

        floor_ids = Floor.objects.filter(
            building_id=building_id
        ).values_list('id', flat=True)

        rooms = Room.objects.filter(
            floor_id__in=floor_ids,
            status="AVAILABLE"
        ).select_related('floor', 'floor__building').order_by('room_number')

        data = [{
            'id': r.pk,
            'room_number': r.room_number,
            'room_name': r.room_name,
            'room_type': r.get_room_type_display(),
            'capacity': r.capacity,
            'building_name': r.floor.building.building_name,
        } for r in rooms]
        return JsonResponse({'success': True, 'rooms': data})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


def _first_overlapping_schedule(qs, requested_dates):
    """
    Return the first schedule where at least one actual date
    matches the requested date.
    """

    requested = set(requested_dates or [])

    for s in qs:

        existing = set(s.dates or [])

        # Only conflict if both schedules have actual dates
        # and at least one date is exactly the same.
        if requested and existing and requested.intersection(existing):
            return s

    return None
############  Rixie code end  ##########

@login_required
def create_schedule(request):
    """Create a new schedule slot via AJAX — supports multiple days"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)

    try:
        data = json.loads(request.body)

        section_id   = data.get('section_id')
        grouped_dates = data.get('grouped_dates', {})

        days = list(grouped_dates.keys())
        
        year         = data.get('year')
        month        = data.get('month')
        start_time   = data.get('start_time')
        end_time     = data.get('end_time')
        room         = data.get('room', '')
        building     = data.get('building', '')
        is_online    = data.get('is_online', False)

        # Backward compat: accept single day_of_week string
        if not days:
            dow = data.get('day_of_week')
            if dow:
                days = [dow]

        # ---------- validation ----------
        if not section_id:
            return JsonResponse({'success': False, 'error': 'Course section is required'}, status=400)
        if not days:
            return JsonResponse({'success': False, 'error': 'At least one day is required'}, status=400)
        if not start_time or not end_time:
            return JsonResponse({'success': False, 'error': 'Start and end time are required'}, status=400)

        valid_day_values = {d[0] for d in Schedule.DAYS_OF_WEEK}
        for d in days:
            if d not in valid_day_values:
                return JsonResponse({'success': False, 'error': f'Invalid day: {d}'}, status=400)

        if 'SUN' in days:
            return JsonResponse({'success': False, 'error': 'Scheduling is only allowed Monday through Saturday. Sunday is not allowed.'}, status=400)

        try:
            start = datetime.strptime(start_time, '%H:%M').time()
            end = datetime.strptime(end_time, '%H:%M').time()
        except (ValueError, TypeError):
            return JsonResponse({'success': False, 'error': 'Invalid time format'}, status=400)

        min_allowed = datetime.strptime('08:00', '%H:%M').time()
        max_allowed = datetime.strptime('17:00', '%H:%M').time()
        if start < min_allowed or end > max_allowed:
            return JsonResponse({'success': False, 'error': 'Time must be between 8:00 AM and 5:00 PM'}, status=400)

        section = get_object_or_404(CourseSection.objects.select_related('course'), section_id=section_id)

        if start >= end:
            return JsonResponse({'success': False, 'error': 'End time must be after start time'}, status=400)

        # ---------- conflict checks ----------
        day_labels = dict(Schedule.DAYS_OF_WEEK)
        instructor = getattr(section, 'faculty_name', '') or ''
        course_obj = section.course
        course_code = course_obj.course_code if course_obj else ''

        for day in days:
            # Requested concrete dates (ISO strings) for this weekday
            requested_dates = set(grouped_dates.get(day) or [])

            # 1) Same section already has a schedule on the requested dates
            section_overlap_qs = Schedule.objects.filter(
                section_id=section,
                day_of_week=day,
            )
            section_overlap = _first_overlapping_schedule(section_overlap_qs, requested_dates)
            if section_overlap:
                return JsonResponse({
                    'success': False,
                    'error': f'This section already has a schedule assigned for {day_labels.get(day, day)}. No additional timing can be allocated for this section.'
                }, status=200)

            # 2) Instructor conflict — same instructor on a DIFFERENT section at overlapping time
            if instructor:
                instructor_conflict_qs = Schedule.objects.filter(
                    day_of_week=day,
                    start_time__lt=end_time,
                    end_time__gt=start_time,
                ).exclude(
                    section_id=section
                ).select_related('section_id', 'section_id__course').filter(
                    section_id__faculty_name=instructor
                )
                instructor_conflict = _first_overlapping_schedule(instructor_conflict_qs, requested_dates)

                if instructor_conflict:
                    other_code = instructor_conflict.section_id.course.course_code if instructor_conflict.section_id.course else 'another course'
                    return JsonResponse({
                        'success': False,
                        'error': f'Instructor conflict: {instructor} is already assigned to {other_code} on {day_labels.get(day, day)} during {start_time} – {end_time}.'
                    }, status=200)

            # 3) Room conflict — same building + room + overlapping time + at least one shared actual date
            if room and building:
                room_conflict_qs = Schedule.objects.filter(
                    room=room,
                    building=building,
                    start_time__lt=end_time,
                    end_time__gt=start_time,
                ).exclude(
                    section_id=section
                ).select_related('section_id', 'section_id__course')

                room_conflict = _first_overlapping_schedule(room_conflict_qs, requested_dates)

                if room_conflict:
                    other_section = room_conflict.section_id
                    other_course = other_section.course
                    other_code = other_course.course_code if other_course else 'another course'
                    return JsonResponse({
                        'success': False,
                        'error': (
                            f'Room conflict: {building} – {room} is already '
                            f'allocated to {other_code} '
                            f'(Section {other_section.section_number}) '
                            f'during {room_conflict.start_time.strftime("%H:%M")} – '
                            f'{room_conflict.end_time.strftime("%H:%M")}.'
                        )
                    }, status=200)

        # ---------- create one Schedule per day ----------
         # Convert selected day numbers into full ISO dates
         ########  Rixie Code ########

        full_dates = []

        if year and month and dates:
         for d in dates:
          try:
            full_date = date(
                int(year),
                int(month) + 1,   # JS months are 0-11
                int(d)
             ).isoformat()

            full_dates.append(full_date)

          except ValueError:
            continue

         
        created_ids = []

        for day, day_dates in grouped_dates.items():

            schedule = Schedule.objects.create(
            section_id=section,
            day_of_week=day,
            dates=day_dates,
            start_time=start_time,
            end_time=end_time,
            room=room,
            building=building,
            is_online=is_online,
           )
            
        # ==========================================
        # AUDIT LOG: SCHEDULE CREATED
        # ==========================================
        after_data = AuditLogger.model_to_dict(
            schedule,
            fields=[
                "schedule_id",
                "day_of_week",
                "dates",
                "start_time",
                "end_time",
                "room",
                "building",
                "is_online",
            ]
        )
        AuditLogger.log(
            request=request,
            action="CREATE",
            module="Academic Schedule",
            object_type="Schedule",
            object_id=schedule.schedule_id,
            description=(
                f'Schedule created for course "{course_code}" '
                f'(Section {section.section_number}) on '
                f'{day_labels.get(day, day)} '
                f'from {start_time} to {end_time}.'
            ),
            before_data=None,
            after_data=after_data,
            status="SUCCESS",
        )

        created_ids.append(schedule.schedule_id)

        # ========== FIX: Update section status to CONFIRMED ==========
        # Only update if status is PENDING or not already CONFIRMED/CONFLICT
        if section.status in ['PENDING', None]:
            section.status = 'CONFIRMED'
            section.save(update_fields=['status'])
            print(f"Section {section.section_id} status updated to CONFIRMED")

            # Create faculty assignment
         # ---------- Create Faculty Assignment ----------
        if section.faculty_name:
        
                        parts = section.faculty_name.split()
        
                        first_name = parts[0]
                        last_name = " ".join(parts[1:])
        
                        faculty = FacultyProfile.objects.select_related("user").filter(
                            Q(preferred_name=section.faculty_name) |
                            Q(
                                user__first_name=first_name,
                                user__last_name=last_name,
                            )
                        ).first()
        
                        print("Faculty Name :", section.faculty_name)
                        print("Faculty Found:", faculty)
        
                        if faculty:
                            FacultyCourseAssignment.objects.update_or_create(
                                faculty=faculty,
                                course_section_id=section.section_id,
                                defaults={
                                    "semester_id": section.semester_id_id,
                                    "role": "INSTRUCTOR",
                                }
                            )
        
                            print("FacultyCourseAssignment created/updated")
                        else:
                            print(" FacultyProfile not found")
        
        created_days = [day_labels.get(d, d) for d in days]

        return JsonResponse({
            'success': True,
            'message': f'Schedule created for {", ".join(created_days)}!',
            'data': {
                'ids': created_ids,
                'course_code': course_code,
                'days': created_days,
                'time': f"{start_time} - {end_time}",
                'section_status': section.status,
            }
        })

    except json.JSONDecodeError:
        return JsonResponse({'success': False, 'error': 'Invalid JSON data'}, status=400)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)

@login_required
def check_section_instructor_conflict(request):
    """Check if same section already has a schedule or same instructor has a conflict on the requested dates."""
    if request.method != 'GET':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)

    section_id = request.GET.get('section_id')
    days = request.GET.get('days', '')
    start_time = request.GET.get('start_time', '')
    end_time = request.GET.get('end_time', '')
    dates_json = request.GET.get('dates', '{}')
    room = request.GET.get('room', '')
    building = request.GET.get('building', '')

    if not section_id or not days or not start_time or not end_time:
        return JsonResponse({'success': False, 'conflict': False})

    try:
        grouped_dates = json.loads(dates_json) if dates_json else {}
    except (json.JSONDecodeError, TypeError):
        grouped_dates = {}

    try:
        section = CourseSection.objects.select_related('course').get(section_id=section_id)
    except CourseSection.DoesNotExist:
        return JsonResponse({'success': False, 'conflict': False})

    instructor = getattr(section, 'faculty_name', '') or ''
    section_number = section.section_number

    day_list = [d.strip() for d in days.split(',') if d.strip()]

    for day in day_list:
        requested_dates = set(grouped_dates.get(day) or [])

        # 1) Same section already has a schedule on the requested dates
        section_qs = Schedule.objects.filter(
            section_id=section,
            day_of_week=day,
        )
        if _first_overlapping_schedule(section_qs, requested_dates):
            return JsonResponse({
                'success': True,
                'conflict': True,
                'error': f'This section already has a schedule assigned for {day}. No additional timing can be allocated for this section.'
            })

        # 2) Instructor conflict — same instructor on a DIFFERENT section at overlapping time on the same dates
        if instructor:
            conflict_qs = Schedule.objects.filter(
                day_of_week=day,
                start_time__lt=end_time,
                end_time__gt=start_time,
            ).exclude(
                section_id=section
            ).select_related('section_id', 'section_id__course').filter(
                section_id__faculty_name=instructor,
            )

            for sched in conflict_qs:
                existing_dates = sched.dates or []
                if not requested_dates or not existing_dates or requested_dates.intersection(existing_dates):
                    other_course = sched.section_id.course
                    other_code = other_course.course_code if other_course else 'another course'
                    return JsonResponse({
                        'success': True,
                        'conflict': True,
                        'error': f'Instructor "{instructor}" is already assigned to {other_code} (Section {sched.section_id.section_number}) on {day} during {start_time} – {end_time}.'
                    })

        # 3) Room conflict — same building + room + overlapping time + at least one shared actual date
        if room and building:
            room_conflict_qs = Schedule.objects.filter(
                day_of_week=day,
                room=room,
                building=building,
                start_time__lt=end_time,
                end_time__gt=start_time,
            ).exclude(
                section_id=section
            ).select_related('section_id', 'section_id__course')

            room_conflict = _first_overlapping_schedule(room_conflict_qs, requested_dates)
            if room_conflict:
                other_course = room_conflict.section_id.course
                other_code = other_course.course_code if other_course else 'another course'
                return JsonResponse({
                    'success': True,
                    'conflict': True,
                    'error': f'Room conflict: {building} – {room} is already allocated to {other_code} (Section {room_conflict.section_id.section_number}) during {start_time} – {end_time}.'
                })

    return JsonResponse({'success': True, 'conflict': False})


@login_required
def get_schedule_detail(request, schedule_id):
    """Get schedule details for the edit modal"""
    try:
        schedule = get_object_or_404(Schedule, schedule_id=schedule_id)
        section = schedule.section_id

        data = {
            'schedule_id': schedule.schedule_id,
            'section_id': section.section_id,
            'course_code': section.course.course_code if section.course else '',
            'course_name': section.course.course_name if section.course else '',
            'day_of_week': schedule.day_of_week,
            'start_time': schedule.start_time.strftime('%H:%M'),
            'end_time': schedule.end_time.strftime('%H:%M'),
            'room': schedule.room,
            'building': schedule.building,
            'is_online': schedule.is_online,
            'section_status': section.status,
            'day_choices': [{'value': k, 'label': v} for k, v in Schedule.DAYS_OF_WEEK],
        }
        return JsonResponse({'success': True, 'data': data})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def update_schedule(request):
    """Update / reschedule via AJAX — reuses conflict detection automatically via save()"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)

    try:
        data = json.loads(request.body)
        schedule_id = data.get('schedule_id')

        schedule = get_object_or_404(Schedule, schedule_id=schedule_id)
        # ==========================================
        # AUDIT: BEFORE DATA
        # ==========================================
        before_data = AuditLogger.model_to_dict(
            schedule,
            fields=[
                "schedule_id",
                "day_of_week",
                "dates",
                "start_time",
                "end_time",
                "room",
                "building",
                "is_online",
            ]
        )
        if data.get('day_of_week'):
            schedule.day_of_week = data['day_of_week']
        if data.get('start_time'):
            schedule.start_time = data['start_time']
        if data.get('end_time'):
            schedule.end_time = data['end_time']

        start = datetime.strptime(str(schedule.start_time)[:5], '%H:%M').time()
        end = datetime.strptime(str(schedule.end_time)[:5], '%H:%M').time()
        min_allowed = datetime.strptime('08:00', '%H:%M').time()
        max_allowed = datetime.strptime('17:00', '%H:%M').time()
        if start < min_allowed or end > max_allowed:
            return JsonResponse({'success': False, 'error': 'Time must be between 8:00 AM and 5:00 PM'}, status=400)
        if start >= end:
            return JsonResponse({'success': False, 'error': 'End time must be after start time'}, status=400)
        if 'room' in data:
            schedule.room = data['room']
        if 'building' in data:
            schedule.building = data['building']

        schedule.save() 
        
        # ==========================================
        # AUDIT: AFTER DATA
        # ==========================================

        after_data = AuditLogger.model_to_dict(
            schedule,
            fields=[
                "schedule_id",
                "day_of_week",
                "dates",
                "start_time",
                "end_time",
                "room",
                "building",
                "is_online",
            ]
        )

        # ==========================================
        # AUDIT LOG
        # ==========================================

        course_code = (
            schedule.section_id.course.course_code
            if schedule.section_id.course
            else "Unknown Course"
        )
        AuditLogger.log(
            request=request,
            action="UPDATE",
            module="Academic Schedule",
            object_type="Schedule",
            object_id=schedule.schedule_id,
            description=(
                f'Schedule updated for course "{course_code}" '
                f'(Section {schedule.section_id.section_number}).'
            ),
            before_data=before_data,
            after_data=after_data,
            status="SUCCESS",
        )
        return JsonResponse({
            'success': True,
            'message': 'Schedule updated successfully!',
            'data': {'id': schedule.schedule_id, 'section_status': schedule.section_id.status}
        })

    except json.JSONDecodeError:
        return JsonResponse({'success': False, 'error': 'Invalid JSON data'}, status=400)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def resolve_schedule_conflict(request, section_id):
    """Manual override — staff clicks 'Fix' and confirms the section anyway"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)
    try:
        section = get_object_or_404(CourseSection, section_id=section_id)
        section.status = 'CONFIRMED'
        section.save(update_fields=['status'])
        return JsonResponse({'success': True, 'message': 'Marked as confirmed'})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


# *********************************** For exams views.py  now *********************************************************

# *********************************** For exams views.py  now *********************************************************

from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.db.models import Count
from django.core.paginator import Paginator
from django.utils import timezone
from django.http import JsonResponse
from django.db import transaction
import json

from Students.models import (
    Exam,
    ExamRoomAllocation,
    ExamInvigilator,
    CourseSection,
    Semester,
    ExamConflict,
    ExamSeat,
    ExamAttendance,
)
from Admin.bela_admin.models import Room
from Faculty.models import FacultyProfile


def _attach_display_fields(exam_list):
    """
    Attaches convenience 'display_*' attributes to each Exam instance,
    computed from real relations (course_section.course, room_allocations,
    invigilators) instead of made-up fields that don't exist on the model.
    Relies on room_allocations/invigilators already being prefetched.
    """
    for exam in exam_list:
        room_allocs = list(exam.room_allocations.all())
        room_alloc = room_allocs[0] if room_allocs else None

        invigs = list(exam.invigilators.all())
        invig = invigs[0] if invigs else None

        course = None
        if exam.course_section_id and exam.course_section.course_id:
            course = exam.course_section.course

        exam.display_course_code = getattr(course, "course_code", None) or "—"
        exam.display_department = getattr(course, "department", "") or ""

        exam.display_room = room_alloc.room.room_number if room_alloc else None
        if room_alloc:
            exam.display_capacity = room_alloc.allocated_capacity
        elif exam.course_section_id:
            exam.display_capacity = exam.course_section.capacity
        else:
            exam.display_capacity = None

        if invig:
            faculty_name = str(invig.faculty)
            exam.display_invigilator = faculty_name
            initials = "".join(part[0] for part in faculty_name.split() if part)[:2].upper()
            exam.display_invigilator_initials = initials or "?"
        else:
            exam.display_invigilator = None
            exam.display_invigilator_initials = None

        exam.display_date = exam.exam_date
        exam.display_time = "{} \u2013 {}".format(
            exam.start_time.strftime("%I:%M %p"),
            exam.end_time.strftime("%I:%M %p"),
        )
        exam.display_section_label = (
            "Sec {}".format(exam.course_section.section_number)
            if exam.course_section_id else ""
        )
    return exam_list


@login_required
def staff_exams(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    try:
        staff = request.user.staff_profile
    except Exception:
        return redirect("login")

    exams_qs = (
        Exam.objects
        .select_related(
            "course_section",
            "semester",
            "created_by",
        )
        .prefetch_related(
            "room_allocations",
            "invigilators",
        )
        .annotate(
            room_count=Count("room_allocations", distinct=True),
            invigilator_count=Count("invigilators", distinct=True),
        )
        .order_by("-exam_date", "-start_time")
    )

    from django.db.models import Q

    search = request.GET.get("search", "").strip()

    if search:
        exams_qs = exams_qs.filter(
            Q(exam_name__icontains=search) |
            Q(course_section__course__course_code__icontains=search)
        )

    total_exams = exams_qs.count()
    scheduled = exams_qs.filter(status="SCHEDULED").count()
    completed = exams_qs.filter(status="COMPLETED").count()
    draft = exams_qs.filter(status="DRAFT").count()
    cancelled = exams_qs.filter(status="CANCELLED").count()
    draft_exams = _attach_display_fields(list(
        exams_qs.filter(status="DRAFT").order_by("exam_date", "start_time")[:20]
    ))
    rooms_for_modal = Room.objects.filter(status="AVAILABLE").order_by("room_number")
    faculty_for_modal = FacultyProfile.objects.filter(employment_status="ACTIVE").select_related("user")

    room_allocated = (
        ExamRoomAllocation.objects
        .values("exam")
        .distinct()
        .count()
    )
    invigilator_assigned = (
        ExamInvigilator.objects
        .values("exam")
        .distinct()
        .count()
    )

    total_rooms = Room.objects.count()
    rooms_available = Room.objects.filter(status="AVAILABLE").count()

    type_counts = (
        exams_qs.values("exam_type")
        .annotate(count=Count("id"))
        .order_by("-count")
    )
    exams_by_type = []
    for row in type_counts:
        pct = round((row["count"] / total_exams) * 100) if total_exams else 0
        exams_by_type.append({
            "type": row["exam_type"],
            "count": row["count"],
            "pct": pct,
        })

    upcoming_exams = (
        exams_qs
        .filter(status="SCHEDULED", exam_date__gte=timezone.now().date())
        .order_by("exam_date", "start_time")[:4]
    )

    pending_room = (
        exams_qs
        .filter(room_count=0)
        .exclude(status__in=["COMPLETED", "CANCELLED"])
        .order_by("exam_date")[:5]
    )
    pending_invigilator = (
        exams_qs
        .filter(invigilator_count=0)
        .exclude(status__in=["COMPLETED", "CANCELLED"])
        .order_by("exam_date")[:5]
    )

    paginator = Paginator(exams_qs, 5)
    page_number = request.GET.get("page", 1)
    page_obj = paginator.get_page(page_number)

    num_pages = paginator.num_pages
    current = page_obj.number
    if num_pages <= 7:
        page_range = list(range(1, num_pages + 1))
    else:
        page_range = [1]
        if current > 3:
            page_range.append("...")
        start = max(2, current - 1)
        end = min(num_pages - 1, current + 1)
        page_range.extend(range(start, end + 1))
        if current < num_pages - 2:
            page_range.append("...")
        page_range.append(num_pages)

    page_obj.object_list = _attach_display_fields(list(page_obj.object_list))
    upcoming_exams = _attach_display_fields(list(upcoming_exams))
    pending_room = _attach_display_fields(list(pending_room))
    pending_invigilator = _attach_display_fields(list(pending_invigilator))

    context = {
        "user": request.user,
        "staff": staff,
        "exams": page_obj,
        "page_obj": page_obj,
        "paginator": paginator,
        "page_range": page_range,
        "total_exams": total_exams,
        "scheduled": scheduled,
        "completed": completed,
        "draft": draft,
        "cancelled": cancelled,
        "room_allocated": room_allocated,
        "invigilator_assigned": invigilator_assigned,
        "total_rooms": total_rooms,
        "rooms_available": rooms_available,
        "exams_by_type": exams_by_type,
        "upcoming_exams": upcoming_exams,
        "pending_room": pending_room,
        "pending_invigilator": pending_invigilator,
        "draft_exams": draft_exams,
        "rooms_for_modal": rooms_for_modal,
        "faculty_for_modal": faculty_for_modal,
    }
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({
            "table": render_to_string(
                "Academicsupport/exams.html",
                context,
                request=request
            )
        })

    return render(
        request,
        "Academicsupport/exams.html",
        context,
    )


@login_required
def new_exam(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    departments = Department.objects.filter(status="ACTIVE").order_by("department_name")
    rooms = Room.objects.filter(status="AVAILABLE").order_by("room_number")
    faculty = FacultyProfile.objects.filter(employment_status="ACTIVE").select_related("user")

    from Admin.Colleges.models import ProgramCourse
    program_course_terms = {}
    for pc in ProgramCourse.objects.filter(status='ACTIVE').values(
        'course_id', 'program_id', 'term', 'study_year'
    ):
        key = f"{pc['course_id']}:{pc['program_id']}"
        program_course_terms.setdefault(key, []).append({
            'term': pc['term'], 'study_year': pc['study_year'],
        })

    context = {
        "user": request.user,
        "departments": departments,
        "rooms": rooms,
        "faculty": faculty,
        "program_course_terms_json": json.dumps(program_course_terms),
    }
    return render(request, "Academicsupport/new_exam.html", context)

@login_required
def get_course_enrolled_students(request):
    """Given a course + term, returns every student enrolled in ANY
    section of that course for the resolved semester — this is who the
    exam applies to, regardless of which section they're in."""
    course_id = request.GET.get('course_id')
    term = request.GET.get('term')

    if not course_id or not term:
        return JsonResponse({'success': False, 'error': 'course_id and term required'}, status=400)

    semester = _resolve_semester(term)
    if not semester:
        return JsonResponse({'success': False, 'error': 'Could not resolve a semester for this term.'}, status=400)

    enrollments = StudentEnrollment.objects.filter(
        section_id__course_id=course_id,
        semester_id=semester.semester_id,
    ).exclude(enrollment_status__in=['WITHDRAWN', 'DROPPED']).select_related(
        'student__user', 'section_id'
    ).order_by('student__user__first_name')

    students = [{
        'id': e.student.id,
        'name': e.student.user.get_full_name(),
        'student_number': e.student.student_number or 'N/A',
        'section': f"Sec {e.section_id.section_number}" if e.section_id else '—',
    } for e in enrollments]

    return JsonResponse({
        'success': True,
        'count': len(students),
        'students': students,
        'semester_id': semester.semester_id,
        'semester_name': str(semester),
    })

@login_required
def create_exam(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request"}, status=400)

    try:
        data = json.loads(request.body)
        course_id = data.get("course_id")
        term = data.get("term")
        exam_name = data.get("exam_name")
        exam_type = data.get("exam_type")
        exam_date = datetime.strptime(data.get("exam_date"), "%Y-%m-%d").date()
        start_time = datetime.strptime(data.get("start_time"), "%H:%M").time()
        end_time = datetime.strptime(data.get("end_time"), "%H:%M").time()
        duration_minutes = data.get("duration_minutes")
        total_marks = data.get("total_marks")
        pass_marks = data.get("pass_marks")
        instructions = data.get("instructions", "")

        if not exam_name:
            return JsonResponse({"success": False, "error": "Exam name is required."})
        if not course_id:
            return JsonResponse({"success": False, "error": "Please select a course."})
        if not term:
            return JsonResponse({"success": False, "error": "Could not determine the term for this course."})

        semester = _resolve_semester(term)
        if not semester:
            return JsonResponse({"success": False, "error": "Could not resolve a semester for this term."})

        sections = CourseSection.objects.filter(
            course_id=course_id, semester_id=semester.semester_id
        )
        if not sections.exists():
            return JsonResponse({
                "success": False,
                "error": "No course sections exist yet for this course in this term. "
                         "Create a schedule/section first, then add the exam."
            })

        created_ids = []
        with transaction.atomic():
            for section in sections:
                exam = Exam.objects.create(
                    course_section=section,
                    semester=semester,
                    exam_name=exam_name,
                    exam_type=exam_type,
                    exam_date=exam_date,
                    start_time=start_time,
                    end_time=end_time,
                    duration_minutes=duration_minutes,
                    total_marks=total_marks,
                    pass_marks=pass_marks,
                    instructions=instructions,
                    created_by=request.user.staff_profile,
                    status="DRAFT",
                )
                after_data = AuditLogger.model_to_dict(
                    exam,
                    fields=[
                        "exam_name",
                        "exam_type",
                        "exam_date",
                        "start_time",
                        "end_time",
                        "duration_minutes",
                        "total_marks",
                        "pass_marks",
                        "instructions",
                        "status",
                        "course_section",
                        "semester",
                    ]
                )
                AuditLogger.log(
                    request=request,
                    action="CREATE",
                    module="Examination",
                    object_type="Exam",
                    object_id=exam.id,
                    description=(
                        f'Created exam "{exam.exam_name}" '
                        f'for section "{exam.course_section}".'
                    ),
                    before_data=None,
                    after_data=after_data,
                    status="SUCCESS",

                )

                created_ids.append(exam.id)

        return JsonResponse({
            "success": True,
            "message": f"Exam created for {sections.count()} section(s).",
            "exam_ids": created_ids,
        })
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})

@login_required
def exam_detail(request, exam_id):
    try:
        exam = get_object_or_404(
            Exam.objects.select_related(
                "course_section", "semester", "created_by",
            ).prefetch_related("invigilators__faculty__user"),
            id=exam_id
        )
        current_invig = exam.invigilators.select_related('faculty__user').first()

        data = {
            "id": exam.id,
            "exam_name": exam.exam_name,
            "exam_type": exam.exam_type,
            "exam_date": exam.exam_date.strftime("%Y-%m-%d"),
            "start_time": exam.start_time.strftime("%H:%M"),
            "end_time": exam.end_time.strftime("%H:%M"),
            "duration_minutes": exam.duration_minutes,
            "total_marks": exam.total_marks,
            "pass_marks": exam.pass_marks,
            "instructions": exam.instructions,
            "status": exam.status,
            "course_section_id": exam.course_section.section_id,
            "semester_id": exam.semester.semester_id,
            "current_invigilator_id": current_invig.faculty_id if current_invig else None,
            "current_invigilator_name": current_invig.faculty.user.get_full_name() if current_invig else None,
        }
        return JsonResponse({"success": True, "data": data})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=500)


from Staff.utils import (
    notify_students_exam_scheduled,
    notify_students_exam_room_assigned,
    notify_faculty_invigilation_assigned,
    notify_exam_cancelled,
    notify_exam_rescheduled,
    notify_faculty_invigilation_removed,
    notify_student_org_request_approved,
    notify_student_added_to_org,
    notify_faculty_assigned_to_org,
)

@login_required
def update_exam(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request"}, status=400)
    try:
        data = json.loads(request.body)
        exam = get_object_or_404(Exam, id=data.get("exam_id"))

        # ==========================================
        # AUDIT: BEFORE DATA
        # MUST BE BEFORE CHANGING ANY FIELD
        # ==========================================
        before_data = AuditLogger.model_to_dict(
            exam,
            fields=[
                "exam_name",
                "exam_type",
                "exam_date",
                "start_time",
                "end_time",
                "duration_minutes",
                "total_marks",
                "pass_marks",
                "instructions",
                "status",
                "course_section",
                "semester",
            ]
        )

        old_status = exam.status
        old_date = exam.exam_date
        old_start = exam.start_time

        exam.course_section = get_object_or_404(CourseSection, section_id=data.get("course_section_id"))
        exam.semester = get_object_or_404(Semester, semester_id=data.get("semester_id"))
        exam.exam_name = data.get("exam_name")
        exam.exam_type = data.get("exam_type")
        exam.exam_date = datetime.strptime(data.get("exam_date"), "%Y-%m-%d").date()
        exam.start_time = datetime.strptime(data.get("start_time"), "%H:%M").time()
        exam.end_time = datetime.strptime(data.get("end_time"), "%H:%M").time()
        exam.duration_minutes = data.get("duration_minutes")
        exam.total_marks = data.get("total_marks")
        exam.pass_marks = data.get("pass_marks")
        exam.instructions = data.get("instructions")
        exam.status = data.get("status")
        exam.save()
        # ==========================================
        # AUDIT: AFTER DATA
        # MUST BE AFTER SAVE
        # ==========================================
        after_data = AuditLogger.model_to_dict(
            exam,
            fields=[
                "exam_name",
                "exam_type",
                "exam_date",
                "start_time",
                "end_time",
                "duration_minutes",
                "total_marks",
                "pass_marks",
                "instructions",
                "status",
                "course_section",
                "semester",
            ]
        )
        # ==========================================
        # AUDIT LOG
        # ==========================================
        AuditLogger.log(
            request=request,
            action="UPDATE",
            module="Examination",
            object_type="Exam",
            object_id=exam.id,
            description=(
                f'Updated exam "{exam.exam_name}".'
            ),
            before_data=before_data,
            after_data=after_data,
            status="SUCCESS",
        )

        if exam.status == "CANCELLED" and old_status != "CANCELLED":
            notify_exam_cancelled(exam)
        elif exam.status != "CANCELLED" and (
            str(old_date) != str(exam.exam_date) or str(old_start) != str(exam.start_time)
        ):
            notify_exam_rescheduled(exam, old_date, old_start)

        return JsonResponse({
            "success": True,
            "message": "Exam updated successfully."
        })
    except Exception as e:
        return JsonResponse({
            "success": False,
            "error": str(e)
        })

@login_required
@transaction.atomic
def reassign_invigilator(request):
    """Replaces the current invigilator on an exam with a different faculty member,
    notifying both the outgoing and incoming faculty."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid Request"}, status=400)
    try:
        data = json.loads(request.body)
        exam = get_object_or_404(Exam, id=data.get("exam_id"))
        new_faculty = get_object_or_404(FacultyProfile, id=data.get("faculty_id"))

        conflict = ExamInvigilator.objects.filter(
            faculty=new_faculty,
            exam__exam_date=exam.exam_date,
            exam__start_time__lt=exam.end_time,
            exam__end_time__gt=exam.start_time
        ).exclude(exam=exam)
        if conflict.exists():
            return JsonResponse({
                "success": False,
                "error": "This faculty member is already invigilating another exam during this time."
            })

        existing = ExamInvigilator.objects.filter(exam=exam).exclude(faculty=new_faculty)
        removed_faculty = [inv.faculty for inv in existing.select_related('faculty__user')]
        existing.delete()

        if ExamInvigilator.objects.filter(exam=exam, faculty=new_faculty).exists():
            return JsonResponse({"success": False, "error": "Faculty already assigned."})

        ExamInvigilator.objects.create(
            exam=exam, faculty=new_faculty, assigned_by=request.user.staff_profile,
        )

        for old_faculty in removed_faculty:
            notify_faculty_invigilation_removed(exam, old_faculty)
        notify_faculty_invigilation_assigned(exam, new_faculty)

        return JsonResponse({"success": True, "message": "Invigilator reassigned successfully."})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=500)
    
@login_required
@transaction.atomic
def delete_exam(request, exam_id):

    if request.method not in ["POST", "DELETE"]:
        return JsonResponse({
            "success": False,
            "error": "Invalid Request"
        }, status=400)

    try:
        exam = Exam.objects.filter(id=exam_id).first()

        if not exam:
            return JsonResponse({
                "success": False,
                "error": "Exam not found"
            }, status=404)

        if exam.status == "COMPLETED":
            return JsonResponse({
                "success": False,
                "error": "Completed exams cannot be deleted."
            })

        # delete related
        ExamRoomAllocation.objects.filter(exam=exam).delete()
        ExamInvigilator.objects.filter(exam=exam).delete()
        ExamSeat.objects.filter(exam=exam).delete()
        ExamAttendance.objects.filter(exam=exam).delete()
        ExamConflict.objects.filter(exam=exam).delete()

        exam.delete()

        return JsonResponse({
            "success": True,
            "message": "Exam deleted successfully."
        })

    except Exception as e:
        print("DELETE ERROR:", str(e))  
        return JsonResponse({
            "success": False,
            "error": str(e)
        }, status=500)


@login_required
@transaction.atomic
def allocate_exam_room(request):
    if request.method != "POST":
        return JsonResponse({
            "success": False,
            "error": "Invalid Request"
        }, status=400)

    try:
        data = json.loads(request.body)
        exam_id = data.get("exam_id")
        room_id = data.get("room_id")
        capacity = data.get("allocated_capacity")
        exam = get_object_or_404(
            Exam,
            id=exam_id
        )
        room = get_object_or_404(
            Room,
            id=room_id
        )
        if ExamRoomAllocation.objects.filter(
            exam=exam,
            room=room
        ).exists():
            return JsonResponse({
                "success": False,
                "error": "This room is already allocated."
            })
        conflict = ExamRoomAllocation.objects.filter(
            room=room,
            exam__exam_date=exam.exam_date,
            exam__start_time__lt=exam.end_time,
            exam__end_time__gt=exam.start_time
        ).exclude(
            exam=exam
        )
        if conflict.exists():
            return JsonResponse({
                "success": False,
                "error": "Room already booked for another exam."
            })
        allocation = ExamRoomAllocation.objects.create(
            exam=exam,
            room=room,
            allocated_capacity=capacity,
            allocated_by=request.user.staff_profile
        )

        notify_students_exam_room_assigned(exam, room)

        return JsonResponse({
            "success": True,
            "message": "Room allocated successfully.",
            "allocation_id": allocation.id
        })
    except Exception as e:
        return JsonResponse({
            "success": False,
            "error": str(e)
        }, status=500)


@login_required
@transaction.atomic
def assign_invigilator(request):
    if request.method != "POST":
        return JsonResponse({
            "success": False,
            "error": "Invalid Request"
        }, status=400)
    try:
        data = json.loads(request.body)
        exam_id = data.get("exam_id")
        faculty_id = data.get("faculty_id")
        exam = get_object_or_404(
            Exam,
            id=exam_id
        )
        faculty = get_object_or_404(
            FacultyProfile,
            id=faculty_id
        )
        if ExamInvigilator.objects.filter(
            exam=exam,
            faculty=faculty
        ).exists():
            return JsonResponse({
                "success": False,
                "error": "Faculty already assigned."
            })
        conflict = ExamInvigilator.objects.filter(
            faculty=faculty,
            exam__exam_date=exam.exam_date,
            exam__start_time__lt=exam.end_time,
            exam__end_time__gt=exam.start_time
        ).exclude(
            exam=exam
        )
        if conflict.exists():
            return JsonResponse({
                "success": False,
                "error": "Faculty already assigned to another exam during this time."
            })
        invigilator = ExamInvigilator.objects.create(
            exam=exam,
            faculty=faculty,
            assigned_by=request.user.staff_profile
        )

        notify_faculty_invigilation_assigned(exam, faculty)

        return JsonResponse({
            "success": True,
            "message": "Invigilator assigned successfully.",
            "invigilator_id": invigilator.id
        })
    except Exception as e:
        return JsonResponse({
            "success": False,
            "error": str(e)
        }, status=500)
    
# ************************************************ Arun Code ********************************************************





#***********************************Rixie code start **********************************************#
from django.db.models import Count
from django.core.paginator import Paginator

def course_section_dashboard(request,uuid):

    section_list = (
        CourseSection.objects
        .select_related("course", "semester_id")
        .order_by("section_id")
    )

    search_query = request.GET.get("search", "").strip()
    if search_query:
        section_list = section_list.filter(
            Q(course__course_code__icontains=search_query) |
            Q(course__course_name__icontains=search_query) |
            Q(section_number__icontains=search_query) |
            Q(section_type__icontains=search_query) |
            Q(semester_id__semester_type__icontains=search_query) |
            Q(semester_id__semester_code__icontains=search_query) |
            Q(semester_id__academic_year__icontains=search_query) |
            Q(capacity__icontains=search_query) |
            Q(faculty_name__icontains=search_query)
        )

    paginator = Paginator(section_list, 10)   # 10 records per page
    page_number = request.GET.get("page")
    sections = paginator.get_page(page_number)

    form = CourseSectionForm()

    departments = Department.objects.filter(status="ACTIVE").order_by("department_name")

    context = {
        "sections": sections,
        "form": form,
        "departments": departments,
        "total_sections": CourseSection.objects.filter(is_active=True).count(),
        "total_courses": Course.objects.count(),
        "total_semesters": Semester.objects.count(),
        "total_section_types": CourseSection.objects.values('section_type').distinct().count(),
        "search_query": search_query,
    }

    return render(
        request,
        "Academicsupport/course_section.html",
        context,
    )


def check_section_number(request, uuid):
    section_number = request.GET.get("section_number", "").strip()
    course_id = request.GET.get("course", "").strip()
    exclude_id = request.GET.get("exclude_id", "")

    if not section_number or not course_id:
        return JsonResponse({"exists": False})

    qs = CourseSection.objects.filter(
        course_id=course_id,
        section_number=section_number,
        is_active=True,
    )

    if exclude_id:
        qs = qs.exclude(section_id=exclude_id)

    return JsonResponse({
        "exists": qs.exists()
    })


def course_section_add(request, uuid):

    if request.method == "POST":

        form = CourseSectionForm(request.POST)

        if form.is_valid():

            obj = form.save(commit=False)
            course = form.cleaned_data["course"]
            section_number = form.cleaned_data["section_number"]
            faculty_name = request.POST.get("faculty_name", "").strip()

            # Same course + same section
            if CourseSection.objects.filter(
                course=course,
                section_number=section_number,
                is_active=True,
            ).exists():

                return JsonResponse({
                    "success": False,
                    "errors": {
                        "section_number": [
                            "This section has already been assigned for this course."
                        ]
                    }
                })

            obj.faculty_name = faculty_name
            obj.save()
            # ==========================================
            # AUDIT: AFTER DATA
            # ==========================================

            after_data = AuditLogger.model_to_dict(
                obj,
                fields=[
                    "section_id",
                    "course",
                    "section_number",
                    "section_type",
                    "semester_id",
                    "capacity",
                    "faculty_name",
                    "is_active",
                ]
            )


            # ==========================================
            # AUDIT LOG: COURSE SECTION CREATED
            # ==========================================

            AuditLogger.log(
                request=request,
                action="CREATE",
                module="Academic Support",
                object_type="Course Section",
                object_id=obj.section_id,
                description=(
                    f'Created course section "{obj.section_number}" '
                    f'for course "{obj.course.course_code} - '
                    f'{obj.course.course_name}".'
                ),
                before_data=None,
                after_data=after_data,
                status="SUCCESS",
            )

            faculty_id = request.POST.get("faculty_id")

            if faculty_id:
                faculty = FacultyProfile.objects.get(user_id=faculty_id)

                FacultyCourseAssignment.objects.update_or_create(
                    faculty=faculty,
                    course_section_id=obj.section_id,
                    defaults={
                        "semester_id": obj.semester_id_id,
                        "role": "INSTRUCTOR",
                    }
                )

            return JsonResponse({
                "success": True,
                "section_id": obj.section_id
            })

        return JsonResponse({
            "success": False,
            "errors": form.errors
        })

    return JsonResponse({
        "success": False
    })

from django.shortcuts import get_object_or_404
from django.http import JsonResponse


def course_section_view(request,uuid, section_id):

    section = get_object_or_404(
        CourseSection.objects.select_related("course", "semester_id"),
        pk=section_id
    )

    return JsonResponse({

        "course": f"{section.course.course_code} - {section.course.course_name}",
        "section_number": section.section_number,
        "section_type": section.get_section_type_display(),
        "semester": str(section.semester_id),
        "capacity": section.capacity,
        "faculty_name": section.faculty_name or "TBA",

    })



from django.shortcuts import get_object_or_404
from django.http import JsonResponse
def course_section_edit(request,uuid, section_id):

    section = get_object_or_404(
        CourseSection,
        section_id=section_id
    )

    if request.method == "POST":
        # ==========================================
        # AUDIT: BEFORE DATA
        # ==========================================

        before_data = AuditLogger.model_to_dict(
            section,
            fields=[
                "section_id",
                "course",
                "section_number",
                "section_type",
                "semester_id",
                "capacity",
                "faculty_name",
                "is_active",
            ]
        )

        print(request.POST)  

        form = CourseSectionForm(
            request.POST,
            instance=section
        )

        if form.is_valid():

            obj = form.save(commit=False)
            course = form.cleaned_data["course"]
            section_number = form.cleaned_data["section_number"]
            faculty_name = request.POST.get("faculty_name", "").strip()

# Same course + same section
        if CourseSection.objects.filter(
           course=course,
           section_number=section_number,
           is_active=True,
           ).exclude(section_id=section.section_id).exists():

            faculty_id = request.POST.get("faculty_id")

            if faculty_id:
                faculty = FacultyProfile.objects.get(user_id=faculty_id)

                FacultyCourseAssignment.objects.update_or_create(
                    faculty=faculty,
                    course_section_id=obj.section_id,
                    defaults={
                        "semester_id": obj.semester_id_id,
                        "role": "INSTRUCTOR",
                    }
                )

                return JsonResponse({
                    "success": True,
                    "section_id": obj.section_id
                })

            return JsonResponse({
        "success": False,
        "errors": {
            "section_number": [
                "This section has already been assigned for this course."
            ]
        }
    })

        obj = form.save(commit=False)
        obj.faculty_name = faculty_name

        is_active = request.POST.get("is_active")
        if is_active is not None:
           obj.is_active = is_active == "true"
           obj.save()
           obj.faculty_name = faculty_name
           is_active = request.POST.get("is_active")
        if is_active is not None:
                obj.is_active = is_active == "true"
                obj.save()

                # ==========================================
                # AUDIT: AFTER DATA
                # ==========================================

                after_data = AuditLogger.model_to_dict(
                    obj,
                    fields=[
                        "section_id",
                        "course",
                        "section_number",
                        "section_type",
                        "semester_id",
                        "capacity",
                        "faculty_name",
                        "is_active",
                    ]
                )


                # ==========================================
                # AUDIT LOG: COURSE SECTION UPDATED
                # ==========================================

                AuditLogger.log(
                    request=request,
                    action="UPDATE",
                    module="Academic Support",
                    object_type="Course Section",
                    object_id=obj.section_id,
                    description=(
                        f'Updated course section "{obj.section_number}" '
                        f'for course "{obj.course.course_code} - '
                        f'{obj.course.course_name}".'
                    ),
                    before_data=before_data,
                    after_data=after_data,
                    status="SUCCESS",
                )

                return JsonResponse({
                "success": True
            })

        print(form.errors) 

        return JsonResponse({
            "success": False,
            "errors": form.errors
        })

    return JsonResponse({
        "section_id": section.section_id,
        "course": section.course_id,
        "section_number": section.section_number,
        "section_type": section.section_type,
        "semester": section.semester_id.get_semester_type_display(),
        "capacity": section.capacity,
        "faculty_name": section.faculty_name or "",
        "department_id": section.course.department_id if section.course else "",
        "program_id": section.course.academic_program_id if section.course and section.course.academic_program_id else "",
        "is_active": section.is_active,
    })


from django.shortcuts import get_object_or_404
from django.http import JsonResponse

def course_section_deactivate(request,uuid, section_id):

    if request.method == "POST":

        section = get_object_or_404(
            CourseSection,
            section_id=section_id
        )

        section.is_active = False
        section.save()

        return JsonResponse({
            "success": True
        })

    return JsonResponse({
        "success": False
    })

from openpyxl import Workbook
from django.http import HttpResponse


def export_course_sections_excel(request,uuid):

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Course Sections"

    headers = [
        "Section ID",
        "Course Code",
        "Course Name",
        "Section Number",
        "Section Type",
        "Semester",
        "Capacity",
        "Instructor",
        
    ]

    for col, header in enumerate(headers, 1):
        worksheet.cell(row=1, column=col).value = header

    sections = (
        CourseSection.objects
        .select_related("course", "semester_id")
        .order_by("course__course_code")
    )
    record_count = sections.count()

    row = 2

    for section in sections:
        worksheet.cell(row=row, column=1).value = section.section_id
        worksheet.cell(row=row, column=2).value = section.course.course_code if section.course else ""
        worksheet.cell(row=row, column=3).value = section.course.course_name if section.course else ""
        worksheet.cell(row=row, column=4).value = section.section_number
        worksheet.cell(row=row, column=5).value = section.get_section_type_display()
        worksheet.cell(row=row, column=6).value = str(section.semester_id)
        worksheet.cell(row=row, column=7).value = section.capacity
        worksheet.cell(row=row, column=8).value = section.faculty_name or "TBA"
       
        row += 1
        # ==========================================
        # AUDIT LOG
        # ==========================================
        AuditLogger.log(
            request=request,
            action="EXPORT",
            module="Academic Support",
            object_type="Course Sections Report",
            object_id="",
            description=(
                f'Exported {record_count} '
                f'course section record(s) to Excel.'
            ),
            after_data={
                "export_format": "EXCEL",
                "record_count": record_count,
                "filename": "Course_Sections.xlsx",
            },
            status="SUCCESS",
        )
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response[
        "Content-Disposition"
    ] = 'attachment; filename="Course_Sections.xlsx"'

    workbook.save(response)

    return response

from django.http import HttpResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle


def export_course_sections_pdf(request,uuid):

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="Course_Sections.pdf"'

    doc = SimpleDocTemplate(response, pagesize=letter)

    data = [
        [
            "Section ID",
            "Course Code",
            "Course Name",
            "Section No",
            "Type",
            "Semester",
            "Capacity",
            "Instructor",
            
        ]
    ]
    

    sections = (
        CourseSection.objects
        .select_related("course", "semester_id")
        .order_by("course__course_code")
    )
    record_count = sections.count()
    for section in sections:
        data.append([
            section.section_id,
            section.course.course_code if section.course else "",
            section.course.course_name if section.course else "",
            section.section_number,
            section.get_section_type_display(),
            str(section.semester_id),
            section.capacity,
            section.faculty_name or "TBA",
           
        ])

    table = Table(data)

    table.setStyle(TableStyle([

        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#2563eb")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),

        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 10),

        ("GRID", (0,0), (-1,-1), 1, colors.black),

        ("BACKGROUND", (0,1), (-1,-1), colors.beige),

        ("ALIGN", (0,0), (-1,-1), "CENTER"),

        ("BOTTOMPADDING", (0,0), (-1,0), 10),

    ]))

    elements = [table]

    doc.build(elements)

         
    # ==========================================
    # AUDIT LOG
    # ==========================================

    AuditLogger.log(

        request=request,

        action="EXPORT",

        module="Academic Support",

        object_type="Course Sections Report",

        object_id="",

        description=(
                f'Exported {record_count} '
                f'course section record(s) to PDF.'
        ),

        after_data={

            "export_format": "PDF",

            "record_count": record_count,

            "filename": "Course_Sections.pdf",

        },

        status="SUCCESS",

    )

    return response


from django.http import JsonResponse

def get_department_courses(request):
    department_id = request.GET.get("department_id")
    courses = Course.objects.filter(
        department_id=department_id,
        status="ACTIVE",
    ).order_by("course_code")
    data = [
        {
            "id": course.course_id,
            "label": f"{course.course_code} - {course.course_name}",
        }
        for course in courses
    ]
    return JsonResponse(data, safe=False)


def get_department_faculty(request):
    from Faculty.models import FacultyProfile
    department_id = request.GET.get("department_id")
    faculty_members = FacultyProfile.objects.filter(
        department_id=department_id,
        employment_status="ACTIVE",
        user__is_faculty=True,
    ).select_related("user")
    data = [
        {
            "id": faculty.user.id,
            "name": faculty.user.full_name,
        }
        for faculty in faculty_members
    ]
    return JsonResponse(data, safe=False)


def get_department_programs(request, department_id):

    programs = AcademicProgram.objects.filter(
        department_id=department_id,
        status="ACTIVE"
    ).order_by("program_name")

    data = [
        {
            "value": program.program_id,
            "label": program.program_name,
        }
        for program in programs
    ]

    return JsonResponse(data, safe=False)



def check_instructor_assignment(request, uuid):
    faculty_name = request.GET.get("faculty_name", "").strip()
    section_number = request.GET.get("section_number", "").strip()
    course_id = request.GET.get("course", "").strip()
    exclude_id = request.GET.get("exclude_id", "")

    if not faculty_name or not section_number:
        return JsonResponse({"exists": False})

    qs = CourseSection.objects.filter(
        faculty_name=faculty_name,
        section_number=section_number,
        is_active=True,
    )

    if course_id:
        qs = qs.exclude(course_id=course_id)

    if exclude_id:
        qs = qs.exclude(section_id=exclude_id)

    return JsonResponse({
        "exists": qs.exists()
    })


def course_view(request, uuid, course_id):

    course = get_object_or_404(
        Course.objects.select_related("department"),
        course_id=course_id
    )

    sections = CourseSection.objects.filter(
        course=course,
        is_active=True
    )

    context = {
        "course": course,
        "sections": sections,
        "total_sections": sections.count(),
    }

    return render(
        request,
        "Academicsupport/course_view.html",
        context
    )

from django.http import JsonResponse
from Admin.Colleges.models import AcademicProgram
 
 
def get_programs_by_department(request):
 
    department_id = request.GET.get("department_id")
 
    programs = AcademicProgram.objects.filter(
        department_id=department_id,
        status="ACTIVE"
    ).order_by("program_name")
 
    return JsonResponse({
        "success": True,
        "programs": [
            {
                "id": p.program_id,
                "name": p.program_name
            }
            for p in programs
        ]
    })

  
from Admin.bela_admin.models import Course, Department
 
 
def get_courses_by_program(request):
 
    program_id = request.GET.get("program_id")
 
    courses = Course.objects.filter(
        academic_program_id=program_id,
        status="ACTIVE"
    ).order_by("course_code")
 
    from Admin.Colleges.models import ProgramCourse
 
    program_terms = {
        pc["course_id"]: pc["term"]
        for pc in ProgramCourse.objects.filter(
            program_id=program_id,
            status="ACTIVE"
        ).values("course_id", "term")
    }
 
    return JsonResponse({
        "success": True,
        "courses": [
            {
                "id": c.course_id,
                "code": c.course_code,
                "name": c.course_name,
                "term": program_terms.get(c.course_id)
            }
            for c in courses
        ]
    })
 

  
def check_section_number_unique(request, uuid):
    course_id = request.GET.get("course")
    section_number = request.GET.get("section_number")
    exclude_id = request.GET.get("exclude_id")
 
    if not course_id or not section_number:
        return JsonResponse({"exists": False})
 
    qs = CourseSection.objects.filter(course_id=course_id, section_number=section_number)
    if exclude_id:
        qs = qs.exclude(section_id=exclude_id)
 
    return JsonResponse({"exists": qs.exists()})
 
 
def check_instructor_unique(request, uuid):
    faculty_name = request.GET.get("faculty_name")
    course_id = request.GET.get("course")
    section_number = request.GET.get("section_number")
    exclude_id = request.GET.get("exclude_id")
 
    if not faculty_name or not course_id or not section_number:
        return JsonResponse({"exists": False})
 
    qs = CourseSection.objects.filter(
        faculty_name=faculty_name,
        section_number=section_number,
    ).exclude(course_id=course_id)
    if exclude_id:
        qs = qs.exclude(section_id=exclude_id)
 
    return JsonResponse({"exists": qs.exists()})


@login_required
def schedule_edit(request, uuid, course_id):
    """Edit all schedule entries for a specific course"""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
 
    from Admin.bela_admin.models import Course
 
    if not course_id:
        return redirect("staff_scheduling", uuid=uuid)
 
    try:
        course = get_object_or_404(
            Course.objects.select_related("department"),
            course_id=course_id,
        )
    except Exception:
        return redirect("staff_scheduling", uuid=uuid)
 
    sections = CourseSection.objects.filter(
        course=course
    ).select_related("semester_id").prefetch_related("schedules").order_by("section_number")
 
    sections_data = []
    all_schedules = []
    for sec in sections:
        enrolled = StudentEnrollment.objects.filter(
            section_id=sec
        ).exclude(enrollment_status__in=["WITHDRAWN"]).count()
 
        sections_data.append({
            "section_id": sec.section_id,
            "section_number": sec.section_number,
            "faculty_name": sec.faculty_name or "Unassigned",
            "room_number": sec.room_number or "—",
            "building_name": sec.building_name or "—",
            "capacity": sec.capacity,
            "enrolled": enrolled,
            "semester": str(sec.semester_id) if sec.semester_id else "—",
            "status": sec.status,
            "status_display": sec.get_status_display(),
        })
 
        for sched in sec.schedules.all():
            all_schedules.append({
                "schedule_id": sched.schedule_id,
                "section_number": sec.section_number,
                "day_of_week": sched.get_day_of_week_display(),
                "start_time": sched.start_time.strftime("%I:%M %p"),
                "end_time": sched.end_time.strftime("%I:%M %p"),
                "room": sched.room or "—",
                "building": sched.building or "—",
                "is_online": sched.is_online,
                "dates": sched.dates or [],
                "section_status": sec.status,
                "section_status_display": sec.get_status_display(),
            })
 
    all_schedules.sort(key=lambda s: s["section_number"])
 
    context = {
        "user": request.user,
        "course": course,
        "sections": sections_data,
        "schedules": all_schedules,
    }
    return render(
    request,
    "Academicsupport/schedule_edit.html",
    context,
)

 
from Students.models import CourseSection
 
 
def view_timetable(request, uuid):
    from Students.models import Schedule as ScheduleModel
    from django.db.models import Subquery, OuterRef
 
    sections = (
        CourseSection.objects
        .select_related("course", "semester_id")
        .annotate(
            display_room=Subquery(
                ScheduleModel.objects.filter(
                    section_id=OuterRef("section_id")
                ).order_by("schedule_id").values("room")[:1]
            ),
            display_building=Subquery(
                ScheduleModel.objects.filter(
                    section_id=OuterRef("section_id")
                ).order_by("schedule_id").values("building")[:1]
            ),
        )
        .all()
        .order_by("section_id")
    )
 
    search = request.GET.get("search", "").strip()
 
    if search:
        sections = sections.filter(
            Q(course__course_name__icontains=search) |
            Q(faculty_name__icontains=search) |
            Q(display_room__icontains=search) |
            Q(room_number__icontains=search)
        )
 
    context = {
        "schedules": sections,
        "search": search
    }
 
    return render(
        request,
        "Academicsupport/view_timetable.html",
        context
    )
 
 
from django.db.models import Count
 
 
def faculty_availability(request, uuid):
 
    faculty_data = (
        CourseSection.objects
        .values("faculty_name")
        .annotate(
            total_classes=Count("section_id")
        )
        .order_by("faculty_name")
    )
 
 
    context = {
        "faculty_data": faculty_data
    }
 
 
    return render(
        request,
        "Academicsupport/faculty_availability.html",
        context
    )
 
from django.db.models import Count
from Students.models import CourseSection
from django.db.models import Count
from Students.models import Schedule
 
 
def room_management(request, uuid):
 
    room_data = (
        Schedule.objects
        .exclude(room__isnull=True)
        .exclude(room="")
        .values(
            "room",
            "building"
        )
        .annotate(
            total_classes=Count("schedule_id")
        )
        .order_by(
            "building",
            "room"
        )
    )
 
 
    return render(
        request,
        "Academicsupport/room_management.html",
        {
            "room_data": room_data
        }
    )
 

#*********************************** Rixie code end **********************************************#
@login_required
def schedule_exam_page(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    exams = (
        Exam.objects
        .filter(status="DRAFT")
        .select_related("course_section", "course_section__course")
        .order_by("exam_date", "start_time")
    )
    exams = _attach_display_fields(list(exams))
    return render(request, "Academicsupport/schedule_exam.html", {
        "user": request.user, "exams": exams,
    })


@login_required
def assign_room_page(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    exams = (
        Exam.objects
        .annotate(room_count=Count("room_allocations", distinct=True))
        .filter(room_count=0)
        .exclude(status__in=["COMPLETED", "CANCELLED"])
        .select_related("course_section", "course_section__course")
        .order_by("exam_date", "start_time")
    )
    exams = _attach_display_fields(list(exams))
    rooms = Room.objects.filter(status="AVAILABLE").order_by("room_number")
    return render(request, "Academicsupport/assign_room.html", {
        "user": request.user, "exams": exams, "rooms": rooms,
    })


@login_required
def assign_invigilator_page(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    exams = (
        Exam.objects
        .annotate(invigilator_count=Count("invigilators", distinct=True))
        .filter(invigilator_count=0)
        .exclude(status__in=["COMPLETED", "CANCELLED"])
        .select_related("course_section", "course_section__course")
        .order_by("exam_date", "start_time")
    )
    exams = _attach_display_fields(list(exams))
    faculty = FacultyProfile.objects.filter(employment_status="ACTIVE").select_related("user")
    return render(request, "Academicsupport/assign_invigilator.html", {
        "user": request.user, "exams": exams, "faculty": faculty,
    })


@login_required
def schedule_exam(request):
    """Sets date/time on a DRAFT exam and flips it to SCHEDULED — doesn't
    touch course_section/semester, so it's lighter than update_exam."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request"}, status=400)
    try:
        data = json.loads(request.body)
        exam = get_object_or_404(Exam, id=data.get("exam_id"))
        exam.exam_date = data.get("exam_date")
        exam.start_time = data.get("start_time")
        exam.end_time = data.get("end_time")
        exam.duration_minutes = data.get("duration_minutes")
        exam.status = "SCHEDULED"
        exam.save()

        notify_students_exam_scheduled(exam)

        return JsonResponse({"success": True, "message": "Exam scheduled successfully."})
    except Exception as e:
        print("ERROR:", str(e))
        traceback.print_exc()
        return JsonResponse({"success": False, "error": str(e)})


from openpyxl import Workbook
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.shortcuts import redirect
from django.utils import timezone

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer
from reportlab.lib.enums import TA_LEFT

TEMPLATE_STYLES = {
    'classic': {
        'accent': colors.HexColor('#2563eb'),
        'ink': colors.HexColor('#111827'),
        'muted': colors.HexColor('#6b7280'),
        'header_bg': colors.HexColor('#2563eb'),
        'header_text': colors.white,
        'row_alt': colors.HexColor('#f3f6fd'),
        'grid_line': colors.HexColor('#e2e8f0'),
        'badge_shape': 'circle',
        'font': 'Helvetica',
        'font_bold': 'Helvetica-Bold',
    },
    'modern': {
        'accent': colors.HexColor('#0e7490'),
        'ink': colors.HexColor('#0f172a'),
        'muted': colors.HexColor('#64748b'),
        'header_bg': colors.HexColor('#0f172a'),
        'header_text': colors.white,
        'row_alt': colors.HexColor('#f1f5f9'),
        'grid_line': colors.HexColor('#e2e8f0'),
        'badge_shape': 'round-square',
        'font': 'Helvetica',
        'font_bold': 'Helvetica-Bold',
    },
    'corporate': {
        'accent': colors.HexColor('#7c2d12'),
        'ink': colors.HexColor('#1f1300'),
        'muted': colors.HexColor('#7a6a5a'),
        'header_bg': colors.HexColor('#7c2d12'),
        'header_text': colors.white,
        'row_alt': colors.HexColor('#fbf1ea'),
        'grid_line': colors.HexColor('#ecdccb'),
        'badge_shape': 'square',
        'font': 'Helvetica',
        'font_bold': 'Helvetica-Bold',
    },
}


def _draw_badge(canvas, style, cx, cy, r):
    """Small monogram mark. Shape varies per template so the three
    styles stay visually distinct without resorting to a loud color band."""
    canvas.setFillColor(style['accent'])
    shape = style['badge_shape']
    if shape == 'circle':
        canvas.circle(cx, cy, r, fill=1, stroke=0)
    elif shape == 'round-square':
        canvas.roundRect(cx - r, cy - r, 2 * r, 2 * r, r * 0.35, fill=1, stroke=0)
    else:
        canvas.rect(cx - r, cy - r, 2 * r, 2 * r, fill=1, stroke=0)
    canvas.setFillColor(colors.white)
    canvas.setFont(style['font_bold'], r * 1.15)
    canvas.drawCentredString(cx, cy - r * 0.38, "W")


def _make_page_decorator(style, title, subtitle):
    """Draws a slim, professional letterhead (monogram + wordmark + title,
    a right-aligned meta block, and a thin rule) plus a matching footer,
    on every page."""
    def _draw(canvas, doc):
        width, height = letter
        top = height - 0.55 * inch
        badge_r = 0.16 * inch
        badge_cx = 0.6 * inch + badge_r
        badge_cy = top - 0.02 * inch

        canvas.saveState()

        _draw_badge(canvas, style, badge_cx, badge_cy, badge_r)

        text_x = badge_cx + badge_r + 0.14 * inch

        canvas.setFillColor(style['muted'])
        canvas.setFont(style['font_bold'], 8)
        canvas.drawString(text_x, badge_cy + 0.045 * inch, "UNIVERSITY OF WISCONSIN")

        canvas.setFillColor(style['ink'])
        canvas.setFont(style['font_bold'], 17)
        canvas.drawString(text_x, badge_cy - 0.22 * inch, title)

        canvas.setFont(style['font'], 8.5)
        canvas.setFillColor(style['muted'])
        canvas.drawRightString(width - 0.6 * inch, top + 0.02 * inch, subtitle)
        canvas.setFont(style['font'], 7.5)
        generated = timezone.now().strftime('%b %d, %Y, %I:%M %p')
        canvas.drawRightString(width - 0.6 * inch, top - 0.16 * inch, f"Generated {generated}")

        rule_y = top - 0.42 * inch
        canvas.setStrokeColor(style['accent'])
        canvas.setLineWidth(1.1)
        canvas.line(0.6 * inch, rule_y, width - 0.6 * inch, rule_y)
        canvas.setStrokeColor(style['grid_line'])
        canvas.setLineWidth(0.5)
        canvas.line(0.6 * inch, rule_y - 2, width - 0.6 * inch, rule_y - 2)

        canvas.setStrokeColor(colors.HexColor('#e5e7eb'))
        canvas.setLineWidth(0.5)
        canvas.line(0.6 * inch, 0.55 * inch, width - 0.6 * inch, 0.55 * inch)
        canvas.setFont(style['font'], 8)
        canvas.setFillColor(style['muted'])
        canvas.drawString(0.6 * inch, 0.4 * inch, "University of Wisconsin — Staff Portal")
        canvas.drawRightString(width - 0.6 * inch, 0.4 * inch, f"Page {doc.page}")

        canvas.restoreState()

    return _draw


def build_pdf_report(response, title, subtitle, headers, rows, template='modern', col_weights=None):
    """
    Builds a styled, letterheaded PDF into an HttpResponse.
    Used by every report export (student records, financial aid, research)
    so all three share the same look for a given template choice.

    col_weights: optional list of relative widths (one per header), e.g.
    [0.15, 0.30, 0.30, 0.10, 0.15]. They don't need to sum to 1 — they're
    normalized automatically. If omitted, columns split the page evenly.
    Without this the table falls back to reportlab's default "shrink to
    fit content" sizing, which is why earlier PDFs looked narrow and
    left-stranded on an otherwise blank page.
    """
    style = TEMPLATE_STYLES.get(template, TEMPLATE_STYLES['modern'])

    left_margin = right_margin = 0.6 * inch
    doc = SimpleDocTemplate(
        response, pagesize=letter,
        topMargin=1.15 * inch, bottomMargin=0.75 * inch,
        leftMargin=left_margin, rightMargin=right_margin,
    )

    available_width = letter[0] - left_margin - right_margin
    if col_weights and sum(col_weights) > 0:
        total_weight = sum(col_weights)
        col_widths = [available_width * (w / total_weight) for w in col_weights]
    else:
        col_widths = [available_width / len(headers)] * len(headers)

    elements = [Spacer(1, 4)]

    if not rows:
        from reportlab.platypus import Paragraph
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        empty_style = ParagraphStyle(
            'Empty', parent=getSampleStyleSheet()['Normal'],
            fontName=style['font'], fontSize=11,
            textColor=colors.HexColor('#6b7280'),
        )
        elements.append(Paragraph("No records found for this report.", empty_style))
        doc.build(
            elements,
            onFirstPage=_make_page_decorator(style, title, subtitle),
            onLaterPages=_make_page_decorator(style, title, subtitle),
        )
        return response

    data = [headers] + rows
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), style['header_bg']),
        ('TEXTCOLOR', (0, 0), (-1, 0), style['header_text']),
        ('FONTNAME', (0, 0), (-1, 0), style['font_bold']),
        ('FONTSIZE', (0, 0), (-1, 0), 9.5),
        ('FONTNAME', (0, 1), (-1, -1), style['font']),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, style['row_alt']]),
        ('GRID', (0, 0), (-1, -1), 0.6, style['grid_line']),
        ('BOX', (0, 0), (-1, -1), 1, style['header_bg']),
        ('LINEBELOW', (0, 0), (-1, 0), 1.4, style['accent']),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (-1, 0), 'LEFT'),
    ]))
    elements.append(table)

    doc.build(
        elements,
        onFirstPage=_make_page_decorator(style, title, subtitle),
        onLaterPages=_make_page_decorator(style, title, subtitle),
    )
    return response


# ============================================================
# STUDENT RECORDS
# ============================================================

@login_required
def export_student_records_excel(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    students = User.objects.filter(is_student=True).select_related(
        'student_profile', 'student_profile__academic_profile__program'
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Student Records"
    ws.append(["Student ID", "Name", "Email", "Program", "GPA", "Status", "Phone"])

    for s in students:
        profile = getattr(s, 'student_profile', None)
        academic = getattr(profile, 'academic_profile', None) if profile else None
        ws.append([
            getattr(profile, 'student_number', 'N/A'),
            s.get_full_name(),
            s.email,
            get_program_name(academic),
            str(get_gpa(academic)) if get_gpa(academic) else '0.00',
            s.account_status,
            s.mobile_number or 'N/A',
        ])

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response['Content-Disposition'] = 'attachment; filename="Student_Records.xlsx"'
    wb.save(response)
    return response


@login_required
def export_student_records_pdf(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    template = request.GET.get('template', 'modern')
    students = User.objects.filter(is_student=True).select_related(
        'student_profile', 'student_profile__academic_profile__program'
    )

    rows = []
    for s in students:
        profile = getattr(s, 'student_profile', None)
        academic = getattr(profile, 'academic_profile', None) if profile else None
        rows.append([
            getattr(profile, 'student_number', 'N/A'),
            s.get_full_name(),
            get_program_name(academic),
            str(get_gpa(academic)) if get_gpa(academic) else '0.00',
            s.account_status,
        ])

    response = HttpResponse(content_type="application/pdf")
    response['Content-Disposition'] = 'attachment; filename="Student_Records.pdf"'
    return build_pdf_report(
        response,
        title="Student Records Report",
        subtitle=f"{students.count()} students",
        headers=["Student ID", "Name", "Program", "GPA", "Status"],
        rows=rows,
        template=template,
        col_weights=[0.15, 0.28, 0.32, 0.10, 0.15],
    )


# ============================================================
# FINANCIAL AID
# ============================================================

@login_required
def export_financial_aid_excel(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    year = request.GET.get('year')  
    aid_records = StudentFinancialAid.objects.select_related('student__user')
    if year:
        aid_records = aid_records.filter(academic_year=year)

    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Aid"
    ws.append(["Student", "Aid Type", "Academic Year", "Amount", "Status", "Applied Date"])

    for aid in aid_records:
        ws.append([
            aid.student.user.get_full_name(),
            aid.aid_type,
            aid.academic_year,
            float(aid.award_amount or 0),
            aid.status,
            aid.applied_date.strftime('%Y-%m-%d') if aid.applied_date else '',
        ])

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response['Content-Disposition'] = 'attachment; filename="Financial_Aid.xlsx"'
    wb.save(response)
    return response


@login_required
def export_financial_aid_pdf(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    template = request.GET.get('template', 'modern')
    year = request.GET.get('year')
    aid_records = StudentFinancialAid.objects.select_related('student__user')
    if year:
        aid_records = aid_records.filter(academic_year=year)

    rows = []
    total_amount = 0
    for aid in aid_records:
        amount = float(aid.award_amount or 0)
        total_amount += amount
        rows.append([
            aid.student.user.get_full_name(),
            aid.aid_type,
            aid.academic_year,
            f"${amount:,.2f}",
            aid.status,
        ])

    subtitle = f"{aid_records.count()} awards"
    if year:
        subtitle += f"  •  {year}"
    subtitle += f"  •  ${total_amount:,.2f} total"

    response = HttpResponse(content_type="application/pdf")
    response['Content-Disposition'] = 'attachment; filename="Financial_Aid.pdf"'
    return build_pdf_report(
        response,
        title="Financial Aid Report",
        subtitle=subtitle,
        headers=["Student", "Aid Type", "Academic Year", "Amount", "Status"],
        rows=rows,
        template=template,
        col_weights=[0.28, 0.20, 0.18, 0.17, 0.17],
    )

@login_required
def export_research_excel(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    students = User.objects.filter(is_student=True).select_related('student_profile')

    wb = Workbook()
    ws = wb.active
    ws.title = "Research Activity"
    ws.append(["Student", "Project Title", "Research Area", "Mentor", "Start Date"])

    for s in students:
        profile = getattr(s, 'student_profile', None)
        if not profile:
            continue
        for r in get_research(profile):
            ws.append([
                s.get_full_name(),
                r['title'],
                r['area'],
                r['mentor'],
                r['start'],
            ])

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response['Content-Disposition'] = 'attachment; filename="Research_Activity.xlsx"'
    wb.save(response)
    return response


@login_required
def export_research_pdf(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    template = request.GET.get('template', 'modern')
    students = User.objects.filter(is_student=True).select_related('student_profile')

    rows = []
    for s in students:
        profile = getattr(s, 'student_profile', None)
        if not profile:
            continue
        for r in get_research(profile):
            rows.append([s.get_full_name(), r['title'], r['area'], r['mentor']])

    response = HttpResponse(content_type="application/pdf")
    response['Content-Disposition'] = 'attachment; filename="Research_Activity.pdf"'
    return build_pdf_report(
        response,
        title="Research Grant Activity",
        subtitle=f"{len(rows)} active records",
        headers=["Student", "Project Title", "Area", "Mentor"],
        rows=rows,
        template=template,
        col_weights=[0.24, 0.36, 0.20, 0.20],
    )




import os
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, FileResponse, Http404
from django.core.files.storage import default_storage
from django.utils import timezone

from Staff.models import Resource

from django.core.paginator import Paginator

@login_required
def staff_downloads(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    resources_qs = Resource.objects.filter(visible_to_staff=True)

    search_query = request.GET.get('search', '').strip()
    type_filter = request.GET.get('type', 'all')
    category_filter = request.GET.get('category', 'all')

    if search_query:
        from django.db.models import Q
        resources_qs = resources_qs.filter(
            Q(title__icontains=search_query) | Q(description__icontains=search_query)
        )
    if type_filter != 'all':
        resources_qs = resources_qs.filter(resource_type=type_filter)
    if category_filter != 'all':
        resources_qs = resources_qs.filter(category=category_filter)

    paginator = Paginator(resources_qs, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    if request.headers.get('x-requested-with') == 'XMLHttpRequest':
        rows = []
        for r in page_obj:
            rows.append({
                "id": r.id,
                "title": r.title,
                "description": r.description or "No description",
                "resource_type": r.resource_type,
                "type_display": r.get_resource_type_display(),
                "category": r.category,
                "ext": r.file_extension,
                "size": r.file_size_display,
                "uploaded_at": r.uploaded_at.strftime("%b %d, %Y"),
                "downloads": r.download_count,
                "download_url": f"/staff/{uuid}/downloads/{r.id}/download/",
            })
        return JsonResponse({
            "rows": rows,
            "page": page_obj.number,
            "num_pages": page_obj.paginator.num_pages,
            "total_count": page_obj.paginator.count,
            "has_previous": page_obj.has_previous(),
            "has_next": page_obj.has_next(),
            "previous_page": page_obj.previous_page_number() if page_obj.has_previous() else None,
            "next_page": page_obj.next_page_number() if page_obj.has_next() else None,
        })

    all_resources = Resource.objects.filter(visible_to_staff=True)
    total_resources = all_resources.count()
    total_downloads = sum(r.download_count for r in all_resources)
    thirty_days_ago = timezone.now() - timezone.timedelta(days=30)
    new_this_month = all_resources.filter(uploaded_at__gte=thirty_days_ago).count()
    storage_bytes = sum(r.file_size for r in all_resources)
    storage_gb = round(storage_bytes / (1024 ** 3), 2)

    alerts_ctx = get_resource_alerts_context(request.user, "visible_to_staff")
    mark_resource_alerts_seen(request.user)

    context = {
        "user": request.user,
        "today": timezone.localdate(),
        "resources": page_obj,
        "page_obj": page_obj,
        "total_resources": total_resources,
        "total_downloads": total_downloads,
        "new_this_month": new_this_month,
        "storage_gb": storage_gb,
        "search_query": search_query,
        "type_filter": type_filter,
        "category_filter": category_filter,
        **alerts_ctx,
    }
    return render(request, "Resources/downloads.html", context)

@login_required
def upload_resource(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False, "message": "Unauthorized"}, status=403)

    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid method"}, status=400)

    files = request.FILES.getlist("files")
    if not files:
        return JsonResponse({"success": False, "message": "No files provided"}, status=400)

    max_size = 50 * 1024 * 1024
    created = []
    skipped = []

    visible_to_staff = request.POST.get("visible_to_staff", "true") == "true"
    visible_to_faculty = request.POST.get("visible_to_faculty") == "true"
    visible_to_students = request.POST.get("visible_to_students") == "true"

    notify_staff = request.POST.get("notify_staff") == "true"
    notify_faculty = request.POST.get("notify_faculty") == "true"
    notify_students = request.POST.get("notify_students") == "true"
    notify_message = request.POST.get("notify_message", "").strip()

    for f in files:
        if f.size > max_size:
            skipped.append(f.name)
            continue

        resource = Resource.objects.create(
            title=request.POST.get("title") or os.path.splitext(f.name)[0],
            description=request.POST.get("description", ""),
            resource_type=request.POST.get("resource_type", "document"),
            category=request.POST.get("category", "student-services"),
            version=request.POST.get("version", "1.0"),
            file=f,
            uploaded_by=request.user,
            form_category=request.POST.get("form_category") or None,
            department_name=request.POST.get("department_name", ""),
            status=request.POST.get("status", "active"),
            expires_at=request.POST.get("expires_at") or None,
            visible_to_staff=visible_to_staff,
            visible_to_faculty=visible_to_faculty,
            visible_to_students=visible_to_students,
        )

        create_resource_alert(
            resource=resource,
            posted_by=request.user,
            notify_staff=notify_staff,
            notify_faculty=notify_faculty,
            notify_students=notify_students,
            message=notify_message,
        )

        created.append({
            "id": resource.id,
            "title": resource.title,
            "type": resource.resource_type,
            "category": resource.category,
            "ext": resource.file_extension,
            "size": resource.file_size_display,
        })

    return JsonResponse({
        "success": True,
        "created": created,
        "skipped": skipped,
        "message": f"{len(created)} file(s) uploaded successfully!" if created else "No files uploaded"
    })


@login_required
def download_resource(request, uuid, resource_id):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    resource = get_object_or_404(Resource, id=resource_id, visible_to_staff=True)

    if not resource.file or not default_storage.exists(resource.file.name):
        raise Http404("File not found")

    resource.download_count += 1
    resource.save(update_fields=["download_count"])

    return FileResponse(
        resource.file.open("rb"),
        as_attachment=True,
        filename=os.path.basename(resource.file.name),
    )

@login_required
def mark_resource_alerts_seen_ajax(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False}, status=403)
    mark_resource_alerts_seen(request.user)
    return JsonResponse({"success": True})

@login_required
def resource_detail_ajax(request, uuid, resource_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({
            "success": False,
            "message": "Unauthorized"
        }, status=403)
    resource = Resource.objects.filter(
        id=resource_id,
        visible_to_staff=True
    ).first()

    if not resource:
        return JsonResponse({
            "success": False,
            "message": f"Resource {resource_id} not found in DB"
        }, status=404)

    return JsonResponse({
        "success": True,
        "data": {
            "id": resource.id,
            "title": resource.title,
            "description": resource.description or "No description provided.",
            "type": resource.get_resource_type_display(),
            "category": resource.get_category_display(),
            "version": resource.version,
            "ext": resource.file_extension,
            "size": resource.file_size_display,
            "downloads": resource.download_count,
            "uploaded_at": resource.uploaded_at.strftime("%b %d, %Y"),
            "uploaded_by": resource.uploaded_by.get_full_name() if resource.uploaded_by else "Unknown",
        }
    })


@login_required
def delete_resource(request, uuid, resource_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False}, status=403)
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid method"}, status=400)

    resource = get_object_or_404(Resource, id=resource_id)
    if resource.file:
        resource.file.delete(save=False)
    resource.delete()
    return JsonResponse({"success": True, "message": "Resource deleted"})



from collections import OrderedDict
from datetime import timedelta
from django.utils import timezone
from django.utils.timesince import timesince
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator
from django.http import JsonResponse
from Staff.models import Notification


TYPE_ICON_MAP = {
    "INFO": "ti ti-info-circle",
    "SUCCESS": "ti ti-circle-check",
    "WARNING": "ti ti-alert-triangle",
    "ERROR": "ti ti-circle-x",
    "REQUEST": "ti ti-clipboard-check",
    "SYSTEM": "ti ti-settings",
}


def _group_label_for_date(n_date, today, yesterday, week_start):
    if n_date == today:
        return "Today"
    elif n_date == yesterday:
        return "Yesterday"
    elif n_date >= week_start:
        return "This Week"
    elif n_date.year == today.year:
        return n_date.strftime("%B %Y")
    else:
        return n_date.strftime("%B %Y, %Y")


def _serialize_notification(n):
    return {
        "id": n.id,
        "title": n.title,
        "message": n.message or "",
        "type": n.notification_type,
        "icon": TYPE_ICON_MAP.get(n.notification_type, "ti ti-info-circle"),
        "is_read": n.is_read,
        "link": n.link or "",
        "time_display": timezone.localtime(n.created_at).strftime("%I:%M %p").lstrip("0"),
        "related_event_id": n.related_event_id,
    }


@login_required
def staff_notifications(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    notifications = Notification.objects.filter(user=request.user).order_by("-created_at")

    paginator = Paginator(notifications, 5)   # 5 per page
    page_obj = paginator.get_page(1)

    today = timezone.localdate()
    yesterday = today - timedelta(days=1)
    week_start = today - timedelta(days=7)

    groups = OrderedDict()
    for n in page_obj.object_list:
        n_date = timezone.localtime(n.created_at).date()
        label = _group_label_for_date(n_date, today, yesterday, week_start)
        groups.setdefault(label, []).append(n)

    context = {
        "user": request.user,
        "grouped_notifications": groups,
        "unread_count": notifications.filter(is_read=False).count(),
        "total_count": notifications.count(),
        "current_page": page_obj.number,
        "total_pages": paginator.num_pages,
        "has_next": page_obj.has_next(),
        "has_previous": page_obj.has_previous(),
    }
    return render(request, "Notifications/notifications.html", context)


@login_required
def staff_notifications_page_ajax(request, uuid):
    """Returns JSON for a specific page of notifications, grouped by date.
    Used by pagination buttons — the client builds the HTML, no partial
    template needed."""
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False}, status=403)

    page_number = request.GET.get("page", 1)

    notifications             = Notification.objects.filter(user=request.user).order_by("-created_at")
    paginator = Paginator(notifications, 5)
    page_obj = paginator.get_page(page_number)

    today = timezone.localdate()
    yesterday = today - timedelta(days=1)
    week_start = today - timedelta(days=7)

    groups = OrderedDict()
    for n in page_obj.object_list:
        n_date = timezone.localtime(n.created_at).date()
        label = _group_label_for_date(n_date, today, yesterday, week_start)
        groups.setdefault(label, []).append(_serialize_notification(n))

    grouped_list = [{"label": label, "items": items} for label, items in groups.items()]

    return JsonResponse({
        "success": True,
        "groups": grouped_list,
        "current_page": page_obj.number,
        "total_pages": paginator.num_pages,
        "has_next": page_obj.has_next(),
        "has_previous": page_obj.has_previous(),
    })


@login_required
def recent_notifications_ajax(request, uuid):
    """Returns the latest 5 notifications for the dropdown modal."""
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False}, status=403)

    notifications = Notification.objects.filter(
        user=request.user
    ).order_by("-created_at")[:5]

    unread_count = Notification.objects.filter(
        user=request.user, is_read=False
    ).count()

    data = [{
        "id": n.id,
        "title": n.title,
        "message": n.message or "",
        "type": n.notification_type,
        "icon": TYPE_ICON_MAP.get(n.notification_type, "ti ti-info-circle"),
        "is_read": n.is_read,
        "link": n.link or "",
        "time_ago": timesince(n.created_at) + " ago",
        
        "borrow_request_uuid": (
            str(n.related_borrow_request_uuid)
            if n.related_borrow_request_uuid
            else None
        ),

         "related_event_id": n.related_event_id,

    } for n in notifications]

    return JsonResponse({
        "success": True,
        "notifications": data,
        "unread_count": unread_count,
    })


@login_required
def mark_notification_read(request, uuid, notification_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False}, status=403)
    if request.method != "POST":
        return JsonResponse({"success": False}, status=400)

    notif = get_object_or_404(Notification, id=notification_id, user=request.user)
    notif.is_read = True
    notif.save(update_fields=["is_read"])

    unread_count = Notification.objects.filter(user=request.user, is_read=False).count()

    return JsonResponse({
        "success": True,
        "unread_count": unread_count,
        "link": notif.link or "",
    })


@login_required
def mark_all_notifications_read(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False}, status=403)
    if request.method != "POST":
        return JsonResponse({"success": False}, status=400)

    Notification.objects.filter(user=request.user, is_read=False).update(is_read=True)
    return JsonResponse({"success": True, "unread_count": 0})


# ************************************************* Communication Code ************************************************

import json
import logging
from datetime import timedelta
from collections import OrderedDict

from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse
from django.db import transaction
from django.db.models import Q, Max
from django.utils import timezone
from django.views.decorators.http import require_http_methods

from Admin.models import User
from Staff.models import (
    StaffProfile,
    MessageThread,
    Message,
    MessageRecipient,
    Announcement,
)

logger = logging.getLogger(__name__)


# AVATAR_PALETTE = [
#     "#2563eb", "#16a34a", "#7c3aed", "#d97706", "#be185d", "#0891b2",
# ]


def _avatar_color(seed_text):
    idx = sum(ord(c) for c in (seed_text or "?")) % len(AVATAR_PALETTE)
    return AVATAR_PALETTE[idx]


def _initials(full_name):
    parts = (full_name or "").split()
    letters = "".join(p[0] for p in parts[:2]).upper()
    return letters or "?"

ALLOWED_ATTACHMENT_TYPES = {
    "image/jpeg", "image/png", "image/gif", "image/webp",
    "video/mp4", "video/quicktime", "video/webm",
    "application/pdf",
}
MAX_ATTACHMENT_SIZE = 25 * 1024 * 1024 
MAX_ATTACHMENTS_PER_MESSAGE = 5


def _save_attachments(message, files):
    """files = request.FILES.getlist('attachments')"""
    saved = []
    for f in files[:MAX_ATTACHMENTS_PER_MESSAGE]:
        if f.size > MAX_ATTACHMENT_SIZE:
            continue
        if f.content_type not in ALLOWED_ATTACHMENT_TYPES:
            continue
        att = MessageAttachment.objects.create(
            message=message,
            file=f,
            original_name=f.name,
            content_type=f.content_type,
            size=f.size,
        )
        saved.append(att)
    return saved


def _serialize_attachments(message):
    return [
        {
            "id": a.id,
            "url": a.file.url,
            "name": a.original_name,
            "content_type": a.content_type,
            "size": a.size,
            "is_image": a.is_image(),
            "is_video": a.is_video(),
        }
        for a in message.attachments.all()
    ]

def _other_party_for_thread(thread, current_user, last_msg):
    """
    Returns the User this thread is 'with', from current_user's point
    of view — the sender if current_user received it, or the
    recipient if current_user sent it.
    """
    if last_msg.sender_id == current_user.id:
        recip = last_msg.recipients.exclude(recipient=current_user).first()
        return recip.recipient if recip else None
    return last_msg.sender


def _build_thread_row(thread, current_user, last_msg, is_archived=False):
    other = _other_party_for_thread(thread, current_user, last_msg)
    other_name = other.get_full_name() if other else "Unknown"

    is_sent_by_me = last_msg.sender_id == current_user.id

    my_recipient_row = None
    if not is_sent_by_me:
        my_recipient_row = last_msg.recipients.filter(recipient=current_user).first()

    unread = bool(my_recipient_row and not my_recipient_row.is_read)

    if is_archived:
        kind = "archived"
    elif is_sent_by_me:
        kind = "sent"
    elif unread:
        kind = "unread"
    else:
        kind = "all"

    avatar_photo = None

    if other and other.profile_photo:
        avatar_photo = other.profile_photo.url

    return {
        "thread_id": thread.id,
        "subject": thread.subject,
        "preview": (last_msg.body or "")[:140],
        "sender_display": f"To: {other_name}" if is_sent_by_me else other_name,

        "avatar_photo": avatar_photo,
        "avatar_bg": _avatar_color(other_name),
        "avatar_letter": _initials(other_name),

        "time": timezone.localtime(last_msg.sent_at).strftime("%b %d, %Y")
                if timezone.localdate(last_msg.sent_at) != timezone.localdate()
                else timezone.localtime(last_msg.sent_at).strftime("%I:%M %p").lstrip("0"),

        "unread": unread,
        "is_sent": is_sent_by_me,
        "is_archived": is_archived,
        "kind": kind,
    }


@login_required
def staff_communications(request, uuid):
    """
    Main Communications dashboard — real KPI counts, real inbox,
    real announcements. Replaces the old hardcoded template.
    """
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    user = request.user

    threads_qs = (
        MessageThread.objects
        .filter(is_announcement=False)
        .filter(Q(messages__sender=user) | Q(messages__recipients__recipient=user))
        .distinct()
        .annotate(last_sent=Max("messages__sent_at"))
        .order_by("-last_sent")
        .prefetch_related("messages__recipients", "messages__sender")
    )

    archived_thread_ids = set(
        MessageRecipient.objects.filter(
            recipient=user, archived=True
        ).values_list("message__thread_id", flat=True)
    )

    thread_rows = []
    unread_count = 0
    week_ago = timezone.now() - timedelta(days=7)

    for thread in threads_qs:
        last_msg = thread.last_message()
        if not last_msg:
            continue
        is_archived = thread.id in archived_thread_ids
        row = _build_thread_row(thread, user, last_msg, is_archived=is_archived)
        thread_rows.append(row)
        if row["unread"] and not is_archived:
            unread_count += 1

    sent_this_week = Message.objects.filter(
        sender=user, sent_at__gte=week_ago
    ).count()

    active_count = sum(1 for r in thread_rows if not r["is_archived"])

    today = timezone.localdate()
    announcements = (
        Announcement.objects
        .filter(Q(expires_at__isnull=True) | Q(expires_at__gte=today))
        .order_by("-posted_at")[:5]
    )
    active_announcements_count = announcements.count()

    context = {
        "user": user,
        "uuid": uuid,
        "thread_rows": thread_rows,
        "announcements": announcements,
        "kpi_inbox": active_count,
        "kpi_unread": unread_count,
        "kpi_sent_week": sent_this_week,
        "kpi_announcements": active_announcements_count,
    }

    return render(request, "Administrator/communications.html", context,)


@login_required
def search_message_recipients(request):
    """
    Autocomplete endpoint for the compose modal's 'To' field.
    Searches real Users so a message can only be sent to someone
    who actually exists in the system.
    """
    q = request.GET.get("q", "").strip()

    if not q:
        users = User.objects.exclude(id=request.user.id).order_by(
            "first_name", "last_name"
        )[:15]
    else:
        users = User.objects.filter(
            Q(first_name__icontains=q) | Q(last_name__icontains=q) | Q(email__icontains=q)
        ).exclude(id=request.user.id)[:15]

    results = []
    for u in users:
        role = "Faculty" if getattr(u, "is_faculty", False) else (
            "Staff" if hasattr(u, "staff_profile") else "Student"
        )
        results.append({
            "id": u.id,
            "name": u.get_full_name() or u.email,
            "email": u.email,
            "role": role,
        })

    return JsonResponse({"results": results})


@login_required
@require_http_methods(["POST"])
def send_message_ajax(request):
    try:
        recipient_id = request.POST.get("recipient_id")
        subject = (request.POST.get("subject") or "").strip()
        body = (request.POST.get("body") or "").strip()
        files = request.FILES.getlist("attachments")

        if not recipient_id:
            return JsonResponse({"success": False, "error": "Please choose a recipient."}, status=400)
        if not body and not files:
            return JsonResponse({"success": False, "error": "Message body or an attachment is required."}, status=400)

        recipient = get_object_or_404(User, id=recipient_id)

        with transaction.atomic():
            thread = MessageThread.objects.create(
                subject=subject or "(No subject)",
                created_by=request.user,
            )
            message = Message.objects.create(
                thread=thread,
                sender=request.user,
                body=body,
            )
            _save_attachments(message, files)

            MessageRecipient.objects.create(message=message, recipient=recipient)
            MessageRecipient.objects.create(message=message, recipient=request.user, is_read=True)

        return JsonResponse({"success": True, "message": "Message sent!", "thread_id": thread.id})

    except Exception as e:
        logger.error(f"Error sending message: {str(e)}")
        return JsonResponse({"success": False, "error": str(e)}, status=500)


@login_required
def thread_detail_ajax(request, uuid, thread_id):
    """
    Returns every message in a thread (oldest -> newest), and marks
    the current user's copy of the latest message as read.
    """
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False, "error": "Unauthorized"}, status=403)

    thread = get_object_or_404(
        MessageThread.objects.filter(
            Q(messages__sender=request.user) | Q(messages__recipients__recipient=request.user)
        ).distinct(),
        id=thread_id,
    )

    messages_qs = thread.messages.select_related("sender").order_by("sent_at")

    MessageRecipient.objects.filter(
        message__thread=thread, recipient=request.user, is_read=False
    ).update(is_read=True, read_at=timezone.now())

    data = []
    for m in messages_qs:
        sender_name = m.sender.get_full_name() or m.sender.email
        data.append({
            "id": m.id,
            "sender_name": sender_name,
            "sender_role": (
                "Faculty" if getattr(m.sender, "is_faculty", False)
                else "Staff" if hasattr(m.sender, "staff_profile") else "Student"
            ),
            "avatar_bg": _avatar_color(sender_name),
            "avatar_letter": _initials(sender_name),
            "is_mine": m.sender_id == request.user.id,
            "body": m.body,
            "sent_at": timezone.localtime(m.sent_at).strftime("%b %d, %Y at %I:%M %p").replace(" 0", " "),
            "attachments": _serialize_attachments(m),
        })

    last_msg = messages_qs.last()
    other = _other_party_for_thread(thread, request.user, last_msg)
    other_name = other.get_full_name() if other else ""
    other_role = (
        "Faculty" if getattr(other, "is_faculty", False)
        else "Staff" if hasattr(other, "staff_profile")
        else "Student"
    ) if other else ""

    other_avatar_letter = _initials(other_name) if other else ""
    other_avatar_bg = _avatar_color(other_name) if other else ""

    other_avatar_photo = None

    if other and other.profile_photo:
        other_avatar_photo = other.profile_photo.url

    return JsonResponse({
        "success": True,
        "thread_id": thread.id,
        "subject": thread.subject,
        "messages": data,

        "reply_to_id": other.id if other else None,
        "reply_to_name": other_name,

        "other_name": other_name,
        "other_role": other_role,

        "other_avatar_photo": other_avatar_photo,
        "other_avatar_letter": other_avatar_letter,
        "other_avatar_bg": other_avatar_bg,
    })


@login_required
@require_http_methods(["POST"])
def reply_message_ajax(request):
    try:
        thread_id = request.POST.get("thread_id")
        recipient_id = request.POST.get("recipient_id")
        body = (request.POST.get("body") or "").strip()
        files = request.FILES.getlist("attachments")

        if not thread_id or not recipient_id or (not body and not files):
            return JsonResponse({"success": False, "error": "Missing required fields."}, status=400)

        thread = get_object_or_404(MessageThread, id=thread_id)
        recipient = get_object_or_404(User, id=recipient_id)

        with transaction.atomic():
            message = Message.objects.create(thread=thread, sender=request.user, body=body)
            _save_attachments(message, files)

            MessageRecipient.objects.create(message=message, recipient=recipient)
            MessageRecipient.objects.create(message=message, recipient=request.user, is_read=True)

        return JsonResponse({"success": True, "message": "Reply sent!"})

    except Exception as e:
        logger.error(f"Error replying to message: {str(e)}")
        return JsonResponse({"success": False, "error": str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def archive_thread_ajax(request, uuid, thread_id):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False}, status=403)

    updated = MessageRecipient.objects.filter(
        message__thread_id=thread_id,
        recipient=request.user
    ).update(archived=True)

    print("UPDATED ROWS =", updated)

    return JsonResponse({
        "success": True,
        "updated": updated,
    })


@login_required
@require_http_methods(["POST"])
def unarchive_thread_ajax(request, uuid, thread_id):
    """
    Restores an archived thread back to the current user's active inbox.
    """
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False, "error": "Unauthorized"}, status=403)

    MessageRecipient.objects.filter(
        message__thread_id=thread_id, recipient=request.user
    ).update(archived=False)

    return JsonResponse({"success": True, "message": "Thread moved back to inbox."})


@login_required
def message_volume_ajax(request, uuid):
    """
    Powers the 'Message Volume (This Week)' widget with real counts
    instead of hardcoded 12/16/9/20/7.
    """
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False}, status=403)

    today = timezone.localdate()
    start_of_week = today - timedelta(days=today.weekday())  

    day_labels = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    counts = []
    for i, label in enumerate(day_labels):
        day = start_of_week + timedelta(days=i)
        count = Message.objects.filter(
            sent_at__date=day
        ).filter(
            Q(sender=request.user) | Q(recipients__recipient=request.user)
        ).distinct().count()
        counts.append({"label": label, "count": count})

    total = sum(c["count"] for c in counts)
    return JsonResponse({"success": True, "days": counts, "total": total})

@login_required
def unread_messages_ajax(request, uuid):
    """
    Lightweight polling endpoint — returns unread message count
    and the latest unread message so we can push a notification.
    """
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({"success": False}, status=403)

    unread_qs = MessageRecipient.objects.filter(
        recipient=request.user,
        is_read=False,
        archived=False,
    ).select_related("message", "message__sender", "message__thread").order_by("-message__sent_at")

    unread_count = unread_qs.count()

    latest = None
    latest_row = unread_qs.first()
    if latest_row:
        msg = latest_row.message
        sender_name = msg.sender.get_full_name() or msg.sender.email
        latest = {
            "thread_id": msg.thread_id,
            "sender_name": sender_name,
            "preview": (msg.body or "")[:120],
            "sent_at": msg.sent_at.isoformat(),
        }

    return JsonResponse({
        "success": True,
        "unread_count": unread_count,
        "latest": latest,
    })

# ************************************************* Communication Code ************************************************

# ============================================================
# EXAMS & ROOMS — EXPORT
# ============================================================

@login_required
def export_exams_excel(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    exams_qs = (
        Exam.objects
        .select_related("course_section", "course_section__course", "semester")
        .prefetch_related("room_allocations__room", "invigilators__faculty")
        .order_by("-exam_date", "-start_time")
    )
    exams = _attach_display_fields(list(exams_qs))

    wb = Workbook()
    ws = wb.active
    ws.title = "Exams"
    ws.append([
        "Exam Name", "Type", "Course", "Section", "Date",
        "Start Time", "End Time", "Room", "Capacity",
        "Invigilator", "Total Marks", "Pass Marks", "Status",
    ])

    for exam in exams:
        ws.append([
            exam.exam_name,
            exam.get_exam_type_display(),
            exam.display_course_code,
            exam.display_section_label,
            exam.exam_date.strftime('%Y-%m-%d') if exam.exam_date else '',
            exam.start_time.strftime('%I:%M %p') if exam.start_time else '',
            exam.end_time.strftime('%I:%M %p') if exam.end_time else '',
            exam.display_room or 'TBD',
            exam.display_capacity if exam.display_capacity is not None else '',
            exam.display_invigilator or 'Unassigned',
            exam.total_marks,
            exam.pass_marks,
            exam.get_status_display(),
        ])

    for i, width in enumerate([22, 12, 14, 10, 12, 12, 12, 14, 10, 20, 12, 12, 12], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename="Exams_Report.xlsx"'
    wb.save(response)
    return response


@login_required
def export_exams_pdf(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    template = request.GET.get('template', 'modern')

    exams_qs = (
        Exam.objects
        .select_related("course_section", "course_section__course", "semester")
        .prefetch_related("room_allocations__room", "invigilators__faculty")
        .order_by("-exam_date", "-start_time")
    )
    exams = _attach_display_fields(list(exams_qs))

    rows = []
    for exam in exams:
        rows.append([
            exam.exam_name,
            exam.display_course_code,
            exam.display_date.strftime('%Y-%m-%d') if exam.display_date else '—',
            exam.display_room or 'TBD',
            exam.display_invigilator or 'Unassigned',
            exam.get_status_display(),
        ])

    response = HttpResponse(content_type="application/pdf")
    response['Content-Disposition'] = 'attachment; filename="Exams_Report.pdf"'
    return build_pdf_report(
        response,
        title="Exams & Rooms Report",
        subtitle=f"{len(rows)} exams",
        headers=["Exam Name", "Course", "Date", "Room", "Invigilator", "Status"],
        rows=rows,
        template=template,
        col_weights=[0.26, 0.16, 0.14, 0.14, 0.18, 0.12],
    )



from django.http import JsonResponse
from django.db import IntegrityError
from Students.models import StudentOrganization, StudentOrganizationMembership
from Admin.Colleges.models import School
from Faculty.models import FacultyProfile


from Admin.bela_admin.models import Department

@login_required
def staff_org_dashboard(request, uuid):
    orgs = StudentOrganization.objects.select_related("school", "department", "advisor__user").order_by("-id")
    if request.method == "POST":
        name = request.POST.get("organization_name", "").strip()

        if not name:
            messages.error(request, "Organization name is required.")
            return redirect("staff_org_dashboard", uuid=uuid)

        if StudentOrganization.objects.filter(organization_name__iexact=name).exists():
            messages.error(request, f'An organization named "{name}" already exists.')
            return redirect("staff_org_dashboard", uuid=uuid)

        # ── validate max_members ──
        max_members_raw = request.POST.get("max_members", "").strip()
        max_members = None
        if max_members_raw:
            try:
                max_members = int(max_members_raw)
                if max_members < 1:
                    messages.error(request, "Max members must be a positive number.")
                    return redirect("staff_org_dashboard", uuid=uuid)
            except ValueError:
                messages.error(request, "Max members must be a valid number.")
                return redirect("staff_org_dashboard", uuid=uuid)

        try:
            org = StudentOrganization.objects.create(
                organization_name=name,
                motto=request.POST.get("motto", "").strip(),
                description=request.POST.get("description", "").strip(),
                school_id=request.POST.get("school") or None,
                department_id=request.POST.get("department") or None,
                advisor_id=request.POST.get("advisor") or None,
                max_members=max_members,
                founded_date=request.POST.get("founded_date") or None,
            )
            if request.FILES.get("cover_image"):
                org.cover_image = request.FILES["cover_image"]
                org.save()
            if org.advisor:                              
                notify_faculty_assigned_to_org(org, org.advisor)
            messages.success(request, f'"{name}" created successfully.')
        except IntegrityError:
            messages.error(request, f'An organization named "{name}" already exists.')

        return redirect("staff_org_dashboard", uuid=uuid)

    pending_requests = StudentOrganizationMembership.objects.filter(
        status="PENDING"
    ).select_related("student__user", "organization").order_by("-membership_id")

    context = {
        "orgs": orgs,
        "schools": School.objects.all(),
        "departments": Department.objects.all(),
        "advisors": FacultyProfile.objects.select_related("user", "department").all(),
        "total_orgs": orgs.count(),
        "pending_requests": pending_requests,
        "total_pending": pending_requests.count(),
        "total_active_members": StudentOrganizationMembership.objects.filter(status="ACTIVE").count(),
    }
    return render(request, "Organizations/org_dashboard.html", context)

@login_required
def staff_departments_by_school_ajax(request, uuid):
    """Filter department dropdown by chosen school"""
    school_id = request.GET.get("school_id")
    qs = Department.objects.all()
    if school_id:
        qs = qs.filter(school_id=school_id)

    return JsonResponse({
        "success": True,
        "departments": [
            {"id": d.department_id, "name": d.department_name}
            for d in qs
        ],
    })

@login_required
def staff_advisors_by_department_ajax(request, uuid):
    """Filter advisor dropdown by chosen department"""
    dept_id = request.GET.get("department_id")
    qs = FacultyProfile.objects.select_related("user")
    if dept_id:
        qs = qs.filter(department_id=dept_id)

    return JsonResponse({
        "success": True,
        "advisors": [
            {"id": f.id, "name": f.user.get_full_name()}
            for f in qs
        ],
    })


@login_required
def staff_org_update_ajax(request, uuid, org_id):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid method"}, status=405)

    org = get_object_or_404(StudentOrganization, id=org_id)
    old_advisor_id = org.advisor_id                   

    name = request.POST.get("organization_name", "").strip()

    if not name:
        return JsonResponse({"success": False, "error": "Organization name is required."})

    duplicate = StudentOrganization.objects.filter(
        organization_name__iexact=name
    ).exclude(id=org.id).exists()
    if duplicate:
        return JsonResponse({"success": False, "error": f'An organization named "{name}" already exists.'})

    max_members_raw = request.POST.get("max_members", "").strip()
    max_members = None
    if max_members_raw:
        try:
            max_members = int(max_members_raw)
            if max_members < 1:
                return JsonResponse({"success": False, "error": "Max members must be a positive number."})
        except ValueError:
            return JsonResponse({"success": False, "error": "Max members must be a valid number."})

    org.organization_name = name
    org.motto = request.POST.get("motto", "").strip()
    org.description = request.POST.get("description", "").strip()
    org.school_id = request.POST.get("school") or None
    org.department_id = request.POST.get("department") or None
    org.advisor_id = request.POST.get("advisor") or None
    org.max_members = max_members
    org.founded_date = request.POST.get("founded_date") or None

    if request.FILES.get("cover_image"):
        org.cover_image = request.FILES["cover_image"]

    org.save()

    if org.advisor_id and org.advisor_id != old_advisor_id: 
        notify_faculty_assigned_to_org(org, org.advisor)        

    return JsonResponse({"success": True, "message": "Organization updated."})


@login_required
def staff_org_detail_ajax(request, uuid, org_id):
    """Fetch a single org's data to prefill the edit modal / show view details"""
    org = get_object_or_404(
        StudentOrganization.objects.select_related("school", "department", "advisor__user"),
        id=org_id
    )
    return JsonResponse({
        "success": True,
        "org": {
            "id": org.id,
            "organization_name": org.organization_name,
            "motto": org.motto or "",
            "description": org.description or "",
            "school_id": org.school_id,
            "school_name": org.school.school_name if org.school else None,
            "department_id": org.department_id,
            "department_name": org.department.department_name if org.department else None,
            "advisor_id": org.advisor_id,
            "advisor_name": org.advisor.user.get_full_name() if org.advisor else None,
            "max_members": org.max_members,
            "founded_date": org.founded_date.isoformat() if org.founded_date else None,
            "cover_image_url": org.cover_image.url if org.cover_image else None,
            "active_member_count": org.active_member_count,
        }
    })


@login_required
def staff_org_delete_ajax(request, uuid, org_id):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid method"}, status=405)

    org = get_object_or_404(StudentOrganization, id=org_id)

    if org.members.filter(status__in=["ACTIVE", "PENDING"]).exists():
        return JsonResponse({
            "success": False,
            "error": f'Cannot delete "{org.organization_name}" — it still has active or pending members.'
        })

    # ==========================================
    # CAPTURE DATA BEFORE DELETE
    # ==========================================

    before_data = AuditLogger.model_to_dict(
        org,
        fields=[
            "organization_name",
            "motto",
            "description",
            "max_members",
            "founded_date",
            "school",
            "department",
            "advisor",
        ]
    )

    name = org.organization_name
    org_id_value = org.pk
    org_name = org.organization_name
    org.delete()


    # ==========================================
    # AUDIT LOG
    # ==========================================

    AuditLogger.log(
        request=request,
        action="DELETE",
        module="Staff",
        object_type="Student Organization",
        object_id=org_id_value,
        description=(
            f'Staff deleted organization '
            f'"{org_name}".'
        ),
        before_data=before_data,
        status="SUCCESS",
    )
    return JsonResponse({"success": True, "message": f'"{name}" deleted.'})


@login_required
def staff_org_approve_ajax(request, uuid, membership_id):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid method"}, status=405)

    membership = get_object_or_404(StudentOrganizationMembership, membership_id=membership_id, status="PENDING")

    # ==========================================
    # BEFORE DATA
    # ==========================================

    before_data = AuditLogger.model_to_dict(
        membership,
        fields=[
            "student",
            "organization",
            "status",
            "role",
            "start_date",
            "end_date",
        ]
    )
    membership.status = "ACTIVE"
    membership.save()

    # ==========================================
    # AFTER DATA
    # ==========================================

    after_data = AuditLogger.model_to_dict(
        membership,
        fields=[
            "student",
            "organization",
            "status",
            "role",
            "start_date",
            "end_date",
        ]
    )
    # ==========================================
    # AUDIT LOG
    # ==========================================
    AuditLogger.log(
        request=request,
        action="APPROVE",
        module="Staff",
        object_type="Organization Membership",
        object_id=membership.membership_id,
        description=(
            f'Staff approved the organization membership '
            f'request for '
            f'"{membership.student.user.get_full_name()}" '
            f'in "{membership.organization.organization_name}".'
        ),
        before_data=before_data,
        after_data=after_data,
        status="SUCCESS",
    )
    notify_student_org_request_approved(membership)

    return JsonResponse({"success": True, "message": "Request approved."})


@login_required
def staff_org_reject_ajax(request, uuid, membership_id):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid method"}, status=405)

    membership = get_object_or_404(StudentOrganizationMembership, membership_id=membership_id , status="PENDING")
    # ==========================================
    # CAPTURE DATA BEFORE DELETE
    # ==========================================
    before_data = AuditLogger.model_to_dict(
        membership,
        fields=[
            "student",
            "organization",
            "status",
            "role",
            "start_date",
            "end_date",
        ]
    )
    membership_id_value = membership.membership_id
    student_name = (
        membership.student.user.get_full_name()
    )
    organization_name = (
        membership.organization.organization_name
    )
    membership.delete()
    # ==========================================
    # AUDIT LOG
    # ==========================================
    AuditLogger.log(
        request=request,
        action="REJECT",
        module="Staff - Organizations",
        object_type="Organization Membership",
        object_id=membership_id_value,
        description=(
            f'Staff rejected the organization membership '
            f'request of "{student_name}" for '
            f'"{organization_name}".'
        ),
        before_data=before_data,
        status="SUCCESS",
    )


    return JsonResponse({"success": True, "message": "Request rejected."})


@login_required
def staff_org_members_ajax(request, uuid, org_id):
    """List current members of an org"""
    org = get_object_or_404(StudentOrganization, id=org_id)
    members = StudentOrganizationMembership.objects.filter(
        organization=org
    ).exclude(status="INACTIVE").select_related("student__user").order_by("student__user__first_name")

    return JsonResponse({
        "success": True,
        "org_name": org.organization_name,
        "members": [
            {
                "membership_id": m.membership_id,
                "name": m.student.user.get_full_name(),
                "email": m.student.university_email,
                "status": m.status,
                "role": m.role or "Member",
            }
            for m in members
        ],
    })


@login_required
def staff_search_students_ajax(request, uuid):
    """Search students to add — excludes anyone already in the target org"""
    q = request.GET.get("q", "").strip()
    org_id = request.GET.get("org_id")

    if not q or len(q) < 2:
        return JsonResponse({"success": True, "students": []})

    already_in = StudentOrganizationMembership.objects.filter(
        organization_id=org_id
    ).exclude(status="INACTIVE").values_list("student_id", flat=True)

    students = StudentProfile.objects.filter(
        Q(user__first_name__icontains=q) |
        Q(user__last_name__icontains=q) |
        Q(university_email__icontains=q) |
        Q(student_number__icontains=q)
    ).exclude(id__in=already_in).select_related("user")[:10]

    return JsonResponse({
        "success": True,
        "students": [
            {
                "id": s.id,
                "name": s.user.get_full_name(),
                "email": s.university_email,
                "student_number": s.student_number,
            }
            for s in students
        ],
    })


@login_required
def staff_add_member_ajax(request, uuid, org_id):
    """Staff directly adds a student — goes straight to ACTIVE, no approval needed"""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid method"}, status=405)

    org = get_object_or_404(StudentOrganization, id=org_id)
    student_id = request.POST.get("student_id")
    role = request.POST.get("role", "").strip() or "Member"

    student = get_object_or_404(StudentProfile, id=student_id)

    existing = StudentOrganizationMembership.objects.filter(
        student=student, organization=org
    ).exclude(status="INACTIVE").first()

    if org.is_full:
        return JsonResponse({
            "success": False,
            "error": f"{org.organization_name} is at capacity ({org.max_members} members)."
        })

    if existing:
        return JsonResponse({"success": False, "error": f"{student.user.get_full_name()} is already a member."})

    membership = StudentOrganizationMembership.objects.create(
        student=student,
        organization=org,
        status="ACTIVE",
        role=role,
        start_date=timezone.now().date(),
    )

    # ==========================================
    # AUDIT LOG
    # ==========================================

    AuditLogger.log(
        request=request,
        action="CREATE",
        module="Staff - Organizations",
        object_type="Organization Membership",
        object_id=membership.membership_id,
        description=(
            f'Staff directly added '
            f'"{student.user.get_full_name()}" '
            f'to "{org.organization_name}" '
            f'as "{role}".'
        ),
        after_data=AuditLogger.model_to_dict(
            membership,
            fields=[
                "student",
                "organization",
                "status",
                "role",
                "start_date",
                "end_date",
            ]
        ),
        status="SUCCESS",
    )


    notify_student_added_to_org(membership)

    return JsonResponse({"success": True, "message": f"{student.user.get_full_name()} added to {org.organization_name}."})


@login_required
def staff_remove_member_ajax(request, uuid, membership_id):
    """Staff removes a member"""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid method"}, status=405)

    membership = get_object_or_404(StudentOrganizationMembership, membership_id=membership_id)
    # ==========================================
    # BEFORE DATA
    # ==========================================
    before_data = AuditLogger.model_to_dict(
        membership,
        fields=[
            "student",
            "organization",
            "status",
            "role",
            "start_date",
            "end_date",
        ]
    )
    student_name = (
        membership.student.user.get_full_name()
    )
    organization_name = (
        membership.organization.organization_name
    )
    membership.status = "INACTIVE"
    membership.end_date = timezone.now().date()
    membership.save()
    # ==========================================
    # AFTER DATA
    # ==========================================

    after_data = AuditLogger.model_to_dict(
        membership,
        fields=[
            "student",
            "organization",
            "status",
            "role",
            "start_date",
            "end_date",
        ]
    )


    # ==========================================
    # AUDIT LOG
    # ==========================================
    AuditLogger.log(
        request=request,
        action="STATUS_CHANGE",
        module="Staff - Organizations",
        object_type="Organization Membership",
        object_id=membership.membership_id,
        description=(
            f'Staff removed "{student_name}" '
            f'from "{organization_name}".'
        ),
        before_data=before_data,
        after_data=after_data,
        status="SUCCESS",
    )

    return JsonResponse({"success": True, "message": "Member removed."})

# ************************************************ Arun Code ********************************************************




#<-----------------------BLAZE CODE START STAFF LEAVE REQUEST(28.07.26)---------------------------->

from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.contrib import messages
from django.utils import timezone
from django.db.models import Q, Sum, Count
from datetime import datetime, timedelta
import json
import csv
from .models import StaffLeaveRequest, StaffLeaveType, StaffLeaveBalance

@login_required
def staff_leave_dashboard(request, uuid):
    """Main staff leave dashboard with stats and analytics"""
    if not hasattr(request.user, 'staff_profile'):
        messages.error(request, 'Access denied. Staff only.')
        return redirect('staff_dashboard', uuid=uuid)

    default_leave_types = [
        {'name': 'Annual Leave', 'code': 'ANNUAL', 'max_days_per_year': 30, 'color': '#3b82f6'},
        {'name': 'Sick Leave', 'code': 'SICK', 'max_days_per_year': 15, 'color': '#22c55e'},
        {'name': 'Casual Leave', 'code': 'CASUAL', 'max_days_per_year': 10, 'color': '#f59e0b'},
        {'name': 'Earned Leave', 'code': 'EARNED', 'max_days_per_year': 5, 'color': '#8b5cf6'},
        {'name': 'Compensatory Leave', 'code': 'COMP', 'max_days_per_year': 5, 'color': '#ec4899'},
        {'name': 'Unpaid Leave', 'code': 'UNPAID', 'max_days_per_year': 0, 'color': '#6b7280'},
    ]
    
    for lt in default_leave_types:
        StaffLeaveType.objects.get_or_create(
            code=lt['code'],
            defaults={
                'name': lt['name'],
                'max_days_per_year': lt['max_days_per_year'],
                'color': lt['color'],
                'is_paid': lt['code'] != 'UNPAID',
                'is_half_day_allowed': lt['code'] != 'UNPAID',
                'is_full_crud': True
            }
        )
    
    leave_types = StaffLeaveType.objects.all()
    leave_requests = StaffLeaveRequest.objects.filter(staff=request.user).order_by('-created_at')
    
    balance, created = StaffLeaveBalance.objects.get_or_create(
        staff=request.user,
        defaults={
            'annual_leave_balance': 30.0,
            'sick_leave_balance': 15.0,
            'casual_leave_balance': 10.0,
            'earned_leave_balance': 0.0,
            'compensatory_leave_balance': 0.0,
        }
    )
    
    total_leaves = leave_requests.count()
    pending_leaves = leave_requests.filter(status='pending').count()
    approved_leaves = leave_requests.filter(status='approved').count()
    rejected_leaves = leave_requests.filter(status='rejected').count()
    cancelled_leaves = leave_requests.filter(status='cancelled').count()
    
    current_year = timezone.now().year
    year_leaves = leave_requests.filter(
        start_date__year=current_year,
        status__in=['approved']
    )

    year_days_used = 0
    for leave in year_leaves:
        year_days_used += leave.days_count

    monthly_data = []
    months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
    for month in range(1, 13):
        month_leaves = StaffLeaveRequest.objects.filter(
            staff=request.user,
            start_date__year=current_year,
            start_date__month=month,
            status__in=['approved']
        )
        total_days = 0
        for leave in month_leaves:
            total_days += leave.days_count
        monthly_data.append(float(total_days))
    
    # Leave type distribution
    type_distribution = []
    for leave_type in leave_types:
        count = StaffLeaveRequest.objects.filter(
            staff=request.user,
            leave_type=leave_type,
            status__in=['approved']
        ).count()
        if count > 0:
            type_distribution.append({
                'leave_type__name': leave_type.name,
                'leave_type__color': leave_type.color,
                'count': count,
            })
    
    type_labels = [item['leave_type__name'] for item in type_distribution]
    type_counts = [item['count'] for item in type_distribution]
    type_colors = [item['leave_type__color'] for item in type_distribution]
    
    context = {
        'leave_types': leave_types,
        'leave_requests': leave_requests,
        'leave_balance': balance,
        'today': timezone.now().date(),
        'uuid': uuid,
        'total_leaves': total_leaves,
        'pending_leaves': pending_leaves,
        'approved_leaves': approved_leaves,
        'rejected_leaves': rejected_leaves,
        'cancelled_leaves': cancelled_leaves,
        'year_days_used': year_days_used,
        'current_year': current_year,
        'monthly_data': json.dumps(monthly_data),
        'months': json.dumps(months),
        'type_labels': json.dumps(type_labels),
        'type_counts': json.dumps(type_counts),
        'type_colors': json.dumps(type_colors),
    }
    return render(request, 'Leave_management/leave_dashboard.html', context)


@login_required
def submit_staff_leave(request, uuid):
    """Submit new staff leave request with attachment"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid request method'})
    
    try:
        data = json.loads(request.body)
        
        leave_type_id = data.get('leave_type')
        start_date = data.get('start_date')
        end_date = data.get('end_date')
        half_day = data.get('half_day', 'full')
        reason = data.get('reason')
        attachment_data = data.get('attachment', None)
        attachment_name = data.get('attachment_name', '')
        
        if not all([leave_type_id, start_date, end_date, reason]):
            return JsonResponse({'success': False, 'message': 'All fields are required'})
        
        leave_type = get_object_or_404(StaffLeaveType, id=leave_type_id)
        
        start = datetime.strptime(start_date, '%Y-%m-%d').date()
        end = datetime.strptime(end_date, '%Y-%m-%d').date()
        
        if start > end:
            return JsonResponse({'success': False, 'message': 'Start date cannot be after end date'})
        
        if start < timezone.now().date():
            return JsonResponse({'success': False, 'message': 'Cannot request leave for past dates'})
        
        # Check overlapping
        overlapping = StaffLeaveRequest.objects.filter(
            staff=request.user,
            status__in=['pending', 'approved'],
            start_date__lte=end,
            end_date__gte=start
        ).exists()
        
        if overlapping:
            return JsonResponse({'success': False, 'message': 'You already have a leave request for these dates'})
        
        # Create leave request
        leave_request = StaffLeaveRequest.objects.create(
            staff=request.user,
            leave_type=leave_type,
            start_date=start,
            end_date=end,
            half_day_type=half_day,
            reason=reason,
            status='pending'
        )


        # from Staff.utils import notify_staff_new_leave_request
        # notify_staff_new_leave_request(leave_request)
        
        # Handle attachment if provided
        if attachment_data and attachment_name:
            import base64
            from django.core.files.base import ContentFile
            
            try:
                format, imgstr = attachment_data.split(';base64,')
                ext = format.split('/')[-1]
                filename = f"{leave_request.id}_{attachment_name}"
                data = ContentFile(base64.b64decode(imgstr), name=filename)
                leave_request.attachment = data
                leave_request.attachment_name = attachment_name
                leave_request.save()
            except:
                pass
        # ==========================================
        # AUDIT: AFTER DATA
        # ==========================================
        after_data = AuditLogger.model_to_dict(
            leave_request,
            fields=[
                "leave_type",
                "start_date",
                "end_date",
                "half_day_type",
                "reason",
                "status",
                "days_count",
            ]
        )
        # ==========================================
        # AUDIT LOG: LEAVE REQUEST CREATED
        # ==========================================
        AuditLogger.log(
            request=request,
            action="CREATE",
            module="Staff Leave",
            object_type="Staff Leave Request",
            object_id=leave_request.id,
            description=(
                f'Staff "{request.user.get_full_name()}" '
                f'submitted leave request #{leave_request.id}.'
            ),
            before_data=None,
            after_data=after_data,
            status="SUCCESS",
        )
        return JsonResponse({
            'success': True,
            'message': 'Leave request submitted successfully!',
            'leave_id': leave_request.id,
            'leave': {
                'id': leave_request.id,
                'type': leave_request.leave_type.name,
                'start_date': leave_request.start_date.strftime('%b %d, %Y'),
                'end_date': leave_request.end_date.strftime('%b %d, %Y'),
                'days': leave_request.days_count,
                'status': leave_request.status,
                'half_day': leave_request.half_day_type,
            }
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})


@login_required
def cancel_staff_leave(request, uuid, leave_id):
    """Cancel a pending staff leave request"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid request method'})
    
    try:
        leave = get_object_or_404(StaffLeaveRequest, id=leave_id, staff=request.user)
        
        if leave.status != 'pending':
            return JsonResponse({'success': False, 'message': 'Only pending requests can be cancelled'})

        # ==========================================
        # BEFORE DATA
        # ==========================================

        before_data = AuditLogger.model_to_dict(
            leave,
            fields=[
                "leave_type",
                "start_date",
                "end_date",
                "days_count",
                "half_day_type",
                "reason",
                "status",
            ]
        )
        
        leave.status = 'cancelled'
        leave.save()
        # ==========================================
        # AFTER DATA
        # ==========================================

        after_data = AuditLogger.model_to_dict(
            leave,
            fields=[
                "leave_type",
                "start_date",
                "end_date",
                "days_count",
                "half_day_type",
                "reason",
                "status",
            ]
        )
        # ==========================================
        # AUDIT LOG
        # ==========================================
        AuditLogger.log(
            request=request,
            action="STATUS_CHANGE",
            module="Staff",
            object_type="Staff Leave Request",
            object_id=leave.id,
            description=(
                f'Staff cancelled a '
                f'{leave.leave_type.name} leave request '
                f'from {leave.start_date} to {leave.end_date}.'
            ),
            before_data=before_data,
            after_data=after_data,
            status="SUCCESS",
        )

        return JsonResponse({'success': True, 'message': 'Leave request cancelled successfully'})
        
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})


@login_required
def get_staff_leave_details(request, uuid, leave_id):
    """Get staff leave request details for editing"""
    try:
        leave = get_object_or_404(StaffLeaveRequest, id=leave_id, staff=request.user)
        
        if leave.status != 'pending':
            return JsonResponse({
                'success': False, 
                'message': 'Only pending requests can be edited'
            })
        
        data = {
            'id': leave.id,
            'leave_type': leave.leave_type_id,
            'start_date': leave.start_date.strftime('%Y-%m-%d'),
            'end_date': leave.end_date.strftime('%Y-%m-%d'),
            'half_day': leave.half_day_type,
            'reason': leave.reason,
            'status': leave.status,
            'days_count': leave.days_count,
        }
        return JsonResponse({'success': True, 'data': data})
        
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})


@login_required
def update_staff_leave(request, uuid, leave_id):
    """Update a pending staff leave request"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid request method'})
    
    try:
        leave = get_object_or_404(StaffLeaveRequest, id=leave_id, staff=request.user)

        before_data = AuditLogger.model_to_dict(
            leave,
            fields=[
                "leave_type",
                "start_date",
                "end_date",
                "days_count",
                "half_day_type",
                "reason",
                "status",
            ]
        )
                
        if leave.status != 'pending':
            return JsonResponse({'success': False, 'message': 'Only pending requests can be updated'})
        
        data = json.loads(request.body)
        
        leave_type_id = data.get('leave_type')
        start_date = data.get('start_date')
        end_date = data.get('end_date')
        half_day = data.get('half_day', 'full')
        reason = data.get('reason')
        
        if not all([leave_type_id, start_date, end_date, reason]):
            return JsonResponse({'success': False, 'message': 'All fields are required'})
        
        leave_type = get_object_or_404(StaffLeaveType, id=leave_type_id)
        
        start = datetime.strptime(start_date, '%Y-%m-%d').date()
        end = datetime.strptime(end_date, '%Y-%m-%d').date()
        
        if start > end:
            return JsonResponse({'success': False, 'message': 'Start date cannot be after end date'})
        
        # Check overlapping (excluding current)
        overlapping = StaffLeaveRequest.objects.filter(
            staff=request.user,
            status__in=['pending', 'approved'],
            start_date__lte=end,
            end_date__gte=start
        ).exclude(id=leave_id).exists()
        
        if overlapping:
            return JsonResponse({'success': False, 'message': 'You already have a leave request for these dates'})
        
        leave.leave_type = leave_type
        leave.start_date = start
        leave.end_date = end
        leave.half_day_type = half_day
        leave.reason = reason
        leave.save()

        after_data = AuditLogger.model_to_dict(
            leave,
            fields=[
                "leave_type",
                "start_date",
                "end_date",
                "days_count",
                "half_day_type",
                "reason",
                "status",
            ]
        )
        AuditLogger.log(
            request=request,
            action="UPDATE",
            module="Staff",
            object_type="Staff Leave Request",
            object_id=leave.id,
            description=(
                f'Staff updated a leave request '
                f'from {leave.start_date} to {leave.end_date}.'
            ),
            before_data=before_data,
            after_data=after_data,
            status="SUCCESS",
        )
        
        return JsonResponse({
            'success': True, 
            'message': 'Leave request updated successfully!',
            'leave': {
                'id': leave.id,
                'type': leave.leave_type.name,
                'start_date': leave.start_date.strftime('%b %d, %Y'),
                'end_date': leave.end_date.strftime('%b %d, %Y'),
                'days': leave.days_count,
            }
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})


@login_required
def get_staff_leave_balance(request, uuid):
    """Get staff leave balance"""
    try:
        balance, created = StaffLeaveBalance.objects.get_or_create(
            staff=request.user,
            defaults={
                'annual_leave_balance': 30.0,
                'sick_leave_balance': 15.0,
                'casual_leave_balance': 10.0,
                'earned_leave_balance': 0.0,
                'compensatory_leave_balance': 0.0,
            }
        )
        data = {
            'annual_leave': balance.annual_leave_balance,
            'sick_leave': balance.sick_leave_balance,
            'casual_leave': balance.casual_leave_balance,
            'earned_leave': balance.earned_leave_balance,
            'compensatory_leave': balance.compensatory_leave_balance,
        }
        return JsonResponse({'success': True, 'data': data})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})


@login_required
def export_staff_leave_report(request, uuid):
    """Export staff leave report as CSV"""
    try:
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="staff_leave_report_{}.csv"'.format(
            timezone.now().strftime('%Y%m%d')
        )
        
        writer = csv.writer(response)
        writer.writerow([
            'S.No', 'Leave Type', 'Start Date', 'End Date', 
            'Days', 'Half Day', 'Reason', 'Status', 
            'Submitted On', 'Approved On'
        ])
        
        leaves = StaffLeaveRequest.objects.filter(staff=request.user).order_by('-created_at')
        for idx, leave in enumerate(leaves, 1):
            writer.writerow([
                idx,
                leave.leave_type.name,
                leave.start_date.strftime('%Y-%m-%d'),
                leave.end_date.strftime('%Y-%m-%d'),
                leave.days_count,
                leave.get_half_day_type_display(),
                leave.reason[:100],
                leave.display_status,
                leave.created_at.strftime('%Y-%m-%d %H:%M'),
                leave.approved_date.strftime('%Y-%m-%d %H:%M') if leave.approved_date else '',
            ])
        # ==========================================
        # AUDIT LOG - STAFF LEAVE CSV EXPORT
        # ==========================================
        AuditLogger.log(
            request=request,
            action="EXPORT",
            module="Staff",
            object_type="Staff Leave Report",
            object_id="",
            description=(
                f'Staff exported {record_count} '
                f'leave request record(s) to CSV.'
            ),
            after_data={
                "export_format": "CSV",
                "record_count": record_count,
                "filename": filename,
            },
            status="SUCCESS",

        )
        
        return response
        
    except Exception as e:
        # ==========================================
        # AUDIT LOG - FAILED EXPORT
        # ==========================================
        AuditLogger.log(
            request=request,
            action="EXPORT",
            module="Staff",
            object_type="Staff Leave Report",
            object_id="",
            description=(
                f'Failed to export staff leave report. '
                f'Error: {str(e)}'
            ),
            status="FAILED",
        )
        messages.error(request, f'Error exporting report: {str(e)}')
        return redirect('Staff:staff_leave_dashboard', uuid=uuid)


#<-----------------------BLAZE CODE END STAFF LEAVE REQUEST(28.07.26)---------------------------->

######################################speed code###############################################
############################# speed code starts#####################################################
 
@login_required
def get_student_academic_profile(request):
    """Return a student's department and academic program for form auto-fill."""
    try:
        student_user_id = request.GET.get('student_id')
        if not student_user_id:
            return JsonResponse({'success': False, 'error': 'student_id required'}, status=400)
 
        student_user = get_object_or_404(User, id=student_user_id, is_student=True)
        student_profile = getattr(student_user, 'student_profile', None)
        academic_profile = getattr(student_profile, 'academic_profile', None) if student_profile else None
 
        department_id = None
        department_name = ''
        program_id = None
        program_name = ''
 
        if academic_profile:
            if hasattr(academic_profile, 'department') and academic_profile.department:
                department_id = academic_profile.department.department_id
                department_name = academic_profile.department.department_name
            if hasattr(academic_profile, 'program') and academic_profile.program:
                program_id = academic_profile.program.program_id
                program_name = academic_profile.program.program_name
 
        return JsonResponse({
            'success': True,
            'department_id': department_id,
            'department_name': department_name,
            'program_id': program_id,
            'program_name': program_name,
        })
################################### Speed code end ##################################################
    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'success': False, 'error': str(e)}, status=500)

 

 # speed code start
def get_bulk_eligible_students(request):
    """Get eligible students for bulk enrollment based on course's department and program."""
    try:
        course_id = request.GET.get('course_id')
        department_id = request.GET.get('department_id')
        program_id = request.GET.get('program_id')
        semester_id = request.GET.get('semester_id')
 
        if not course_id or not str(course_id).isdigit():
            return JsonResponse({'success': False, 'error': 'course_id required'}, status=400)
 
        student_users = User.objects.filter(is_student=True, account_status='ACTIVE')
 
        if department_id:
            student_users = student_users.filter(
                student_profile__academic_profile__department_id=department_id
            )
 
        if program_id and program_id.isdigit():
            student_users = student_users.filter(
                student_profile__academic_profile__program_id=program_id
            )
 
        # Exclude students already enrolled in this course for the current semester
        current_semester = (
            _resolve_semester(semester_id)
            if semester_id else Semester.objects.filter(is_current=True).first()
        )
        all_course_ids = {int(course_id)}
        raw_course_ids = request.GET.get('course_ids', '')
        if raw_course_ids:
            for cid in raw_course_ids.split(','):
                cid = cid.strip()
                if cid.isdigit():
                    all_course_ids.add(int(cid))

        already_enrolled_student_ids = []
        already_enrolled_by_student = {}
        if current_semester:
            enrolled_rows = StudentEnrollment.objects.filter(
                section_id__course_id__in=all_course_ids,
                semester_id=current_semester.semester_id,
            ).values_list('student_id', 'section_id__course_id')
            already_enrolled_student_ids = list({r[0] for r in enrolled_rows})
            for student_pk, cid in enrolled_rows:
                already_enrolled_by_student.setdefault(student_pk, set()).add(cid)
 
        current_section_by_student = {}
        current_course_by_student = {}
        student_course_count = {}
        student_total_credits = {}
        if current_semester:
            current_enrollments = list(StudentEnrollment.objects.filter(
                student_id__in=student_users.values_list('student_profile__id', flat=True),
                semester_id=current_semester.semester_id,
            ).select_related('section_id').order_by('student_id', '-id'))
            section_ids = set()
            for enrollment in current_enrollments:
                if enrollment.section_id:
                    section_ids.add(enrollment.section_id.pk)
            section_course_map = {}
            if section_ids:
                sections = CourseSection.objects.filter(pk__in=section_ids).select_related('course')
                for s in sections:
                    if s.course:
                        section_course_map[s.pk] = f"{s.course.course_code} — {s.course.course_name}"
            for enrollment in current_enrollments:
                sid = enrollment.student_id
                sec = enrollment.section_id
                if sid not in current_section_by_student:
                    current_section_by_student[sid] = (
                        f"Section {sec.section_number}" if sec else 'Not assigned'
                    )
                if sid not in current_course_by_student:
                    course_info = section_course_map.get(sec.pk, '—') if sec else '—'
                    current_course_by_student[sid] = course_info
                student_course_count[sid] = student_course_count.get(sid, 0) + 1
                student_total_credits[sid] = student_total_credits.get(sid, 0) + (enrollment.credit_load or 0)
 
        students = []
        for user in student_users.select_related('student_profile').order_by('first_name', 'last_name'):
            sp = getattr(user, 'student_profile', None)
            if sp:
                already = sp.id in already_enrolled_student_ids
                enrolled_cids = list(already_enrolled_by_student.get(sp.id, set()))
                students.append({
                    'id': user.id,
                    'first_name': user.first_name or '',
                    'last_name': user.last_name or '',
                    'student_number': sp.student_number or 'N/A',
                    'academic_level': sp.get_academic_level_display() if sp.academic_level else '—',
                    'status': user.get_account_status_display() if user.account_status else 'Active',
                    'section': current_section_by_student.get(sp.id, 'Not assigned'),
                    'current_course': current_course_by_student.get(sp.id, ''),
                    'current_courses_count': student_course_count.get(sp.id, 0),
                    'total_credits': student_total_credits.get(sp.id, 0),
                    'already_enrolled': already,
                    'already_enrolled_course_ids': enrolled_cids,
                })
 
        return JsonResponse({'success': True, 'students': students})
 
    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'success': False, 'error': str(e)}, status=500)
    
def _student_admission_year(student_profile):
    """Best-effort cohort year: prefer admission_date, fall back to the first
    4-digit group found in the student number."""
    sp = student_profile
    if sp is None:
        return None
    if getattr(sp, 'admission_date', None):
        try:
            return int(sp.admission_date.year)
        except (TypeError, ValueError):
            pass
    match = re.search(r'(\d{4})', str(sp.student_number or ''))
    return int(match.group(1)) if match else None


def _cohort_study_years(course_ids, program_id=None):
    """Study years a set of courses belongs to in the program curriculum."""
    from Admin.Colleges.models import ProgramCourse

    pc_qs = ProgramCourse.objects.filter(
        course_id__in=list(set(course_ids)),
        status='ACTIVE',
    )
    if program_id:
        pc_qs = pc_qs.filter(program_id=program_id)
    return sorted({int(y) for y in pc_qs.values_list('study_year', flat=True) if y})


def _cohort_matches(admission_year, study_years, academic_year):
    """True when admission_year + (study_year - 1) equals the semester's
    academic year for any of the curriculum study years.  None means we cannot
    judge (missing admission year or curriculum mapping) so treat as eligible."""
    if not admission_year or not academic_year or not study_years:
        return None
    return any(int(admission_year) + (int(y) - 1) == int(academic_year) for y in study_years)

 



@login_required
def bulk_enroll_students(request):
    """Bulk enroll multiple students into a course section."""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'}, status=400)
 
    try:
        data = json.loads(request.body)
        student_ids = data.get('student_ids', [])
        course_section_id = data.get('course_section_id')
        semester_id = data.get('semester_id')
        enrollment_status = data.get('enrollment_status', 'FULL_TIME')
        credit_load = data.get('credit_load', 3)
 
        if not student_ids:
            return JsonResponse({'success': False, 'error': 'No students selected'}, status=400)
        if not course_section_id:
            return JsonResponse({'success': False, 'error': 'Course section required'}, status=400)
        if not semester_id:
            return JsonResponse({'success': False, 'error': 'Semester required'}, status=400)
 
        semester = _resolve_semester(semester_id)
        if not semester:
            return JsonResponse({'success': False, 'error': 'Invalid semester or term.'}, status=400)
 
        section = get_object_or_404(CourseSection, section_id=course_section_id)
        if str(semester_id).isalpha() and _term_of_semester(section.semester_id) != str(semester_id).upper():
            return JsonResponse({
                'success': False,
                'error': 'Please select a section for the selected term.',
                'field': 'section',
            }, status=400)

        # Cohort gate: a course scheduled for study_year N of the curriculum
        # belongs to batch (academic_year - N + 1).  Block students from other
        # batches unless the caller explicitly overrode the warning.
        cohort_override = bool(data.get('cohort_override'))
        program_id_raw = data.get('program_id')
        mismatched_students = []
        if student_ids:
            cohort_program_id = program_id_raw if str(program_id_raw or '').isdigit() else None
            try:
                cohort_study_years = _cohort_study_years([section.course_id], cohort_program_id)
            except Exception:
                cohort_study_years = []
            if cohort_study_years:
                profile_by_user_id = {
                    sp.user_id: sp
                    for sp in StudentProfile.objects.filter(user_id__in=student_ids)
                }
                for uid in student_ids:
                    sp = profile_by_user_id.get(uid)
                    admission_year = _student_admission_year(sp)
                    if _cohort_matches(
                        admission_year, cohort_study_years, semester.academic_year
                    ) is False:
                        name = (
                            sp.user.get_full_name()
                            if sp and sp.user else f'Student #{uid}'
                        )
                        mismatched_students.append({
                            'student_id': uid,
                            'name': name,
                            'admission_year': admission_year,
                        })
        if mismatched_students and not cohort_override:
            names = ', '.join(m['name'] for m in mismatched_students[:5])
            more = '' if len(mismatched_students) <= 5 else f" and {len(mismatched_students) - 5} more"
            return JsonResponse({
                'success': False,
                'error': (
                    f"Cohort mismatch: {names}{more} "
                    f"{'is' if len(mismatched_students) == 1 else 'are'} not in the expected "
                    f"batch for this course's study year. Enable \"Include other-batch students\" "
                    f"to enroll them anyway."
                ),
                'cohort_mismatch': mismatched_students,
            }, status=400)

        results = []
        success_count = 0
        fail_count = 0
 
        for uid in student_ids:
            try:
                student_user = User.objects.get(id=uid, is_student=True)
                student_profile = StudentProfile.objects.get(user=student_user)

                existing = StudentEnrollment.objects.filter(
                    student=student_profile,
                    semester_id=semester.semester_id,
                     section_id__course=section.course,
                ).first()
                if existing:
                    existing.section_id = section
                    existing.enrollment_status = enrollment_status
                    existing.credit_load = credit_load
                    existing.save()
                    results.append({
                        'student_id': uid,
                        'student_name': student_user.get_full_name(),
                        'student_number': student_profile.student_number,
                        'success': True,
                        'updated': True,
                    })
                    success_count += 1
                    continue
 
                StudentEnrollment.objects.create(
                    student=student_profile,
                    semester=semester,
                    section_id=section,
                    enrollment_status=enrollment_status,
                    credit_load=credit_load,
                    academic_standing='GOOD_STANDING',
                )
                results.append({
                    'student_id': uid,
                    'student_name': student_user.get_full_name(),
                    'student_number': student_profile.student_number,
                    'success': True,
                })
                success_count += 1
 
            except User.DoesNotExist:
                results.append({'student_id': uid, 'student_name': f'User #{uid}', 'student_number': '-', 'success': False, 'error': 'User not found'})
                fail_count += 1
            except StudentProfile.DoesNotExist:
                results.append({'student_id': uid, 'student_name': f'User #{uid}', 'student_number': '-', 'success': False, 'error': 'No student profile'})
                fail_count += 1
            except Exception as e:
                results.append({'student_id': uid, 'student_name': f'User #{uid}', 'student_number': '-', 'success': False, 'error': str(e)})
                fail_count += 1
 
        return JsonResponse({
            'success': True,
            'data': {
                'success_count': success_count,
                'fail_count': fail_count,
                'results': results,
            }
        })
 
    except json.JSONDecodeError:
        return JsonResponse({'success': False, 'error': 'Invalid JSON data'}, status=400)
    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'success': False, 'error': str(e)}, status=500)



# speed code ends
 

@login_required
def bulk_enrollment(request, uuid):
    """Standalone multi-step page for enrolling several students at once."""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
 
    from Admin.bela_admin.models import Department
    from Admin.Colleges.models import ProgramCourse

    program_course_terms = {}
    for pc in ProgramCourse.objects.filter(status='ACTIVE').values('course_id', 'program_id', 'term', 'study_year'):
        key = f"{pc['course_id']}:{pc['program_id']}"
        program_course_terms.setdefault(key, []).append({
            'term': pc['term'],
            'study_year': pc['study_year']
        })

    context = {
        'user': request.user,
        'departments': Department.objects.all().order_by('department_name'),
        'courses': Course.objects.select_related('department', 'academic_program').filter(
            status='ACTIVE'
        ).order_by('course_code'),
        'semesters': Semester.objects.order_by('-is_current', '-academic_year', '-start_date'),
        'programs': AcademicProgram.objects.filter(status='ACTIVE').select_related('department').order_by('program_name'),
        'program_course_terms_json': json.dumps(program_course_terms),
        'today': timezone.now().date(),
    }
    return render(request, 'Student_Services/Student_enrollment/bulk_enrollment.html', context)
 #speed code ends##########################################################################################################


@login_required
@require_http_methods(["POST"])
def remove_section_student(request):
    try:
        data = json.loads(request.body or '{}')
        section_id = data.get('section_id')
        student_id = data.get('student_id')
        enrollment_filter = {
            'section_id': section_id,
            'student__user_id': student_id,
        }
        if data.get('semester_id'):
            sem_ids = _resolve_semester_ids(data.get('semester_id'))
            if sem_ids:
                enrollment_filter['semester_id__in'] = sem_ids
        deleted, _ = StudentEnrollment.objects.filter(**enrollment_filter).delete()
        return JsonResponse({'success': deleted > 0})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)

@login_required
def staff_enrollment_student_view(request, uuid, enrollment_id):
    """Separate page to view a student's enrollment details, linked courses and schedules."""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    try:
        enrollment = get_object_or_404(StudentEnrollment, id=enrollment_id)
        student = enrollment.student
        user = student.user

        # Academic profile — safe to handle missing profile
        try:
            academic_profile = student.academic_profile
        except Exception:
            academic_profile = None

        department = getattr(academic_profile, 'department', None)
        program    = getattr(academic_profile, 'program', None)
        degree     = getattr(academic_profile, 'degree', None)

        # ----------------------------------------------------------------
        # Fetch ALL enrollments for this student — simple ordering to avoid
        # NULL JOIN issues with multi-level order_by on nullable FKs.
        # ----------------------------------------------------------------
        raw_enrollments = (
            StudentEnrollment.objects
            .filter(student=student)
            .select_related(
                'section_id',
                'section_id__course',
                'section_id__semester_id',
                'semester',
            )
            .order_by('-id')          # newest enrollment first, always safe
        )

        course_cards = []

        # Program curriculum term -> term label (curriculum is the source of
        # truth for which term a course is taken in a given program, using the
        # same term values as ProgramCourse).
        from Admin.Colleges.models import ProgramCourse
        program_course_terms = {}
        for pc in ProgramCourse.objects.filter(status='ACTIVE').values('course_id', 'program_id', 'term'):
            key = f"{pc['course_id']}:{pc['program_id']}"
            program_course_terms.setdefault(key, set()).add(pc['term'])
        program_course_terms = {k: sorted(v) for k, v in program_course_terms.items()}

        def curriculum_term_label(course_id, program_id):
            terms = program_course_terms.get(f"{course_id}:{program_id}")
            if not terms:
                return None
            return terms[0].title()

        for enroll in raw_enrollments:
            try:
                section = enroll.section_id
                if not section:
                    continue
                course = section.course
                if not course:
                    continue

                # Fetch schedules separately — avoids prefetch_related issues
                schedules = []
                try:
                    for sched in section.schedules.all():
                        schedules.append({
                            'day':        sched.get_day_of_week_display(),
                            'day_short':  sched.day_of_week,
                            'start_time': sched.start_time.strftime('%I:%M %p') if sched.start_time else '—',
                            'end_time':   sched.end_time.strftime('%I:%M %p')   if sched.end_time   else '—',
                            'room':       sched.room     or '—',
                            'building':   sched.building or '—',
                            'is_online':  sched.is_online,
                        })
                except Exception as sched_err:
                    print(f"[enrollment_view] schedule fetch error: {sched_err}")

                status_map = {
                    'FULL_TIME': ('enrolled', 'Full-Time'),
                    'PART_TIME': ('waitlist', 'Part-Time'),
                    'WITHDRAWN': ('dropped',  'Withdrawn'),
                    'LEAVE':     ('pending',  'Leave of Absence'),
                }
                status_class, status_label = status_map.get(
                    enroll.enrollment_status,
                    ('pending', enroll.enrollment_status or 'Pending'),
                )

                # Semester label — try enrollment.semester first, fall back to section.semester_id
                try:
                    sem_label = str(enroll.semester) if enroll.semester else (
                        str(section.semester_id) if section.semester_id else '—'
                    )
                except Exception:
                    sem_label = '—'

                # Curriculum is the source of truth: override the stored semester
                # when the program curriculum defines a term for this course.
                cur_term_label = curriculum_term_label(
                    course.course_id, program.program_id if program else None
                )
                if cur_term_label:
                    sem_label = cur_term_label

                # Section type display
                try:
                    sec_type = section.get_section_type_display()
                except Exception:
                    sec_type = getattr(section, 'section_type', '—')

                course_cards.append({
                    'enrollment_id': enroll.id,
                    'course_code':   course.course_code  or '—',
                    'course_name':   course.course_name  or '—',
                    'credits':       course.credits      or 0,
                    'section_number': section.section_number or '—',
                    'section_type':   sec_type,
                    'faculty_name':   section.faculty_name   or 'Unassigned',
                    'room_number':    section.room_number    or '—',
                    'building_name':  section.building_name  or '—',
                    'semester':       sem_label,
                    'status':         enroll.enrollment_status,
                    'status_class':   status_class,
                    'status_label':   status_label,
                    'schedules':      schedules,
                })

            except Exception as row_err:
                print(f"[enrollment_view] row error (enrollment id={enroll.pk}): {row_err}")
                continue

        total_credits = sum(card['credits'] for card in course_cards)

        print(f"[enrollment_view] student={user.get_full_name()}, course_cards={len(course_cards)}, total_credits={total_credits}")

        context = {
            'user':             request.user,
            'student':          student,
            'student_user':     user,
            'academic_profile': academic_profile,
            'department':       department,
            'program':          program,
            'degree':           degree,
            'enrollment':       enrollment,
            'course_cards':     course_cards,
            'total_credits':    total_credits,
            'today':            timezone.now().date(),
        }
        return render(request, 'Student_Services/Student_enrollment/student_enrollment_view.html', context)

    except Exception as e:
        import traceback; traceback.print_exc()
        return redirect('staff_enrollment', uuid=uuid)




SEMESTER_TYPE_DISPLAY = {'FA': 'Fall', 'SP': 'Spring', 'SU': 'Summer', 'WI': 'Winter'}


@login_required
def get_enrollment_list(request):
    """Return one enrollment summary per student for client-side search and pagination."""
    enrollments = StudentEnrollment.objects.select_related(
        'student', 'student__user', 'student__program',
        'student__program__department',
        'section_id__course', 'semester',
    ).all().order_by('-section_id__semester_id__start_date', 'student__user__first_name')

    students = {}
    for enrollment in enrollments:
        student = enrollment.student
        user = student.user
        section = enrollment.section_id
        course = section.course if section else None
        semester = enrollment.semester

        # Derive term/year from semester
        term = SEMESTER_TYPE_DISPLAY.get(getattr(semester, 'semester_type', None), '') if semester else ''
        year = getattr(semester, 'academic_year', None)
        term_year = f"{term} {year}" if term and year else ''

        if student.id not in students:
            photo = getattr(user, 'profile_photo', None)
            admission_year = None
            if getattr(student, 'admission_date', None):
                admission_year = student.admission_date.year

            # Resolve program and department from StudentProfile -> AcademicProgram
            program = getattr(student, 'program', None)
            program_name = getattr(program, 'program_name', '') or ''
            dept = getattr(program, 'department', None) if program else None
            department_name = getattr(dept, 'department_name', '') or ''

            students[student.id] = {
                # The queryset is newest first, so retain this record for View/Edit.
                'id': enrollment.id,
                'student_name': user.get_full_name() or 'Unknown Student',
                'student_number': student.student_number or 'N/A',
                'status': enrollment.enrollment_status or 'PENDING',
                'status_display': enrollment.get_enrollment_status_display() or enrollment.enrollment_status or 'Pending',
                'photo_url': photo.url if photo else '',
                'initials': f"{(user.first_name or 'U')[:1]}{(user.last_name or '')[:1]}",
                'course_ids': set(),
                'course_names': set(),
                'course_credits': {},
                'statuses': set(),
                'program_name': program_name,
                'department_name': department_name,
                'terms': set(),
                'admission_year': admission_year,
            }

        record = students[student.id]
        record['statuses'].add(enrollment.enrollment_status or 'PENDING')
        if term_year:
            record['terms'].add(term_year)
        if course:
            record['course_ids'].add(course.pk)
            record['course_names'].add(course.course_name or '')
            record['course_credits'][course.pk] = course.credits or 0

    results = []
    for record in students.values():
        record['course_count'] = len(record.pop('course_ids'))
        record['course_names'] = sorted(record['course_names'])
        record['credits'] = sum(record.pop('course_credits').values())
        record['statuses'] = list(record['statuses'])
        record['terms'] = sorted(record['terms'])
        results.append(record)
    return JsonResponse({'success': True, 'enrollments': results})

# Term codes mirror ProgramCourse.term (FALL/SPRING/SUMMER/WINTER) rather than
# the Semester model's types (FA/SP/SU/WI). The enrollment UI now uses term
# codes, so these helpers bridge the two.
TERM_TO_SEMESTER_TYPE = {'FALL': 'FA', 'SPRING': 'SP', 'SUMMER': 'SU', 'WINTER': 'WI'}
SEMESTER_TYPE_TO_TERM = {'FA': 'FALL', 'SP': 'SPRING', 'SU': 'SUMMER', 'WI': 'WINTER'}


def _resolve_semester_ids(value):
    """Resolve a numeric Semester id or a term code (FALL/SPRING/SUMMER/WINTER)
    to a list of matching Semester ids. Returns None when no filter is given,
    [] for a term with no matching semester."""
    if value is None or value == '':
        return None
    s = str(value).strip()
    if not s:
        return None
    if s.isdigit():
        return [int(s)]
    stype = TERM_TO_SEMESTER_TYPE.get(s.upper())
    if not stype:
        return []
    current = Semester.objects.filter(semester_type=stype, is_current=True).first()
    if current:
        return [current.semester_id]
    latest = Semester.objects.filter(semester_type=stype).order_by('-academic_year', '-semester_id').first()
    if latest:
        return [latest.semester_id]
    return []

def _resolve_semester(value):
    """Resolve a numeric Semester id or term code to a single Semester row,
    preferring the current semester of a term, then the newest. None when
    unresolvable."""
    if value is None or value == '':
        return None
    s = str(value).strip()
    if not s:
        return None
    if s.isdigit():
        return Semester.objects.filter(semester_id=int(s)).first()
    stype = TERM_TO_SEMESTER_TYPE.get(s.upper())
    if not stype:
        return None
    return Semester.objects.filter(semester_type=stype).order_by(
        '-is_current', '-academic_year', '-semester_id'
    ).first()


def _term_of_semester(sem):
    return SEMESTER_TYPE_TO_TERM.get(getattr(sem, 'semester_type', None)) if sem else None


@login_required
def course_allocation(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
        
    from Admin.bela_admin.models import Department, Course
    from Admin.Colleges.models import ProgramCourse, AcademicProgram, AcademicTerm
    from django.utils import timezone
    
    courses = Course.objects.select_related('department').filter(status='ACTIVE').order_by('course_code')
    departments = Department.objects.prefetch_related('programs').filter(status='ACTIVE').order_by('department_name')
    academic_terms = AcademicTerm.objects.filter(status='ACTIVE').order_by('academic_year', 'start_date')
    allocated_course_ids = set(ProgramCourse.objects.filter(status='ACTIVE').values_list('course_id', flat=True))
    
    study_years = [{'value': y[0], 'label': y[1]} for y in ProgramCourse.YEAR_CHOICES]
    term_choices = [{'value': t[0], 'label': t[1]} for t in ProgramCourse.TERM_CHOICES]
    
    context = {
        'user': request.user,
        'courses': courses,
        'departments': departments,
        'academic_terms': academic_terms,
        'study_years': study_years,
        'term_choices': term_choices,
        'allocated_course_ids': allocated_course_ids,
        'today': timezone.now().date(),
    }
    return render(request, 'Student_Services/Student_enrollment/course_allocation.html', context)


@login_required
def get_course_allocations(request):
    from Admin.Colleges.models import ProgramCourse
    course_id = request.GET.get('course_id')
    if not course_id:
        return JsonResponse({'success': False, 'error': 'course_id required'}, status=400)
        
    allocations_qs = ProgramCourse.objects.filter(course_id=course_id).select_related('program__department')
    allocations = []
    
    user_name = request.user.get_full_name() or request.user.username or 'Admin User'
    
    for pc in allocations_qs:
        allocations.append({
            'id': pc.program_course_id,
            'department_id': pc.program.department.department_id,
            'department_name': pc.program.department.department_name,
            'program_id': pc.program.program_id,
            'program_name': pc.program.program_name,
            'study_year': pc.study_year,
            'study_year_display': f"Year {pc.study_year}",
            'term': pc.term,
            'term_display': pc.get_term_display(),
            'is_elective': pc.is_elective,
            'status': pc.status,
            'allocated_on': pc.created_at.strftime('%b %d, %Y') if pc.created_at else '—',
            'allocated_by': user_name
        })
        
    return JsonResponse({'success': True, 'allocations': allocations})


@login_required
@require_http_methods(["POST"])
def save_course_allocations(request):
    import json
    import traceback
    from django.db import transaction
    from Admin.Colleges.models import ProgramCourse, AcademicProgram
    from Admin.bela_admin.models import Course
    
    try:
        data = json.loads(request.body or '{}')
        course_id = data.get('course_id')
        term = data.get('term')
        study_year = data.get('study_year')
        program_ids = data.get('program_ids', [])
        
        if not course_id or not term or not study_year:
            return JsonResponse({'success': False, 'error': 'Missing required fields'}, status=400)
            
        course = get_object_or_404(Course, pk=course_id)
        study_year_int = int(study_year)
        
        with transaction.atomic():
            existing_pcs = ProgramCourse.objects.filter(
                course=course,
                term=term,
                study_year=study_year_int
            )
            
            existing_program_ids = set(existing_pcs.values_list('program_id', flat=True))
            target_program_ids = set(int(pid) for pid in program_ids)
            
            existing_pcs.exclude(program_id__in=target_program_ids).delete()
            existing_pcs.filter(program_id__in=target_program_ids).update(status='ACTIVE')
            
            for prog_id in target_program_ids - existing_program_ids:
                program = get_object_or_404(AcademicProgram, pk=prog_id)
                ProgramCourse.objects.create(
                    course=course,
                    program=program,
                    term=term,
                    study_year=study_year_int,
                    status='ACTIVE'
                )
                
        return JsonResponse({'success': True, 'message': 'Allocations saved successfully'})
    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def delete_course_allocation(request):
    import json
    import traceback
    from Admin.Colleges.models import ProgramCourse
    try:
        data = json.loads(request.body or '{}')
        allocation_id = data.get('allocation_id')
        if not allocation_id:
            return JsonResponse({'success': False, 'error': 'allocation_id required'}, status=400)
            
        pc = get_object_or_404(ProgramCourse, pk=allocation_id)
        pc.delete()
        return JsonResponse({'success': True, 'message': 'Allocation deleted successfully'})
    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'success': False, 'error': str(e)}, status=500)

@login_required
def staff_view_students(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    context = {
        'user': request.user,
        'today': timezone.now().date(),
    }
    return render(request, "Student_Services/Student_enrollment/enrollment_view_students.html", context)

@login_required
def staff_enrollment_student_edit(request, uuid, enrollment_id):
    """Separate edit page for one student's enrollments — Only Course and Section are editable
    every other field is read-only. Each enrolled course is shown as a button,
    and only the selected course's enrollment is changed."""
    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    try:
        from Admin.Colleges.models import ProgramCourse  # noqa: F401

        enrollment = get_object_or_404(StudentEnrollment, id=enrollment_id)
        student = enrollment.student
        user = student.user

        try:
            academic_profile = student.academic_profile
        except Exception:
            academic_profile = None

        department = getattr(academic_profile, 'department', None)
        program = getattr(academic_profile, 'program', None)
        degree = getattr(academic_profile, 'degree', None)

        # ---- Fixed curriculum position per course (ProgramCourse.study_year).
        # This never changes: CS101 is always a Year-1 course no matter which
        # academic year the student actually took it in. ----
        curriculum_year_map = {}
        if program:
            for cid, sy in ProgramCourse.objects.filter(program=program).values_list('course_id', 'study_year'):
                curriculum_year_map.setdefault(cid, set()).add(sy)

        # ---- All enrolled courses of this student (one button each) ----
        raw_enrollments = (
            StudentEnrollment.objects.filter(student=student)
            .select_related(
                'section_id',
                'section_id__course',
                'section_id__semester_id',
                'semester',
            )
            .order_by('id')
        )

        enrolled_courses = []
        for e in raw_enrollments:
            sec = e.section_id
            c = sec.course if sec else None
            try:
                sem_label = str(e.semester) if e.semester else (
                    str(sec.semester_id) if sec and sec.semester_id else '—'
                )
            except Exception:
                sem_label = '—'
            try:
                sem_type = getattr(
                    e.semester or (sec.semester_id if sec else None),
                    'semester_type',
                    None,
                )
            except Exception:
                sem_type = None
            try:
                sec_type = sec.get_section_type_display() if sec else '—'
            except Exception:
                sec_type = getattr(sec, 'section_type', '—')
            try:
                academic_year = getattr(
                    e.semester or (sec.semester_id if sec else None),
                    'academic_year',
                    None,
                )
            except Exception:
                academic_year = None

            enrolled_courses.append({
                'enrollment_id': e.id,
                'course_id': c.course_id if c else None,
                'course_code': (c.course_code or '—') if c else '—',
                'course_name': (c.course_name or 'No Course') if c else 'No Course',
                'credits': c.credits if c else 0,
                'section_id': sec.section_id if sec else None,
                'section_number': (sec.section_number or '—') if sec else '—',
                'section_type': sec_type,
                'faculty_name': (sec.faculty_name or 'Unassigned') if sec else 'Unassigned',
                'semester_id': e.semester_id or (sec.semester_id if sec else None),
                'semester': sem_label,
                'semester_type': sem_type,
                'academic_year': academic_year,
                'study_year': (
                    min(curriculum_year_map[c.course_id])
                    if c and curriculum_year_map.get(c.course_id) else None
                ),
                'status': e.enrollment_status or 'PENDING',
                'status_display': e.get_enrollment_status_display() or 'Pending',
            })

        # ---- Editable course options: department + program courses, always
        # including every course the student is currently enrolled in ----
        department_course_ids = []
        if department:
            department_course_ids = list(
                Course.objects.filter(department=department).values_list('course_id', flat=True)
            )
        enrolled_course_ids = [ec['course_id'] for ec in enrolled_courses if ec['course_id']]
        allowed_ids = set(department_course_ids) | set(enrolled_course_ids)

        course_qs = Course.objects.filter(course_id__in=allowed_ids) if allowed_ids else Course.objects.none()
        if program:
            course_qs = course_qs.filter(
                Q(academic_program=program)
                | Q(academic_program__isnull=True)
                | Q(course_id__in=enrolled_course_ids)
            )
        course_qs = course_qs.order_by('course_code')

        course_options = [
            {'id': c.course_id, 'code': c.course_code, 'name': c.course_name, 'credits': c.credits}
            for c in course_qs
        ]

        # ---- Term TYPE each course is offered in ('FA'/'SP'/'SU'/'WI'): the
        # edit page locks the enrollment's semester, so the course dropdown may
        # only contain courses offered in that term. Matching by TYPE (not by
        # Semester row id) keeps this correct across academic years. Same rule
        # as the new-enrollment page: ProgramCourse curriculum allocations are
        # the source of truth; courses with no allocation fall back to the
        # terms where they actually have sections. ----
        TERM_TO_SEM_TYPE = {'FALL': 'FA', 'SPRING': 'SP', 'SUMMER': 'SU', 'WINTER': 'WI'}

        pc_qs = ProgramCourse.objects.filter(course_id__in=allowed_ids)
        if program:
            pc_qs = pc_qs.filter(program=program)
        curriculum_types = {}
        for cid, term in pc_qs.values_list('course_id', 'term'):
            ttype = TERM_TO_SEM_TYPE.get(term)
            if ttype:
                curriculum_types.setdefault(str(cid), set()).add(ttype)

        section_types = {}
        for cid, stype in (
            CourseSection.objects.filter(course_id__in=allowed_ids)
            .values_list('course_id', 'semester_id__semester_type')
        ):
            if stype:
                section_types.setdefault(str(cid), set()).add(stype)

        course_semesters = {}
        for cid in set(curriculum_types) | set(section_types):
            course_semesters[cid] = sorted(
                curriculum_types.get(cid) or section_types.get(cid) or []
            )

        # ---- Specific semester IDs each course has sections in (not just term
        # type). Used client-side to restrict the dropdown to courses actually
        # offered in the *same* semester as the enrollment being edited, so a
        # Fall 2026 course can only be replaced by another Fall 2026 course. ----
        course_semester_ids = {}
        for cid, sem_id in (
            CourseSection.objects.filter(course_id__in=allowed_ids)
            .values_list('course_id', 'semester_id')
        ):
            course_semester_ids.setdefault(str(cid), set()).add(sem_id)
        course_semester_ids = {k: sorted(v) for k, v in course_semester_ids.items()}

        context = {
            'user': request.user,
            'student': student,
            'student_user': user,
            'academic_profile': academic_profile,
            'department': department,
            'program': program,
            'degree': degree,
            'enrollment': enrollment,
            'enrolled_courses': enrolled_courses,
            'enrolled_courses_json': json.dumps(enrolled_courses),
            'course_options_json': json.dumps(course_options),
            'course_semesters_json': json.dumps(course_semesters),
            'course_semester_ids_json': json.dumps(course_semester_ids),
            'today': timezone.now().date(),
        }
        response = render(request, 'Student_Services/Student_enrollment/enrollment_edit_student.html', context)
        # Never let a cached copy of this page show stale course lists.
        response['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        return response

    except Exception:
        import traceback; traceback.print_exc()
        return redirect('staff_view_students', uuid=uuid)



#speed code end ##########################################################################################################

# *******************Bela code statrt**********************
@login_required
def staff_advising_department(request, uuid, department_id):

    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    department = get_object_or_404(
        Department,
        department_id=department_id,
        status="ACTIVE"
    )

    programs = AcademicProgram.objects.filter(
        department=department,
        status="ACTIVE"
    ).order_by("program_name")

    # Get active student counts from StudentAcademicProfile
    student_counts = (
        StudentAcademicProfile.objects
        .filter(
            program_id__in=programs.values("program_id"),
            student__current_status="ACTIVE"
        )
        .values("program_id")
        .annotate(total=Count("student"))
    )

    count_map = {
        item["program_id"]: item["total"]
        for item in student_counts
    }

    for program in programs:
        program.student_count = count_map.get(
            program.program_id,
            0
        )

    context = {
        "user": request.user,
        "department": department,
        "programs": programs,
        "today": timezone.localdate(),
    }

    return render(
        request,
        "Student_Services/advising_program.html",
        context
    )

@login_required
def staff_advising_program(request, uuid, department_id, program_id):

    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    # =====================================================
    # DEPARTMENT
    # =====================================================

    department = get_object_or_404(
        Department,
        department_id=department_id,
        status="ACTIVE"
    )

    # =====================================================
    # PROGRAM
    # =====================================================

    program = get_object_or_404(
        AcademicProgram,
        program_id=program_id,
        department=department,
        status="ACTIVE"
    )

    # =====================================================
    # ACTIVE STUDENTS
    # =====================================================

    students = (
        StudentAcademicProfile.objects
        .filter(
            program=program,
            student__current_status="ACTIVE"
        )
        .select_related(
            "student",
            "student__user"
        )
        .order_by(
            "student__student_number"
        )
    )

    # =====================================================
    # ACTIVE ADVISOR ASSIGNMENTS
    # =====================================================

    advisor_assignments = (
        AdvisorAssignment.objects
        .filter(
            program=program,
            is_active=True
        )
        .select_related(
            "advisor",
            "advisor__user"
        )
        .order_by(
            "admission_batch",
            "roll_number_from"
        )
    )

    # =====================================================
    # GROUP STUDENTS BY ADMISSION BATCH
    # =====================================================

    batches = {}

    for academic_profile in students:

        student = academic_profile.student

        if not student.admission_date:
            continue

        batch_year = student.admission_date.year

        if batch_year not in batches:

            batches[batch_year] = {
                "batch": batch_year,
                "students": [],
                "student_count": 0,

                # All advisors for this batch
                "advisor_assignments": [],

                # Number of students already covered
                "assigned_count": 0,

                # Number of students not yet covered
                "unassigned_count": 0,
            }

        batches[batch_year]["students"].append(student)

        batches[batch_year]["student_count"] += 1

    # =====================================================
    # ATTACH ADVISOR ASSIGNMENTS TO EACH BATCH
    # =====================================================

    for batch in batches.values():

        batch_year = batch["batch"]

        batch_assignments = [
            assignment
            for assignment in advisor_assignments
            if assignment.admission_batch == batch_year
        ]

        batch["advisor_assignments"] = batch_assignments

        # Unique advisors for display
        unique_advisors = []
        seen_advisors = set()

        for assignment in batch_assignments:
            if assignment.advisor_id not in seen_advisors:
                unique_advisors.append(assignment.advisor)
                seen_advisors.add(assignment.advisor_id)

        batch["advisors"] = unique_advisors

        # -------------------------------------------------
        # FIND HOW MANY STUDENTS ARE ALREADY ASSIGNED
        # -------------------------------------------------

        assigned_students = set()

        for assignment in batch_assignments:

            roll_from = (
                assignment.roll_number_from
                .strip()
                .upper()
            )

            roll_to = (
                assignment.roll_number_to
                .strip()
                .upper()
            )

            for student in batch["students"]:

                student_number = (
                    student.student_number
                    .strip()
                    .upper()
                )

                if roll_from <= student_number <= roll_to:
                    assigned_students.add(
                        student.student_number
                    )

        batch["assigned_count"] = len(
            assigned_students
        )

        batch["unassigned_count"] = (
            batch["student_count"]
            - batch["assigned_count"]
        )

    # =====================================================
    # SORT BATCHES
    # =====================================================

    batches = sorted(
        batches.values(),
        key=lambda item: item["batch"]
    )

    # =====================================================
    # ACTIVE ADVISORS
    # =====================================================

    advisors = (
        FacultyProfile.objects
        .filter(
            employment_status="ACTIVE"
        )
        .select_related(
            "user",
            "department"
        )
        .order_by(
            "preferred_name",
            "employee_id"
        )
    )

    # =====================================================
    # TOTAL ASSIGNED STUDENTS
    # =====================================================

    total_assigned_students = sum(
        batch["assigned_count"]
        for batch in batches
    )

    total_unassigned_students = sum(
        batch["unassigned_count"]
        for batch in batches
    )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        "user": request.user,

        "department": department,

        "program": program,

        "batches": batches,

        "total_students": students.count(),

        "total_assigned_students": total_assigned_students,

        "total_unassigned_students": total_unassigned_students,

        "advisors": advisors,

        "today": timezone.localdate(),
    }

    return render(
        request,
        "Student_Services/advising_students.html",
        context
    )

@login_required
def staff_assign_advisor(request, uuid):

    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    if request.method != "POST":
        return redirect(
            "staff_advising",
            uuid=uuid
        )

    # =====================================================
    # GET FORM DATA
    # =====================================================

    program_id = request.POST.get("program_id")
    admission_batch = request.POST.get("admission_batch")

    roll_number_from = (
        request.POST.get("roll_number_from", "")
        .strip()
        .upper()
    )

    roll_number_to = (
        request.POST.get("roll_number_to", "")
        .strip()
        .upper()
    )

    advisor_id = request.POST.get("advisor_id")


    # =====================================================
    # BASIC VALIDATION
    # =====================================================

    if not all([
        program_id,
        admission_batch,
        roll_number_from,
        roll_number_to,
        advisor_id,
    ]):

        messages.error(
            request,
            "Please fill in all advisor assignment fields."
        )

        return redirect(
            request.META.get(
                "HTTP_REFERER",
                reverse(
                    "staff_advising",
                    kwargs={"uuid": uuid}
                )
            )
        )


    # =====================================================
    # VALIDATE BATCH
    # =====================================================

    try:

        admission_batch = int(admission_batch)

    except (TypeError, ValueError):

        messages.error(
            request,
            "Invalid admission batch."
        )

        return redirect(
            request.META.get("HTTP_REFERER")
        )


    # =====================================================
    # GET PROGRAM
    # =====================================================

    program = get_object_or_404(
        AcademicProgram,
        program_id=program_id,
        status="ACTIVE"
    )


    # =====================================================
    # GET ADVISOR
    # =====================================================

    advisor = get_object_or_404(
        FacultyProfile,
        pk=advisor_id,
        employment_status="ACTIVE"
    )


    # =====================================================
    # RULE 1
    # FACULTY MUST BELONG TO SAME DEPARTMENT
    # =====================================================

    if (
        advisor.department_id
        and advisor.department_id != program.department_id
    ):

        messages.error(
            request,
            "This faculty member belongs to a different department "
            "and cannot be assigned to this program."
        )

        return redirect(
            request.META.get("HTTP_REFERER")
        )


    # =====================================================
    # RULE 2
    # FACULTY CANNOT HAVE ANOTHER ACTIVE BATCH
    # =====================================================

    existing_faculty_assignment = (
        AdvisorAssignment.objects
        .filter(
            advisor=advisor,
            is_active=True
        )
        .exclude(
            program=program,
            admission_batch=admission_batch
        )
        .select_related("program")
        .first()
    )


    if existing_faculty_assignment:

        existing_program = (
            existing_faculty_assignment.program.program_name
        )

        advisor_name = (
            advisor.preferred_name
            or advisor.user.get_full_name()
            or advisor.employee_id
        )

        messages.error(
            request,
            f"{advisor_name} is already assigned to "
            f"{existing_program} "
            f"({existing_faculty_assignment.admission_batch} batch). "
            f"This advisor cannot be assigned to another active "
            f"batch until the current assignment is completed."
        )

        return redirect(
            request.META.get("HTTP_REFERER")
        )


    # =====================================================
    # RULE 3
    # FIND EXISTING ASSIGNMENTS FOR SAME PROGRAM + BATCH
    # =====================================================

    existing_assignments = AdvisorAssignment.objects.filter(
        program=program,
        admission_batch=admission_batch,
        is_active=True
    ).select_related(
        "advisor"
    )


    # =====================================================
    # RULE 4
    # CHECK ROLL RANGE OVERLAP
    # =====================================================

    for assignment in existing_assignments:

        existing_from = (
            assignment.roll_number_from
            .strip()
            .upper()
        )

        existing_to = (
            assignment.roll_number_to
            .strip()
            .upper()
        )


        # Range overlap:
        #
        # New FROM <= Existing TO
        # AND
        # New TO >= Existing FROM

        if (
            roll_number_from <= existing_to
            and roll_number_to >= existing_from
        ):

            existing_advisor = assignment.advisor

            existing_advisor_name = (
                existing_advisor.preferred_name
                or existing_advisor.user.get_full_name()
                or existing_advisor.employee_id
            )

            messages.error(
                request,
                f"This roll-number range overlaps with an "
                f"existing assignment "
                f"({existing_from} - {existing_to}) "
                f"assigned to {existing_advisor_name}."
            )

            return redirect(
                request.META.get("HTTP_REFERER")
            )


    # =====================================================
    # RULE 5
    # CHECK ACTUAL STUDENTS
    #
    # This protects against assigning a student that is
    # already assigned, even if the range check somehow
    # does not catch it.
    # =====================================================

    students = StudentAcademicProfile.objects.filter(
        program=program,
        student__current_status="ACTIVE",
        student__admission_date__year=admission_batch
    ).select_related(
        "student"
    )


    already_assigned_students = []


    for academic_profile in students:

        student = academic_profile.student

        student_number = (
            student.student_number
            .strip()
            .upper()
        )


        # Is this student inside the new range?

        if (
            roll_number_from
            <= student_number
            <= roll_number_to
        ):

            # Check whether this student is already covered
            # by an existing assignment.

            for assignment in existing_assignments:

                existing_from = (
                    assignment.roll_number_from
                    .strip()
                    .upper()
                )

                existing_to = (
                    assignment.roll_number_to
                    .strip()
                    .upper()
                )

                if (
                    existing_from
                    <= student_number
                    <= existing_to
                ):

                    existing_advisor = assignment.advisor

                    advisor_name = (
                        existing_advisor.preferred_name
                        or existing_advisor.user.get_full_name()
                        or existing_advisor.employee_id
                    )

                    already_assigned_students.append(
                        f"{student.student_number} → {advisor_name}"
                    )

                    break


    # =====================================================
    # STOP IF STUDENTS ARE ALREADY ASSIGNED
    # =====================================================

    if already_assigned_students:

        preview = ", ".join(
            already_assigned_students[:5]
        )

        if len(already_assigned_students) > 5:
            preview += " ..."

        messages.error(
            request,
            "Some students in this range are already assigned: "
            f"{preview}"
        )

        return redirect(
            request.META.get("HTTP_REFERER")
        )


    # =====================================================
    # CREATE ASSIGNMENT
    # =====================================================

    AdvisorAssignment.objects.create(
        program=program,
        admission_batch=admission_batch,
        roll_number_from=roll_number_from,
        roll_number_to=roll_number_to,
        advisor=advisor,
        is_active=True,
    )


    # =====================================================
    # SUCCESS
    # =====================================================

    advisor_name = (
        advisor.preferred_name
        or advisor.user.get_full_name()
        or advisor.employee_id
    )

    messages.success(
        request,
        f"{advisor_name} has been assigned to "
        f"{program.program_name} "
        f"{admission_batch} batch "
        f"({roll_number_from} - {roll_number_to})."
    )


    return redirect(
        "staff_advising_program",
        uuid=uuid,
        department_id=program.department_id,
        program_id=program.program_id,
    ) 

@login_required
def staff_advising_batch(
    request,
    uuid,
    department_id,
    program_id,
    batch_year
):

    if str(request.user.uuid) != str(uuid):
        return redirect("login")

    department = get_object_or_404(
        Department,
        department_id=department_id,
        status="ACTIVE"
    )

    program = get_object_or_404(
        AcademicProgram,
        program_id=program_id,
        department=department,
        status="ACTIVE"
    )
    advisors = (
    FacultyProfile.objects
    .filter(
        employment_status="ACTIVE",
        department=department
    )
    .select_related("user", "department")
    .order_by("preferred_name", "employee_id")
)

    students = StudentAcademicProfile.objects.filter(
        program=program,
        student__current_status="ACTIVE",
        student__admission_date__year=batch_year
    ).select_related(
        "student",
        "student__user"
    ).order_by(
        "student__student_number"
    )
  
    advisor_assignments = AdvisorAssignment.objects.filter(
        program=program,
        admission_batch=batch_year,
        is_active=True
    ).select_related(
        "advisor",
        "advisor__user"
    )

    # Attach the matching advisor to each student
    for academic_profile in students:

        student = academic_profile.student

        student.advisor_assignment = None

        student_number = student.student_number.strip().upper()

        for assignment in advisor_assignments:

            roll_from = assignment.roll_number_from.strip().upper()
            roll_to = assignment.roll_number_to.strip().upper()

            if roll_from <= student_number <= roll_to:

                student.advisor_assignment = assignment

                break

    context = {
        "user": request.user,
        "department": department,
        "program": program,
        "batch_year": batch_year,
        "students": students,
        "advisors": advisors,
        "today": timezone.localdate(),
    }

    return render(
        request,
        "Student_Services/advising_batch.html",
        context
    )

# ******************************* Bela code start *******************************************************



from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.db.models import Q
from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse
from django.views.decorators.http import require_http_methods

from Admin.models import User, RoomAllocation, Hostel, Room
from Students.models import MaintenanceRequest


@login_required
def staff_housing_view(request, uuid):
    user = get_object_or_404(User, uuid=uuid)
    
    if str(request.user.uuid) != str(uuid):
        return redirect("login")
    
    current_allocation = RoomAllocation.objects.filter(
        student=user,
        status='ACTIVE'
    ).select_related('room', 'room__hostel').first()
    
    history_allocations = RoomAllocation.objects.filter(
        student=user
    ).exclude(
        id=current_allocation.id if current_allocation else None
    ).select_related('room', 'room__hostel').order_by('-allocated_date')[:5]
    
    roommate_info = None
    if current_allocation:
        roommates = RoomAllocation.objects.filter(
            room=current_allocation.room,
            status='ACTIVE'
        ).exclude(student=user).select_related('student')
        
        if roommates.exists():
            roommate = roommates.first().student
            roommate_info = {
                'name': roommate.full_name,
                'email': roommate.email,
                'username': roommate.username,
                'gender': getattr(roommate, 'gender', 'N/A'),
                'program': 'Unknown',
            }
    
    amenities_list = []
    if current_allocation and current_allocation.room.amenities:
        amenities_list = [a.strip() for a in current_allocation.room.amenities.split(',') if a.strip()]
    
    duration_days = None
    if current_allocation and current_allocation.move_in_date:
        today = timezone.now().date()
        duration_days = (today - current_allocation.move_in_date).days
    
    status_filter = request.GET.get('status', 'all')
    priority_filter = request.GET.get('priority', 'all')
    search = request.GET.get('search', '').strip()
    
    maintenance_requests = MaintenanceRequest.objects.filter(
        student=user
    ).select_related('room', 'room__hostel').order_by('-created_at')
    
    if status_filter != 'all':
        maintenance_requests = maintenance_requests.filter(status=status_filter.upper())
    
    if priority_filter != 'all':
        maintenance_requests = maintenance_requests.filter(priority=priority_filter.upper())
    
    if search:
        maintenance_requests = maintenance_requests.filter(
            Q(room__room_number__icontains=search) |
            Q(issue_type__icontains=search) |
            Q(description__icontains=search)
        )
    
    context = {
        'user': user,
        'current_allocation': current_allocation,
        'history_allocations': history_allocations,
        'roommate_info': roommate_info,
        'amenities_list': amenities_list,
        'duration_days': duration_days,
        'has_housing': current_allocation is not None,
        'maintenance_requests': maintenance_requests,
        'status_filter': status_filter,
        'priority_filter': priority_filter,
        'search': search,
        'active_sb': 'housing',
        'page_title': 'Staff Housing',
    }
    
    return render(request, 'Leave_management/staff_housing.html', context)


@login_required
def staff_get_available_rooms_json(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'error': 'Unauthorized'}, status=401)
    
    try:
        current_allocation = RoomAllocation.objects.filter(
            student=request.user,
            status='ACTIVE'
        ).select_related('room', 'room__hostel').first()
        
        if not current_allocation:
            return JsonResponse([], safe=False)
        
        rooms = Room.objects.filter(
            hostel=current_allocation.room.hostel
        ).values(
            'id', 'room_number', 'capacity', 'current_occupancy', 'status'
        ).order_by('room_number')
        
        room_list = []
        for room in rooms:
            room_list.append({
                'id': room['id'],
                'room_number': room['room_number'],
                'capacity': room['capacity'],
                'current_occupancy': room['current_occupancy'],
                'status': room['status'],
                'available_slots': room['capacity'] - room['current_occupancy'],
                'is_available': room['current_occupancy'] < room['capacity'],
                'is_current_room': room['id'] == current_allocation.room.id
            })
        
        return JsonResponse(room_list, safe=False)
        
    except Exception as e:
        import logging
        logging.error(f"Error loading rooms: {e}")
        return JsonResponse([], safe=False)


@login_required
@require_http_methods(["POST"])
def staff_maintenance_request_submit(request, uuid):
    if str(request.user.uuid) != str(uuid):
        return JsonResponse({'success': False, 'error': 'Unauthorized'})
    
    staff = get_object_or_404(User, uuid=uuid)
    
    room_id = request.POST.get('room_id')
    issue_type = request.POST.get('issue_type')
    description = request.POST.get('description')
    priority = request.POST.get('priority', 'MEDIUM')
    preferred_time = request.POST.get('preferred_time', '')
    
    if not room_id or not issue_type or not description:
        return JsonResponse({'success': False, 'error': 'All required fields must be filled'})
    
    try:
        room = get_object_or_404(Room, id=room_id)
        
        has_allocation = RoomAllocation.objects.filter(
            student=staff,
            room=room,
            status='ACTIVE'
        ).exists()
        
        if not has_allocation:
            return JsonResponse({'success': False, 'error': 'You are not allocated to this room'})
        
        MaintenanceRequest.objects.create(
            student=staff,
            room=room,
            issue_type=issue_type,
            description=description,
            priority=priority,
            preferred_time=preferred_time,
            status='PENDING'
        )
        
        return JsonResponse({
            'success': True,
            'message': 'Maintenance request submitted successfully!'
        })
        
    except Room.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Room not found'})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})

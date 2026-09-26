import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..constants import UW_RED_HEX


class ExcelWriter:

    def __init__(self):
        self.workbook = openpyxl.Workbook()
        self.worksheet = self.workbook.active
        self.header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        self.header_fill = PatternFill(start_color=UW_RED_HEX, end_color=UW_RED_HEX, fill_type="solid")
        self.header_alignment = Alignment(horizontal="center", vertical="center")
        self.data_font = Font(name="Calibri", size=10)
        self.data_alignment = Alignment(horizontal="center", vertical="center")
        self.thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

    def set_title(self, title):
        self.worksheet.title = title[:31]

    def write_headers(self, headers):
        for col_idx, header_text in enumerate(headers, 1):
            cell = self.worksheet.cell(row=1, column=col_idx, value=header_text)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.thin_border

    def write_row(self, row_idx, values):
        for col_idx, value in enumerate(values, 1):
            cell = self.worksheet.cell(row=row_idx, column=col_idx, value=value)
            cell.font = self.data_font
            cell.alignment = self.data_alignment
            cell.border = self.thin_border

    def set_column_widths(self, widths):
        for col_idx, width in enumerate(widths, 1):
            self.worksheet.column_dimensions[get_column_letter(col_idx)].width = width

    def get_response(self, filename):
        output = io.BytesIO()
        self.workbook.save(output)
        output.seek(0)
        return output, filename

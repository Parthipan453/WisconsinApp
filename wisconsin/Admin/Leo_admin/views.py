## Leo's Code Start ##
from datetime import date
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Avg, Count, F, Prefetch, Q
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.contrib.contenttypes.models import ContentType
from django.db import transaction
from Students.Leo_Student.models import DoctoralCommittee, PhDStudent, PhDVideoEvidence
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import csv
from Faculty.leo.forms import (
    PreliminaryExaminationForm,
    PreliminaryExamEvaluationForm,
)

from .forms import (
    AcademicStandingForm,
    DoctoralCommitteeForm,
    PhDProgramForm,
    PhDStudentForm,
)

from Faculty.leo.models import PreliminaryExamEvaluation
from Faculty.models import FacultyProfile

from Students.models import (
    CommitteeMember,
    DoctoralCommittee,
    PhD,
    PhDStudent,
    PreliminaryExamination,
    StudentProfile,
)

from Students.Leo_Student.models import (
    DoctoralCommittee,
    PhD,
    PhDStudent,
)

from .models import AcademicStanding
from Staff.models import Notification
from notifications.utils import send_notification

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.contenttypes.models import ContentType
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string

from openpyxl import Workbook

from Faculty.models import FacultyProfile
from Students.Leo_Student.models import PhD, PhDStudent, PhDVideoEvidence

from .forms import PhDStudentForm

#######--ACADEMIC STANDINGS VIEWS START--########


# Academic Standings List
@login_required
def academic_standing_list(request):

    search = request.GET.get(
        "search",
        "",
    ).strip()

    sort = request.GET.get(
        "sort",
        "gpa",
    )

    order = request.GET.get(
        "order",
        "desc",
    )

    standings = AcademicStanding.objects.all()

    if search:
        standings = standings.filter(
            standing_name__icontains=search,
        )

    allowed_sort_fields = {
        "name": "standing_name",
        "gpa": "minimum_gpa",
        "description": "description",
    }

    sort_field = allowed_sort_fields.get(
        sort,
        "minimum_gpa",
    )

    if order == "desc":
        sort_field = f"-{sort_field}"

    standings = standings.order_by(
        sort_field,
    )

    paginator = Paginator(
        standings,
        10,
    )

    page_number = request.GET.get("page")

    standings = paginator.get_page(
        page_number,
    )

    if request.headers.get("x-requested-with") == "XMLHttpRequest":

        data = []

        for standing in standings:
            data.append(
                {
                    "id": standing.standing_id,
                    "name": standing.standing_name,
                    "gpa": str(standing.minimum_gpa),
                    "description": standing.description,
                    "deactivate_url": reverse(
                        "academic_standing_deactivate",
                        args=[standing.standing_id],
                    ),
                    "activate_url": reverse(
                        "academic_standing_activate",
                        args=[standing.standing_id],
                    ),
                    "is_active": standing.is_active,
                }
            )

        return JsonResponse(
            {
                "results": data,
                "start_index": standings.start_index(),
                "end_index": standings.end_index(),
                "total": paginator.count,
                "current_page": standings.number,
                "total_pages": paginator.num_pages,
                "has_previous": standings.has_previous(),
                "has_next": standings.has_next(),
                "previous_page": (
                    standings.previous_page_number()
                    if standings.has_previous()
                    else None
                ),
                "next_page": (
                    standings.next_page_number() if standings.has_next() else None
                ),
                "page_range": list(
                    standings.paginator.page_range,
                ),
                "sort": sort,
                "order": order,
            }
        )

    context = {
        "standings": standings,
        "form": AcademicStandingForm(),
        "search": search,
        "sort": sort,
        "order": order,
    }

    return render(
        request,
        "Leo_admin/academic_standing.html",
        context,
    )


# Create function for Academic Standings
@login_required
def academic_standing_create(request):
    if request.method == "POST":

        form = AcademicStandingForm(request.POST)

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Academic Standing created successfully.",
            )

            return redirect("academic_standing_list")

        standings = AcademicStanding.objects.all().order_by(
            "standing_name",
        )

        paginator = Paginator(
            standings,
            10,
        )

        page_number = request.GET.get("page")

        standings = paginator.get_page(page_number)

        context = {
            "standings": standings,
            "form": form,
            "show_modal": True,
        }

        return render(
            request,
            "Leo_admin/academic_standing.html",
            context,
        )

    return redirect("academic_standing_list")


# Academic standings update function
@login_required
def academic_standing_update(request, standing_id):
    standing = get_object_or_404(
        AcademicStanding,
        pk=standing_id,
    )

    if request.method == "POST":

        form = AcademicStandingForm(
            request.POST,
            instance=standing,
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Academic Standing updated successfully.",
            )

            return redirect("academic_standing_list")

        standings = AcademicStanding.objects.all().order_by(
            "standing_name",
        )

        paginator = Paginator(
            standings,
            10,
        )

        page_number = request.GET.get("page")

        standings = paginator.get_page(page_number)

        context = {
            "standings": standings,
            "form": form,
            "edit_standing": standing,
            "show_modal": True,
        }

        return render(
            request,
            "Leo_admin/academic_standing.html",
            context,
        )

    return redirect("academic_standing_list")


# Academic Standings Deactivate
@login_required
def academic_standing_deactivate(request, standing_id):
    standing = get_object_or_404(
        AcademicStanding,
        pk=standing_id,
    )

    if request.method != "POST":
        return redirect("academic_standing_list")

    standing.is_active = False
    standing.save()

    return JsonResponse(
        {
            "success": True,
            "message": "Academic Standing deactivated successfully.",
            "is_active": False,
        }
    )


# Academic Standings Activate
@login_required
def academic_standing_activate(request, standing_id):
    standing = get_object_or_404(
        AcademicStanding,
        pk=standing_id,
    )

    if request.method != "POST":
        return redirect("academic_standing_list")

    standing.is_active = True
    standing.save()

    return JsonResponse(
        {
            "success": True,
            "message": "Academic Standing activated successfully.",
            "is_active": True,
        }
    )


# Export function for Academic Standings
@login_required
def academic_standing_export(request):

    standings = AcademicStanding.objects.all().order_by("-minimum_gpa")

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="academic_standings.csv"'

    writer = csv.writer(response)

    writer.writerow(
        [
            "S.No",
            "Standing Name",
            "Minimum GPA",
            "Description",
            "Status",
        ]
    )

    for index, standing in enumerate(standings, start=1):
        writer.writerow(
            [
                index,
                standing.standing_name,
                standing.minimum_gpa,
                standing.description,
                "Active" if standing.is_active else "Inactive",
            ]
        )

    return response


# Print function for Academic Standings
@login_required
def academic_standing_print(request):

    standings = AcademicStanding.objects.all().order_by(
        "-minimum_gpa",
    )

    context = {
        "standings": standings,
    }

    return render(
        request,
        "Leo_admin/academic_standing_print.html",
        context,
    )


#######--ACADEMIC STANDINGS VIEWS END--########


#######--PHD PROGRAM VIEWS START--########
# PhD Program List


@login_required
def phd_program_list(request):

    search = request.GET.get(
        "search",
        "",
    ).strip()

    programs = PhD.objects.all()

    department_count = (
        PhD.objects.values(
            "department",
        )
        .distinct()
        .count()
    )

    average_credits = (
        PhD.objects.aggregate(
            Avg(
                "total_credits_required",
            )
        )["total_credits_required__avg"]
        or 0
    )

    average_duration = (
        PhD.objects.aggregate(
            Avg(
                "duration_years",
            )
        )["duration_years__avg"]
        or 0
    )

    if search:
        programs = programs.filter(
            Q(program_name__icontains=search)
            | Q(department__department_name__icontains=search)
            | Q(degree_type__icontains=search)
            | Q(total_credits_required__icontains=search)
            | Q(residency_requirement__icontains=search)
            | Q(duration_years__icontains=search)
            | Q(program_description__icontains=search)
            | Q(status__icontains=search)
        )

    programs = programs.order_by(
        "program_name",
    )

    paginator = Paginator(
        programs,
        10,
    )

    page_number = request.GET.get(
        "page",
    )

    programs = paginator.get_page(
        page_number,
    )

    if request.headers.get("x-requested-with") == "XMLHttpRequest":

        data = []

        for program in programs:

            data.append(
                {
                    "phd_program_id": program.phd_program_id,
                    "program_name": program.program_name,
                    "department": str(program.department),
                    "department_id": program.department.pk,
                    "degree_type": program.degree_type,
                    "total_credits_required": program.total_credits_required,
                    "residency_requirement": program.residency_requirement,
                    "duration_years": program.duration_years,
                    "program_description": program.program_description,
                    "status": program.status,
                    "deactivate_url": reverse(
                        "phd_program_deactivate",
                        args=[program.phd_program_id],
                    ),
                    "activate_url": reverse(
                        "phd_program_activate",
                        args=[program.phd_program_id],
                    ),
                }
            )

        return JsonResponse(
            {
                "results": data,
                "start_index": programs.start_index(),
                "end_index": programs.end_index(),
                "total": paginator.count,
                "current_page": programs.number,
                "total_pages": paginator.num_pages,
                "has_previous": programs.has_previous(),
                "has_next": programs.has_next(),
                "previous_page": (
                    programs.previous_page_number() if programs.has_previous() else None
                ),
                "next_page": (
                    programs.next_page_number() if programs.has_next() else None
                ),
                "page_range": list(
                    programs.paginator.page_range,
                ),
            }
        )

    context = {
        "programs": programs,
        "form": PhDProgramForm(),
        "search": search,
        "department_count": department_count,
        "average_credits": round(average_credits),
        "average_duration": round(average_duration, 1),
    }

    return render(
        request,
        "Leo_admin/phd_program.html",
        context,
    )


# PhD Program Create
@login_required
def phd_program_create(request):

    if request.method == "POST":

        form = PhDProgramForm(request.POST)

        print("=" * 60)
        print("POST DATA:", request.POST)
        print("VALID:", form.is_valid())
        print("ERRORS:", form.errors)
        print("=" * 60)

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "PhD Program created successfully.",
            )

            return redirect(
                "phd_program_list",
            )

        programs = PhD.objects.all().order_by(
            "program_name",
        )

        paginator = Paginator(
            programs,
            10,
        )

        page_number = request.GET.get(
            "page",
        )

        programs = paginator.get_page(
            page_number,
        )

        context = {
            "programs": programs,
            "form": form,
            "show_modal": True,
        }

        return render(
            request,
            "Leo_admin/phd_program.html",
            context,
        )

    return redirect(
        "phd_program_list",
    )


# PhD Program Update
@login_required
def phd_program_update(request, phd_program_id):

    program = get_object_or_404(
        PhD,
        pk=phd_program_id,
    )

    if request.method == "POST":

        form = PhDProgramForm(
            request.POST,
            instance=program,
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "PhD Program updated successfully.",
            )

            return redirect(
                "phd_program_list",
            )

        programs = PhD.objects.all().order_by(
            "program_name",
        )

        paginator = Paginator(
            programs,
            10,
        )

        page_number = request.GET.get(
            "page",
        )

        programs = paginator.get_page(
            page_number,
        )

        context = {
            "programs": programs,
            "form": form,
            "edit_program": program,
            "show_modal": True,
        }

        return render(
            request,
            "Leo_admin/phd_program.html",
            context,
        )

    return redirect(
        "phd_program_list",
    )


# PhD Program Deactivate
@login_required
def phd_program_deactivate(request, phd_program_id):

    program = get_object_or_404(
        PhD,
        pk=phd_program_id,
    )

    if request.method != "POST":
        return redirect(
            "phd_program_list",
        )

    program.status = "INACTIVE"

    program.save()

    return JsonResponse(
        {
            "success": True,
            "message": "PhD Program deactivated successfully.",
            "status": "INACTIVE",
        }
    )


# PhD Program Activate
@login_required
def phd_program_activate(request, phd_program_id):

    program = get_object_or_404(
        PhD,
        pk=phd_program_id,
    )

    if request.method != "POST":
        return redirect(
            "phd_program_list",
        )

    program.status = "ACTIVE"

    program.save()

    return JsonResponse(
        {
            "success": True,
            "message": "PhD Program activated successfully.",
            "status": "ACTIVE",
        }
    )


# PhD Program Export
@login_required
def phd_program_export(request):

    programs = PhD.objects.all().order_by(
        "program_name",
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "PhD Programs"

    headers = [
        "S.No",
        "Program Name",
        "Department",
        "Degree Type",
        "Credits Required",
        "Residency Requirement",
        "Duration (Years)",
        "Status",
    ]

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="C5050C",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    header_alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    for column, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=1,
            column=column,
        )

        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    for row, program in enumerate(
        programs,
        start=2,
    ):
        worksheet.cell(row=row, column=1).value = row - 1

        worksheet.cell(row=row, column=2).value = program.program_name
        worksheet.cell(row=row, column=3).value = str(program.department)
        worksheet.cell(row=row, column=4).value = program.degree_type
        worksheet.cell(row=row, column=5).value = program.total_credits_required
        worksheet.cell(row=row, column=6).value = program.residency_requirement
        worksheet.cell(row=row, column=7).value = program.duration_years
        worksheet.cell(row=row, column=8).value = program.status

        center_alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        for column in [1, 4, 5, 6, 7, 8]:
            worksheet.cell(
                row=row,
                column=column,
            ).alignment = center_alignment

    for column_cells in worksheet.columns:

        length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)

        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = (
            length + 4
        )

    worksheet.freeze_panes = "A2"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    response["Content-Disposition"] = 'attachment; filename="phd_programs.xlsx"'

    workbook.save(response)

    return response


# PhD Program Print
@login_required
def phd_program_print(request):

    programs = PhD.objects.all().order_by(
        "program_name",
    )

    context = {
        "programs": programs,
    }

    return render(
        request,
        "Leo_admin/phd_program_print.html",
        context,
    )


#######--PHD PROGRAM VIEWS END--########


#######--PHD STUDENTS VIEWS START--########
def _get_assignment_videos(phd_student):
    videos = PhDVideoEvidence.objects.filter(
        phd_student=phd_student,
        evidence_type__in=[
            "PROGRAM_ASSIGNMENT",
            "ADVISOR_ASSIGNMENT",
        ],
    ).order_by(
        "-uploaded_at",
        "-video_evidence_id",
    )

    return {
        "program_assignment_videos": videos.filter(evidence_type="PROGRAM_ASSIGNMENT"),
        "advisor_assignment_videos": videos.filter(evidence_type="ADVISOR_ASSIGNMENT"),
    }


def _delete_assignment_videos(phd_student, evidence_type):
    videos = PhDVideoEvidence.objects.filter(
        phd_student=phd_student,
        evidence_type=evidence_type,
    )

    for evidence in videos:
        if evidence.video:
            evidence.video.delete(save=False)

        evidence.delete()


def _save_assignment_video(
    phd_student,
    evidence_type,
    uploaded_file,
    user,
):
    if not uploaded_file:
        return None

    _delete_assignment_videos(
        phd_student,
        evidence_type,
    )

    content_type = ContentType.objects.get_for_model(PhDStudent)

    return PhDVideoEvidence.objects.create(
        phd_student=phd_student,
        evidence_type=evidence_type,
        target_content_type=content_type,
        target_object_id=phd_student.pk,
        video=uploaded_file,
        uploaded_by=user,
    )


@login_required
def phd_student_list(request):
    search = request.GET.get(
        "search",
        "",
    ).strip()

    status_filter = request.GET.get(
        "status",
        "",
    )

    program_filter = request.GET.get(
        "program",
        "",
    )

    advisor_filter = request.GET.get(
        "advisor",
        "",
    )

    phd_students = (
        PhDStudent.objects.select_related(
            "student",
            "student__user",
            "phd_program",
            "advisor",
            "advisor__user",
        )
        .all()
        .order_by(
            "student__user__first_name",
            "student__user__last_name",
        )
    )

    if search:
        phd_students = phd_students.filter(
            Q(student__user__first_name__icontains=search)
            | Q(student__user__last_name__icontains=search)
            | Q(student__user__email__icontains=search)
            | Q(phd_program__program_name__icontains=search)
            | Q(advisor__user__first_name__icontains=search)
            | Q(advisor__user__last_name__icontains=search)
            | Q(cohort_year__icontains=search)
        )

    if status_filter:
        phd_students = phd_students.filter(current_status=status_filter)

    if program_filter:
        phd_students = phd_students.filter(phd_program_id=program_filter)

    if advisor_filter:
        phd_students = phd_students.filter(advisor_id=advisor_filter)

    paginator = Paginator(
        phd_students,
        10,
    )

    page_number = request.GET.get("page")

    page_obj = paginator.get_page(page_number)

    active_programs = PhD.objects.filter(status="ACTIVE").order_by("program_name")

    context = {
        "page_obj": page_obj,
        "form": PhDStudentForm(),
        "search": search,
        "status_filter": status_filter,
        "program_filter": program_filter,
        "advisor_filter": advisor_filter,
        "programs": active_programs,
        "advisors": (
            FacultyProfile.objects.select_related("user").order_by("user__first_name")
        ),
        "has_active_phd_program": active_programs.exists(),
        "total_students": PhDStudent.objects.count(),
        "active_students": PhDStudent.objects.filter(current_status="ACTIVE").count(),
        "completed_students": PhDStudent.objects.filter(
            current_status="COMPLETED"
        ).count(),
        "on_hold_students": PhDStudent.objects.filter(current_status="ON_HOLD").count(),
        "withdrawn_students": PhDStudent.objects.filter(
            current_status="WITHDRAWN"
        ).count(),
    }

    return render(
        request,
        "Leo_admin/phd_student.html",
        context,
    )


def _create_phd_notification(user, title, message, notification_type="INFO", link=""):
    if not user:
        return

    notification = Notification.objects.create(
        user=user,
        title=title,
        message=message,
        notification_type=notification_type,
        link=link or "",
        is_read=False,
    )

    payload = {
        "id": notification.pk,
        "title": notification.title,
        "message": notification.message,
        "type": notification.notification_type,
        "link": notification.link,
        "is_read": notification.is_read,
        "created_at": notification.created_at.isoformat(),
    }

    transaction.on_commit(
        lambda user_id=user.pk, payload=payload: send_notification(
            user_id,
            payload,
            event="phd_assignment",
        )
    )


def _create_doctoral_committee_notification(
    user, title, message, notification_type="INFO", link=""
):
    if not user:
        return

    notification = Notification.objects.create(
        user=user,
        title=title,
        message=message,
        notification_type=notification_type,
        link=link or "",
        is_read=False,
    )

    payload = {
        "id": notification.pk,
        "title": notification.title,
        "message": notification.message,
        "type": notification.notification_type,
        "link": notification.link,
        "is_read": notification.is_read,
        "created_at": notification.created_at.isoformat(),
    }

    transaction.on_commit(
        lambda user_id=user.pk, payload=payload: send_notification(
            user_id,
            payload,
            event="doctoral_committee_assignment",
        )
    )


@login_required
def phd_student_create(request):
    if request.method != "POST":
        return redirect("phd_student_list")

    if not PhD.objects.filter(status="ACTIVE").exists():
        messages.error(
            request,
            "No active PhD program is available. Please create a PhD program before adding a PhD student.",
        )
        return redirect("phd_student_list")

    form = PhDStudentForm(
        request.POST,
        request.FILES,
    )

    if form.is_valid():
        selected_program = form.cleaned_data.get("phd_program")
        selected_advisor = form.cleaned_data.get("advisor")
        program_video = form.cleaned_data.get("program_assignment_video")
        advisor_video = form.cleaned_data.get("advisor_assignment_video")

        if not selected_program or selected_program.status != "ACTIVE":
            messages.error(
                request,
                "An active PhD program is required before creating the PhD student assignment.",
            )
            return redirect("phd_student_list")

        if advisor_video and not selected_advisor:
            messages.error(
                request,
                "A Research Advisor must be assigned before uploading Advisor Assignment Video.",
            )
            return redirect("phd_student_list")

        with transaction.atomic():
            phd_student = form.save()

            if program_video:
                _save_assignment_video(
                    phd_student=phd_student,
                    evidence_type="PROGRAM_ASSIGNMENT",
                    uploaded_file=program_video,
                    user=request.user,
                )

            if advisor_video:
                _save_assignment_video(
                    phd_student=phd_student,
                    evidence_type="ADVISOR_ASSIGNMENT",
                    uploaded_file=advisor_video,
                    user=request.user,
                )

            student_user = None
            advisor_user = None

            if phd_student.student_id and phd_student.student.user_id:
                student_user = phd_student.student.user

            if phd_student.advisor_id and phd_student.advisor.user_id:
                advisor_user = phd_student.advisor.user

            program_name = str(selected_program)

            if student_user and selected_program:
                _create_phd_notification(
                    user=student_user,
                    title="PhD Program Assigned",
                    message=f"Your PhD program has been assigned: {program_name}.",
                    notification_type="SUCCESS",
                )

            if student_user and selected_advisor:
                advisor_name = selected_advisor.user.get_full_name().strip()
                if not advisor_name:
                    advisor_name = selected_advisor.user.username

                _create_phd_notification(
                    user=student_user,
                    title="Research Advisor Assigned",
                    message=f"Your Research Advisor has been assigned: {advisor_name}.",
                    notification_type="SUCCESS",
                )

            if advisor_user:
                student_name = phd_student.student.user.get_full_name().strip()
                if not student_name:
                    student_name = phd_student.student.user.username

                _create_phd_notification(
                    user=advisor_user,
                    title="New PhD Student Assigned",
                    message=f"A new PhD student has been assigned to you: {student_name}.",
                    notification_type="INFO",
                )

        messages.success(
            request,
            "PhD Student created successfully.",
        )

    else:
        for field in form:
            for error in field.errors:
                messages.error(
                    request,
                    f"{field.label}: {error}",
                )

        for error in form.non_field_errors():
            messages.error(
                request,
                error,
            )

    return redirect("phd_student_list")


@login_required
def phd_student_update(request, phd_student_id):
    phd_student = get_object_or_404(
        PhDStudent.objects.select_related(
            "student",
            "student__user",
            "phd_program",
            "advisor",
            "advisor__user",
        ),
        pk=phd_student_id,
    )

    video_context = _get_assignment_videos(phd_student)

    if request.method == "GET" and request.GET.get("ajax") == "1":
        form = PhDStudentForm(instance=phd_student)

        context = {
            "form": form,
            "student": phd_student,
            **video_context,
        }

        html = render_to_string(
            "Leo_admin/partials/phd_student_form.html",
            context,
            request=request,
        )

        return JsonResponse(
            {
                "form": html,
            }
        )

    if request.method == "POST":
        old_program_id = phd_student.phd_program_id
        old_advisor_id = phd_student.advisor_id

        form = PhDStudentForm(
            request.POST,
            request.FILES,
            instance=phd_student,
        )

        if form.is_valid():
            selected_program = form.cleaned_data.get("phd_program")
            selected_advisor = form.cleaned_data.get("advisor")
            program_video = form.cleaned_data.get("program_assignment_video")
            advisor_video = form.cleaned_data.get("advisor_assignment_video")
            remove_program_video = form.cleaned_data.get(
                "remove_program_assignment_video"
            )
            remove_advisor_video = form.cleaned_data.get(
                "remove_advisor_assignment_video"
            )

            if not selected_program or selected_program.status != "ACTIVE":
                error_message = (
                    "An active PhD program is required before assigning a mentor."
                )

                if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                    return JsonResponse(
                        {
                            "success": False,
                            "errors": {"phd_program": [error_message]},
                        },
                        status=400,
                    )

                messages.error(
                    request,
                    error_message,
                )

                return redirect("phd_student_list")

            if advisor_video and not selected_advisor:
                error_message = "A Research Advisor must be assigned before uploading Advisor Assignment Video."

                if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                    return JsonResponse(
                        {
                            "success": False,
                            "errors": {"advisor": [error_message]},
                        },
                        status=400,
                    )

                messages.error(
                    request,
                    error_message,
                )

                return redirect("phd_student_list")

            with transaction.atomic():
                phd_student = form.save()

                if program_video:
                    _save_assignment_video(
                        phd_student=phd_student,
                        evidence_type="PROGRAM_ASSIGNMENT",
                        uploaded_file=program_video,
                        user=request.user,
                    )
                elif remove_program_video:
                    _delete_assignment_videos(
                        phd_student,
                        "PROGRAM_ASSIGNMENT",
                    )

                if advisor_video:
                    _save_assignment_video(
                        phd_student=phd_student,
                        evidence_type="ADVISOR_ASSIGNMENT",
                        uploaded_file=advisor_video,
                        user=request.user,
                    )
                elif remove_advisor_video:
                    _delete_assignment_videos(
                        phd_student,
                        "ADVISOR_ASSIGNMENT",
                    )

                program_changed = old_program_id != phd_student.phd_program_id
                advisor_changed = old_advisor_id != phd_student.advisor_id

                student_user = None

                if phd_student.student_id and phd_student.student.user_id:
                    student_user = phd_student.student.user

                if program_changed and student_user and selected_program:
                    program_name = str(selected_program)

                    _create_phd_notification(
                        user=student_user,
                        title="PhD Program Changed",
                        message=f"Your PhD program has been changed to: {program_name}.",
                        notification_type="INFO",
                    )

                if advisor_changed and student_user:
                    if selected_advisor:
                        advisor_name = selected_advisor.user.get_full_name().strip()

                        if not advisor_name:
                            advisor_name = selected_advisor.user.username

                        _create_phd_notification(
                            user=student_user,
                            title="Research Advisor Changed",
                            message=f"Your Research Advisor has been changed to: {advisor_name}.",
                            notification_type="INFO",
                        )
                    else:
                        _create_phd_notification(
                            user=student_user,
                            title="Research Advisor Removed",
                            message="Your Research Advisor assignment has been removed.",
                            notification_type="WARNING",
                        )

                if advisor_changed and selected_advisor:
                    advisor_user = None

                    if selected_advisor.user_id:
                        advisor_user = selected_advisor.user

                    if advisor_user:
                        student_name = phd_student.student.user.get_full_name().strip()

                        if not student_name:
                            student_name = phd_student.student.user.username

                        _create_phd_notification(
                            user=advisor_user,
                            title="PhD Student Assigned",
                            message=f"A PhD student has been assigned to you: {student_name}.",
                            notification_type="INFO",
                        )

            if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                return JsonResponse(
                    {
                        "success": True,
                        "message": "PhD Student assignment updated successfully.",
                    }
                )

            messages.success(
                request,
                "PhD Student assignment updated successfully.",
            )

            return redirect("phd_student_list")

        if request.headers.get("X-Requested-With") == "XMLHttpRequest":
            video_context = _get_assignment_videos(phd_student)

            context = {
                "form": form,
                "student": phd_student,
                **video_context,
            }

            html = render_to_string(
                "Leo_admin/partials/phd_student_form.html",
                context,
                request=request,
            )

            return JsonResponse(
                {
                    "success": False,
                    "form": html,
                },
                status=400,
            )

        for field in form:
            for error in field.errors:
                messages.error(
                    request,
                    f"{field.label}: {error}",
                )

        for error in form.non_field_errors():
            messages.error(
                request,
                error,
            )

    return redirect("phd_student_list")


@login_required
def phd_student_activate(request, phd_student_id):
    phd_student = get_object_or_404(
        PhDStudent,
        pk=phd_student_id,
    )

    phd_student.current_status = "ACTIVE"

    phd_student.save(
        update_fields=[
            "current_status",
        ]
    )

    messages.success(
        request,
        "PhD Student activated successfully.",
    )

    return redirect("phd_student_list")


@login_required
def phd_student_deactivate(request, phd_student_id):
    phd_student = get_object_or_404(
        PhDStudent,
        pk=phd_student_id,
    )

    phd_student.current_status = "ON_HOLD"

    phd_student.save(
        update_fields=[
            "current_status",
        ]
    )

    messages.success(
        request,
        "PhD Student placed On Hold successfully.",
    )

    return redirect("phd_student_list")


@login_required
def phd_student_export(request):
    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "PhD Students"

    headers = [
        "Student",
        "Program",
        "Advisor",
        "Admission Date",
        "Cohort Year",
        "Status",
        "Expected Graduation",
    ]

    for column, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=1,
            column=column,
        ).value = header

    phd_students = PhDStudent.objects.select_related(
        "student",
        "student__user",
        "phd_program",
        "advisor",
        "advisor__user",
    ).order_by("student__user__first_name")

    row = 2

    for student in phd_students:
        worksheet.cell(
            row=row,
            column=1,
        ).value = str(student.student)

        worksheet.cell(
            row=row,
            column=2,
        ).value = student.phd_program.program_name

        worksheet.cell(
            row=row,
            column=3,
        ).value = (
            str(student.advisor) if student.advisor else "-"
        )

        worksheet.cell(
            row=row,
            column=4,
        ).value = student.admission_date.strftime("%d-%m-%Y")

        worksheet.cell(
            row=row,
            column=5,
        ).value = student.cohort_year

        worksheet.cell(
            row=row,
            column=6,
        ).value = student.get_current_status_display()

        worksheet.cell(
            row=row,
            column=7,
        ).value = (
            student.expected_graduation_date.strftime("%d-%m-%Y")
            if student.expected_graduation_date
            else "-"
        )

        row += 1

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = 'attachment; filename="phd_students.xlsx"'

    workbook.save(response)

    return response


@login_required
def phd_student_print(request):
    phd_students = PhDStudent.objects.select_related(
        "student",
        "student__user",
        "phd_program",
        "advisor",
        "advisor__user",
    ).order_by("student__user__first_name")

    context = {
        "phd_students": phd_students,
    }

    return render(
        request,
        "Leo_admin/phd_student_print.html",
        context,
    )


from django.http import JsonResponse

@login_required
def phd_student_view(request, pk):
    student = get_object_or_404(
        PhDStudent.objects.select_related(
            "student__user",
            "phd_program",
            "advisor__user",
        ),
        pk=pk,
    )

    video_context = _get_assignment_videos(student)

    if request.GET.get("ajax"):
        return JsonResponse(
            {
                "success": True,
                "student": {
                    "id": student.pk,
                    "first_name": student.student.user.first_name,
                    "last_name": student.student.user.last_name,
                    "email": student.student.user.email,
                    "program": (
                        student.phd_program.program_name
                        if student.phd_program
                        else ""
                    ),
                    "advisor": (
                        f"{student.advisor.user.first_name} "
                        f"{student.advisor.user.last_name}"
                        if student.advisor
                        else ""
                    ),
                    "admission_date": (
                        student.admission_date.strftime("%b %d, %Y")
                        if student.admission_date
                        else ""
                    ),
                    "cohort_year": student.cohort_year or "",
                    "expected_graduation_date": (
                        student.expected_graduation_date.strftime("%b %d, %Y")
                        if student.expected_graduation_date
                        else ""
                    ),
                    "status": student.current_status,
                },
                "videos": {
                    "program": [
                        {
                            "url": evidence.video.url,
                            "name": evidence.video.name.split("/")[-1],
                            "uploaded_at": evidence.uploaded_at.strftime(
                                "%b %d, %Y %H:%M"
                            ),
                        }
                        for evidence in video_context.get(
                            "program_assignment_videos",
                            [],
                        )
                    ],
                    "advisor": [
                        {
                            "url": evidence.video.url,
                            "name": evidence.video.name.split("/")[-1],
                            "uploaded_at": evidence.uploaded_at.strftime(
                                "%b %d, %Y %H:%M"
                            ),
                        }
                        for evidence in video_context.get(
                            "advisor_assignment_videos",
                            [],
                        )
                    ],
                },
            }
        )

    return redirect("phd_student_list")

#######--PHD STUDENTS VIEWS END--########
#######--DOCTORAL COMMITTEE VIEWS START--########


def doctoral_committee_form_errors(form):
    errors = {}
    messages = []

    for field, field_errors in form.errors.as_data().items():
        field_messages = []

        for error in field_errors:
            for message in error.messages:
                field_messages.append(str(message))
                messages.append(str(message))

        if field_messages:
            errors[field] = field_messages

    if "__all__" in errors:
        errors["non_field_errors"] = errors.pop("__all__")

    main_message = messages[0] if messages else "Please correct the errors below."

    return {
        "success": False,
        "message": main_message,
        "errors": errors,
    }


def _committee_video_queryset(phd_student):
    content_type = ContentType.objects.get_for_model(DoctoralCommittee)

    committee = DoctoralCommittee.objects.filter(phd_student=phd_student).first()

    if not committee:
        return PhDVideoEvidence.objects.none()

    return PhDVideoEvidence.objects.filter(
        phd_student=phd_student,
        evidence_type="COMMITTEE_ASSIGNMENT",
        target_content_type=content_type,
        target_object_id=committee.pk,
    )


@login_required
def doctoral_committee_list(request):

    search = request.GET.get(
        "search",
        "",
    ).strip()

    approval_status = (
        request.GET.get(
            "approval_status",
            "",
        )
        .strip()
        .upper()
    )
    committees = DoctoralCommittee.objects.select_related(
        "phd_student__student__user",
        "phd_student__phd_program",
        "phd_student__advisor__user",
        "chair_faculty__user",
    )

    if search:
        search_terms = search.split()

        for term in search_terms:
            committees = committees.filter(
                Q(phd_student__student__user__first_name__icontains=term)
                | Q(phd_student__student__user__last_name__icontains=term)
                | Q(phd_student__student__preferred_name__icontains=term)
                | Q(phd_student__student__student_number__icontains=term)
                | Q(chair_faculty__user__first_name__icontains=term)
                | Q(chair_faculty__user__last_name__icontains=term)
                | Q(chair_faculty__preferred_name__icontains=term)
                | Q(chair_faculty__employee_id__icontains=term)
                | Q(approval_status__icontains=term)
            )

    if approval_status:
        committees = committees.filter(
            approval_status=approval_status,
        )
    committees = committees.order_by(
        "-committee_id",
    )

    paginator = Paginator(
        committees,
        10,
    )

    page_number = request.GET.get(
        "page",
    )

    committees_page = paginator.get_page(
        page_number,
    )

    if request.headers.get("x-requested-with") == "XMLHttpRequest":

        data = []

        for committee in committees_page:

            committee_videos = _committee_video_queryset(committee.phd_student)

            data.append(
                {
                    "committee_id": committee.committee_id,
                    "phd_student": str(committee.phd_student),
                    "phd_student_id": committee.phd_student.pk,
                    "student_number": (committee.phd_student.student.student_number),
                    "preferred_name": (committee.phd_student.student.preferred_name),
                    "program": str(committee.phd_student.phd_program),
                    "advisor": (
                        str(committee.phd_student.advisor)
                        if committee.phd_student.advisor
                        else "Not Assigned"
                    ),
                    "admission_date": (
                        committee.phd_student.admission_date.strftime("%Y-%m-%d")
                        if committee.phd_student.admission_date
                        else ""
                    ),
                    "cohort_year": (committee.phd_student.cohort_year),
                    "expected_graduation": (
                        committee.phd_student.expected_graduation_date.strftime(
                            "%Y-%m-%d"
                        )
                        if committee.phd_student.expected_graduation_date
                        else ""
                    ),
                    "current_status": (
                        committee.phd_student.get_current_status_display()
                    ),
                    "chair_faculty": str(committee.chair_faculty),
                    "chair_faculty_id": (committee.chair_faculty.pk),
                    "formation_date": (
                        committee.formation_date.strftime("%Y-%m-%d")
                        if committee.formation_date
                        else ""
                    ),
                    "formation_date_display": (
                        committee.formation_date.strftime("%d %b %Y")
                        if committee.formation_date
                        else ""
                    ),
                    "approval_status": committee.approval_status,
                    "approval_status_display": (
                        committee.get_approval_status_display()
                    ),
                    "video_evidence": [
                        {
                            "id": video.video_evidence_id,
                            "url": video.video.url if video.video else "",
                            "filename": (
                                video.video.name.split("/")[-1] if video.video else ""
                            ),
                            "uploaded_at": (
                                video.uploaded_at.strftime("%Y-%m-%d %H:%M")
                                if video.uploaded_at
                                else ""
                            ),
                        }
                        for video in committee_videos
                        if video.video
                    ],
                    "has_video_evidence": committee_videos.exists(),
                }
            )

        return JsonResponse(
            {
                "success": True,
                "results": data,
                "start_index": committees_page.start_index(),
                "end_index": committees_page.end_index(),
                "total": paginator.count,
                "current_page": committees_page.number,
                "total_pages": paginator.num_pages,
                "has_previous": committees_page.has_previous(),
                "has_next": committees_page.has_next(),
                "previous_page": (
                    committees_page.previous_page_number()
                    if committees_page.has_previous()
                    else None
                ),
                "next_page": (
                    committees_page.next_page_number()
                    if committees_page.has_next()
                    else None
                ),
                "page_range": list(committees_page.paginator.page_range),
            }
        )

    current_year = date.today().year

    active_programs = PhD.objects.filter(
        status="ACTIVE",
    )

    phd_students = PhDStudent.objects.all()

    students_with_active_program = phd_students.filter(
        phd_program__status="ACTIVE",
        current_status="ACTIVE",
    )

    students_with_mentor = students_with_active_program.filter(
        advisor__isnull=False,
    )

    eligible_students = students_with_mentor.filter(
        doctoral_committee__isnull=True,
    )

    total_committees = DoctoralCommittee.objects.count()

    approved_committees = DoctoralCommittee.objects.filter(
        approval_status="APPROVED",
    ).count()

    pending_committees = DoctoralCommittee.objects.filter(
        approval_status="PENDING",
    ).count()

    rejected_committees = DoctoralCommittee.objects.filter(
        approval_status="REJECTED",
    ).count()

    chair_faculty_count = (
        DoctoralCommittee.objects.exclude(chair_faculty=None)
        .values("chair_faculty")
        .distinct()
        .count()
    )

    advisor_count = (
        PhDStudent.objects.exclude(advisor=None).values("advisor").distinct().count()
    )

    committees_this_year = DoctoralCommittee.objects.filter(
        formation_date__year=current_year,
    ).count()

    approval_rate = (
        round(
            (approved_committees / total_committees) * 100,
            1,
        )
        if total_committees
        else 0
    )

    context = {
        "committees": committees_page,
        "form": DoctoralCommitteeForm(),
        "search": search,
        "total_committees": total_committees,
        "approved_committees": approved_committees,
        "pending_committees": pending_committees,
        "rejected_committees": rejected_committees,
        "chair_faculty_count": chair_faculty_count,
        "advisor_count": advisor_count,
        "committees_this_year": committees_this_year,
        "approval_rate": approval_rate,
        "current_year": current_year,
        "has_active_phd_program": active_programs.exists(),
        "has_phd_student": phd_students.exists(),
        "has_active_phd_student": (students_with_active_program.exists()),
        "has_mentor_assigned_student": (students_with_mentor.exists()),
        "has_eligible_committee_student": (eligible_students.exists()),
        "approval_status": approval_status,
    }

    return render(
        request,
        "Leo_admin/doctoral_committee.html",
        context,
    )


@login_required
def doctoral_committee_create(request):

    if request.method != "POST":
        return JsonResponse(
            {
                "success": False,
                "message": ("Invalid request method. " "Please submit the form."),
                "errors": {},
            },
            status=400,
        )

    if not PhD.objects.filter(status="ACTIVE").exists():
        return JsonResponse(
            {
                "success": False,
                "message": (
                    "No active PhD program is available. "
                    "Please create a PhD program before creating "
                    "a doctoral committee."
                ),
                "errors": {},
            },
            status=400,
        )

    if not PhDStudent.objects.filter(
        phd_program__status="ACTIVE",
        current_status="ACTIVE",
    ).exists():
        return JsonResponse(
            {
                "success": False,
                "message": (
                    "No active PhD student is available. "
                    "Please create a PhD student before creating "
                    "a doctoral committee."
                ),
                "errors": {},
            },
            status=400,
        )

    if not PhDStudent.objects.filter(
        phd_program__status="ACTIVE",
        current_status="ACTIVE",
        advisor__isnull=False,
        doctoral_committee__isnull=True,
    ).exists():
        return JsonResponse(
            {
                "success": False,
                "message": (
                    "No eligible PhD student is available. "
                    "Please assign a research advisor to a PhD "
                    "student before creating a doctoral committee."
                ),
                "errors": {},
            },
            status=400,
        )

    form = DoctoralCommitteeForm(
        request.POST,
        request.FILES,
    )

    if not form.is_valid():
        return JsonResponse(
            doctoral_committee_form_errors(form),
            status=400,
        )

    selected_student = form.cleaned_data.get("phd_student")
    selected_chair = form.cleaned_data.get("chair_faculty")
    uploaded_video = form.cleaned_data.get("committee_assignment_video")

    if not selected_student:
        return JsonResponse(
            {
                "success": False,
                "message": "Please select a PhD student.",
                "errors": {"phd_student": ["Please select a PhD student."]},
            },
            status=400,
        )

    if not selected_student.phd_program:
        return JsonResponse(
            {
                "success": False,
                "message": ("The PhD student must be assigned " "to a PhD program."),
                "errors": {
                    "phd_student": [
                        "The PhD student must be assigned " "to a PhD program."
                    ]
                },
            },
            status=400,
        )

    if selected_student.phd_program.status != "ACTIVE":
        return JsonResponse(
            {
                "success": False,
                "message": ("The PhD student's PhD program " "must be active."),
                "errors": {
                    "phd_student": ["The PhD student's PhD program " "must be active."]
                },
            },
            status=400,
        )

    if selected_student.current_status != "ACTIVE":
        return JsonResponse(
            {
                "success": False,
                "message": (
                    "The PhD student must be active "
                    "before creating a doctoral committee."
                ),
                "errors": {
                    "phd_student": [
                        "The PhD student must be active "
                        "before creating a doctoral committee."
                    ]
                },
            },
            status=400,
        )

    if not selected_student.advisor:
        return JsonResponse(
            {
                "success": False,
                "message": (
                    "Please assign a research advisor to the "
                    "PhD student before creating the doctoral committee."
                ),
                "errors": {
                    "phd_student": [
                        "Please assign a research advisor to the "
                        "PhD student before creating the doctoral committee."
                    ]
                },
            },
            status=400,
        )

    if selected_chair and (selected_student.advisor_id == selected_chair.pk):
        message = (
            "The selected Chair Faculty cannot be the "
            "PhD student's research advisor. Please select "
            "a different Chair Faculty."
        )

        return JsonResponse(
            {
                "success": False,
                "message": message,
                "errors": {"chair_faculty": [message]},
            },
            status=400,
        )

    try:
        with transaction.atomic():
            committee = form.save(commit=True)

            if uploaded_video:
                content_type = ContentType.objects.get_for_model(DoctoralCommittee)

                PhDVideoEvidence.objects.create(
                    phd_student=committee.phd_student,
                    evidence_type="COMMITTEE_ASSIGNMENT",
                    target_content_type=content_type,
                    target_object_id=committee.pk,
                    video=uploaded_video,
                    uploaded_by=request.user,
                    is_full_crud=True,
                )

            student_user = (
                committee.phd_student.student.user
                if committee.phd_student.student_id
                and committee.phd_student.student.user_id
                else None
            )

            advisor_user = (
                committee.phd_student.advisor.user
                if committee.phd_student.advisor_id
                and committee.phd_student.advisor.user_id
                else None
            )

            chair_user = (
                committee.chair_faculty.user
                if committee.chair_faculty_id and committee.chair_faculty.user_id
                else None
            )

            student_name = (
                committee.phd_student.student.user.get_full_name().strip()
                if student_user
                else str(committee.phd_student)
            )

            if not student_name:
                student_name = (
                    committee.phd_student.student.user.username
                    if student_user
                    else str(committee.phd_student)
                )

            chair_name = (
                committee.chair_faculty.user.get_full_name().strip()
                if chair_user
                else str(committee.chair_faculty)
            )

            if not chair_name and chair_user:
                chair_name = committee.chair_faculty.user.username

            if student_user:
                _create_doctoral_committee_notification(
                    user=student_user,
                    title="Doctoral Committee Assigned",
                    message=(
                        f"Your Doctoral Committee has been assigned. "
                        f"Chair: {chair_name}."
                    ),
                    notification_type="SUCCESS",
                )

            if advisor_user:
                _create_doctoral_committee_notification(
                    user=advisor_user,
                    title="Doctoral Committee Assigned",
                    message=(
                        f"A Doctoral Committee has been assigned "
                        f"for your PhD advisee: {student_name}."
                    ),
                    notification_type="INFO",
                )

            if chair_user:
                _create_doctoral_committee_notification(
                    user=chair_user,
                    title="Doctoral Committee Assigned",
                    message=(
                        f"You have been assigned as Chair for "
                        f"PhD student: {student_name}."
                    ),
                    notification_type="INFO",
                )

    except ValidationError as error:
        error_messages = error.messages

        return JsonResponse(
            {
                "success": False,
                "message": (
                    error_messages[0]
                    if error_messages
                    else "Unable to create the doctoral committee."
                ),
                "errors": {"non_field_errors": error_messages},
            },
            status=400,
        )

    except Exception as error:
        return JsonResponse(
            {
                "success": False,
                "message": ("Unable to create the doctoral committee."),
                "errors": {"non_field_errors": [str(error)]},
            },
            status=500,
        )

    return JsonResponse(
        {
            "success": True,
            "message": ("Doctoral Committee created successfully."),
            "committee_id": committee.committee_id,
            "phd_student": str(committee.phd_student),
            "phd_student_id": (committee.phd_student.pk),
            "student_number": (committee.phd_student.student.student_number),
            "phd_program": str(committee.phd_student.phd_program),
            "advisor": (
                str(committee.phd_student.advisor)
                if committee.phd_student.advisor
                else "Not Assigned"
            ),
            "chair_faculty": str(committee.chair_faculty),
            "chair_faculty_id": (committee.chair_faculty.pk),
            "formation_date": (
                committee.formation_date.strftime("%Y-%m-%d")
                if committee.formation_date
                else ""
            ),
            "formation_date_display": (
                committee.formation_date.strftime("%d %b %Y")
                if committee.formation_date
                else ""
            ),
            "approval_status": (committee.approval_status),
            "approval_status_display": (committee.get_approval_status_display()),
            "has_video_evidence": bool(uploaded_video),
        }
    )


@login_required
def doctoral_committee_update(
    request,
    committee_id,
):

    committee = get_object_or_404(
        DoctoralCommittee,
        pk=committee_id,
    )

    if request.method == "GET":

        form = DoctoralCommitteeForm(
            instance=committee,
        )

        videos = PhDVideoEvidence.objects.filter(
            phd_student=committee.phd_student,
            evidence_type="COMMITTEE_ASSIGNMENT",
            target_content_type=ContentType.objects.get_for_model(DoctoralCommittee),
            target_object_id=committee.pk,
        )

        html = render_to_string(
            "Leo_admin/partials/committee_form.html",
            {
                "form": form,
                "is_edit": True,
                "committee": committee,
                "committee_videos": videos,
            },
            request=request,
        )

        return JsonResponse(
            {
                "success": True,
                "html": html,
            }
        )

    if request.method != "POST":
        return JsonResponse(
            {
                "success": False,
                "message": ("Invalid request method. " "Please submit the form."),
                "errors": {},
            },
            status=400,
        )

    old_student_id = committee.phd_student_id
    old_chair_id = committee.chair_faculty_id

    form = DoctoralCommitteeForm(
        request.POST,
        request.FILES,
        instance=committee,
    )

    if not form.is_valid():
        return JsonResponse(
            doctoral_committee_form_errors(form),
            status=400,
        )

    selected_student = form.cleaned_data.get("phd_student")
    selected_chair = form.cleaned_data.get("chair_faculty")
    uploaded_video = form.cleaned_data.get("committee_assignment_video")
    remove_video = form.cleaned_data.get("remove_committee_assignment_video")

    if not selected_student:
        return JsonResponse(
            {
                "success": False,
                "message": "Please select a PhD student.",
                "errors": {"phd_student": ["Please select a PhD student."]},
            },
            status=400,
        )

    if not selected_student.phd_program:
        return JsonResponse(
            {
                "success": False,
                "message": ("The PhD student must be assigned " "to a PhD program."),
                "errors": {
                    "phd_student": [
                        "The PhD student must be assigned " "to a PhD program."
                    ]
                },
            },
            status=400,
        )

    if selected_student.phd_program.status != "ACTIVE":
        return JsonResponse(
            {
                "success": False,
                "message": ("The PhD student's PhD program " "must be active."),
                "errors": {
                    "phd_student": ["The PhD student's PhD program " "must be active."]
                },
            },
            status=400,
        )

    if selected_student.current_status != "ACTIVE":
        return JsonResponse(
            {
                "success": False,
                "message": (
                    "The PhD student must be active "
                    "before updating the doctoral committee."
                ),
                "errors": {
                    "phd_student": [
                        "The PhD student must be active "
                        "before updating the doctoral committee."
                    ]
                },
            },
            status=400,
        )

    if not selected_student.advisor:
        return JsonResponse(
            {
                "success": False,
                "message": (
                    "Please assign a research advisor to the "
                    "PhD student before updating the doctoral committee."
                ),
                "errors": {
                    "phd_student": [
                        "Please assign a research advisor to the "
                        "PhD student before updating the doctoral committee."
                    ]
                },
            },
            status=400,
        )

    if selected_chair and (selected_student.advisor_id == selected_chair.pk):
        message = (
            "The selected Chair Faculty cannot be the "
            "PhD student's research advisor. Please select "
            "a different Chair Faculty."
        )

        return JsonResponse(
            {
                "success": False,
                "message": message,
                "errors": {"chair_faculty": [message]},
            },
            status=400,
        )

    content_type = ContentType.objects.get_for_model(DoctoralCommittee)

    existing_videos = PhDVideoEvidence.objects.filter(
        phd_student=committee.phd_student,
        evidence_type="COMMITTEE_ASSIGNMENT",
        target_content_type=content_type,
        target_object_id=committee.pk,
    )

    try:
        with transaction.atomic():
            committee = form.save(commit=True)

            if remove_video:
                for evidence in existing_videos:
                    if evidence.video:
                        evidence.video.delete(save=False)
                    evidence.delete()

                existing_videos = PhDVideoEvidence.objects.none()

            if uploaded_video:
                for evidence in existing_videos:
                    if evidence.video:
                        evidence.video.delete(save=False)
                    evidence.delete()

                PhDVideoEvidence.objects.create(
                    phd_student=committee.phd_student,
                    evidence_type="COMMITTEE_ASSIGNMENT",
                    target_content_type=content_type,
                    target_object_id=committee.pk,
                    video=uploaded_video,
                    uploaded_by=request.user,
                    is_full_crud=True,
                )

            student_changed = old_student_id != committee.phd_student_id
            chair_changed = old_chair_id != committee.chair_faculty_id

            student_user = (
                committee.phd_student.student.user
                if committee.phd_student.student_id
                and committee.phd_student.student.user_id
                else None
            )

            advisor_user = (
                committee.phd_student.advisor.user
                if committee.phd_student.advisor_id
                and committee.phd_student.advisor.user_id
                else None
            )

            chair_user = (
                committee.chair_faculty.user
                if committee.chair_faculty_id and committee.chair_faculty.user_id
                else None
            )

            student_name = (
                committee.phd_student.student.user.get_full_name().strip()
                if student_user
                else str(committee.phd_student)
            )

            if not student_name:
                student_name = (
                    committee.phd_student.student.user.username
                    if student_user
                    else str(committee.phd_student)
                )

            chair_name = (
                committee.chair_faculty.user.get_full_name().strip()
                if chair_user
                else str(committee.chair_faculty)
            )

            if not chair_name and chair_user:
                chair_name = committee.chair_faculty.user.username

            if student_changed:
                if student_user:
                    _create_doctoral_committee_notification(
                        user=student_user,
                        title="Doctoral Committee Assigned",
                        message=(
                            f"Your Doctoral Committee has been assigned. "
                            f"Chair: {chair_name}."
                        ),
                        notification_type="SUCCESS",
                    )

                if advisor_user:
                    _create_doctoral_committee_notification(
                        user=advisor_user,
                        title="Doctoral Committee Assigned",
                        message=(
                            f"A Doctoral Committee has been assigned "
                            f"for your PhD student: {student_name}."
                        ),
                        notification_type="INFO",
                    )

                if chair_user:
                    _create_doctoral_committee_notification(
                        user=chair_user,
                        title="Doctoral Committee Assigned",
                        message=(
                            f"You have been assigned as Chair for "
                            f"PhD student: {student_name}."
                        ),
                        notification_type="INFO",
                    )

            elif chair_changed:
                if student_user:
                    _create_doctoral_committee_notification(
                        user=student_user,
                        title="Doctoral Committee Updated",
                        message=(
                            f"Your Doctoral Committee has been updated. "
                            f"New Chair: {chair_name}."
                        ),
                        notification_type="INFO",
                    )

                if advisor_user:
                    _create_doctoral_committee_notification(
                        user=advisor_user,
                        title="Doctoral Committee Updated",
                        message=(
                            f"The Doctoral Committee for your PhD student "
                            f"{student_name} has been updated. "
                            f"New Chair: {chair_name}."
                        ),
                        notification_type="INFO",
                    )

                if chair_user:
                    _create_doctoral_committee_notification(
                        user=chair_user,
                        title="Doctoral Committee Assigned",
                        message=(
                            f"You have been assigned as Chair for "
                            f"PhD student: {student_name}."
                        ),
                        notification_type="INFO",
                    )

    except ValidationError as error:
        error_messages = error.messages

        return JsonResponse(
            {
                "success": False,
                "message": (
                    error_messages[0]
                    if error_messages
                    else "Unable to update the doctoral committee."
                ),
                "errors": {"non_field_errors": error_messages},
            },
            status=400,
        )

    except Exception as error:
        return JsonResponse(
            {
                "success": False,
                "message": ("Unable to update the doctoral committee."),
                "errors": {"non_field_errors": [str(error)]},
            },
            status=500,
        )

    committee.refresh_from_db()

    current_video = PhDVideoEvidence.objects.filter(
        phd_student=committee.phd_student,
        evidence_type="COMMITTEE_ASSIGNMENT",
        target_content_type=content_type,
        target_object_id=committee.pk,
    ).first()

    return JsonResponse(
        {
            "success": True,
            "message": ("Doctoral Committee updated successfully."),
            "committee_id": committee.committee_id,
            "phd_student": str(committee.phd_student),
            "phd_student_id": (committee.phd_student.pk),
            "student_number": (committee.phd_student.student.student_number),
            "phd_program": str(committee.phd_student.phd_program),
            "advisor": (
                str(committee.phd_student.advisor)
                if committee.phd_student.advisor
                else "Not Assigned"
            ),
            "chair_faculty": str(committee.chair_faculty),
            "chair_faculty_id": (committee.chair_faculty.pk),
            "formation_date": (
                committee.formation_date.strftime("%Y-%m-%d")
                if committee.formation_date
                else ""
            ),
            "formation_date_display": (
                committee.formation_date.strftime("%d %b %Y")
                if committee.formation_date
                else ""
            ),
            "approval_status": (committee.approval_status),
            "approval_status_display": (committee.get_approval_status_display()),
            "video_evidence": (
                {
                    "id": current_video.video_evidence_id,
                    "url": (current_video.video.url if current_video.video else ""),
                    "filename": (
                        current_video.video.name.split("/")[-1]
                        if current_video.video
                        else ""
                    ),
                }
                if current_video
                else None
            ),
            "has_video_evidence": bool(current_video),
        }
    )


@login_required
def doctoral_committee_delete(
    request,
    committee_id,
):

    if request.method != "POST":
        return JsonResponse(
            {
                "success": False,
                "message": "Invalid request method.",
                "errors": {},
            },
            status=400,
        )

    committee = get_object_or_404(
        DoctoralCommittee,
        pk=committee_id,
    )

    content_type = ContentType.objects.get_for_model(DoctoralCommittee)

    try:
        with transaction.atomic():

            videos = PhDVideoEvidence.objects.filter(
                phd_student=committee.phd_student,
                evidence_type="COMMITTEE_ASSIGNMENT",
                target_content_type=content_type,
                target_object_id=committee.pk,
            )

            for evidence in videos:
                if evidence.video:
                    evidence.video.delete(save=False)
                evidence.delete()

            committee.delete()

    except Exception as error:
        return JsonResponse(
            {
                "success": False,
                "message": ("Unable to delete the doctoral committee."),
                "errors": {"non_field_errors": [str(error)]},
            },
            status=500,
        )

    return JsonResponse(
        {
            "success": True,
            "message": (
                "Doctoral Committee and its video evidence "
                "were deleted successfully."
            ),
        }
    )


@login_required
def doctoral_committee_get_form(request):

    committee_id = request.GET.get("committee_id")

    committee = None

    if committee_id:
        committee = get_object_or_404(
            DoctoralCommittee,
            pk=committee_id,
        )

        form = DoctoralCommitteeForm(instance=committee)

        is_edit = True

        content_type = ContentType.objects.get_for_model(DoctoralCommittee)

        committee_videos = PhDVideoEvidence.objects.filter(
            phd_student=committee.phd_student,
            evidence_type="COMMITTEE_ASSIGNMENT",
            target_content_type=content_type,
            target_object_id=committee.pk,
        )

    else:
        form = DoctoralCommitteeForm()

        is_edit = False

        committee_videos = PhDVideoEvidence.objects.none()

    try:
        html = render_to_string(
            "Leo_admin/partials/committee_form.html",
            {
                "form": form,
                "is_edit": is_edit,
                "committee": committee,
                "committee_videos": committee_videos,
            },
            request=request,
        )

    except Exception as error:
        return JsonResponse(
            {
                "success": False,
                "message": ("Unable to load the doctoral committee form."),
                "errors": {"non_field_errors": [str(error)]},
            },
            status=500,
        )

    return JsonResponse(
        {
            "success": True,
            "html": html,
        }
    )


#######--DOCTORAL COMMITTEE VIEWS END--########

## Leo's Code End ##

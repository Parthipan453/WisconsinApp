from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.http import HttpResponse

import csv
from .forms import DepartmentForm
from .forms import DepartmentForm, CourseSectionForm
from .forms import DepartmentEditForm
from .models import Department, Course
from Admin.Colleges.models import School, AcademicProgram, Degree, University
from .forms import CourseForm, CourseEditForm
from openpyxl import Workbook
from django.http import HttpResponse
from django.db.models import Count
from openpyxl import Workbook
from django.http import HttpResponse
from django.db.models import Count
from django.http import HttpResponse
from reportlab.platypus import SimpleDocTemplate, Table
from .models import Department, Course, CourseDesignation, CourseDesignationType,CourseRequirement,CourseOffering,CourseLearningOutcome
from .forms import CourseForm, CourseEditForm, CourseDesignationTypeForm, CourseDesignationTypeEditForm,CourseOfferingForm,CourseLearningOutcomeForm
from django.http import JsonResponse
import json
from .forms import CourseForm, CourseEditForm ,CourseRequirementTypeForm
from django.http import JsonResponse
from .models import Course, CourseOffering
from .forms import CourseOfferingForm
from .forms import CourseLearningOutcomeForm
from .models import CourseLearningOutcome
from Students.models import CourseSection, Semester

from Admin.audit import AuditLogger


# ─── Department Dashboard ────────────────────────────────────────────────────

def department_dashboard(request):

    departments = (
        Department.objects
        .select_related('school')
        .annotate(course_count=Count('courses'))
        .order_by('department_name')
    )

    total_departments = Department.objects.count()

    active_departments = Department.objects.filter(
        status='ACTIVE'
    ).count()

    inactive_departments = Department.objects.filter(
        status='INACTIVE'
    ).count()

    total_schools = School.objects.count()

    total_courses = Course.objects.count()

    total_programs = AcademicProgram.objects.count()

    active_pct = (
        round(
            (active_departments / total_departments) * 100,
            1
        )
        if total_departments else 0
    )

    schools = School.objects.order_by(
        'school_name'
    )
    universities = University.objects.all()
    no_school = not schools.exists()

    schools_with_counts = School.objects.annotate(
        dept_count=Count('departments')
    )

    recent_programs = (
        AcademicProgram.objects
        .select_related(
            'department',
            'degree'
        )[:5]
    )

    paginator = Paginator(
        departments,
        10
    )

    page_number = request.GET.get('page')

    page_obj = paginator.get_page(
        page_number
    )

    context = {
        'departments': page_obj,

        'page_obj': page_obj,
        'paginator': paginator,
        'is_paginated': page_obj.has_other_pages(),

        'total_departments': total_departments,
        'active_departments': active_departments,
        'inactive_departments': inactive_departments,
        
        'total_schools': total_schools,
        'total_courses': total_courses,
        'total_programs': total_programs,

        'active_pct': active_pct,

        'schools': schools,
        'schools_with_counts': schools_with_counts,
        'no_school': no_school,
         "no_university": not universities.exists(),

        'recent_programs': recent_programs,
    }

    return render(
        request,
        'department/department_home.html',
        context
    )

def department_add(request):

    schools = School.objects.filter(status="ACTIVE").order_by("school_name")

    # Check whether any schools exist
    no_school = not schools.exists()
    universities = University.objects.all()
    no_university = not universities.exists()

    if request.method == "POST" and not no_school:

        print("POST DATA:")
        print(request.POST)

        form = DepartmentForm(request.POST)

        if form.is_valid():

            print("FORM VALID")

            department = form.save(commit=False)

            department.status = "ACTIVE"

            department.save()

            AuditLogger.log(
                 request=request,
                 action="CREATE",
                 module="Department",
                 object_type="Department",
                 object_id=department.department_id,
                 description=f"Created department '{department.department_name}'.",
                 after_data=AuditLogger.model_to_dict(
                 department,
                    [
                    "department_code",
                     "department_name",
                     "school",
                     "status",
                    ]
               ),
                status="SUCCESS",
                  )

            print("SAVED:", department)
            messages.success(
                request,
                "Department created successfully."
            )

            return redirect("department_dashboard")

        else:

            print("FORM ERRORS:")
            print(form.errors.as_json())

    else:

        form = DepartmentForm()

    context = {
        "form": form,
        "page_title": "Add Department",
        "schools": schools,
        "no_school": no_school,
        "no_university": no_university,
    }

    return render(
        request,
        "department/department_add.html",
        context
    )


def department_edit(request, department_uuid):

    department = get_object_or_404(
        Department,
        department_uuid=department_uuid
    )

    if request.method == "POST":

        form = DepartmentEditForm(
            request.POST,
            instance=department
        )

        if form.is_valid():
            before = AuditLogger.model_to_dict(
                department,
                [
                    "department_code",
                    "department_name",
                    "school",
                    "status",
                ]
            )
            department = form.save()

            AuditLogger.log(
                request=request,
                action="UPDATE",
                module="Department",
                object_type="Department",
                object_id=department.department_id,
                description=f"Updated department '{department.department_name}'.",
                before_data=before,
                after_data=AuditLogger.model_to_dict(
                    department,
                    [
                        "department_code",
                        "department_name",
                        "school",
                        "status",
                    ]
                ),
                status="SUCCESS",
            )

            messages.success(
                request,
                "Department updated successfully."
            )

            return redirect("department_dashboard")

    else:

        form = DepartmentEditForm(
            instance=department
        )

    context = {
        "form": form,
        "department": department,
        "schools": School.objects.filter(status="ACTIVE").order_by("school_name"),
        "page_title": "Edit Department"
    }

    return render(
        request,
        "department/edit_department.html",
        context
    )


def department_inactivate(request, department_uuid):

    department = get_object_or_404(
        Department,
        department_uuid=department_uuid
    )

    if request.method == "POST":
        before = AuditLogger.model_to_dict(
            department,
            [
                "department_code",
                "department_name",
                "school",
                "status",
            ]
        )
        department.status = "INACTIVE"

        department.save()
        AuditLogger.log(
            request=request,
            action="DELETE",
            module="Department",
            object_type="Department",
            object_id=department.department_id,
            description=f"Inactivated department '{department.department_name}'.",
            before_data=before,
            after_data=AuditLogger.model_to_dict(
                department,
                [
                    "department_code",
                    "department_name",
                    "school",
                    "status",
                ]
            ),
            status="SUCCESS",
        )

        messages.success(
            request,
            "Department marked as inactive."
        )

    return redirect("department_dashboard")


def department_view(request, department_uuid):
 
    department = get_object_or_404(
        Department,
        department_uuid=department_uuid
    )
  
    context = {
 
        "department": department,
 
        "course_count": department.courses.count(),
 
        "program_count":
        department.programs.count()
 
    }
 
    return render(
        request,
        "department/department_view.html",
        context
    )
 





def course_dashboard(request):

    courses = (
        Course.objects
        .select_related(
            'department',
            'department__school'
        )
        .order_by('course_name')
    )

    total_courses = Course.objects.count()

    active_courses = Course.objects.filter(
        status='ACTIVE'
    ).count()

    inactive_courses = Course.objects.filter(
        status='INACTIVE'
    ).count()

    active_pct = (
        round(
            (active_courses / total_courses) * 100,
            1
        )
        if total_courses else 0
    )

    departments = Department.objects.order_by(
        'department_name'
    )
    schools = School.objects.all()       
    no_department = not departments.exists()

    departments_with_counts = Department.objects.annotate(
        course_count=Count('courses')
    )

    total_departments = Department.objects.count()

    paginator = Paginator(
        courses,
        10
    )

    page_number = request.GET.get('page')

    page_obj = paginator.get_page(
        page_number
    )

    context = {

        'courses': page_obj,

        'page_obj': page_obj,
        'paginator': paginator,
        'is_paginated': page_obj.has_other_pages(),

        'total_courses': total_courses,
        'active_courses': active_courses,
        'inactive_courses': inactive_courses,

        'active_pct': active_pct,

        'departments': departments,
        "no_school": not schools.exists(),
        "no_department": no_department,
        'departments_with_counts': departments_with_counts,
        'total_departments': total_departments,
    }

    return render(
        request,
        'course/course_home.html',
        context
    )


def course_inactivate(request, course_uuid):

    course = get_object_or_404(
        Course,
        course_uuid=course_uuid
    )

    if request.method == "POST":
        before = AuditLogger.model_to_dict(
            course,
            [
                "course_code",
                "course_name",
                "department",
                "academic_program",
                "degree",
                "status",
            ]
        )

        course.status = "INACTIVE"

        course.save()
        AuditLogger.log(
            request=request,
            action="DELETE",
            module="Course",
            object_type="Course",
            object_id=course.course_id,
            description=f"Inactivated course '{course.course_name}' ({course.course_code}).",
            before_data=before,
            after_data=AuditLogger.model_to_dict(
                course,
                [
                    "course_code",
                    "course_name",
                    "department",
                    "academic_program",
                    "degree",
                    "status",
                ]
            ),
            status="SUCCESS",
        )

        messages.success(
            request,
            "Course marked as inactive."
        )

    return redirect("course_dashboard")


def course_add(request):

    if request.method == "POST":

        form = CourseForm(request.POST)

        if form.is_valid():

            

            course = form.save(commit=False)
            course.status = "ACTIVE"
            course.save()
            AuditLogger.log(
                request=request,
                action="CREATE",
                module="Course",
                object_type="Course",
                object_id=course.course_id,
                description=f"Created course '{course.course_name}' ({course.course_code}).",
                after_data=AuditLogger.model_to_dict(
                    course,
                    [
                        "course_code",
                        "course_name",
                        "department",
                        "academic_program",
                        "degree",
                        "status",
                    ]
                ),
                status="SUCCESS",
            )

        

            messages.success(
                request,
                "Course created successfully."
            )

            return redirect("course_dashboard")

        else:

            print("FORM ERRORS:", form.errors)

    else:

        form = CourseForm()

    return render(
        request,
        "course/course_add.html",
        {
            "form": form,
            "departments": Department.objects.filter(status="ACTIVE"),
            "academic_programs": AcademicProgram.objects.filter(status="ACTIVE"),
            "degrees": Degree.objects.filter(status="ACTIVE"),
        },
    )


# Speed code: course_edit view 



def course_edit(request, course_uuid):

    course = get_object_or_404(
        Course,
        course_uuid=course_uuid
    )
    current_course_number = int(course.course_code)

    available_courses = [
        c for c in Course.objects.filter(
            status="ACTIVE",
            department=course.department,
            academic_program=course.academic_program
        ).exclude(
            course_uuid=course.course_uuid
        ).order_by("course_code")
        if int(c.course_code) < current_course_number
]

    #ranganayagi code satrt
    offering_form = CourseOfferingForm()
    learning_outcome_form = CourseLearningOutcomeForm()
    learning_outcomes = course.learning_outcomes.all()
    if request.method == "POST":

        form = CourseEditForm(
            request.POST,
            instance=course
        )
        before = AuditLogger.model_to_dict(
            course,
            [
                "course_code",
                "course_name",
                "department",
                "academic_program",
                "degree",
                "status",
            ]
        )
        if form.is_valid():

            course = form.save()

            AuditLogger.log(
                request=request,
                action="UPDATE",
                module="Course",
                object_type="Course",
                object_id=course.course_id,
                description=f"Updated course '{course.course_name}' ({course.course_code}).",
                before_data=before,
                after_data=AuditLogger.model_to_dict(
                    course,
                    [
                        "course_code",
                        "course_name",
                        "department",
                        "academic_program",
                        "degree",
                        "status",
                    ]
                ),
                status="SUCCESS",
            )

            messages.success(
                request,
                "Course updated successfully."
            )

            return redirect("course_dashboard")

    else:

        form = CourseEditForm(
            instance=course
        )
    #ranganayagi code start
    offerings = course.offerings.all()
    last_taught = CourseOffering.objects.filter(
        course=course
    ).order_by("-year", "-created_at").first()

    learning_outcomes = CourseLearningOutcome.objects.filter(
        course=course
    ).order_by("display_order")    

    context = {

        # Existing
        "form": form,
        "course": course,
        "page_title": "Edit Course",

        # Dropdowns
        "departments": Department.objects.filter(
            status="ACTIVE"
        ),

        "requirement_types": CourseRequirementType.objects.filter(
            status="ACTIVE"
        ).order_by("name"),

        "designation_types": CourseDesignationType.objects.filter(
            status="ACTIVE"
        ).order_by("designation_name"),

        "available_courses": available_courses,
        

        # Existing linked data
        "requirements": course.requirements.select_related(
            "requirement_type",
            "related_course"
        ).order_by("display_order"),

        "designations": course.designations.select_related(
            "designation_type"
        ).order_by("display_order"),
        "departments": Department.objects.filter(status="ACTIVE"),
        "page_title": "Edit Course",
        "offering_form": offering_form,
    "learning_outcome_form": learning_outcome_form,
      "last_taught": last_taught,
    "learning_outcomes": learning_outcomes,
     "learning_outcomes": learning_outcomes,
    "offerings": offerings,
        "academic_programs": AcademicProgram.objects.filter(status="ACTIVE"),
        "degrees": Degree.objects.filter(status="ACTIVE"),
        "page_title": "Edit Course"
    }

    return render(
        request,
        "course/edit_course.html",
        context
    )


def course_view(request, course_uuid):

    course = get_object_or_404(
        Course.objects.select_related(
            'department__school__university',
            'academic_program',
            'degree'
        ),
        course_uuid=course_uuid
    )

    context = {
        "course": course,
        "requirements": course.requirements.select_related("requirement_type", "related_course"),
        "designations": course.designations.select_related("designation_type"),
        "offerings": course.offerings.all(),
        "learning_outcomes": course.learning_outcomes.all(),
    }

    return render(
        request,
        "course/course_view.html",
        context
    )

def course_designation(request):
 
    types_qs = CourseDesignationType.objects.annotate(
        linked_count=Count("course_designations")
    ).order_by("designation_name")
 
    paginator = Paginator(types_qs, 10)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)
 
    types_list = []
    for t in types_qs:
        types_list.append({
            "id": t.designation_type_id,
            "designation_name": t.designation_name,
            "description": t.description or "",
            "linked_count": t.linked_count,
            "status": t.status,
           
        })
 
    total_linked = CourseDesignation.objects.count()
    total_types = types_qs.count()
    active_types = types_qs.filter(status="ACTIVE").count()
    inactive_types = total_types - active_types
 
    context = {
        "types": page_obj,
        "page_obj": page_obj,
        "paginator": paginator,
        "is_paginated": page_obj.has_other_pages(),
        "types_json": json.dumps(types_list),
        "total_types": types_qs.count(),
        "total_linked": total_linked,
        "total_records": types_qs.count(),
        "active_main_nav": "course_designation",
        "active_types": active_types,
        "inactive_types": inactive_types,
    }
    return render(request, 'course/course_designation.html', context)
 

def designation_type_add(request):
    if request.method == "POST":
        data = request.POST.copy()
        data["status"] = "ACTIVE"
        form = CourseDesignationTypeForm(data)
        if form.is_valid():
            obj = form.save()
            return JsonResponse({"success": True, "id": obj.pk, "message": "Designation type created successfully."})
        else:
            return JsonResponse({"success": False, "errors": form.errors}, status=400)
    return JsonResponse({"error": "Invalid request"}, status=400)
 
 
def designation_type_edit(request, pk):
    obj = get_object_or_404(CourseDesignationType, pk=pk)
    if request.method == "POST":
        form = CourseDesignationTypeEditForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            return JsonResponse({"success": True, "message": "Designation type updated successfully."})
        else:
            return JsonResponse({"success": False, "errors": form.errors}, status=400)
    return JsonResponse({"error": "Invalid request"}, status=400)
 
 
 
def course_designation_dashboard(request):
 
    search = request.GET.get("search", "")
    designation_type = request.GET.get("designation_type", "")
 
    designations = CourseDesignation.objects.select_related(
        "course",
        "designation_type"
    ).order_by(
        "designation_type__designation_name",
        "display_order"
    )
 
    designation_chart = (
        CourseDesignationType.objects
        .annotate(total=Count("course_designations"))
        .order_by("-total")
    )
 
    chart_labels = [
        item.designation_name
        for item in designation_chart
    ]
 
    chart_values = [
        item.total
        for item in designation_chart
    ]
 
    from django.db.models import Count
 
    designation_stats = (
        CourseDesignationType.objects
        .annotate(total_records=Count("course_designations"))
        .order_by("-total_records")
    )
 
    if search:
        designations = designations.filter(
            Q(course__course_code__icontains=search) |
            Q(course__course_name__icontains=search)
        )
 
    if designation_type:
        designations = designations.filter(
            designation_type_id=designation_type
        )
 
    paginator = Paginator(designations, 15)
 
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)
 
    context = {
        "page_obj": page_obj,
        "designation_types": CourseDesignationType.objects.all(),
        "selected_designation": designation_type,
        "search": search,
         "designation_stats": designation_stats,
        "total_designations": CourseDesignation.objects.count(),
        "total_types": CourseDesignationType.objects.count(),
        "total_courses": CourseDesignation.objects.values("course").distinct().count(),
         "chart_labels": json.dumps(chart_labels),
    "chart_values": json.dumps(chart_values),
    }
 
    return render(
        request,
        "dashboard/course_designation.html",
        context
    )
 
 
 #ranganayagi code start
def export_designations_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Designation Types"

    ws.append(["#", "Designation Type", "Description", "Linked Courses", "Status"])

    types = CourseDesignationType.objects.annotate(
        linked_count=Count("course_designations")
    ).order_by("designation_name")

    for idx, t in enumerate(types, start=1):
        ws.append([
            idx,
            t.designation_name,
            t.description or "",
            t.linked_count,
            t.status
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=designation_types.xlsx"
    wb.save(response)
    return response


def export_designations_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.units import inch
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import TableStyle, Paragraph

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="designation_types.pdf"'

    pw, ph = landscape(A4)
    doc = SimpleDocTemplate(response, pagesize=landscape(A4),
                            leftMargin=30, rightMargin=30,
                            topMargin=30, bottomMargin=30)

    styles = getSampleStyleSheet()
    styleN = styles["Normal"]
    styleN.fontSize = 8
    styleN.leading = 11
    styleN.spaceBefore = 0
    styleN.spaceAfter = 0

    data = [["#", "Designation Type", "Description", "Linked", "Status"]]

    types = CourseDesignationType.objects.annotate(
        linked_count=Count("course_designations")
    ).order_by("designation_name")

    for idx, t in enumerate(types, start=1):
        data.append([
            str(idx),
            Paragraph(t.designation_name, styleN),
            Paragraph(t.description or "—", styleN),
            str(t.linked_count),
            Paragraph(t.status, styleN),
        ])

    col_widths = [30, round(pw * 0.18), round(pw * 0.52), round(pw * 0.07), round(pw * 0.07)]

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563eb")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (3, 0), (4, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))

    doc.build([table])
    return response
 
 
def designation_type_delete(request, pk):
    if request.method == "POST":
        obj = get_object_or_404(CourseDesignationType, pk=pk)
        name = obj.designation_name
        obj.delete()
        return JsonResponse({"success": True, "message": f'Designation type "{name}" deleted.'})
    return JsonResponse({"error": "Invalid request"}, status=400)


def designation_type_inactivate(request, pk):
    obj = get_object_or_404(CourseDesignationType, pk=pk)
    if request.method != "POST":
        return redirect("course_designation")
    if obj.status != "INACTIVE":
        obj.status = "INACTIVE"
        obj.save(update_fields=["status"])
        return JsonResponse({"success": True, "message": f'"{obj.designation_name}" marked as inactive.'})
    return JsonResponse({"success": False, "message": "Already inactive."}, status=400)


####################################bela code start ######################################
from .models import CourseRequirementType
def course_requirement_type_dashboard(request):

    qs = (
        CourseRequirementType.objects
        .annotate(
            usage_count=Count("course_requirements")
        )
        .order_by("name")
    )

    total = qs.count()
    active = qs.filter(status="ACTIVE").count()
    inactive = total - active

    add_form = CourseRequirementTypeForm()

    context = {
        "requirement_types": qs,
        "total": total,
        "active": active,
        "inactive": inactive,
        "add_form": add_form,
    }

    return render(
        request,
        "course/courserequirement_type_dashboard.html",
        context,
    )
 

def course_requirement_type_add(request):

    if request.method != "POST":
        return redirect("course_requirement_type_dashboard")

    data = request.POST.copy()
    data["status"] = "ACTIVE"

    form = CourseRequirementTypeForm(data)

    if form.is_valid():

        form.save()

        messages.success(
            request,
            "Requirement type added successfully."
        )

        return redirect("course_requirement_type_dashboard")

    # If validation fails, reload dashboard with the invalid form
    qs = (
        CourseRequirementType.objects
        .annotate(
            usage_count=Count("course_requirements")
        )
        .order_by("name")
    )

    total = qs.count()
    active = qs.filter(status="ACTIVE").count()
    inactive = total - active

    context = {
        "requirement_types": qs,
        "total": total,
        "active": active,
        "inactive": inactive,
        "add_form": form,          # Invalid form with errors
        "show_add_modal": True,    # Reopen modal
    }

    return render(
        request,
        "course/courserequirement_type_dashboard.html",
        context,
    )
 
def course_requirement_type_edit(request, pk):

    obj = get_object_or_404(
        CourseRequirementType,
        pk=pk
    )

    if request.method != "POST":
        return redirect(
            "course_requirement_type_dashboard"
        )

    form = CourseRequirementTypeForm(
        request.POST,
        instance=obj
    )

    if form.is_valid():

        updated = form.save()

        messages.success(
            request,
            f'"{updated.name}" updated successfully.'
        )

    else:

        messages.error(
            request,
            "Please fix the errors below."
        )

    return redirect(
        "course_requirement_type_dashboard"
    )
 
 
def course_requirement_type_inactivate(request, pk):

    obj = get_object_or_404(
        CourseRequirementType,
        pk=pk
    )

    if request.method != "POST":
        return redirect(
            "course_requirement_type_dashboard"
        )

    if obj.status != "INACTIVE":

        obj.status = "INACTIVE"

        obj.save(update_fields=["status"])

        messages.success(
            request,
            f'"{obj.name}" marked as inactive.'
        )

    return redirect(
        "course_requirement_type_dashboard"
    )



#ranganayagi code start

from django.views.decorators.http import require_POST
 
# ── Add Requirement ──────────────────────────────────────────
@require_POST
def course_requirement_add(request, course_uuid):
    course = get_object_or_404(Course, course_uuid=course_uuid)
 
    type_id   = request.POST.get('requirement_type_id')
    rc_id     = request.POST.get('related_course_id') or None
    note      = request.POST.get('note', '').strip() or None
    order     = request.POST.get('display_order', 1)
 
    if not type_id:
        return JsonResponse({'success': False, 'error': 'Requirement type is required.'})
 
    req_type = get_object_or_404(CourseRequirementType, pk=type_id)
 
    related_course = None
    if rc_id:
        related_course = get_object_or_404(Course, pk=rc_id)

    if CourseRequirement.objects.filter(
        course=course,
        requirement_type=req_type,
        related_course=related_course
    ).exists():

        return JsonResponse({
            "success": False,
            "error": "This prerequisite is already linked to this course."
        })
 
    req = CourseRequirement.objects.create(
        course=course,
        requirement_type=req_type,
        related_course=related_course,
        note=note,
        display_order=int(order),
    )
 
    return JsonResponse({
        'success': True,
        'requirement_id':       req.requirement_id,
        'requirement_type_name': req_type.name,
        'related_course_code':  related_course.course_code if related_course else None,
        'related_course_name':  related_course.course_name if related_course else None,
        'note':                 req.note or '',
        'display_order':        req.display_order,
    })


def course_requirement_type_view(request, pk):

    obj = get_object_or_404(
        CourseRequirementType,
        pk=pk
    )

    linked_courses = (
        CourseRequirement.objects
        .filter(
            requirement_type=obj,
            related_course__isnull=False
        )
        .select_related("related_course")
        .order_by("related_course__course_code")
    )

    return JsonResponse({

        "name": obj.name,
        "description": obj.description or "",
        "status": obj.status,
        "created": obj.created_at.strftime("%b %d, %Y %I:%M %p"),
        "updated": obj.updated_at.strftime("%b %d, %Y %I:%M %p"),

        "linked_courses": [
            {
                "code": req.related_course.course_code,
                "name": req.related_course.course_name,
            }
            for req in linked_courses
        ]

    })
 
# ── Delete Requirement ───────────────────────────────────────
@require_POST
def course_requirement_delete(request, requirement_id):
    req = get_object_or_404(CourseRequirement, pk=requirement_id)
    req.delete()
    return JsonResponse({'success': True})
 
 
# ── Add Designation ──────────────────────────────────────────
@require_POST
def course_designation_add(request, course_uuid):
    course = get_object_or_404(Course, course_uuid=course_uuid)
 
    type_id = request.POST.get('designation_type_id')
    value   = request.POST.get('designation_value', '').strip() or None
    order   = request.POST.get('display_order', 1)
 
    if not type_id:
        return JsonResponse({'success': False, 'error': 'Designation type is required.'})
 
    desg_type = get_object_or_404(CourseDesignationType, pk=type_id)

    # Prevent duplicate designation
    if CourseDesignation.objects.filter(
        course=course,
        designation_type=desg_type
    ).exists():

        return JsonResponse({
            "success": False,
            "error": "This designation is already linked to this course."
        })
 
    desg = CourseDesignation.objects.create(
        course=course,
        designation_type=desg_type,
        designation_value=value,
        display_order=int(order),
    )
 
    return JsonResponse({
        'success': True,
        'designation_id':        desg.designation_id,
        'designation_type_name': desg_type.designation_name,
        'designation_value':     desg.designation_value or '',
        'display_order':         desg.display_order,
    })
 
 
# ── Delete Designation ───────────────────────────────────────
@require_POST
def course_designation_delete(request, designation_id):
    desg = get_object_or_404(CourseDesignation, pk=designation_id)
    desg.delete()
    return JsonResponse({'success': True})


@require_POST
def course_offering_delete(request, offering_id):
    try:
        offering = get_object_or_404(CourseOffering, pk=offering_id)
        offering.delete()
        return JsonResponse({"success": True})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({"success": False, "error": str(e)}, status=500)


@require_POST
def course_learning_outcome_delete(request, outcome_id):
    try:
        outcome = get_object_or_404(CourseLearningOutcome, pk=outcome_id)
        outcome.delete()
        return JsonResponse({"success": True})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({"success": False, "error": str(e)}, status=500)


def course_offering_add(request, course_uuid):
    try:
        course = get_object_or_404(
            Course,
            course_uuid=course_uuid
        )

        if request.method != "POST":
            return JsonResponse({"success": False})

        form = CourseOfferingForm(request.POST)

        if form.is_valid():

            offering = form.save(commit=False)
            offering.course = course
            offering.save()

            return JsonResponse({
                "success": True,
                "id": offering.offering_id,
                "term": offering.term,
                "year": offering.year,
            })

        return JsonResponse({
            "success": False,
            "errors": form.errors,
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({"success": False, "error": str(e)}, status=500)


def course_learning_outcome_add(request, course_uuid):
    try:
        course = get_object_or_404(
            Course,
            course_uuid=course_uuid
        )

        if request.method != "POST":
            return JsonResponse({"success": False})

        form = CourseLearningOutcomeForm(request.POST)

        if form.is_valid():

            outcome = form.save(commit=False)
            outcome.course = course
            outcome.save()

            return JsonResponse({
                "success": True,
                "id": outcome.outcome_id,
                "outcome": outcome.outcome,
            })

        return JsonResponse({
            "success": False,
            "errors": form.errors,
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({"success": False, "error": str(e)}, status=500)



from django.db.models import Count
from django.core.paginator import Paginator

def course_section_dashboard(request):

    section_list = (
        CourseSection.objects
        .select_related("course", "semester_id")
        .order_by("section_id")
    )

    paginator = Paginator(section_list, 10)   # 10 records per page
    page_number = request.GET.get("page")
    sections = paginator.get_page(page_number)

    form = CourseSectionForm()

    departments = Department.objects.filter(status="ACTIVE").order_by("department_name")

    context = {
        "sections": sections,
        "form": form,
        "departments": departments,
        "total_sections": CourseSection.objects.count(),
        "total_courses": Course.objects.count(),
        "total_semesters": Semester.objects.count(),
        "total_section_types": CourseSection.objects.values('section_type').distinct().count(),
    }

    return render(
        request,
        "course/course_section.html",
        context,
    )


def course_section_add(request):

    if request.method == "POST":

        form = CourseSectionForm(request.POST)

        if form.is_valid():

            obj = form.save(commit=False)
            faculty_name = request.POST.get("faculty_name", "").strip()
            obj.faculty_name = faculty_name
            obj.save()

            return JsonResponse({
                "success": True,
                "section_id": obj.section_id
            })

        return JsonResponse({
            "success": False,
            "errors": form.errors
        })

    return JsonResponse({
        "success": False
    })


from django.shortcuts import get_object_or_404
from django.http import JsonResponse


def course_section_view(request, section_id):

    section = get_object_or_404(
        CourseSection.objects.select_related("course", "semester_id"),
        pk=section_id
    )

    return JsonResponse({

        "course": f"{section.course.course_code} - {section.course.course_name}",
        "section_number": section.section_number,
        "section_type": section.get_section_type_display(),
        "semester": str(section.semester_id),
        "capacity": section.capacity,
        "faculty_name": section.faculty_name or "TBA",

    })



from django.shortcuts import get_object_or_404
from django.http import JsonResponse
def course_section_edit(request, section_id):

    section = get_object_or_404(
        CourseSection,
        section_id=section_id
    )

    if request.method == "POST":

        print(request.POST)  

        form = CourseSectionForm(
            request.POST,
            instance=section
        )

        if form.is_valid():

            print("VALID")   

            obj = form.save(commit=False)
            faculty_name = request.POST.get("faculty_name", "").strip()
            obj.faculty_name = faculty_name
            obj.save()

            return JsonResponse({
                "success": True
            })

        print(form.errors) 

        return JsonResponse({
            "success": False,
            "errors": form.errors
        })

    return JsonResponse({
        "section_id": section.section_id,
        "course": section.course_id,
        "section_number": section.section_number,
        "section_type": section.section_type,
        "semester": section.semester_id.get_semester_type_display(),
        "capacity": section.capacity,
        "faculty_name": section.faculty_name or "",
        "department_id": section.course.department_id if section.course else "",
    })


from django.shortcuts import get_object_or_404
from django.http import JsonResponse

def course_section_delete(request, section_id):

    print(request.method, section_id) 

    if request.method == "POST":

        section = get_object_or_404(
            CourseSection,
            section_id=section_id
        )

        section.delete()

        return JsonResponse({
            "success": True
        })

    return JsonResponse({
        "success": False
    })

from openpyxl import Workbook
from django.http import HttpResponse


def export_course_sections_excel(request):

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Course Sections"

    headers = [
        "Course",
        "Section Number",
        "Section Type",
        "Semester",
        "Capacity",
    ]

    for col, header in enumerate(headers, 1):
        worksheet.cell(row=1, column=col).value = header

    sections = (
        CourseSection.objects
        .select_related("course", "semester_id")
        .order_by("course__course_code")
    )

    row = 2

    for section in sections:

        worksheet.cell(row=row, column=1).value = (
            section.course.course_code if section.course else ""
        )

        worksheet.cell(row=row, column=2).value = (
            section.section_number
        )

        worksheet.cell(row=row, column=3).value = (
            section.get_section_type_display()
        )

        worksheet.cell(row=row, column=4).value = (
            str(section.semester_id)
        )

        worksheet.cell(row=row, column=5).value = (
            section.capacity
        )

        row += 1

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response[
        "Content-Disposition"
    ] = 'attachment; filename="Course_Sections.xlsx"'

    workbook.save(response)

    return response

from django.http import HttpResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle


def export_course_sections_pdf(request):

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="Course_Sections.pdf"'

    doc = SimpleDocTemplate(response, pagesize=letter)

    data = [
        [
            "Course",
            "Section Number",
            "Section Type",
            "Semester",
            "Capacity"
        ]
    ]

    sections = (
        CourseSection.objects
        .select_related("course", "semester_id")
        .order_by("course__course_code")
    )

    for section in sections:

        data.append([
            section.course.course_code if section.course else "",
            section.section_number,
            section.get_section_type_display(),
            str(section.semester_id),
            section.capacity,
        ])

    table = Table(data)

    table.setStyle(TableStyle([

        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#2563eb")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),

        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 10),

        ("GRID", (0,0), (-1,-1), 1, colors.black),

        ("BACKGROUND", (0,1), (-1,-1), colors.beige),

        ("ALIGN", (0,0), (-1,-1), "CENTER"),

        ("BOTTOMPADDING", (0,0), (-1,0), 10),

    ]))

    elements = [table]

    doc.build(elements)

    return response


from django.http import JsonResponse

def get_department_courses(request):
    department_id = request.GET.get("department_id")
    courses = Course.objects.filter(
        department_id=department_id,
        status="ACTIVE",
    ).order_by("course_code")
    data = [
        {
            "id": course.course_id,
            "label": f"{course.course_code} - {course.course_name}",
        }
        for course in courses
    ]
    return JsonResponse(data, safe=False)


def get_department_faculty(request):
    from Faculty.models import FacultyProfile
    department_id = request.GET.get("department_id")
    faculty_members = FacultyProfile.objects.filter(
        department_id=department_id,
        employment_status="ACTIVE",
        user__is_faculty=True,
    ).select_related("user")
    data = [
        {
            "id": faculty.user.id,
            "name": faculty.user.full_name,
        }
        for faculty in faculty_members
    ]
    return JsonResponse(data, safe=False)


def get_department_programs(request, department_id):

    programs = AcademicProgram.objects.filter(
        department_id=department_id,
        status="ACTIVE"
    ).order_by("program_name")

    data = [
        {
            "value": program.program_id,
            "label": program.program_name,
        }
        for program in programs
    ]

    return JsonResponse(data, safe=False)














#<-----------------Blaze App Start(23.09.26)---------------->




from rest_framework.decorators import api_view
from rest_framework.response import Response
from Admin.Colleges.models import AcademicProgram
from django.db.models import Q



@api_view(['GET'])
def programs_api(request):
    """
    Returns list of active academic programs.
    Supports filtering via query params:
      - search: filter by name/code/department
      - program_type: FULL_TIME / PART_TIME / ONLINE
      - interest: filter by area of interest ID (can be multiple)
    """
    try:
        programs = AcademicProgram.objects.select_related(
            'department', 'degree'
        ).filter(status='ACTIVE').order_by('program_name')
        
        # Search filter
        search = request.GET.get('search', '').strip()
        if search:
            programs = programs.filter(
                Q(program_name__icontains=search) |
                Q(program_code__icontains=search) |
                Q(department__department_name__icontains=search)
            )
        
        # Program type filter
        program_type = request.GET.get('program_type', '').strip()
        if program_type and program_type != 'all':
            programs = programs.filter(program_type=program_type.upper())
        
        # Interests filter (multiple)
        interests = request.GET.getlist('interest')
        if interests:
            programs = programs.filter(
                areas_of_interest__id__in=interests
            ).distinct()
        
        data = []
        for p in programs:
            data.append({
                'id': p.program_id,
                'code': p.program_code or '',
                'name': p.program_name,
                'department': p.department.department_name if p.department else 'N/A',
                'department_code': p.department.department_code if p.department else 'N/A',
                'degree': p.degree.degree_name if p.degree else 'N/A',
                'degree_code': p.degree.degree_code if p.degree else 'N/A',
                'degree_level': p.degree.level if p.degree else 'N/A',
                'program_type': p.get_program_type_display(),
                'program_type_raw': p.program_type,  # For filtering
                'duration': p.duration,
                'total_credits': p.total_credits,
                'description': p.description or '',
                'status': p.status,
            })
        
        return Response(data)
    
    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({'error': str(e)}, status=500)




# ═══════════════════════════════════════════════════════════════
# MOBILE APP APIs — Areas of Interest
# ═══════════════════════════════════════════════════════════════

from Admin.Colleges.models import AreaOfInterest

@api_view(['GET'])
def interests_api(request):
    """
    Returns list of active areas of interest.
    Used by React Native mobile app for filtering.
    """
    try:
        interests = AreaOfInterest.objects.filter(
            status='ACTIVE'
        ).order_by('interest_name')
        
        data = []
        for i in interests:
            data.append({
                'id': i.id,
                'name': i.interest_name,
            })
        
        return Response(data)
    
    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({'error': str(e)}, status=500)

#<-----------------Blaze App End(23.09.26)---------------->
